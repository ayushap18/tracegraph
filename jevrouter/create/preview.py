"""A light preview of a created file for the file card (FilePreview in web/src/protocol.ts): the Markdown text, the
outline of a PDF, DOCX or PPTX (headings or slide titles), or the first rows of each sheet."""
import datetime
import io
import math

MD_CHARS = 200_000
SHEET_ROWS = 20


def _value(v):
    if v is None or isinstance(v, str):
        return v
    if isinstance(v, bool):
        return str(v).upper()
    if isinstance(v, (int, float)):
        return v if math.isfinite(v) else None
    if isinstance(v, (datetime.date, datetime.time)):
        return v.isoformat()
    return str(v)


def preview(fmt: str, data: bytes) -> dict:
    if fmt == 'md':
        return {'kind': 'markdown', 'text': data.decode('utf-8', 'replace')[:MD_CHARS]}
    if fmt == 'pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        items = []

        def walk(nodes, depth):
            for n in nodes:
                if isinstance(n, list):
                    walk(n, depth + 1)
                else:
                    items.append({'level': depth, 'text': str(n.title)})
        walk(reader.outline, 0)
        return {'kind': 'outline', 'items': items, 'pages': len(reader.pages), 'slides': None}
    if fmt == 'docx':
        from docx import Document
        doc = Document(io.BytesIO(data))
        items = []
        for p in doc.paragraphs:
            name = p.style.name if p.style is not None else ''
            if name == 'Title':
                items.append({'level': 0, 'text': p.text})
            elif name.startswith('Heading ') and name[8:].isdigit():
                items.append({'level': int(name[8:]), 'text': p.text})
        return {'kind': 'outline', 'items': items, 'pages': None, 'slides': None}
    if fmt == 'pptx':
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        items = []
        for i, slide in enumerate(prs.slides):
            title = slide.shapes.title.text_frame.text if slide.shapes.title is not None else f'Slide {i + 1}'
            items.append({'level': 0 if i == 0 else 1, 'text': title})
        return {'kind': 'outline', 'items': items, 'pages': None, 'slides': len(items)}
    if fmt == 'xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True)
        sheets = []
        try:
            for ws in wb.worksheets:
                rows = [list(r) for r in ws.iter_rows(max_row=SHEET_ROWS + 1, values_only=True)]
                head = rows[0] if rows else []
                total = max((ws.max_row or len(rows)) - 1, 0)
                sheets.append({'name': ws.title, 'columns': ['' if v is None else str(_value(v)) for v in head],
                               'rows': [[_value(v) for v in r] for r in rows[1:]], 'total_rows': total})
        finally:
            wb.close()
        return {'kind': 'sheets', 'sheets': sheets}
    raise ValueError(f'no preview for {fmt!r}')
