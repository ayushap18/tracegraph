"""The ruleset for created files (docs/RULES-files.md): the rule list served by GET /api/rules, the RuleResult each
check produces, and verify(), which reopens a rendered file with its own library and checks it (V1-V4, L4, plus the
per-format F, X and A rules that can be read back from the file). Given the request's brief
(docs/PLAN-accuracy-v2.md C7), it also checks the file against it: pages or slides (V5), images (V6), font (V7), theme
and the black-and-white scan (V8) and diagrams (V9), and every embedded image's origin and credit (X6)."""
import io
import re
import zipfile
from dataclasses import asdict, dataclass

MAX_BYTES = 15 * 1024 * 1024


@dataclass
class RuleResult:
    id: str
    severity: str  # block | fix | warn
    ok: bool
    note: str

    def to_dict(self) -> dict:
        return asdict(self)


class SpecError(ValueError):
    """A block rule failed: the file is not produced. rule_id says which rule."""

    def __init__(self, rule_id: str, message: str):
        super().__init__(f'{rule_id}: {message}')
        self.rule_id, self.message = rule_id, message


_GROUPS = {'S': 'Content spec', 'L': 'Size limits', 'F': 'Structure and style', 'X': 'Safety',
           'V': 'Verification', 'A': 'Accessibility', 'D': 'Design'}


def _rule(rid: str, severity: str | None, text: str) -> dict:
    return {'id': rid, 'group': _GROUPS[rid[0]], 'text': text, 'severity': severity, 'enforced': True}


RULES: list[dict] = [
    _rule('S1', 'fix', 'Each block is repaired to the DocSpec schema: known aliases are mapped (list, ul, ol to bullets; '
                       'content or body to text; points or lines to items), a missing type is inferred from its keys, and '
                       'a string where a list is expected is split into lines. A block that still cannot be read keeps its '
                       'text as a paragraph; a block with no text is left out. Every change is noted.'),
    _rule('S2', 'block', 'The file has something to show: at least one section with a heading or a non-empty block, not '
                         'counting page breaks and figures that found no image. A spec with no text at all is refused.'),
    _rule('S3', 'fix', 'A title of 1 to 120 characters. Missing: taken from the first heading or the request.'),
    _rule('S4', 'fix', 'Text is plain: no HTML, no raw Markdown syntax inside paragraphs (converted to plain text, with '
                       'bold and italic kept as runs where the format supports them).'),
    _rule('S5', 'fix', 'Numbers in tables and charts are numbers, not strings ("1,200" becomes 1200; "$1,200.50" and "12%" '
                       'become numbers that keep their money or percent format). Non-numeric chart '
                       'values drop that point.'),
    _rule('S6', 'fix', 'Every table row has as many cells as there are columns (short rows padded, long rows trimmed). '
                       'Diagrams are well formed: one root, no loops or cycles, known node names. Column names are at '
                       'most 80 characters, cut at a word.'),
    _rule('S7', 'fix', 'Nothing is fetched: link, image and path keys (url, src, href, image, link, path and the like) are '
                       'removed from the spec with a note, and the file holds only the content given.'),
    _rule('S8', 'fix', 'Every section has a heading: a missing one is taken from its first line of text, or becomes '
                       '"Section N" (slides: the previous slide\'s heading with "(cont.)"). A section given as plain text '
                       'becomes a heading-only section, or a paragraph when it is longer than 80 characters.'),
    _rule('L1', 'fix', 'At most 200 KB of spec JSON (table rows are limited by L2), 40 sections and 30 blocks per section. '
                       'Blocks past 30 continue in a section headed "<heading> (cont.)"; sections past 40, or past 200 KB, '
                       'are cut from the end and the note says how many.'),
    _rule('L2', 'fix', 'Tables: at most 2,000 rows and 30 columns in XLSX; 200 rows in PDF, DOCX and MD (the rest noted); '
                       '12 rows per slide (split across slides). Diagrams: at most 30 timeline events, 40 tree nodes in '
                       '4 levels, 12 flow steps.'),
    _rule('L3', 'fix', 'Slides: at most 6 bullets per slide and 18 words per bullet; longer content moves to the slide '
                       'notes.'),
    _rule('L4', 'block', 'Output file at most 15 MB.'),
    _rule('L5', None, 'The model is asked for the spec once per request; conversions to another format reuse the stored '
                      'spec (0 LLM tokens).'),
    _rule('L6', 'block', 'A model reply larger than 2 MB is not read as a file spec.'),
    _rule('F1', 'fix', 'PDF: A4 (Letter when asked for), 2 cm margins, page numbers, the title in the document '
                       'properties, headings as PDF outline entries.'),
    _rule('F2', 'fix', 'DOCX: built-in Heading 1 to 3 styles (so a table of contents works), real Word tables with a '
                       'header row, core properties set. The subject property holds at most 255 characters of the '
                       'subtitle, cut at a word.'),
    _rule('F3', 'fix', 'PPTX: 16:9; a title slide, then one slide per section; native charts (not pictures); speaker '
                       'notes for overflow text. The subject property holds at most 255 characters of the subtitle, cut '
                       'at a word.'),
    _rule('F4', 'fix', 'XLSX: one sheet per table (sheet names at most 31 characters, unique, none of []:*?/\\), a bold '
                       'frozen header row, column widths fitted, numbers stored as numbers, native charts next to '
                       'their data.'),
    _rule('F5', 'fix', 'Markdown: CommonMark with GFM tables; one # title; headings never skip a level.'),
    _rule('F6', 'fix', 'One theme per file (font, colours); body text at least 10 pt in PDF and DOCX and 18 pt on '
                       'slides.'),
    _rule('F7', 'fix', 'File names: lowercase words joined by hyphens from the title, the right extension, at most 80 '
                       'characters, no path separators.'),
    _rule('X1', 'block', 'No macros or active content: never .docm, .xlsm or .pptm, no embedded scripts, no OLE objects.'),
    _rule('X2', 'fix', 'Spreadsheet formula injection: a cell whose text starts with =, +, -, @, tab or CR is stored as '
                       'text with a leading apostrophe, unless it is a formula the spec marked as one, built only from '
                       'cell references and SUM, AVERAGE, MIN, MAX, COUNT or ROUND.'),
    _rule('X3', 'fix', 'No remote references; images only as embedded bytes from the asset cache. No external links in '
                       'formulas, no linked objects. Web addresses appear only as text.'),
    _rule('X4', 'block', 'Content Jev would block is not written to a file: the request goes through the same safety '
                         'check as any question, and so does the spec text.'),
    _rule('X5', None, 'Created files are stored like uploads (by id, never by user path). Sandbox files stay in memory '
                      'and disappear with the sandbox.'),
    _rule('X6', 'fix', 'Every embedded image is a PNG or JPEG from the local asset cache, re-encoded, found under an '
                       'allowed licence (public domain, CC0, CC BY, CC BY-SA), with a credit line naming its title, '
                       'author, licence and source. An image without a full credit is left out, never embedded.'),
    _rule('V1', 'block', 'The file reopens with its own library (pypdf, python-docx, python-pptx, openpyxl) or parses as '
                         'Markdown.'),
    _rule('V2', 'warn', 'The title and every section heading appear in the reopened file\'s text.'),
    _rule('V3', 'warn', 'Page, slide or sheet counts match the spec (after the L2 and L3 splits).'),
    _rule('V4', 'warn', 'Every chart and diagram in the spec exists in PPTX or XLSX, is drawn in PDF, or appears as a '
                        'table or picture in DOCX and MD.'),
    _rule('V5', 'warn', 'Pages (or slides) are within the count the request asked for.'),
    _rule('V6', 'warn', 'When images were asked for, at least one licensed image is embedded.'),
    _rule('V7', 'warn', 'The requested font was used, or the answer names the font used instead.'),
    _rule('V8', 'warn', 'The file uses the theme the request asked for; a black and white PDF draws only greys, and its '
                        'images are greyscale.'),
    _rule('V9', 'warn', 'At least as many diagrams are drawn as the request asked for (2 for "multiple diagrams").'),
    _rule('V10', 'warn', 'When a design file was applied, the file uses its background and text colours (or the answer '
                         'says what could not be used).'),
    _rule('V11', 'fix', 'A section the renderer cannot lay out in this format is retried as plain text, then left out, '
                        'and the answer names it.'),
    _rule('V12', 'warn', 'The file holds every planned section, or the answer lists the missing ones and offers Resume.'),
    _rule('A1', 'fix', 'Headings are real headings (styles, outline entries, slide titles), never bold paragraphs.'),
    _rule('A2', 'fix', 'Tables have a header row marked as a header.'),
    _rule('A3', 'warn', 'Charts carry a text title, and a one-line summary of what they show appears next to them.'),
    _rule('A4', 'fix', 'Colour is never the only way information is shown; theme colours meet 4.5:1 contrast for text.'),
    # docs/PLAN-designer.md 9.8: Studio's visual QA, read from the design report of a designed file (absent otherwise)
    _rule('D1', 'warn', 'Overflow: no text is taller than its box (0.5 pt of slack) and no line is wider than its box.'),
    _rule('D2', 'warn', 'Overlap: no two boxes overlap by more than 1 square point, unless the upper one is meant to sit '
                        'on top (text on an overlay, art behind type).'),
    _rule('D3', 'warn', 'Readability: text is at least the minimum size (slides 18 pt, captions 12 pt; print 10 pt, '
                        'captions 8 pt) and meets 4.5:1 contrast (3:1 for large text), measured against the photo under '
                        'it too.'),
    _rule('D4', 'warn', 'Density: a slide holds at most 40 words of bullets or 60 of prose, and 25 to 60 percent of each '
                        'page is white space (15 to 60 percent in print).'),
    _rule('D5', 'warn', 'Balance: the visual weight of a page sits in its middle third both ways (asymmetric layouts, '
                        'covers and freeform pages are skipped).'),
    _rule('D6', 'warn', 'Consistency: text sizes are on the type scale, text left edges sit on grid columns, and the file '
                        'uses one image treatment.'),
    _rule('D7', 'warn', 'Variety: no layout is used on more than 3 slides in a row, and a deck of 8 or more slides uses at '
                        'least 4 layouts.'),
    _rule('D8', 'warn', 'Images: no picture is enlarged more than 1.5 times (at 96 dpi on slides, 150 dpi in print) and '
                        'its focal point stays inside the crop.'),
]
DESIGN_RULES = ('D1', 'D2', 'D3', 'D4', 'D5', 'D6', 'D7', 'D8')


# ---------- verify ----------


def _norm(text: str) -> str:
    s = re.sub(r'[\u200b-\u200d\ufeff]', '', str(text)).replace('**', '').replace('*', '')
    return re.sub(r'\s+', ' ', s).strip().lower()


def _found(needle: str, corpus: str) -> bool:
    n = _norm(needle)
    if not n or n in corpus:
        return True
    squashed = corpus.replace(' ', '')
    return n.replace(' ', '') in squashed  # PDF text extraction can drop or add spaces at line breaks


def _zip_checks(data: bytes) -> tuple[list[str], list[str]]:
    """(active content found, external references found) inside an OOXML package (X1, X3)."""
    active, external = [], []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for n in z.namelist():
            low = n.lower()
            if 'vbaproject' in low or 'activex' in low or 'oleobject' in low or ('embeddings/' in low and low.endswith('.bin')):
                active.append(n)
            if low.endswith('.rels'):
                rels = z.read(n).decode('utf-8', 'replace')
                if 'TargetMode="External"' in rels:
                    external.append(n)
        types = z.read('[Content_Types].xml').decode('utf-8', 'replace') if '[Content_Types].xml' in z.namelist() else ''
        if 'macroEnabled' in types or 'vbaProject' in types:
            active.append('[Content_Types].xml')
    return active, external


def _count_charts(data: bytes, folder: str) -> tuple[int, list[str]]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = [n for n in z.namelist() if re.match(rf'{folder}/charts/chart\d+\.xml$', n)]
        titles = [z.read(n).decode('utf-8', 'replace') for n in names]
    return len(names), titles


_COLOR_OP = re.compile(rb'(?<![\w.])(-?\d*\.?\d+)\s+(-?\d*\.?\d+)\s+(-?\d*\.?\d+)\s+(rg|RG)(?![\w])')


def pdf_scan(data: bytes, reader=None) -> dict:
    """What a PDF draws, read from its pages: {images: [colour space of each image XObject], fonts: [BaseFont names],
    colors: [(r, g, b) of every rg/RG operator], gray: every colour is a grey and every image DeviceGray}."""
    from pypdf import PdfReader
    reader = reader or PdfReader(io.BytesIO(data))
    images, fonts, colors, seen = [], set(), [], set()

    def resources(res):
        res = res.get_object() if res is not None else None
        if not res:
            return
        for name, f in (res.get('/Font') or {}).items():
            f = f.get_object()
            if f.get('/BaseFont'):
                fonts.add(str(f['/BaseFont']).lstrip('/'))
        for name, x in (res.get('/XObject') or {}).items():
            ref = getattr(x, 'idnum', None)
            x = x.get_object()
            if ref is not None:
                if ref in seen:
                    continue
                seen.add(ref)
            if x.get('/Subtype') == '/Image':
                cs = x.get('/ColorSpace')
                cs = cs.get_object() if hasattr(cs, 'get_object') else cs
                images.append(str(cs[0] if isinstance(cs, list) else cs))
            elif x.get('/Subtype') == '/Form':
                content(x.get_data())
                resources(x.get('/Resources'))

    def content(raw: bytes):
        for m in _COLOR_OP.finditer(raw or b''):
            colors.append(tuple(float(v) for v in m.groups()[:3]))
    for page in reader.pages:
        c = page.get_contents()
        content(c.get_data() if c is not None else b'')
        resources(page.get('/Resources'))
    grey = all(abs(r - g) < 1e-3 and abs(g - b) < 1e-3 for r, g, b in colors)
    return {'images': images, 'fonts': sorted(fonts), 'colors': colors,
            'gray': grey and all(cs == '/DeviceGray' for cs in images)}


def grey_scan(data: bytes) -> tuple[bool, str]:
    """V8 for a black and white PDF: every rg/RG colour has r == g == b and every image is DeviceGray."""
    scan = pdf_scan(data)
    bad = [c for c in scan['colors'] if not (abs(c[0] - c[1]) < 1e-3 and abs(c[1] - c[2]) < 1e-3)]
    colour_images = [cs for cs in scan['images'] if cs != '/DeviceGray']
    if not bad and not colour_images:
        return True, f'only greys drawn; {len(scan["images"])} greyscale image{"" if len(scan["images"]) == 1 else "s"}'
    parts = []
    if bad:
        parts.append(f'{len(bad)} coloured fills or strokes (for example rgb {", ".join(f"{v:.2f}" for v in bad[0])})')
    if colour_images:
        parts.append(f'{len(colour_images)} images in colour ({colour_images[0]})')
    return False, '; '.join(parts)


def _pdf_info(data: bytes, spec: dict) -> dict:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages = len(reader.pages)
    text = '\n'.join((p.extract_text() or '') for p in reader.pages)

    def walk(items, depth=0):
        for it in items:
            if isinstance(it, list):
                yield from walk(it, depth + 1)
            else:
                yield depth, str(it.title)
    outline = list(walk(reader.outline))
    meta = reader.metadata or {}
    box = reader.pages[0].mediabox
    # the structure only: compressed page and image data (ASCII85 text) can hold "/JS" by chance
    raw = re.sub(rb'(?<!end)stream\r?\n.*?endstream', b'', data, flags=re.S)
    return {'text': text + '\n' + '\n'.join(t for _, t in outline), 'pages': pages, 'outline': outline,
            'title': str(meta.get('/Title') or ''), 'size': (float(box.width), float(box.height)),
            'first_page': reader.pages[0].extract_text() or '',
            'active': re.findall(rb'/(JavaScript|JS|Launch|EmbeddedFiles?|RichMedia|XFA)\b', raw) + _active_names(reader),
            'external': re.findall(rb'/(URI|GoToR|SubmitForm|ImportData)\b', raw), 'reader': reader}


ACTIVE = {'/JavaScript', '/JS', '/Launch', '/EmbeddedFiles', '/EmbeddedFile', '/RichMedia', '/XFA'}


def _active_names(reader) -> list[bytes]:
    """Active content found by walking the document's objects (the catalog, its name tree and actions, each page's
    actions and annotations), which also finds what an object stream would hide from the byte scan."""
    found: set[str] = set()

    def look(obj, depth=0):
        if depth > 6:
            return
        try:
            obj = obj.get_object()
        except Exception:
            pass
        if isinstance(obj, dict):
            for key, value in obj.items():
                if key in ACTIVE:
                    found.add(key)
                if key == '/S' and str(value) in ACTIVE:
                    found.add(str(value))
                if key in ('/OpenAction', '/AA', '/A', '/Names', '/AcroForm', '/Next') or key in ACTIVE:
                    look(value, depth + 1)
        elif isinstance(obj, list):
            for value in obj:
                look(value, depth + 1)
    try:
        look(reader.trailer['/Root'])
        for page in reader.pages:
            look(page.get('/AA'))
            for annot in page.get('/Annots') or []:
                look(annot)
    except Exception:
        pass
    return [name.lstrip('/').encode() for name in sorted(found)]


def _docx_info(data: bytes) -> dict:
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(io.BytesIO(data))
    paras = [(p.style.name if p.style is not None else '', p.text) for p in doc.paragraphs]
    cells = []
    header_marked = []
    for t in doc.tables:
        for row in t.rows:
            cells.extend(c.text for c in row.cells)
        tr_pr = t.rows[0]._tr.trPr if len(t.rows) else None
        header_marked.append(tr_pr is not None and tr_pr.find(qn('w:tblHeader')) is not None)
    normal = doc.styles['Normal'].font.size
    pictures = [str(s._inline.docPr.get('descr') or '') for s in doc.inline_shapes]
    return {'text': '\n'.join([t for _, t in paras] + cells), 'paras': paras, 'tables': len(doc.tables),
            'header_marked': header_marked, 'title': doc.core_properties.title or '',
            'normal_pt': normal.pt if normal is not None else None, 'pictures': pictures,
            'diagrams': sum(p.startswith('Diagram:') for p in pictures),
            'images': sum(not p.startswith('Diagram:') for p in pictures)}


def _walk(shapes):
    """Every shape on a slide, the ones inside group shapes too."""
    for shape in shapes:
        yield shape
        if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
            yield from _walk(shape.shapes)


def _pptx_info(data: bytes) -> dict:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    slides, texts, body_sizes, tables_first_row, notes = [], [], [], [], []
    diagrams, images = 0, 0
    for slide in prs.slides:
        title = slide.shapes.title.text_frame.text if slide.shapes.title is not None else ''
        slides.append(title)
        for shape in _walk(slide.shapes):
            if shape.shape_type == 6 and shape.name.startswith('Diagram:'):
                diagrams += 1
            if shape.shape_type == 13 and shape.name == 'Image':  # MSO_SHAPE_TYPE.PICTURE
                images += 1
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
                if shape.name == 'Body':
                    body_sizes += [r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs
                                   if r.font.size is not None]
            if getattr(shape, 'has_table', False) and shape.has_table:
                tables_first_row.append(shape.table.first_row)
                texts += [c.text for row in shape.table.rows for c in row.cells]
            if getattr(shape, 'has_chart', False) and shape.has_chart and shape.chart.has_title:
                texts.append(shape.chart.chart_title.text_frame.text)
        notes.append(slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else '')
    return {'text': '\n'.join(slides + texts), 'slides': slides, 'size': (prs.slide_width, prs.slide_height),
            'body_sizes': body_sizes, 'tables_first_row': tables_first_row, 'notes': notes,
            'title': prs.core_properties.title or '', 'diagrams': diagrams, 'images': images}


def _xlsx_info(data: bytes) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    sheets, texts, cells = [], [wb.properties.title or ''], []
    for ws in wb.worksheets:
        sheets.append({'name': ws.title, 'freeze': ws.freeze_panes, 'a1_bold': bool(ws['A1'].font and ws['A1'].font.b),
                       'ws': ws})
        texts.append(ws.title)
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    cells.append(c)
                    if isinstance(c.value, str):
                        texts.append(c.value)
    return {'text': '\n'.join(texts), 'sheets': sheets, 'cells': cells, 'title': wb.properties.title or '', 'wb': wb}


def _md_info(data: bytes) -> dict:
    import html
    text = data.decode('utf-8')
    heads, fence, prose = [], None, []
    for ln in text.split('\n'):
        m = re.match(r'^(`{3,}|~{3,})', ln)
        if m and (fence is None or ln.startswith(fence)):
            fence = None if fence else m.group(1)
            continue
        if fence:
            continue
        prose.append(ln)
        h = re.match(r'^(#{1,6}) (.*)$', ln)
        if h:
            heads.append((len(h.group(1)), h.group(2)))
    # the text as a reader sees it (escapes and entities resolved), for finding titles and headings
    shown = html.unescape(re.sub(r'\\([\\`*_{}\[\]()#+\-.!|>])', r'\1', text))
    return {'text': shown, 'heads': heads, 'prose': '\n'.join(prose)}


# Markdown outside code blocks that a viewer would load or run: raw HTML tags, images (inline or by reference) and
# reference definitions pointing at a web address.
_MD_REMOTE = re.compile(r'!\[[^\]]*\]\(\s*<?(?:https?:)?//|!\[[^\]]*\]\[[^\]]*\]|^\s{0,3}\[[^\]]+\]:\s*<?(?:https?:)?//|'
                        r'<\s*(?:img|image|iframe|object|embed|link|svg|video|audio|source|picture|input|form)\b', re.I | re.M)
_MD_ACTIVE = re.compile(r'<\s*(?:script|style|iframe|object|embed)\b|<[A-Za-z][^<>]*\son[a-z]+\s*=|'
                        r'<[A-Za-z][^<>]*(?:javascript|vbscript):', re.I)


def _embedded_formulas(data: bytes) -> list[str]:
    """Formulas inside the workbooks a PPTX embeds for its charts (X2): python-pptx writes each chart's data there, and
    any "=..." label written as-is would become a live formula when someone picks Edit Data."""
    found = []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for n in z.namelist():
            if re.match(r'ppt/embeddings/.*\.xlsx$', n, re.I):
                with zipfile.ZipFile(io.BytesIO(z.read(n))) as inner:
                    for m in inner.namelist():
                        if m.startswith('xl/worksheets/') and m.endswith('.xml'):
                            found += re.findall(r'<f(?:\s[^>]*)?>([^<]*)</f>', inner.read(m).decode('utf-8', 'replace'))
    return found


def _lost_chars(spec: dict, fonts) -> tuple[int, int, str]:
    """(characters the PDF fonts could not draw, all non-space characters, a few examples) over the spec's text."""
    texts = [spec['title'], spec.get('subtitle') or '']
    for s in spec['sections']:
        texts.append(s['heading'])
        for b in s['blocks']:
            texts += [str(b.get(k) or '') for k in ('text', 'title', 'by')] + [str(x) for x in b.get('items') or []]
            texts += [str(x) for x in (b.get('columns') or []) + (b.get('labels') or [])]
            texts += [str(v) for r in b.get('rows') or [] for v in r if isinstance(v, str)]
            texts += [str(x.get('name') or '') for x in b.get('series') or []]
    lost, total, examples = 0, 0, []
    for t in texts:
        drawn, _ = fonts.fit(t)
        for a, d in zip(t, drawn):
            if a.isspace():
                continue
            total += 1
            if a != d and d == '?':
                lost += 1
                if a not in examples and len(examples) < 5:
                    examples.append(a)
    return lost, total, ''.join(examples)


def _planned_pages(report: dict | None) -> int:
    """The pages or slides a Studio design report says were laid out (the sum of its layout counts); 0 without one."""
    if not isinstance(report, dict) or not isinstance(report.get('layouts'), dict):
        return 0
    return sum(v for v in report['layouts'].values() if isinstance(v, int) and not isinstance(v, bool) and v > 0)


def design_checks(report: dict | None) -> list[RuleResult]:
    """D1-D8 from a Studio design report (docs/PLAN-designer.md 9.8): one warn RuleResult per check, in D1..D8 order;
    [] for a file made without Studio (no report). A check the report does not list counts as passed."""
    if not isinstance(report, dict):
        return []
    by = {c.get('id'): c for c in report.get('checks') or [] if isinstance(c, dict)}
    out = []
    for rid in DESIGN_RULES:
        c = by.get(rid)
        if c is None:
            out.append(RuleResult(rid, 'warn', True, 'not checked'))
            continue
        ok = bool(c.get('ok'))
        n = c.get('failures') if isinstance(c.get('failures'), int) else 0
        note = ' '.join(str(c.get('note') or '').split())[:200]
        if not ok:
            head = f'{n} issue{"" if n == 1 else "s"}' if n else 'failed'
            note = f'{head}: {note}' if note else head
        out.append(RuleResult(rid, 'warn', ok, note or 'passed'))
    return out


def verify(spec: dict, fmt: str, data: bytes, brief=None, extra=(), design_report: dict | None = None) -> list[RuleResult]:
    """Reopen the rendered file and check it. Raises SpecError for V1 (does not reopen), L4 (too big), X1 (active
    content) and X6 (an image not from the asset cache, or without a credit); everything else comes back as one
    RuleResult per rule. With the request's brief (create/brief.py) the file is also checked against it (V5-V9), and
    with a design in the spec against the design (V10). `extra` takes the results render_safe (V11) and the create
    agent (V12) made: when render_safe drew some sections as plain text or left them out, the file is checked against
    what was drawn. They are not repeated in the list returned. With a Studio design report (a designed file), its QA
    checks come back as D1-D8 (design_checks)."""
    from . import themes
    from .render import (_Fonts, body_font, chart_summary, safe_formula, sheet_name, sheet_theme, theme_for,
                         theme_name, xlsx_plan)
    from .spec import DIAGRAMS, FORMATS, normalize

    if fmt not in FORMATS:
        raise SpecError('X1', f'Files are made as {", ".join(FORMATS)} only, not {str(fmt)[:12]!r}.')
    drawn = next((getattr(r, 'rendered', None) for r in extra or () if getattr(r, 'rendered', None) is not None), None)
    spec, _ = normalize(drawn if drawn is not None else spec, fmt)
    size = len(data)
    if size > MAX_BYTES:
        raise SpecError('L4', f'The file is {size / 1e6:.1f} MB; the limit is {MAX_BYTES // (1024 * 1024)} MB.')
    try:
        info = {'pdf': lambda: _pdf_info(data, spec), 'docx': lambda: _docx_info(data), 'pptx': lambda: _pptx_info(data),
                'xlsx': lambda: _xlsx_info(data), 'md': lambda: _md_info(data)}[fmt]()
    except Exception as e:
        raise SpecError('V1', f'The {fmt.upper()} file does not reopen: {str(e)[:160]}')
    out = [RuleResult('V1', 'block', True, f'reopened with {_LIB[fmt]}'), RuleResult('L4', 'block', True, f'{size:,} bytes')]
    sections = spec['sections']
    headed = [s for s in sections if s['heading']]
    charts = [b for s in sections for b in s['blocks'] if b['type'] == 'chart']
    tables = [b for s in sections for b in s['blocks'] if b['type'] == 'table']
    corpus = _norm(info['text'])
    theme = theme_name(spec, None)
    look = theme_for(spec, theme)  # the theme dict the file was drawn with (a design file's, when there is one)
    if fmt == 'xlsx' and spec.get('design'):
        look = sheet_theme(look)  # sheets stay white: the colours the sheet really uses
    pdf_fonts = _Fonts(look, body_font(spec, 'pdf', theme, look)) if fmt == 'pdf' else None
    fix = (lambda s: pdf_fonts.fit(s)[0]) if pdf_fonts else (lambda s: s)  # what the PDF could draw

    # V2: title and headings in the text, as written (a PDF that drew them as "????" has lost them; a PDF heading is
    # drawn at most MAX_PDF_HEADING characters long)
    from .render import MAX_PDF_HEADING, _short
    shown = (lambda h: _short(h, MAX_PDF_HEADING)) if fmt == 'pdf' else (lambda h: h)
    missing = [x for x in [spec['title']] + [s['heading'] for s in headed] if not _found(shown(x), corpus)]
    lost_note = ''
    if pdf_fonts is not None:
        lost, total, examples = _lost_chars(spec, pdf_fonts)
        if lost and lost * 2 >= total:
            raise SpecError('V1', f'The PDF fonts on this server cannot draw {lost:,} of the {total:,} characters '
                                  f'(for example {examples}), so most of the text would print as "?". Make it as Word, '
                                  f'PowerPoint, Excel or Markdown, or set TRACEGRAPH_PDF_FONT to a font that has them.')
        if lost:
            lost_note = f'{lost:,} characters could not be drawn in the PDF font and show as "?" (for example {examples})'
            missing = missing or ['']
    if fmt == 'xlsx':
        names = {sh['name'].lower() for sh in info['sheets']}
        missing = [x for x in missing if sheet_name(x, set()).lower() not in names]
    note = 'not found in the file: ' + '; '.join(m[:60] for m in missing[:5] if m) if any(missing) else ''
    out.append(RuleResult('V2', 'warn', not missing, 'title and headings found' if not missing else
                          '; '.join(x for x in (note, lost_note) if x)))

    # V3: counts; V4: charts; format rules
    if fmt == 'pdf':
        want = 1 + len(headed)
        got = len(info['outline'])
        out.append(RuleResult('V3', 'warn', info['pages'] >= 1 and got == want,
                              f'{info["pages"]} pages, {got} outline entries for {want} headings'))
        lost = [b['title'] for b in charts if not _found(fix(b['title']), corpus)]
        out.append(RuleResult('V4', 'warn', not lost, f'{len(charts)} charts drawn' if not lost else
                              'charts not found: ' + '; '.join(lost[:3])))
        w, h = info['size']
        paper = (612, 792) if spec.get('paper') == 'letter' else (595.27, 841.89)
        problems = []
        if abs(w - paper[0]) > 2 or abs(h - paper[1]) > 2:
            problems.append(f'page size {w:.0f}x{h:.0f}')
        if _norm(info['title']) != _norm(fix(spec['title'])):
            problems.append('title missing from the document properties')
        if not info['outline']:
            problems.append('no outline')
        if 'page 1' not in _norm(info['first_page']):
            problems.append('no page number')
        out.append(RuleResult('F1', 'fix', not problems, '; '.join(problems) or
                              f'{"Letter" if spec.get("paper") == "letter" else "A4"}, page numbers, outline, title set'))
        out.append(RuleResult('F6', 'fix', True, 'one theme; body text 10.5 pt'))
        if info['active']:
            raise SpecError('X1', 'The PDF holds active content: ' + ', '.join(sorted({a.decode() for a in info['active']})))
        out.append(RuleResult('X1', 'block', True, 'no scripts or embedded files'))
        out.append(RuleResult('X3', 'fix', not info['external'], 'no links or remote actions' if not info['external']
                              else 'external actions found'))
        out.append(RuleResult('A1', 'fix', got == want, 'headings are outline entries'))
        out.append(RuleResult('A2', 'fix', True, 'table header rows repeat on every page'))
    elif fmt == 'docx':
        heads = [(int(st.split()[-1]), t) for st, t in info['paras'] if re.match(r'Heading [1-9]$', st)]
        from .render import chart_rows, formats_of, table_groups
        font = 'Times-Roman' if theme_for(spec, theme, paper=True)['font'] == 'Georgia' else 'Helvetica'
        want_tables = sum(len(table_groups(b['columns'], b['rows'], 'docx', spec.get('paper'), formats_of(b), font))
                          for b in tables)
        want_tables += sum(len(table_groups(*chart_rows(b), 'docx', spec.get('paper'), None, font)) for b in charts)
        out.append(RuleResult('V3', 'warn', len(heads) == len(headed) and info['tables'] == want_tables,
                              f'{len(heads)} headings for {len(headed)} sections, {info["tables"]} tables for '
                              f'{want_tables}'))
        lost = [b['title'] for b in charts if not _found(b['title'], corpus)]
        out.append(RuleResult('V4', 'warn', not lost and info['tables'] >= want_tables,
                              f'{len(charts)} charts as tables' if not lost else 'charts not found: ' + '; '.join(lost)))
        has_title = any(st == 'Title' for st, _ in info['paras'])
        ok = has_title and len(heads) == len(headed) and bool(info['title'])
        out.append(RuleResult('F2', 'fix', ok, 'Title and Heading styles, core properties set' if ok else
                              'heading styles or core properties missing'))
        pt = info['normal_pt'] or 0
        out.append(RuleResult('F6', 'fix', pt >= 10, f'body text {pt:g} pt'))
        active, external = _zip_checks(data)
        if active:
            raise SpecError('X1', 'The file holds active content: ' + ', '.join(active[:3]))
        out.append(RuleResult('X1', 'block', True, 'no macros, scripts or OLE objects'))
        out.append(RuleResult('X3', 'fix', not external, 'no external references' if not external else
                              'external references in ' + ', '.join(external[:3])))
        out.append(RuleResult('A1', 'fix', len(heads) == len(headed), 'headings use Heading styles'))
        out.append(RuleResult('A2', 'fix', all(info['header_marked']), 'every table has a marked header row'
                              if all(info['header_marked']) else 'a table has no marked header row'))
    elif fmt == 'pptx':
        # a designed deck has the slides its plan laid out (covers, dividers, splits and a closing slide included)
        planned = _planned_pages(design_report)
        want = planned if planned else 1 + len(sections)
        out.append(RuleResult('V3', 'warn', len(info['slides']) == want, f'{len(info["slides"])} slides for {want}'))
        n, xml = _count_charts(data, 'ppt')
        out.append(RuleResult('V4', 'warn', n == len(charts), f'{n} native charts for {len(charts)}'))
        w, h = info['size']
        notes_ok = all(not s.get('notes') or info['notes'][i + 1].strip() for i, s in enumerate(sections)
                       if i + 1 < len(info['notes'])) if not planned else \
            not any(s.get('notes') for s in sections) or any(n.strip() for n in info['notes'])
        ok = abs(w / h - 16 / 9) < 0.01 and _norm(info['slides'][0]) == _norm(spec['title']) and notes_ok
        out.append(RuleResult('F3', 'fix', ok, '16:9, title slide, one slide per section, notes kept' if ok else
                              'slide size, title slide or notes are off'))
        small = [s for s in info['body_sizes'] if s < 18]
        out.append(RuleResult('F6', 'fix', not small, 'body text at least 18 pt' if not small else
                              f'{len(small)} body runs under 18 pt'))
        active, external = _zip_checks(data)
        if active:
            raise SpecError('X1', 'The file holds active content: ' + ', '.join(active[:3]))
        out.append(RuleResult('X1', 'block', True, 'no macros, scripts or OLE objects'))
        formulas = _embedded_formulas(data)
        out.append(RuleResult('X2', 'fix', not formulas, 'no formulas in the chart data' if not formulas else
                              'formulas in the chart data: ' + ', '.join(f[:40] for f in formulas[:3])))
        out.append(RuleResult('X3', 'fix', not external, 'no external references' if not external else
                              'external references in ' + ', '.join(external[:3])))
        if planned:  # every section heading is the real title of a slide somewhere in the designed deck
            titles = {_norm(t) for t in info['slides']}
            titled = all(_norm(s['heading']) in titles for s in sections if s['heading'])
        else:
            titled = all(_norm(t) == _norm(s['heading']) for t, s in zip(info['slides'][1:], sections))
        out.append(RuleResult('A1', 'fix', titled, 'every slide has a real title' if titled else
                              'a slide title does not match its section'))
        out.append(RuleResult('A2', 'fix', all(info['tables_first_row']), 'table header rows marked'))
    elif fmt == 'xlsx':
        plan = xlsx_plan(spec)
        got = [sh['name'] for sh in info['sheets']]
        out.append(RuleResult('V3', 'warn', got == [p['name'] for p in plan], f'{len(got)} sheets for {len(plan)}'))
        n, xml = _count_charts(data, 'xl')
        out.append(RuleResult('V4', 'warn', n == len(charts), f'{n} native charts for {len(charts)}'))
        problems = []
        for p, sh in zip(plan, info['sheets']):
            if p['kind'] == 'table' and (sh['freeze'] != 'A2' or not sh['a1_bold']):
                problems.append(f'{sh["name"]}: header not bold and frozen')
            if p['kind'] == 'table':
                ws, b = sh['ws'], p['table']
                for i, row in enumerate(b['rows'][:50], 2):
                    for j, v in enumerate(row, 1):
                        if isinstance(v, (int, float)) and not isinstance(v, bool) and \
                                not isinstance(ws.cell(row=i, column=j).value, (int, float)):
                            problems.append(f'{sh["name"]}: a number stored as text')
                            break
        bad_names = [g for g in got if len(g) > 31 or re.search(r'[\[\]:*?/\\]', g)]
        problems += [f'bad sheet name {g}' for g in bad_names]
        out.append(RuleResult('F4', 'fix', not problems, '; '.join(problems[:3]) or
                              'bold frozen headers, numbers as numbers, valid sheet names'))
        active, external = _zip_checks(data)
        if active:
            raise SpecError('X1', 'The file holds active content: ' + ', '.join(active[:3]))
        out.append(RuleResult('X1', 'block', True, 'no macros, scripts or OLE objects'))
        risky = []
        for c in info['cells']:
            v = c.value
            if c.data_type == 'f' or (isinstance(v, str) and v.startswith('=')):
                if not safe_formula(str(v)):
                    risky.append(c.coordinate)
            elif isinstance(v, str) and v.lstrip(' \u00a0')[:1] in ('=', '+', '-', '@', '\t', '\r'):
                risky.append(c.coordinate)
        out.append(RuleResult('X2', 'fix', not risky, 'no cell can run as a formula' if not risky else
                              'cells that could run as formulas: ' + ', '.join(risky[:5])))
        out.append(RuleResult('X3', 'fix', not external, 'no external references' if not external else
                              'external references in ' + ', '.join(external[:3])))
        heads = all(sh['a1_bold'] for p, sh in zip(plan, info['sheets']) if p['kind'] == 'table')
        out.append(RuleResult('A2', 'fix', heads, 'header rows bold and frozen'))
    else:  # md
        heads = info['heads']
        titles = [h for h in heads if h[0] == 1]
        out.append(RuleResult('V3', 'warn', len(heads) == 1 + len(headed),
                              f'{len(heads)} headings for {1 + len(headed)}'))
        lost = [b['title'] for b in charts if f'chart: {_norm(b["title"])}' not in corpus]
        out.append(RuleResult('V4', 'warn', not lost, f'{len(charts)} charts as tables' if not lost else
                              'charts not found: ' + '; '.join(lost)))
        skips = [t for (a, _), (b2, t) in zip(heads, heads[1:]) if b2 > a + 1]
        ok = len(titles) == 1 and heads and heads[0][0] == 1 and not skips
        out.append(RuleResult('F5', 'fix', bool(ok), 'one # title, no skipped heading levels' if ok else
                              'title or heading levels are off'))
        if _MD_ACTIVE.search(info['prose']):
            raise SpecError('X1', 'The Markdown file holds raw HTML that could run in a viewer.')
        out.append(RuleResult('X1', 'block', True, 'no raw HTML'))
        remote = _MD_REMOTE.findall(info['prose'])
        out.append(RuleResult('X3', 'fix', not remote, 'no remote images or raw HTML' if not remote else
                              'remote images or raw HTML: ' + ', '.join(r[:30] for r in remote[:3])))
        out.append(RuleResult('A1', 'fix', len(heads) == 1 + len(headed), 'headings are Markdown headings'))

    # A3: chart titles and summaries next to them; A4: theme contrast
    if charts:
        if fmt == 'xlsx':
            _, xml = _count_charts(data, 'xl')
            titled = sum('<c:title>' in x or '<title>' in x for x in xml) >= len(charts)
        elif fmt == 'pptx':
            _, xml = _count_charts(data, 'ppt')
            titled = sum('<c:title>' in x for x in xml) >= len(charts)
        else:
            titled = all(_found(fix(b['title']), corpus) for b in charts)
        summaries = all(_found(fix(chart_summary(b)), corpus) for b in charts)
        out.append(RuleResult('A3', 'warn', titled and summaries, 'every chart has a title and a summary line'
                              if titled and summaries else 'a chart is missing its title or summary line'))
    ratio, pair = themes.worst_contrast(look if spec.get('design') else theme)
    out.append(RuleResult('A4', 'fix', ratio >= 4.5, f'lowest text contrast {ratio:.1f}:1 ({pair})'))
    if spec.get('design'):
        out.append(_v10(spec, fmt, info, data, look))

    # V4 counts diagrams too; X6 whenever the file holds images; V5-V9 against the brief
    diagrams = [b for s in sections for b in s['blocks'] if b['type'] in DIAGRAMS]
    images = [b for s in sections for b in s['blocks'] if b['type'] == 'image']
    drawn = _diagrams_drawn(fmt, info, data, diagrams, lambda t: _found(fix(t), corpus))
    if diagrams:
        v4 = next(r for r in out if r.id == 'V4')
        v4.ok = v4.ok and drawn >= len(diagrams)
        v4.note = f'{v4.note}; {drawn} of {len(diagrams)} diagrams drawn'
    embedded = _images_embedded(fmt, info)
    if images or brief is not None:
        out.append(_x6(fmt, data, images, embedded))
    if brief is not None:
        out += _brief_checks(spec, fmt, info, brief, theme, drawn, charts, embedded, data)
    out += design_checks(design_report)
    return out


def _v10(spec: dict, fmt: str, info: dict, data: bytes, t: dict) -> RuleResult:
    """V10: the design's background and text colours are in the file (PDF fills, the Word page background and heading
    style, the slide background, the sheet header fill). Markdown carries no colours and says so."""
    name = (spec.get('design') or {}).get('name') or 'the design file'
    if fmt == 'md':
        return RuleResult('V10', 'warn', True, 'A Markdown file carries no colours or fonts; convert it to PDF, Word or '
                                               'PowerPoint to see the design (0 tokens).')
    bg, text, heading = t['bg'].upper(), t['text'].upper(), t['heading'].upper()
    missing = []
    if fmt == 'pdf':
        seen = {''.join(f'{round(c * 255):02X}' for c in rgb) for rgb in pdf_scan(b'', info['reader'])['colors']}
        missing += [f'{what} #{c}' for what, c in (('background', bg), ('text', text)) if c not in seen and
                    not (what == 'background' and c == 'FFFFFF')]
    else:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            names = set(z.namelist())
            read = lambda n: z.read(n).decode('utf-8', 'replace') if n in names else ''  # noqa: E731
            if fmt == 'docx':
                doc, styles = read('word/document.xml'), read('word/styles.xml')
                if bg != 'FFFFFF' and not re.search(rf'<w:background [^>]*w:color="{bg}"', doc, re.I):
                    missing.append(f'page background #{bg}')
                block = re.search(r'<w:style [^>]*w:styleId="Heading1".*?</w:style>', styles, re.S)
                if not block or not re.search(rf'<w:color [^>]*w:val="{heading}"', block.group(0), re.I):
                    missing.append(f'heading colour #{heading}')
            elif fmt == 'pptx':
                slide = read('ppt/slides/slide2.xml') or read('ppt/slides/slide1.xml')
                fill = re.search(r'<p:bg>.*?</p:bg>', slide, re.S)
                if not fill or not re.search(rf'val="{bg}"', fill.group(0), re.I):
                    missing.append(f'slide background #{bg}')
            else:  # xlsx: sheets stay white; the table header uses the design's header colours
                sheet = next((sh for sh in info['sheets'] if sh['a1_bold'] and sh['freeze'] == 'A2'), None)
                if sheet is not None:
                    fill = sheet['ws']['A1'].fill
                    got = str(getattr(fill.fgColor, 'rgb', '') or '')[-6:].upper()
                    if got != t['header_bg'].upper():
                        missing.append(f'table header #{t["header_bg"].upper()}')
    if missing:
        return RuleResult('V10', 'warn', False, f'{name}: not found in the file: ' + ', '.join(missing))
    if fmt == 'xlsx':
        return RuleResult('V10', 'warn', True, f'{name}: table headers #{t["header_bg"].upper()}, text #{text}; sheets '
                                               'stay white')
    return RuleResult('V10', 'warn', True, f'{name}: background #{bg} and text #{text} used')


def _diagrams_drawn(fmt: str, info: dict, data: bytes, diagrams: list[dict], found) -> int:
    if fmt == 'pdf':
        return sum(1 for b in diagrams if found(b['title']))
    if fmt in ('docx', 'pptx'):
        return info['diagrams']
    if fmt == 'md':
        return data.decode('utf-8', 'replace').count('```mermaid\n')
    return sum(1 for line in info['text'].split('\n') if line.startswith('Diagram: '))


def _images_embedded(fmt: str, info: dict) -> int:
    if fmt == 'pdf':
        return len(pdf_scan(b'', info['reader'])['images'])
    if fmt in ('docx', 'pptx'):
        return info['images']
    return 0


def _x6(fmt: str, data: bytes, images: list[dict], embedded: int) -> RuleResult:
    """X6: every image block is a re-encoded PNG in the asset cache with a credit, and the file embeds only PNG or JPEG
    pictures. Raises SpecError when that does not hold."""
    from . import assets
    for b in images:
        path = assets.CACHE / f'{b["asset"]}.png'
        try:
            head = path.read_bytes()[:8]
        except OSError:
            raise SpecError('X6', 'An image in the file is not in the local asset cache.')
        if head != b'\x89PNG\r\n\x1a\n':
            raise SpecError('X6', 'An image in the asset cache is not a re-encoded PNG.')
        if not str(b.get('credit') or '').strip():
            raise SpecError('X6', 'An image in the file has no credit line (title, author, licence and source).')
    if fmt in ('docx', 'pptx'):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            media = [n for n in z.namelist() if re.match(r'(word|ppt)/media/', n)]
        odd = [n for n in media if not re.search(r'\.(png|jpe?g)$', n, re.I)]
        if odd:
            raise SpecError('X6', 'The file embeds a picture that is not PNG or JPEG: ' + ', '.join(odd[:3]))
    if not images:
        return RuleResult('X6', 'fix', True, 'no images')
    shown = embedded if fmt in ('pdf', 'docx', 'pptx') else 0
    where = f'{shown} embedded' if fmt in ('pdf', 'docx', 'pptx') else f'not embedded in {fmt.upper()}; credits kept'
    return RuleResult('X6', 'fix', True, f'{len(images)} image{"" if len(images) == 1 else "s"} from the asset cache, '
                                           f'each credited ({where})')


def span(lo_hi) -> str:
    lo, hi = lo_hi
    return f'{lo}' if lo == hi else f'{lo}-{hi}'


def _brief_checks(spec, fmt, info, brief, theme, drawn, charts, embedded, data) -> list[RuleResult]:
    from . import fonts as font_mod
    from .brief import WORDS_PER_PAGE, diagrams_min
    out = []
    # V5: pages or slides within the range asked for
    if brief.slides and fmt == 'pptx':
        n = len(info['slides'])
        lo, hi = brief.slides
        out.append(RuleResult('V5', 'warn', lo <= n <= hi, f'{n} slides, asked for {span(brief.slides)}' if lo <= n <= hi
                              else f'asked for {span(brief.slides)} slides, made {n}'))
    elif brief.pages and fmt == 'pdf':
        n = info['pages']
        lo, hi = brief.pages
        out.append(RuleResult('V5', 'warn', lo <= n <= hi, f'{n} pages, asked for {span(brief.pages)}' if lo <= n <= hi
                              else f'asked for {span(brief.pages)} pages, made {n}'))
    elif brief.pages and fmt in ('docx', 'md'):
        words = len(re.findall(r'\w+', info['text']))
        n = max(1, round(words / WORDS_PER_PAGE))
        lo, hi = brief.pages
        ok = lo - 1 <= n <= hi + 1
        out.append(RuleResult('V5', 'warn', ok, f'about {n} pages ({words:,} words), asked for {span(brief.pages)}'
                              if ok else f'asked for {span(brief.pages)} pages, made about {n} ({words:,} words)'))
    # V6: images
    if brief.images:
        ok = embedded >= 1
        note = (f'{embedded} image{"" if embedded == 1 else "s"} embedded' if ok else
                f'no images embedded; a {fmt.upper()} file can\'t hold them' if fmt in ('md', 'xlsx') else
                'no images embedded (asked for images)')
        out.append(RuleResult('V6', 'warn', ok, note))
    # V7: the font
    if brief.font:
        choice = font_mod.resolve(brief.font, fmt, theme=theme)
        shown = font_mod.display(brief.font)
        if choice.embedded and fmt == 'pdf':
            ps = _postscript(choice.regular)
            used = any(ps and f.split('+')[-1] == ps for f in pdf_scan(b'', info['reader'])['fonts'])
            out.append(RuleResult('V7', 'warn', used, f'{shown} embedded' if used else
                                  f'{shown} was chosen but is not embedded in the PDF'))
        elif choice.note:
            out.append(RuleResult('V7', 'warn', True, f'{shown} not used; the answer names what the file uses instead'
                                  if fmt == 'pdf' else f'{shown} named, not embedded' if fmt != 'md' else
                                  'Markdown carries no font; the answer says so'))
        else:
            out.append(RuleResult('V7', 'warn', True, f'{choice.used or shown} used'))
    # V8: the theme, and for a black and white PDF every colour and image
    if brief.theme:
        ok, note = theme == brief.theme, f'theme {theme}' + ('' if theme == brief.theme else f', asked for {brief.theme}')
        if ok and brief.theme == 'mono' and fmt == 'pdf':
            ok, scan = grey_scan(data)
            note = f'{note}; {scan}'
        out.append(RuleResult('V8', 'warn', ok, note))
    # V9: diagrams (charts count only when no diagram kind was named: "a report with charts")
    if brief.diagrams:
        want = diagrams_min(brief)
        got = drawn + (len(charts) if not brief.diagram_kinds else 0)
        out.append(RuleResult('V9', 'warn', got >= want, f'{got} diagram{"" if got == 1 else "s"} drawn, asked for at '
                                                          f'least {want}'))
    return out


def _postscript(path: str | None) -> str | None:
    if not path:
        return None
    try:
        from reportlab.pdfbase.ttfonts import TTFontFile
        name = TTFontFile(path).name
        return name.decode('latin-1') if isinstance(name, bytes) else str(name)
    except Exception:
        return None


_LIB = {'pdf': 'pypdf', 'docx': 'python-docx', 'pptx': 'python-pptx', 'xlsx': 'openpyxl', 'md': 'a UTF-8 Markdown read'}
