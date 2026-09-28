"""Evals: run evals/cases.jsonl (plus the git-ignored evals/cases.local.jsonl of promoted labels and runs) through the
real pipeline and score each final run record.

A case passes when every expectation holds: each `expect_agents` entry was chosen for some subtask and no other agent
was (guards aside), the outcome (clarify / blocked / unsupported / answer / error) matches, the steps and the plan have
the expected shape, and the final answer matches `must_match`, every `must_mention` and not `must_not_match`
(case-insensitive regexes; `@name` expands to a named matcher from evals/matchers.json). `silent_wrong` counts failures
where every subtask still reported ok=true: the answers a user would trust without a warning.

Optional case keys (docs/PLAN-speed-evals-chat.md, "Eval cases", and docs/PLAN-accuracy-v2.md D1-D8): `split` (dev or
holdout), `turns` (a multi-turn conversation run in one session, each turn with its own expectations), `files`
(fixtures under evals/fixtures/, uploaded through the same store path as a user's upload and attached to every turn),
`judge` + `judge_min` (a rubric scored by an LLM judge, see jevrouter/judge.py; without a judge the case is unjudged),
`max_ms` + `strict_ms` (a latency budget), `paraphrases` (extra phrasings, each its own attempt), `chat` ({mode, agent,
style}, also per turn), `plan` (a pinned plan replayed instead of calling the planner), `expect_steps` (+ `unordered`),
`expect_plan`, `forbid_agents`, `must_mention`, `no_repeat` and `budget`. An eval can repeat every attempt (1-5
times); a case passes only when all its attempts pass, and is flaky when the repeats of one phrasing disagree.

Two modes: `full` runs every agent; `route` (D3) stops each run after routing (the pipeline's dry-run hook), with Jev
answers from a cassette (jevrouter/cassette.py), and reports a confusion matrix, per-agent precision and recall,
failure stages and calibration.

CLI: .venv/bin/python -m jevrouter.evals [--mode full|route] [--jev live|replay|record] [--engine NAME|none]
     [--examples on|off] [--split dev|holdout|all] [--tags a,b] [--repeat N] [--judge NAME|auto] [--matrix] [--json]
     [--gates evals/gates.json] [--strict-budgets] [--cases PATH] [--save-baseline]
     .venv/bin/python -m jevrouter.evals calibrate | harvest | flaky | paraphrase | judge-calibrate  (see --help)
It prints a table and exits 1 when a gate fails (or, with no gates, when a tag drops below the baseline).
"""
import argparse
import asyncio
import hashlib
import itertools
import json
import math
import os
import re
import sys
import time
import uuid
import weakref
from pathlib import Path

from . import judge as judge_mod
from .config import GUARDS, MAX_SUBTASKS, MODES as CHAT_MODES, ROOT, STYLES as CHAT_STYLES

EVALS_DIR = ROOT / 'evals'
CASES = EVALS_DIR / 'cases.jsonl'
LOCAL_CASES = EVALS_DIR / 'cases.local.jsonl'  # labels and runs promoted to cases; personal, so git-ignored
FIXTURES = EVALS_DIR / 'fixtures'
BASELINE = EVALS_DIR / 'baseline.json'
MATCHERS = EVALS_DIR / 'matchers.json'
GATES = EVALS_DIR / 'gates.json'
BUDGETS = EVALS_DIR / 'budgets.json'
JUDGE_GOLD = EVALS_DIR / 'judge_gold.jsonl'
CASSETTE = EVALS_DIR / 'cassettes' / 'jev.jsonl'
CASSETTE_META = EVALS_DIR / 'cassettes' / 'jev.meta.json'  # the suite_sha the cassette was recorded for
CONCURRENCY = 2
ROUTE_CONCURRENCY = 4  # route mode runs no agents, so more cases can share Jev at once
MAX_REPEAT = 5
MAX_UNRECORDED = 0.02  # route replay fails its gate when more than this share of cases met a cassette miss
SPLITS = ('dev', 'holdout', 'all')
EVAL_MODES = ('full', 'route')
JEV_SOURCES = ('live', 'replay', 'record')
OUTCOMES = ('clarify', 'blocked', 'unsupported', 'answer', 'error')
ROUTE_EXPECT = ('expect_agents', 'expect_outcome', 'expect_steps', 'expect_plan', 'forbid_agents')
ANSWER_EXPECT = ('must_match', 'must_not_match', 'must_mention', 'expect_file', 'no_repeat')
EXPECT_KEYS = (*ROUTE_EXPECT, *ANSWER_EXPECT)
FILE_FORMATS = ('pdf', 'docx', 'pptx', 'xlsx', 'md')  # jevrouter/create/spec.py FORMATS
FILE_SOURCES = ('answer', 'llm', 'table', 'convert')
FILE_KEYS = {'format', 'contains', 'min_pages', 'max_pages', 'slides_min', 'sheets', 'charts_min', 'rules_ok',
             # accuracy v2, D2
             'pages', 'slides', 'words_min', 'headings_min', 'images_min', 'diagrams_min', 'fonts', 'grayscale',
             'files_exact', 'only_formats', 'source', 'not_contains', 'credits',
             # files that always build (docs/PLAN-files-robust.md 7.2)
             'design_bg', 'partial_ok'}
HEX_COLOR = re.compile(r'^#?[0-9A-Fa-f]{6}$')
DIAGRAM_BLOCKS = ('timeline', 'tree', 'flow')
# Always on for every created file (D2): a file whose body is the create agent's own reply line is not a document.
TEMPLATE_BODY = re.compile(r'^Created \*\*.+\*\*, \d|No format was named', re.M)
# X2: the only functions a created spreadsheet may compute with; anything else is an injected formula.
SAFE_FORMULA = re.compile(r'^=(?:[A-Z]+\d+|[A-Z]+\d+:[A-Z]+\d+|(?:SUM|AVERAGE|MIN|MAX|COUNT|ROUND)\(|[\s\d.,+\-*/()])+$', re.I)
STEP_KEYS = {'agent', 'not_agent', 'depends_on', 'text', 'args', 'format', 'forced'}
ARG_KEYS = {'amount', 'from', 'to', 'city'}
PLAN_EXPECT_KEYS = {'min_steps', 'max_steps', 'file_steps', 'no_fragment', 'max_clarify'}
BUDGET_KEYS = {'ms', 'p95_ms', 'llm_in', 'llm_out', 'jev_in', 'file_tokens', 'plan_ms', 'planner'}
PLANNERS = ('single', 'heuristic', 'llm')
TURN_KEYS = {'query', 'chat', 'plan', 'unordered', *EXPECT_KEYS}
CASE_KEYS = {'id', 'query', 'tags', 'split', 'turns', 'files', 'judge', 'judge_min', 'max_ms', 'strict_ms', 'paraphrases',
             'note', 'chat', 'plan', 'unordered', 'budget', 'run_on', *EXPECT_KEYS}
# `run_on`: where a case is meant to run: keyless, cli or api engines, or route (routing-only, which also runs every
# keyless case: it routes as a keyless run does). A case without run_on runs everywhere.
ENGINE_CLASSES = ('keyless', 'cli', 'api', 'route')
CASE_ID = re.compile(r'^[A-Za-z0-9][A-Za-z0-9_.-]{0,63}$')
AGENT_NAME = re.compile(r'^@?[a-z][a-z0-9_-]{1,23}$')
MATCHER_TOKEN = re.compile(r'(?<![\w\\@])@([a-z][a-z0-9_]*)\b')
CODES = ('wrong_agent', 'extra_agent', 'missing_agent', 'forbidden_agent', 'wrong_outcome', 'plan_shape', 'wrong_deps',
         'answer_regex', 'forbidden_regex', 'file', 'judge', 'budget', 'unjudged', 'run_status', 'unrecorded')
ROUTE_CODES = {'wrong_agent', 'extra_agent', 'missing_agent', 'forbidden_agent', 'wrong_outcome', 'plan_shape',
               'wrong_deps', 'run_status'}
ANSWER_CODES = {'answer_regex', 'forbidden_regex', 'file', 'judge'}
STAGES = ('plan_text', 'plan_shape', 'jev_pick', 'clarity', 'confidence', 'gate_missing_detail', 'gate_cant',
          'gate_wants_file', 'unsupported', 'forced', 'frame', 'agent_answer', 'file', 'judge', 'budget')
# The failure stage a policy rule stands for (jevrouter/policy.py RULES, plus the follow-up `frame`).
RULE_STAGE = {'blocked': 'jev_pick', 'blocked_dependency': 'plan_shape', 'unsupported': 'unsupported',
              'cant_do': 'gate_cant', 'forced': 'forced', 'create_demote': 'gate_wants_file',
              'file_intent': 'gate_wants_file', 'keyless_contract': 'gate_missing_detail', 'confirmed': 'clarity',
              'attached_file': 'clarity', 'refers_back_file': 'clarity', 'advice': 'jev_pick',
              'time_sensitive': 'jev_pick', 'clarify_policy': 'clarity', 'missing_slot': 'gate_missing_detail',
              'mode_research': 'jev_pick', 'ambiguous_term': 'clarity', 'frame': 'frame'}
CODE_STAGE = {'answer_regex': 'agent_answer', 'forbidden_regex': 'agent_answer', 'file': 'file', 'judge': 'judge',
              'unjudged': 'judge', 'budget': 'budget', 'plan_shape': 'plan_shape', 'wrong_deps': 'plan_shape'}
AUTO_TAG = re.compile(r'^(mode|forced):')
# Summary fields beyond the evals table's own columns, kept as JSON in its `extra` column.
EXTRA_KEYS = ('split', 'repeat', 'judge', 'tags', 'by_tag', 'p50_ms', 'p95_ms', 'flaky', 'judge_mean', 'judge_errors',
              'mode', 'jev', 'suite_sha', 'route_pass', 'answer_pass', 'by_tag_route', 'unjudged', 'unrecorded',
              'confusion', 'per_agent', 'stages', 'calibration', 'gates', 'tokens', 'budget_warnings')


class CaseError(ValueError):
    pass


# ---------- named matchers (D1) ----------

_matchers: dict | None = None


def matchers() -> dict[str, str]:
    """evals/matchers.json: {name: regex}. Read once."""
    global _matchers
    if _matchers is None:
        try:
            d = json.loads(MATCHERS.read_text())
        except (OSError, ValueError):
            d = {}
        _matchers = {k: v for k, v in d.items() if isinstance(k, str) and isinstance(v, str)}
    return _matchers


def expand(pattern: str) -> str:
    """A case regex with every `@name` token replaced by that matcher's regex (as a group). Raises CaseError for an
    unknown name."""
    known = matchers()

    def sub(m):
        name = m.group(1)
        if name not in known:
            raise CaseError(f'unknown matcher @{name} (evals/matchers.json)')
        return f'(?:{known[name]})'
    return MATCHER_TOKEN.sub(sub, pattern)


def search(pattern: str, text: str):
    return re.search(expand(pattern), text or '', re.I)


# ---------- cases ----------

def check_regex(where: str, key: str, pattern) -> list[str]:
    if not isinstance(pattern, str) or not pattern:
        return [f'{where}{key} must be a regex']
    try:
        re.compile(expand(pattern))
    except CaseError as e:
        return [f'{where}{key}: {e}']
    except re.error as e:
        return [f'{where}{key} is not a valid regex ({e})']
    return []


REFUSAL_MATCHERS = ('@cant', '@unknowable', '@needs_engine')
# an alternative that accepts a sourced answer: a web engine may answer these, so route mode can't require a refusal
SOURCED_ALTERNATIVE = re.compile(r'https\?|as of|sources\?|approx|fluctuat')


def refusal_only(expect: dict) -> bool:
    """True when the case's answer must be a refusal (its must_match leads with @cant, @unknowable or @needs_engine and
    accepts no sourced answer)."""
    m = expect.get('must_match')
    return isinstance(m, str) and m.startswith(REFUSAL_MATCHERS) and not SOURCED_ALTERNATIVE.search(m)


def route_term(tasks: list[dict]) -> bool:
    """A route-only step that found a lone term ("Mercury") whose meanings only a real run looks up."""
    from .pipeline import ROUTE_TERM
    return any(e.get('rule') == 'ambiguous_term' and e.get('why') == ROUTE_TERM for t in tasks for e in t.get('trace') or [])


def as_outcomes(v) -> list[str]:
    return [v] if isinstance(v, str) else list(v or [])


def check_expectations(d: dict, where: str) -> list[str]:
    errors = []
    agents = d.get('expect_agents')
    if agents is not None and not (isinstance(agents, list) and agents and all(isinstance(a, str) and a for a in agents)):
        errors.append(f'{where}expect_agents must be a non-empty list of agent names')
    forbid = d.get('forbid_agents')
    if forbid is not None and not (isinstance(forbid, list) and forbid and all(isinstance(a, str) and a for a in forbid)):
        errors.append(f'{where}forbid_agents must be a non-empty list of agent names')
    if 'expect_outcome' in d:
        v = d['expect_outcome']
        if not (v in OUTCOMES or (isinstance(v, list) and v and all(isinstance(x, str) and x in OUTCOMES for x in v))):
            errors.append(f'{where}expect_outcome must be one of {", ".join(OUTCOMES)} (or a list of them)')
    for k in ('must_match', 'must_not_match'):
        if k in d:
            if isinstance(d[k], str):
                errors += check_regex(where, k, d[k])
            else:
                errors.append(f'{where}{k} is not a valid regex (not a string)')
    if 'must_mention' in d:
        v = d['must_mention']
        if not (isinstance(v, list) and v):
            errors.append(f'{where}must_mention must be a non-empty list of regexes or @matchers')
        else:
            for x in v:
                errors += check_regex(where, 'must_mention', x)
    if 'no_repeat' in d and not isinstance(d['no_repeat'], bool):
        errors.append(f'{where}no_repeat must be true or false')
    if 'expect_file' in d:
        errors += check_file_expectation(d['expect_file'], where)
    if 'expect_steps' in d:
        errors += check_steps_expectation(d['expect_steps'], where)
    if 'unordered' in d and (not isinstance(d['unordered'], bool) or 'expect_steps' not in d):
        errors.append(f'{where}unordered must be true or false and needs expect_steps')
    if 'expect_plan' in d:
        errors += check_plan_expectation(d['expect_plan'], where)
    if 'chat' in d:
        errors += check_chat(d['chat'], where)
    if 'plan' in d:
        errors += check_pinned_plan(d['plan'], where)
    return errors


def check_steps_expectation(steps, where: str) -> list[str]:
    if not (isinstance(steps, list) and 1 <= len(steps) <= 8 and all(isinstance(s, dict) for s in steps)):
        return [f'{where}expect_steps must be a list of 1 to 8 objects']
    errors = []
    for i, s in enumerate(steps):
        at = f'{where}expect_steps[{i}] '
        if unknown := sorted(set(s) - STEP_KEYS):
            errors.append(f'{at}has unknown key {unknown[0]!r}')
        for k in ('agent', 'not_agent', 'text'):
            if k in s:
                errors += check_regex(at, k, s[k])
        if 'depends_on' in s and not (isinstance(s['depends_on'], list)
                                      and all(type(x) is int and 0 <= x < i for x in s['depends_on'])):
            errors.append(f'{at}depends_on must list earlier step numbers (from 0)')
        if 'args' in s and not (isinstance(s['args'], dict) and set(s['args']) <= ARG_KEYS
                                and all(isinstance(v, (str, int, float)) and not isinstance(v, bool)
                                        for v in s['args'].values())):
            errors.append(f'{at}args must be an object with amount, from, to or city')
        if 'format' in s and s['format'] not in FILE_FORMATS:
            errors.append(f'{at}format must be one of {", ".join(FILE_FORMATS)}')
        if 'forced' in s and not isinstance(s['forced'], bool):
            errors.append(f'{at}forced must be true or false')
    return errors


def check_plan_expectation(p, where: str) -> list[str]:
    if not isinstance(p, dict) or not p:
        return [f'{where}expect_plan must be an object']
    errors = []
    if unknown := sorted(set(p) - PLAN_EXPECT_KEYS):
        errors.append(f'{where}expect_plan has unknown key {unknown[0]!r}')
    for k in ('min_steps', 'max_steps', 'file_steps', 'max_clarify'):
        if k in p and not (type(p[k]) is int and p[k] >= 0):
            errors.append(f'{where}expect_plan {k} must be a whole number')
    if 'no_fragment' in p and not isinstance(p['no_fragment'], bool):
        errors.append(f'{where}expect_plan no_fragment must be true or false')
    return errors


def check_chat(chat, where: str) -> list[str]:
    """The run's chat options, by the rules app.chat_options applies to a request (the agent must also be one the run
    can pick, which app.check_agent decides when the case runs)."""
    if not isinstance(chat, dict):
        return [f'{where}chat must be an object with mode, agent or style']
    errors = []
    if unknown := sorted(set(chat) - {'mode', 'agent', 'style'}):
        errors.append(f'{where}chat has unknown key {unknown[0]!r}')
    if 'mode' in chat and chat['mode'] not in CHAT_MODES:
        errors.append(f'{where}chat mode must be one of {", ".join(CHAT_MODES)}')
    if 'style' in chat and chat['style'] not in CHAT_STYLES:
        errors.append(f'{where}chat style must be one of {", ".join(CHAT_STYLES)}')
    if 'agent' in chat:
        a = chat['agent']
        if not (isinstance(a, str) and AGENT_NAME.match(a)) or a.lstrip('@') in GUARDS:
            errors.append(f'{where}chat agent must be an agent name (guards cannot be picked)')
    return errors


def check_pinned_plan(p, where: str) -> list[str]:
    if not isinstance(p, dict) or set(p) - {'subtasks', 'deps'}:
        return [f'{where}plan must be {{subtasks, deps}}']
    subs, deps = p.get('subtasks'), p.get('deps', [[] for _ in p.get('subtasks') or []])
    if not (isinstance(subs, list) and 1 <= len(subs) <= MAX_SUBTASKS and all(isinstance(s, str) and s.strip() for s in subs)):
        return [f'{where}plan subtasks must be 1 to {MAX_SUBTASKS} non-empty strings']
    if not (isinstance(deps, list) and len(deps) == len(subs)
            and all(isinstance(d, list) and all(type(x) is int and 0 <= x < i for x in d) for i, d in enumerate(deps))):
        return [f'{where}plan deps must list, for each subtask, the earlier subtasks it needs (from 0)']
    return []


def check_budget(b) -> list[str]:
    if not isinstance(b, dict) or not b:
        return ['budget must be an object']
    errors = []
    if unknown := sorted(set(b) - BUDGET_KEYS):
        errors.append(f'budget has unknown key {unknown[0]!r}')
    for k in BUDGET_KEYS - {'planner'}:
        if k in b and not (type(b[k]) is int and b[k] > 0):
            errors.append(f'budget {k} must be a positive whole number')
    if 'planner' in b and not (isinstance(b['planner'], list) and b['planner'] and set(b['planner']) <= set(PLANNERS)):
        errors.append(f'budget planner must list planners from {", ".join(PLANNERS)}')
    return errors


def check_file_expectation(f, where: str) -> list[str]:
    """expect_file is false (no file may be created) or an object of FILE_KEYS with a format."""
    if f is False:
        return []
    if not isinstance(f, dict):
        return [f'{where}expect_file must be an object or false']
    errors = []
    if unknown := sorted(set(f) - FILE_KEYS):
        errors.append(f'{where}expect_file has unknown key {unknown[0]!r}')
    if f.get('format') not in FILE_FORMATS:
        errors.append(f'{where}expect_file format must be one of {", ".join(FILE_FORMATS)}')
    for k in ('contains', 'sheets'):
        v = f.get(k, [])
        if not (isinstance(v, list) and all(isinstance(x, str) and x for x in v)):
            errors.append(f'{where}expect_file {k} must be a list of regexes')
            continue
        for x in v:
            try:
                re.compile(x)
            except re.error as e:
                errors.append(f'{where}expect_file {k} has a bad regex {x!r} ({e})')
    for k in ('min_pages', 'max_pages', 'slides_min', 'charts_min', 'words_min', 'headings_min', 'images_min',
              'diagrams_min', 'files_exact'):
        if k in f and not (type(f[k]) is int and f[k] >= 0):
            errors.append(f'{where}expect_file {k} must be a whole number')
    for k in ('pages', 'slides'):
        if k in f and not (isinstance(f[k], list) and len(f[k]) == 2 and all(type(x) is int and x >= 0 for x in f[k])
                           and f[k][0] <= f[k][1]):
            errors.append(f'{where}expect_file {k} must be [lo, hi]')
    for k in ('rules_ok', 'grayscale', 'credits', 'partial_ok'):
        if k in f and not isinstance(f[k], bool):
            errors.append(f'{where}expect_file {k} must be true or false')
    if 'design_bg' in f and not (isinstance(f['design_bg'], str) and HEX_COLOR.match(f['design_bg'])):
        errors.append(f'{where}expect_file design_bg must be a colour like F7F4ED')
    if 'fonts' in f:
        try:
            re.compile(f['fonts'])
        except (re.error, TypeError) as e:
            errors.append(f'{where}expect_file fonts is not a valid regex ({e})')
    if 'only_formats' in f and not (isinstance(f['only_formats'], list) and f['only_formats']
                                    and set(f['only_formats']) <= set(FILE_FORMATS)):
        errors.append(f'{where}expect_file only_formats must list formats from {", ".join(FILE_FORMATS)}')
    if 'source' in f and f['source'] not in FILE_SOURCES:
        errors.append(f'{where}expect_file source must be one of {", ".join(FILE_SOURCES)}')
    if 'not_contains' in f and not (isinstance(f['not_contains'], list)
                                    and all(isinstance(x, str) and x for x in f['not_contains'])):
        errors.append(f'{where}expect_file not_contains must be a list of strings')
    return errors


def validate_case(c, strict: bool = False) -> list[str]:
    """Everything wrong with one case, as plain sentences; [] when it's valid. strict (used for the committed suite's
    own test) also wants every case to expect something, so no case passes by merely finishing."""
    if not isinstance(c, dict):
        return ['a case must be a JSON object']
    errors = []
    if not (isinstance(c.get('id'), str) and CASE_ID.match(c['id'])):
        errors.append('id must be 1-64 letters, digits, _ . or -')
    if unknown := sorted(set(c) - CASE_KEYS):
        errors.append(f'unknown key {unknown[0]!r}')
    if not (isinstance(c.get('tags', []), list) and all(isinstance(t, str) and t for t in c.get('tags', []))):
        errors.append('tags must be a list of strings')
    if c.get('split', 'dev') not in ('dev', 'holdout'):
        errors.append('split must be dev or holdout')
    turns = c.get('turns')
    if turns is not None:
        if not (isinstance(turns, list) and len(turns) >= 2 and all(isinstance(t, dict) for t in turns)):
            errors.append('turns must be a list of at least two objects')
        else:
            for i, t in enumerate(turns, 1):
                if unknown := sorted(set(t) - TURN_KEYS):
                    errors.append(f'turn {i}: unknown key {unknown[0]!r}')
                if not (isinstance(t.get('query'), str) and t['query'].strip()):
                    errors.append(f'turn {i}: query must be a non-empty string')
                errors += check_expectations(t, f'turn {i}: ')
            if strict and not any(k in t for t in turns for k in EXPECT_KEYS) and 'judge' not in c:
                errors.append('at least one turn needs an expectation')
        if any(k in c for k in (*EXPECT_KEYS, 'plan', 'unordered')):
            errors.append('a multi-turn case keeps its expectations on its turns')
        if 'paraphrases' in c:
            errors.append('a multi-turn case cannot have paraphrases')
        if 'query' in c and not (isinstance(c['query'], str) and c['query'].strip()):
            errors.append('query must be a non-empty string')
        if 'chat' in c:
            errors += check_chat(c['chat'], '')
    else:
        if not (isinstance(c.get('query'), str) and c['query'].strip()):
            errors.append('query must be a non-empty string')
        errors += check_expectations(c, '')
        if strict and not any(k in c for k in (*EXPECT_KEYS, 'judge')):
            errors.append('a case needs at least one expectation (expect_agents, expect_outcome, must_match, '
                          'must_not_match or judge)')
    files = c.get('files')
    if files is not None:
        if not (isinstance(files, list) and files and all(isinstance(f, str) for f in files)):
            errors.append('files must be a non-empty list of fixture names')
        else:
            for f in files:
                if '/' in f or '\\' in f or f.startswith('.') or not (FIXTURES / f).is_file():
                    errors.append(f'no fixture evals/fixtures/{f}')
    if 'judge' in c and not (isinstance(c['judge'], str) and c['judge'].strip()):
        errors.append('judge must be the rubric text')
    if 'judge_min' in c:
        if 'judge' not in c:
            errors.append('judge_min needs a judge rubric')
        elif not (isinstance(c['judge_min'], (int, float)) and not isinstance(c['judge_min'], bool) and 1 <= c['judge_min'] <= 5):
            errors.append('judge_min must be a number from 1 to 5')
    if 'max_ms' in c and not (type(c['max_ms']) is int and c['max_ms'] > 0):
        errors.append('max_ms must be a positive whole number of milliseconds')
    if 'strict_ms' in c and (not isinstance(c['strict_ms'], bool) or 'max_ms' not in c):
        errors.append('strict_ms must be true or false and needs max_ms')
    if 'budget' in c:
        errors += check_budget(c['budget'])
    if 'run_on' in c and not (isinstance(c['run_on'], list) and c['run_on'] and set(c['run_on']) <= set(ENGINE_CLASSES)):
        errors.append(f'run_on must list engine kinds from {", ".join(ENGINE_CLASSES)}')
    p = c.get('paraphrases')
    if p is not None and not (isinstance(p, list) and p and all(isinstance(x, str) and x.strip() for x in p)
                              and len(set(p)) == len(p) and c.get('query') not in p):
        errors.append('paraphrases must be a list of distinct non-empty strings, other than the query')
    return errors


def auto_tags(c: dict) -> dict:
    """The case with `mode:<m>` and `forced:<agent>` tags added from its chat options (case or turns)."""
    chats = [c.get('chat') or {}, *((t.get('chat') or {}) for t in c.get('turns') or [])]
    extra = [f"mode:{ch['mode']}" for ch in chats if ch.get('mode')] + \
            [f"forced:{ch['agent'].lstrip('@')}" for ch in chats if ch.get('agent')]
    tags = list(c.get('tags', []))
    new = [t for t in dict.fromkeys(extra) if t not in tags]
    return {**c, 'tags': tags + new} if new else c


def read_cases(path) -> list[dict]:
    """A cases file, every line validated. Raises CaseError naming the first bad line and what's wrong with it."""
    out, seen = [], set()
    for n, line in enumerate(open(path), 1):
        if not line.strip():
            continue
        try:
            c = json.loads(line)
        except json.JSONDecodeError as e:
            raise CaseError(f'{path.name if hasattr(path, "name") else path} line {n}: not JSON ({e.msg})')
        errors = validate_case(c)
        if not errors and c['id'] in seen:
            errors = [f'duplicate id {c["id"]!r}']
        if errors:
            cid = c.get('id') if isinstance(c, dict) else None
            raise CaseError(f'{getattr(path, "name", path)} line {n}{f" ({cid})" if cid else ""}: {"; ".join(errors)}')
        seen.add(c['id'])
        out.append(auto_tags(c))
    return out


def load_cases(path=None) -> list[dict]:
    """The committed cases, then the local ones. The local file is optional and a bad line in it is skipped. When an id
    appears twice the first one wins, so a local case can never replace a committed one."""
    if path is not None:
        return read_cases(path)
    cases, seen = [], set()
    for c in [*read_cases(CASES), *local_cases()]:
        if c['id'] not in seen:
            seen.add(c['id'])
            cases.append(c)
    return cases


def local_cases() -> list[dict]:
    if not LOCAL_CASES.exists():
        return []
    out = []
    for line in LOCAL_CASES.read_text().splitlines():
        try:
            c = json.loads(line)
        except json.JSONDecodeError:
            continue
        if isinstance(c, dict) and isinstance(c.get('id'), str) and isinstance(c.get('query'), str) and c['query'].strip():
            c = {**c, 'tags': c.get('tags') if isinstance(c.get('tags'), list) else []}
            if not validate_case(c):
                out.append(auto_tags(c))
    return out


def select(cases: list[dict], split: str = 'all', tags=None) -> list[dict]:
    """Cases in `split` (dev, holdout or all) that carry any of `tags` (None, empty or ['all']: every tag)."""
    wanted = set(tags or ()) - {'all'}
    return [c for c in cases if (split == 'all' or c.get('split', 'dev') == split)
            and (not wanted or wanted & set(c.get('tags', [])))]


def runs_on(c: dict, eclass: str) -> bool:
    """Whether the case is meant for an eval of this kind (keyless, cli, api or route); every kind by default."""
    where = c.get('run_on') or ENGINE_CLASSES
    return eclass in where or (eclass == 'route' and 'keyless' in where)


def routable(c: dict) -> bool:
    """A case route mode can score: it (or one of its turns) has a routing expectation."""
    return any(k in t for t in [c, *(c.get('turns') or [])] for k in ROUTE_EXPECT)


def case_query(c: dict) -> str:
    return c.get('query') or c['turns'][0]['query']


def case_kind(c: dict) -> str:
    return 'multi_turn' if c.get('turns') else 'file' if c.get('files') else 'judge' if c.get('judge') else 'single'


def suite_sha(path=None) -> str:
    """sha256 of the committed suite's sorted case lines: two evals with the same suite_sha ran the same cases."""
    try:
        lines = sorted(line.strip() for line in open(path or CASES) if line.strip())
    except OSError:
        return ''
    return hashlib.sha256('\n'.join(lines).encode('utf-8')).hexdigest()


def route_sha(path=None) -> str:
    """sha256 of the sorted lines of the cases route mode runs (routable and meant for route). The Jev cassette is
    checked against this, so adding a case that only runs on cli or api engines doesn't invalidate it. A line that
    doesn't parse is kept, which errs toward saying the cassette is stale."""
    try:
        raw = [line.strip() for line in open(path or CASES) if line.strip()]
    except OSError:
        return ''
    kept = []
    for line in raw:
        try:
            c = json.loads(line)
        except ValueError:
            kept.append(line)
            continue
        if not isinstance(c, dict) or (routable(c) and runs_on(c, 'route')):
            kept.append(line)
    return hashlib.sha256('\n'.join(sorted(kept)).encode('utf-8')).hexdigest()


def case_sha(c: dict) -> str:
    plain = {**c, 'tags': [t for t in c.get('tags', []) if not AUTO_TAG.match(t)]}
    return hashlib.sha256(json.dumps(plain, sort_keys=True, ensure_ascii=False).encode('utf-8')).hexdigest()[:12]


# ---------- scoring ----------

def outcome(rec: dict, mode: str = 'full') -> str:
    agents = [t.get('agent') for t in rec.get('tasks', [])]
    for guard in ('blocked', 'clarify', 'unsupported'):
        if guard in agents:
            return guard
    if mode == 'route':  # nothing ran: a routed step that isn't a guard would have answered
        return 'answer' if any(agents) else 'error'
    return 'answer' if any(t.get('ok') for t in rec.get('tasks', [])) else 'error'


def answer_of(rec: dict) -> str:
    return (rec.get('merged') or {}).get('answer') or rec.get('error') or ''


def step_index(tid: str) -> int | None:
    """'2741.2' -> 1."""
    try:
        return int(str(tid).rsplit('.', 1)[1]) - 1
    except (IndexError, ValueError):
        return None


def fullmatch(pattern: str, value) -> bool:
    return bool(value) and re.fullmatch(expand(pattern), str(value), re.I) is not None


STOP = {'a', 'an', 'the', 'and', 'or', 'of', 'to', 'in', 'on', 'for', 'it', 'that', 'this', 'from', 'with', 'then',
        'also', 'is', 'be', 'by', 'at', 'as', 'its'}


def fragment(text: str) -> bool:
    """A plan step too thin to act on: under 3 words, or no content word at all."""
    ws = re.findall(r"[A-Za-z][A-Za-z'-]*|\d+", text or '')
    return len(ws) < 3 or not any(w.lower() not in STOP and not w.isdigit() and len(w) > 2 for w in ws)


def repeated_sentence(answer: str) -> str | None:
    """The first sentence (4 words or more) that appears twice in the answer, else None."""
    seen = set()
    for s in re.split(r'(?<=[.!?])\s+|\n+', answer or ''):
        norm = re.sub(r'[^a-z0-9 ]', '', re.sub(r'\*\*[a-z_]+\*\*:', '', s.lower())).strip()
        norm = re.sub(r'\s+', ' ', norm)
        if len(norm.split()) < 4:
            continue
        if norm in seen:
            return s.strip()
        seen.add(norm)
    return None


def step_args(t: dict) -> dict:
    """What the step's keyless parser took from its text: the stored frame's slots, else gate.parse_for when present."""
    slots = ((t.get('frame') or {}).get('slots')) or {}
    if slots:
        return slots
    try:
        from . import gate
        parse_for = getattr(gate, 'parse_for', None)
        p = parse_for(t.get('agent'), t.get('input') or t.get('text') or '') if parse_for else None
        return dict(getattr(p, 'slots', None) or {})
    except Exception:
        return {}


def same_arg(want, got) -> bool:
    if got is None:
        return False
    try:
        return math.isclose(float(want), float(str(got).replace(',', '')), rel_tol=1e-6)
    except (TypeError, ValueError):
        return str(want).strip().casefold() == str(got).strip().casefold()


def step_misses(i: int, e: dict, t: dict, mode: str) -> list[tuple[str, str, dict]]:
    """(code, reason, fields) for each way the step task `t` misses expected step `e` (number i)."""
    out = []
    agent = t.get('agent')
    if 'agent' in e and not fullmatch(e['agent'], agent):
        out.append(('wrong_agent', f"step {i}: expected agent /{e['agent']}/, got {agent or 'none'}",
                    {'step': i, 'want': e['agent'], 'got': agent}))
    if 'not_agent' in e and fullmatch(e['not_agent'], agent):
        out.append(('wrong_agent', f"step {i}: agent {agent} is not allowed (/{e['not_agent']}/)",
                    {'step': i, 'want': f"not {e['not_agent']}", 'got': agent}))
    if 'forced' in e and bool(t.get('forced')) != e['forced']:
        out.append(('wrong_agent', f"step {i}: expected {'a forced' if e['forced'] else 'an unforced'} step",
                    {'step': i, 'want': 'forced' if e['forced'] else 'not forced',
                     'got': 'forced' if t.get('forced') else 'not forced', 'stage': 'forced'}))
    if 'depends_on' in e:
        got = sorted(x for x in (step_index(d) for d in t.get('depends_on') or []) if x is not None)
        if got != sorted(e['depends_on']):
            out.append(('wrong_deps', f"step {i}: depends on {got}, expected {sorted(e['depends_on'])}",
                        {'step': i, 'want': str(sorted(e['depends_on'])), 'got': str(got), 'stage': 'plan_shape'}))
    if 'text' in e and not search(e['text'], t.get('input') or t.get('text') or ''):
        out.append(('plan_shape', f"step {i}: text does not match /{e['text']}/",
                    {'step': i, 'want': e['text'], 'got': (t.get('input') or t.get('text') or '')[:80],
                     'stage': 'plan_text'}))
    if e.get('args'):
        got = step_args(t)
        for k, v in e['args'].items():
            if not same_arg(v, got.get(k)):
                out.append(('plan_shape', f"step {i}: {k} is {got.get(k)!r}, expected {v!r}",
                            {'step': i, 'want': f'{k}={v}', 'got': f'{k}={got.get(k)}', 'stage': 'plan_text'}))
    if 'format' in e and mode == 'full':
        fmts = [f.get('format') for f in t.get('created_files') or [] if isinstance(f, dict)]
        if e['format'] not in fmts:
            out.append(('file', f"step {i}: expected a {e['format']} file, got {', '.join(fmts) or 'none'}",
                        {'step': i, 'want': e['format'], 'got': ', '.join(fmts) or None}))
    return out


def align(expect_steps: list[dict], tasks: list[dict], unordered: bool, mode: str) -> list[tuple[int, int]]:
    """Pairs (expected step, task) to compare. In order, or with `unordered` the assignment where the most expected
    steps pass (fewest misses breaking ties); an exhaustive search, fine for n <= 4 (larger plans keep the order)."""
    n, m = len(expect_steps), len(tasks)
    k = min(n, m)
    if not unordered or max(n, m) > 4:
        return [(i, i) for i in range(k)]
    best, best_key = None, None
    for perm in itertools.permutations(range(m), k) if n <= m else itertools.permutations(range(n), k):
        pairs = list(zip(range(k), perm)) if n <= m else list(zip(perm, range(k)))
        misses = [len(step_misses(i, expect_steps[i], tasks[j], mode)) for i, j in pairs]
        key = (-sum(x == 0 for x in misses), sum(misses))
        if best_key is None or key < best_key:
            best, best_key = pairs, key
    return sorted(best or [])


def check(expect: dict, rec: dict, mode: str = 'full') -> tuple[list[str], list[dict]]:
    """Why a run record misses these expectations, as (reasons, codes); both empty when it meets them. In route mode
    only the routing checks run (there is no answer to check)."""
    reasons, codes = [], []

    def add(code, reason, **fields):
        reasons.append(reason)
        codes.append({'code': code, **{k: v for k, v in fields.items() if v is not None}})

    tasks = rec.get('tasks', [])
    agents = [t['agent'] for t in tasks if t.get('agent')]
    answer = answer_of(rec)
    if rec.get('status') != 'done':
        add('run_status', f"run ended {rec.get('status')}")
    want_agents = expect.get('expect_agents') or []
    for a in want_agents:
        if a not in agents:
            add('missing_agent', f'expected agent {a}, got {", ".join(agents) or "none"}', want=a,
                got=', '.join(agents) or None)
    if want_agents:
        for a in dict.fromkeys(agents):
            if a not in want_agents and a not in GUARDS:
                add('extra_agent', f'unexpected agent {a} (expected only {", ".join(want_agents)})', got=a,
                    want=', '.join(want_agents))
    for a in expect.get('forbid_agents') or []:
        if a in agents:
            add('forbidden_agent', f'agent {a} must not be used', got=a)
    if expect.get('expect_outcome'):
        want = as_outcomes(expect['expect_outcome'])
        got = outcome(rec, mode)
        if mode == 'route' and refusal_only(expect) and 'unsupported' in want:
            # route mode has no answer to match @cant against, so an honest refusal must show in the routing: an
            # unsupported, clarify or blocked step, or the policy's cant_do answer. Any other agent would have answered.
            want = [o for o in want if o in ('unsupported', 'clarify', 'blocked')]
            if got == 'answer' and any(e.get('rule') == 'cant_do' for t in tasks for e in t.get('trace') or []):
                got = 'unsupported'
        if mode == 'route' and got == 'answer' and 'error' in want:
            got = 'error'  # route mode can't tell an answer from an error; either was allowed
        if mode == 'route' and got == 'answer' and 'clarify' in want and route_term(tasks):
            got = 'clarify'  # a lone term the run would look up and may ask about; route mode checks it was found
        if got not in want:
            add('wrong_outcome', f"expected outcome {'/'.join(want)}, got {got}", want='/'.join(want), got=got)
    steps = expect.get('expect_steps')
    if steps:
        if len(tasks) != len(steps):
            add('plan_shape', f'expected {len(steps)} steps, got {len(tasks)}', want=str(len(steps)),
                got=str(len(tasks)), stage='plan_shape')
        for i, j in align(steps, tasks, bool(expect.get('unordered')), mode):
            for code, reason, fields in step_misses(i, steps[i], tasks[j], mode):
                add(code, reason, **fields)
    p = expect.get('expect_plan')
    if p:
        subs = [s.get('text') or '' for s in ((rec.get('plan') or {}).get('subtasks') or tasks)]
        if 'min_steps' in p and len(subs) < p['min_steps']:
            add('plan_shape', f"planned {len(subs)} steps, fewer than {p['min_steps']}", stage='plan_shape')
        if 'max_steps' in p and len(subs) > p['max_steps']:
            add('plan_shape', f"planned {len(subs)} steps, more than {p['max_steps']}", stage='plan_shape')
        if 'file_steps' in p:
            from .planner import is_file_request
            n = sum(bool(is_file_request(s)) for s in subs)
            if n != p['file_steps']:
                add('plan_shape', f"planned {n} file steps, expected {p['file_steps']}", stage='plan_shape')
        if p.get('no_fragment'):
            for i, s in enumerate(subs):
                if fragment(s):
                    add('plan_shape', f'step {i} is a fragment: {s[:60]!r}', step=i, stage='plan_text')
        if 'max_clarify' in p and agents.count('clarify') > p['max_clarify']:
            add('wrong_outcome', f"{agents.count('clarify')} steps asked to clarify, at most {p['max_clarify']} may",
                stage='clarity')
    if mode == 'full':
        if expect.get('must_match') and not search(expect['must_match'], answer):
            add('answer_regex', f"answer does not match /{expect['must_match']}/")
        if expect.get('must_not_match') and search(expect['must_not_match'], answer):
            add('forbidden_regex', f"answer matches forbidden /{expect['must_not_match']}/")
        for m in expect.get('must_mention') or []:
            if not search(m, answer):
                add('answer_regex', f'answer does not mention /{m}/')
        if expect.get('no_repeat') and (s := repeated_sentence(answer)):
            add('answer_regex', f'answer repeats a sentence: {s[:80]!r}')
    return reasons, codes


def score(case: dict, rec: dict, mode: str = 'full') -> dict:
    """One scored single-run case, in the shape GET /api/evals/{id} returns."""
    tasks = rec.get('tasks', [])
    reasons, codes = check(case, rec, mode)
    return {'id': case['id'], 'query': case_query(case), 'tags': case.get('tags', []), 'pass': not reasons,
            'reasons': reasons, 'codes': codes, 'agents': [t['agent'] for t in tasks if t.get('agent')],
            'answer': answer_of(rec), 'ms': rec.get('total_ms'), 'qid': rec.get('qid'),
            'expect_agents': case.get('expect_agents'),
            'expect_outcome': '/'.join(as_outcomes(case['expect_outcome'])) if case.get('expect_outcome') else None,
            'all_ok': bool(tasks) and all(t.get('ok') for t in tasks),
            'jev_tokens': (rec.get('tokens') or {}).get('jev_in', 0)}


# ---------- where a step went wrong (D3 stages) ----------

def expected_agents(expect: dict, rec: dict) -> list[str | None]:
    """The expected agent per task of a run, or None where the case says nothing about that step. From expect_steps
    (a regex: a plain name, or the alternative the run matched), or for a one-step run a single expect_agents entry or
    a guard outcome. Never over a bag of several expected agents."""
    tasks = rec.get('tasks', [])
    out: list[str | None] = [None] * len(tasks)
    steps = expect.get('expect_steps')
    if steps:
        for i, j in align(steps, tasks, bool(expect.get('unordered')), 'route'):
            pat = steps[i].get('agent')
            if not pat:
                continue
            got = tasks[j].get('agent')
            if fullmatch(pat, got):
                out[j] = got
            else:
                first = re.split(r'\|', expand(pat).strip('^$()?:'))[0]
                out[j] = first if re.fullmatch(r'[a-z_-]+', first or '') else pat
        return out
    if len(tasks) == 1:
        agents = expect.get('expect_agents') or []
        outs = as_outcomes(expect.get('expect_outcome'))
        if len(agents) == 1:
            out[0] = agents[0]
        elif len(outs) == 1 and outs[0] in ('clarify', 'blocked', 'unsupported'):
            out[0] = outs[0]
    return out


def initial_agent(t: dict) -> str | None:
    """Jev's own decision before the policy: jev.decide over the step's stored scores."""
    from types import SimpleNamespace as NS
    from . import jev as jev_mod
    if not t.get('pick'):
        return t.get('agent')
    try:
        return jev_mod.decide(NS(choice=t['pick'], confidence=float(t.get('confidence') or 0),
                                 probabilities=t.get('probabilities') or {}),
                              float(t.get('unsafe') or 0), float(t.get('clear') or 0))[0]
    except Exception:
        return t['pick']


def stage_of(t: dict, expected: str) -> str | None:
    """The failure stage of a step whose final agent isn't `expected` (None when it is): the first trace entry that
    moved the step away from its expected agent, else what Jev itself decided."""
    ok = lambda a: bool(a) and (a == expected or fullmatch(expected, a))
    if ok(t.get('agent')):
        return None
    trace = t.get('trace') or []
    if t.get('forced') or any(e.get('rule') == 'forced' for e in trace):
        return 'forced'
    prev = initial_agent(t)
    for e in trace:
        if ok(prev) and not ok(e.get('agent')):
            return RULE_STAGE.get(e.get('rule'), 'jev_pick')
        prev = e.get('agent') or prev
    if ok(t.get('pick')):  # Jev picked right; its own clarity or confidence veto turned it away
        from .config import MIN_CLEAR
        return 'clarity' if float(t.get('clear') or 0) < MIN_CLEAR else 'confidence'
    return 'jev_pick'


def step_routes(expect: dict, rec: dict) -> list[dict]:
    """StepRoute per task: its routing and, where the case expects an agent, that agent and the failure stage."""
    tasks = rec.get('tasks', [])
    want = expected_agents(expect, rec)
    out = []
    for t, e in zip(tasks, want):
        out.append({'tid': t.get('tid'), 'agent': t.get('agent'), 'pick': t.get('pick'),
                    'confidence': t.get('confidence'), 'clear': t.get('clear'), 'expected': e,
                    'stage': stage_of(t, e) if e else None, 'trace': t.get('trace') or [],
                    # what the calibration sweep needs to decide the step again
                    'probabilities': t.get('probabilities') or {}, 'unsafe': t.get('unsafe'),
                    'signals': t.get('signals') or {}, 'input': t.get('input') or t.get('text') or '',
                    'depends_on': t.get('depends_on') or [],
                    **({'forced': True} if t.get('forced') else {}), **({'bound': True} if t.get('bound') else {})})
    return out


def guard_stage(steps: list[dict]) -> str | None:
    """The stage of the first trace entry, in step order, that turned a step Jev routed to an agent into a guard."""
    for t in steps:
        prev = initial_agent(t)
        for e in t.get('trace') or []:
            if e.get('agent') in GUARDS and prev not in GUARDS and e.get('rule') in RULE_STAGE:
                return RULE_STAGE[e['rule']]
            prev = e.get('agent') or prev
    return None


def with_stages(codes: list[dict], steps: list[dict]) -> list[dict]:
    """Each code with a stage: its own, the failing step's, or the stage its kind stands for. When no step has an
    expected agent to judge it by, a routing miss is put on the first policy rule that turned a step into a guard
    ("time in the capital" clarified by missing_slot), and on Jev's pick only when no rule did."""
    lone = next((s['stage'] for s in steps if s.get('stage')), None) or guard_stage(steps)
    out = []
    for c in codes:
        if 'stage' not in c and c['code'] == 'wrong_outcome' and c.get('got') in GUARDS:
            # a guard the case didn't want: where did the step Jev picked an agent for turn into it?
            s = next((s for s in steps if s.get('agent') == c['got'] and s.get('pick') and s['pick'] not in GUARDS), None)
            c = {**c, 'stage': (stage_of(s, s['pick']) if s else None) or 'jev_pick'}
        if 'stage' not in c:
            stage = CODE_STAGE.get(c['code'])
            if stage is None and c['code'] in ROUTE_CODES - {'run_status'}:
                at = c.get('step')
                stage = (steps[at]['stage'] if at is not None and at < len(steps) and steps[at].get('stage')
                         else lone or 'jev_pick')
            c = {**c, **({'stage': stage} if stage else {})}
        out.append(c)
    return out


# ---------- created files (docs/PLAN-files.md, "Evals"; accuracy v2 D2) ----------

def created_files(rec: dict) -> list[dict]:
    """Every CreatedFile the run's steps made, in step order."""
    return [f for t in rec.get('tasks', []) for f in (t.get('created_files') or []) if isinstance(f, dict)]


def read_created(store, fid: str) -> bytes:
    """A created file's bytes from the store, by id only (X5): the store's own reader when it has one, else the plan's
    data/created/<id> next to the uploads."""
    for name in ('created_raw', 'created_bytes', 'get_created_bytes'):
        if callable(getattr(store, name, None)):
            return getattr(store, name)(fid)
    if not CASE_ID.match(fid):
        raise ValueError(f'bad created file id {fid!r}')
    base = getattr(store, 'created_dir', None) or store.files_dir.parent / 'created'
    return (Path(base) / fid).read_bytes()


def spec_reader(store):
    """fid -> the stored spec (or None), when the store keeps specs."""
    get = getattr(store, 'created_spec', None)
    if not callable(get):
        return None

    def read(fid):
        try:
            return get(fid)
        except Exception:
            return None
    return read


GRAY_OP = re.compile(rb'(-?[\d.]+)\s+(-?[\d.]+)\s+(-?[\d.]+)\s+(rg|RG)\b')
CMYK_OP = re.compile(rb'(-?[\d.]+)\s+(-?[\d.]+)\s+(-?[\d.]+)\s+(-?[\d.]+)\s+(k|K)\b')


def pdf_scan(r) -> dict:
    """What the pages hold: image XObjects (with their colour spaces), BaseFont names, and whether every colour set in
    a content stream is a grey (r == g == b, or CMYK with no colour) and every image is DeviceGray (the V8 scan)."""
    from pypdf.generic import IndirectObject

    def deref(o):
        return o.get_object() if isinstance(o, IndirectObject) else o

    images, fonts, gray, seen = 0, set(), True, set()

    def colours_gray(data: bytes) -> bool:
        for m in GRAY_OP.finditer(data):
            try:
                a, b, c = (float(x) for x in m.groups()[:3])
            except ValueError:
                continue
            if max(a, b, c) - min(a, b, c) > 1e-3:
                return False
        for m in CMYK_OP.finditer(data):
            try:
                c, mm, y = (float(x) for x in m.groups()[:3])
            except ValueError:
                continue
            if max(c, mm, y) > 1e-3:
                return False
        return True

    def walk(resources, depth=0):
        nonlocal images, gray
        resources = deref(resources) or {}
        for f in (deref(resources.get('/Font')) or {}).values():
            name = deref(f).get('/BaseFont') if hasattr(deref(f), 'get') else None
            if name:
                fonts.add(str(name).lstrip('/').split('+')[-1])
        for x in (deref(resources.get('/XObject')) or {}).values():
            ref = id(deref(x))
            if ref in seen:
                continue
            seen.add(ref)
            obj = deref(x)
            sub = obj.get('/Subtype')
            if sub == '/Image':
                images += 1
                cs = deref(obj.get('/ColorSpace'))
                if cs != '/DeviceGray' and not (obj.get('/ImageMask')):
                    gray = False
            elif sub == '/Form' and depth < 3:
                try:
                    if not colours_gray(obj.get_data()):
                        gray = False
                except Exception:
                    pass
                walk(obj.get('/Resources'), depth + 1)

    for page in r.pages:
        try:
            content = page.get_contents()
            if content is not None and not colours_gray(content.get_data()):
                gray = False
        except Exception:
            pass
        walk(page.get('/Resources'))
    return {'images': images, 'fonts': sorted(fonts), 'gray': gray}


def outline_count(r) -> int:
    def count(items):
        return sum(count(i) if isinstance(i, list) else 1 for i in items)
    try:
        return count(r.outline)
    except Exception:
        return 0


def reopen(fmt: str, data: bytes) -> dict:
    """A created file reopened with its own library (V1): its text, and its pages, slides, sheets and charts, plus (D2)
    words, headings, images, fonts and, for a PDF, whether it is all grey. Raises when the file does not open."""
    import io
    buf = io.BytesIO(data)
    out = {'text': '', 'pages': None, 'slides': None, 'sheets': [], 'charts': 0, 'formulas': [], 'headings': 0,
           'images': 0, 'fonts': [], 'gray': None, 'page_texts': []}
    if fmt == 'md':
        out['text'] = data.decode('utf-8')
        out['headings'] = len(re.findall(r'^#{1,6}\s', out['text'], re.M))
        out['images'] = len(re.findall(r'!\[[^\]]*\]\(', out['text']))
    elif fmt == 'pdf':
        from pypdf import PdfReader
        r = PdfReader(buf)
        out['pages'] = len(r.pages)
        title = (r.metadata or {}).get('/Title') or ''
        out['page_texts'] = [(p.extract_text() or '') for p in r.pages]
        out['text'] = '\n'.join([str(title), *out['page_texts']])
        out['headings'] = outline_count(r)
        scan = pdf_scan(r)
        out['images'], out['fonts'], out['gray'] = scan['images'], scan['fonts'], scan['gray']
    elif fmt == 'docx':
        import docx
        d = docx.Document(buf)
        parts = [d.core_properties.title or '', *(p.text for p in d.paragraphs)]
        parts += [c.text for t in d.tables for row in t.rows for c in row.cells]
        out['text'] = '\n'.join(parts)
        out['headings'] = sum(1 for p in d.paragraphs if (p.style.name or '').startswith(('Heading', 'Title')))
        out['images'] = len(d.inline_shapes)
        out['fonts'] = sorted({r.font.name for p in d.paragraphs for r in p.runs if r.font.name})
    elif fmt == 'pptx':
        from pptx import Presentation
        prs = Presentation(buf)
        out['slides'] = len(prs.slides)
        parts = []
        for slide in prs.slides:
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame and slide.shapes.title.text:
                out['headings'] += 1
            for sh in slide.shapes:
                if sh.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                    out['images'] += 1
                if sh.has_text_frame:
                    parts.append(sh.text_frame.text)
                if getattr(sh, 'has_table', False) and sh.has_table:
                    parts += [c.text for row in sh.table.rows for c in row.cells]
                if getattr(sh, 'has_chart', False) and sh.has_chart:
                    out['charts'] += 1
                    ch = sh.chart
                    if ch.has_title and ch.chart_title.has_text_frame:
                        parts.append(ch.chart_title.text_frame.text)
                    parts += [str(c) for c in ch.plots[0].categories] if len(ch.plots) else []
            if slide.has_notes_slide:
                parts.append(slide.notes_slide.notes_text_frame.text)
        out['text'] = '\n'.join(parts)
        # X2 in slides too: each chart's data is an embedded workbook, where an "=..." label would be a live formula
        from .create.rules import _embedded_formulas
        out['formulas'] += ['=' + f for f in _embedded_formulas(data)]
    elif fmt == 'xlsx':
        import openpyxl
        wb = openpyxl.load_workbook(buf)  # formulas as written, so an injected one shows up
        out['sheets'] = wb.sheetnames
        parts = []
        for ws in wb.worksheets:
            out['charts'] += len(getattr(ws, '_charts', []))
            out['images'] += len(getattr(ws, '_images', []))
            parts.append(ws.title)
            for row in ws.iter_rows():
                for c in row:
                    if c.value is None:
                        continue
                    if c.data_type == 'f':
                        out['formulas'].append(str(c.value))
                    parts.append(str(c.value))
        out['text'] = '\n'.join(parts)
    else:
        raise ValueError(f'unknown format {fmt!r}')
    out['words'] = len(re.findall(r'\b\w+\b', out['text']))
    return out


def diagrams_of(meta: dict, spec: dict | None) -> int:
    """Drawn diagrams: the file's own count when its meta has one, else the stored spec's diagram blocks."""
    if type(meta.get('diagrams')) is int:
        return meta['diagrams']
    n = 0

    def walk(x):
        nonlocal n
        if isinstance(x, dict):
            if x.get('type') in DIAGRAM_BLOCKS:
                n += 1
            for v in x.values():
                walk(v)
        elif isinstance(x, list):
            for v in x:
                walk(v)
    walk(spec or {})
    return n


def pick_primary(files: list[dict], fmt: str) -> dict | None:
    """The file a case scores: the primary one in its format when a file says so, else the first in that format."""
    of_fmt = [f for f in files if f.get('format') == fmt]
    return next((f for f in of_fmt if f.get('role') == 'primary'), of_fmt[0] if of_fmt else None)


def body_checks(files: list[dict], read, opened: dict) -> list[str]:
    """Always-on (D2): no created file's body may be the create agent's reply line."""
    reasons = []
    for f in files:
        name = f.get('name') or f.get('id')
        got = opened.get(f.get('id'))
        if got is None:
            try:
                got = opened[f.get('id')] = reopen(f.get('format'), read(f['id']))
            except Exception:
                continue  # the scored file reports its own reopen failure; an extra that won't open is not a body
        if TEMPLATE_BODY.search(got['text']):
            reasons.append(f'{name} file body is a create reply template')
    return reasons


def score_files(expect, rec: dict, read, spec_of=None, default_exact: int | None = None) -> tuple[list[str], dict | None]:
    """(reasons, FileScore) for the files a run created against expect_file. read(fid) returns a file's bytes and
    spec_of(fid) its stored spec. Every file is checked for a reply-template body; the scored file is the one in the
    expected format and the others count towards files_exact and only_formats. expect_file false means no file may be
    created; None checks only the always-on rules."""
    files = created_files(rec)
    if expect is False:
        return ([f"expected no file, got {', '.join(f.get('name', '?') for f in files)}"] if files else []), None
    opened: dict = {}
    if expect is None:
        return body_checks(files, read, opened), None
    fmt = expect['format']
    f = pick_primary(files, fmt)
    if f is None:
        got = ', '.join(f.get('name') or f.get('format', '?') for f in files) or 'none'
        return [f'expected a {fmt} file, got {got}'], None
    name = f.get('name') or f.get('id')
    try:
        got = opened[f['id']] = reopen(fmt, read(f['id']))
    except Exception as e:
        return [f'{name} does not reopen ({str(e)[:120]})'], None
    reasons = []
    for pat in expect.get('contains', []):
        if not re.search(pat, got['text'], re.I):
            reasons.append(f'{name} does not contain /{pat}/')
    pages = got['pages'] if got['pages'] is not None else f.get('pages')
    if 'min_pages' in expect and (pages or 0) < expect['min_pages']:
        reasons.append(f"{name} has {pages} pages, fewer than {expect['min_pages']}")
    if 'max_pages' in expect and pages is not None and pages > expect['max_pages']:
        reasons.append(f"{name} has {pages} pages, more than {expect['max_pages']}")
    if 'pages' in expect:
        lo, hi = expect['pages']
        if pages is None or not lo <= pages <= hi:
            reasons.append(f'{name} has {pages} pages, outside {lo}-{hi}')
    slides = got['slides'] if got['slides'] is not None else f.get('slides')
    if 'slides_min' in expect and (got['slides'] or 0) < expect['slides_min']:
        reasons.append(f"{name} has {got['slides'] or 0} slides, fewer than {expect['slides_min']}")
    if 'slides' in expect:
        lo, hi = expect['slides']
        if slides is None or not lo <= slides <= hi:
            reasons.append(f'{name} has {slides} slides, outside {lo}-{hi}')
    for pat in expect.get('sheets', []):
        if not any(re.search(pat, s, re.I) for s in got['sheets']):
            reasons.append(f"{name} has no sheet like /{pat}/ (sheets: {', '.join(got['sheets']) or 'none'})")
    if 'charts_min' in expect and got['charts'] < expect['charts_min']:
        reasons.append(f"{name} has {got['charts']} charts, fewer than {expect['charts_min']}")
    if 'words_min' in expect and got['words'] < expect['words_min']:
        reasons.append(f"{name} has {got['words']} words, fewer than {expect['words_min']}")
    if 'headings_min' in expect and got['headings'] < expect['headings_min']:
        reasons.append(f"{name} has {got['headings']} headings, fewer than {expect['headings_min']}")
    if 'images_min' in expect and got['images'] < expect['images_min']:
        reasons.append(f"{name} has {got['images']} images, fewer than {expect['images_min']}")
    spec = spec_of(f['id']) if spec_of else None
    diagrams = diagrams_of(f, spec)
    if 'diagrams_min' in expect and diagrams < expect['diagrams_min']:
        reasons.append(f"{name} has {diagrams} diagrams, fewer than {expect['diagrams_min']}")
    if 'fonts' in expect and not any(re.search(expect['fonts'], x, re.I) for x in got['fonts']):
        reasons.append(f"{name} uses no font like /{expect['fonts']}/ (fonts: {', '.join(got['fonts']) or 'none'})")
    if expect.get('grayscale') and got['gray'] is not True:
        reasons.append(f'{name} is not all grey' if got['gray'] is False else f'{name} cannot be checked for grey')
    if 'source' in expect and f.get('source') != expect['source']:
        reasons.append(f"{name} came from {f.get('source')}, expected {expect['source']}")
    for s in expect.get('not_contains', []):
        if s.casefold() in got['text'].casefold():
            reasons.append(f'{name} contains {s!r}')
    if expect.get('credits') and got['images']:
        credits = f.get('credits')
        if not re.search(r'image credits', got['text'], re.I) or (isinstance(credits, list) and len(credits) < got['images']):
            reasons.append(f"{name} has {got['images']} images without an image credits section")
    if 'design_bg' in expect:
        want = expect['design_bg'].lstrip('#').upper()
        design = f.get('design') if isinstance(f.get('design'), dict) else None
        bg = str(((design or {}).get('colors') or {}).get('bg') or '').lstrip('#').upper()
        if design is None:
            reasons.append(f'{name} has no design applied (expected background {want})')
        elif bg != want:
            reasons.append(f"{name} has background {bg or 'none'} from its design, expected {want}")
    if expect.get('partial_ok') is False and isinstance(f.get('partial'), dict):
        p = f['partial']
        missing = ', '.join(str(h) for h in (p.get('missing') or [])[:5]) or 'unnamed sections'
        reasons.append(f"{name} is partial: {p.get('written', '?')} of {p.get('planned', '?')} written "
                       f"(missing {missing})")
    # X2 always: a formula outside the allow-list means a cell's text was written as a live formula
    if bad := [x for x in got['formulas'] if not SAFE_FORMULA.match(x)]:
        reasons.append(f'{name} has an unsafe formula {bad[0][:40]!r}')
    if expect.get('rules_ok'):
        failed = [r.get('id', '?') for r in f.get('rules') or []
                  if not r.get('ok') and r.get('severity') in ('block', 'warn')]
        if failed:
            reasons.append(f"{name} failed rules {', '.join(failed)}")
    exact = expect.get('files_exact', default_exact)
    if exact is not None and len(files) != exact:
        extras = [x for x in files if x is not f]
        if len(files) > exact and extras:
            reasons += [f"extra file .{x.get('format')} ({x.get('name') or x.get('id')})" for x in extras[:max(1, len(files) - exact)]]
        else:
            reasons.append(f'made {len(files)} files, expected {exact}')
    if 'only_formats' in expect:
        for x in files:
            if x.get('format') not in expect['only_formats']:
                reasons.append(f"{x.get('name') or x.get('id')} is .{x.get('format')}, not {'/'.join(expect['only_formats'])}")
    reasons += [r for r in body_checks(files, read, opened) if r not in reasons]
    fs = {'format': fmt, 'pages': pages, 'slides': slides, 'words': got['words'], 'headings': got['headings'],
          'images': got['images'], 'diagrams': diagrams, 'fonts': got['fonts'], 'grayscale': got['gray'],
          'source': f.get('source'), 'files': len(files), 'id': f.get('id'), 'name': name,
          '_sample': sample_text(got)}
    return reasons, fs


def check_file(expect, rec: dict, read, spec_of=None) -> list[str]:
    """Why the files a run created miss expect_file ([] when they meet it). See score_files."""
    return score_files(expect, rec, read, spec_of)[0]


def sample_text(got: dict, n: int = 1500) -> str:
    """About n chars from the first, middle and last pages (or the start, middle and end of the text)."""
    pages = got.get('page_texts') or []
    if len(pages) >= 3:
        parts = [pages[0], pages[len(pages) // 2], pages[-1]]
    else:
        t = got.get('text') or ''
        third = max(1, len(t) // 3)
        parts = [t[:third], t[third:2 * third], t[2 * third:]]
    each = n // 3
    return '\n...\n'.join(p.strip()[:each] for p in parts if p.strip())


def file_summary(fs: dict | None, headings: list[str] | None = None) -> str | None:
    """The judge's FILE: block (D5) for the scored file."""
    if not fs:
        return None
    size = f"{fs['pages']} pages" if fs.get('pages') is not None else f"{fs['slides']} slides" if fs.get('slides') is not None else ''
    lines = [f"format {fs['format']}{', ' + size if size else ''}, {fs['words']} words, {fs['headings']} headings",
             f"images {fs['images']}, diagrams {fs['diagrams']}, fonts: {', '.join(fs['fonts']) or 'unknown'}"]
    if headings:
        lines.append('headings: ' + '; '.join(headings[:20]))
    if fs.get('_sample'):
        lines.append('sample:\n' + fs['_sample'])
    return '\n'.join(lines)


def public_file(fs: dict | None) -> dict | None:
    return {k: v for k, v in fs.items() if not k.startswith('_')} if fs else None


def file_expected(expect: dict, case: dict) -> int | None:
    """files_exact's default (D2): 1 for a turn with expect_file, unless its expect_steps name several file steps."""
    if not isinstance(expect.get('expect_file'), dict):
        return None
    steps = expect.get('expect_steps') or []
    return None if sum('format' in s for s in steps) > 1 else 1


# ---------- one turn, one attempt, one case ----------

def run_tokens(rec: dict) -> dict:
    t = rec.get('tokens') or {}
    return {k: int(t.get(k) or 0) for k in ('jev_in', 'llm_in', 'llm_out')}


def turn_result(t: dict, rec: dict, read=None, spec_of=None, mode: str = 'full', case: dict | None = None) -> dict:
    s = score({'id': '', **t}, rec, mode)
    fs = None
    files = created_files(rec)
    if mode == 'full' and ('expect_file' in t or files):
        if read is None:
            if 'expect_file' in t:
                s['reasons'].append('created files cannot be checked here')
                s['codes'].append({'code': 'file'})
        else:
            file_reasons, fs = score_files(t.get('expect_file'), rec, read, spec_of, file_expected(t, case or t))
            s['reasons'] += file_reasons
            s['codes'] += [{'code': 'file'}] * len(file_reasons)
        s['pass'] = not s['reasons']
    steps = step_routes(t, rec)
    out = {k: s[k] for k in ('query', 'pass', 'reasons', 'answer', 'agents', 'ms', 'qid', 'all_ok', 'jev_tokens')}
    out['codes'] = with_stages(s['codes'], steps)
    out['steps'] = steps
    out['file'] = fs
    out['tokens'] = run_tokens(rec)
    out['timings'] = rec.get('timings') or {}
    out['file_tokens'] = sum(int(f.get('tokens') or 0) for f in files)
    # the engines that wrote this answer (steps and merge), so the judge can be one that didn't
    out['engines'] = sorted({e for e in [*(t.get('engine') for t in rec.get('tasks', [])),
                                         (rec.get('merged') or {}).get('engine')] if e and e != 'keyless'})
    return out


def attempt_label(n: int, phrasing: int, repeat: int, phrasings: int) -> str:
    parts = ([f'run {n + 1}'] if repeat > 1 else []) + ([f'paraphrase {phrasing}'] if phrasing else
                                                          ['original'] if phrasings > 1 else [])
    return ', '.join(parts)


def tri(values) -> bool | None:
    """False when any is False, else None when any is None, else True (an empty list is None)."""
    values = list(values)
    if not values:
        return None
    return False if any(v is False for v in values) else None if any(v is None for v in values) else True


def has_route_checks(case: dict) -> bool:
    return routable(case)


def has_answer_checks(case: dict) -> bool:
    return bool(case.get('judge')) or any(k in t for t in [case, *(case.get('turns') or [])] for k in ANSWER_EXPECT)


def attempt_passes(a: dict, case: dict, mode: str) -> tuple[bool | None, bool | None]:
    """(route_pass, answer_pass) of one attempt: None where the case has no such checks (or, in route mode, answers)."""
    kinds = {c['code'] for c in a.get('codes', [])}
    route = None if not has_route_checks(case) else not (kinds & ROUTE_CODES)
    if mode == 'route' or not has_answer_checks(case):
        answer = None
    elif a.get('unjudged') and not (kinds & ANSWER_CODES):
        answer = None
    else:
        answer = not (kinds & ANSWER_CODES)
    if a.get('unrecorded'):
        route = answer = None
    return route, answer


def combine(case: dict, attempts: list[dict], mode: str = 'full') -> dict:
    """One case's result from its attempts. The shown answer, agents and turns are the first failing attempt's (or the
    first attempt's when all passed); the case passes only when every attempt passed. A case with a rubric no judge
    scored, or that met a cassette miss, and nothing else wrong, has pass None: it is left out of passed and total."""
    shown = next((a for a in attempts if a['pass'] is False), attempts[0])
    turns = shown['turns']
    reasons = []
    for a in attempts:
        for r in a['reasons']:
            r = f"{a['label']}: {r}" if a['label'] else r
            if r not in reasons:
                reasons.append(r)
    codes = []
    for a in attempts:
        for c in a.get('codes', []):
            if c not in codes:
                codes.append(c)
    route_pass = tri(attempt_passes(a, case, mode)[0] for a in attempts if attempt_passes(a, case, mode)[0] is not None)
    answer_pass = tri(attempt_passes(a, case, mode)[1] for a in attempts if attempt_passes(a, case, mode)[1] is not None)
    tokens = {k: sum(t.get('tokens', {}).get(k, 0) for t in turns) for k in ('jev_in', 'llm_in', 'llm_out')}
    fs = next((t['file'] for t in reversed(turns) if t.get('file')), None)
    out = {'id': case['id'], 'query': case_query(case), 'tags': case.get('tags', []),
           'pass': tri(a['pass'] for a in attempts),
           'reasons': reasons, 'agents': [a for t in turns for a in t['agents']], 'answer': turns[-1]['answer'],
           'ms': sum(t['ms'] or 0 for t in turns), 'qid': turns[-1]['qid'],
           'expect_agents': case.get('expect_agents'),
           'expect_outcome': '/'.join(as_outcomes(case['expect_outcome'])) if case.get('expect_outcome') else None,
           'all_ok': all(t['all_ok'] for t in turns), 'jev_tokens': sum(t['jev_tokens'] or 0 for t in turns),
           'kind': case_kind(case), 'split': case.get('split', 'dev'),
           'judge': shown['judge'] if shown['judge'] else next((a['judge'] for a in attempts if a['judge']), None),
           'max_ms': case.get('max_ms'), 'over_budget': any(a['over_budget'] for a in attempts),
           # flaky: the same text passed on one run and failed on another (paraphrases are compared to themselves)
           'flaky': any(len({a['pass'] for a in attempts if a.get('phrasing', 0) == i and a['pass'] is not None}) > 1
                        for i in {a.get('phrasing', 0) for a in attempts}),
           'route_pass': route_pass, 'answer_pass': answer_pass, 'codes': codes,
           'steps': [s for t in turns for s in t.get('steps', [])], 'file': public_file(fs), 'tokens': tokens,
           'file_tokens': sum(t.get('file_tokens', 0) for t in turns), 'case_sha': case_sha(case),
           'timings': [t.get('timings') or {} for t in turns] if len(turns) > 1 else (turns[0].get('timings') or {})}
    if any(a.get('unjudged') for a in attempts):
        out['unjudged'] = True
    if any(a.get('unrecorded') for a in attempts):
        out['unrecorded'] = True
    chat = case.get('chat') or next((t.get('chat') for t in case.get('turns') or [] if t.get('chat')), None)
    if chat:
        out['chat'] = chat
    warnings = [w for a in attempts for w in a.get('budget_warnings', [])]
    if warnings:
        out['budget_warnings'] = list(dict.fromkeys(warnings))
    if err := next((a['judge_error'] for a in attempts if a.get('judge_error')), None):
        out['judge_error'] = err
    if case.get('turns'):
        keep = ('query', 'pass', 'reasons', 'answer', 'agents', 'ms', 'qid')
        out['turns'] = [{k: t[k] for k in keep} for t in turns]
    if sum(1 for a in attempts if a.get('phrasing', 0) == 0) > 1:  # repeat > 1
        out['attempts'] = [a['pass'] for a in attempts]
    return out


def percentile(values: list, p: float) -> int | None:
    """Nearest-rank percentile of the numbers in `values` (None when there are none)."""
    v = sorted(x for x in values if isinstance(x, (int, float)))
    return round(v[max(0, math.ceil(p / 100 * len(v)) - 1)]) if v else None


# ---------- route report (D3, D4) ----------

def confusion_of(cases: list[dict]) -> dict[str, dict[str, int]]:
    """{expected: {got: n}} over every step with an expected agent. A route-mode lone term ("Mercury") counts as the
    clarify it would ask, as check() scores it: only a real run looks its meanings up."""
    out: dict[str, dict[str, int]] = {}
    for c in cases:
        if c.get('unrecorded'):
            continue
        for s in c.get('steps') or []:
            if s.get('expected'):
                row = out.setdefault(s['expected'], {})
                got = 'clarify' if s['expected'] == 'clarify' and route_term([s]) else s.get('agent') or 'none'
                row[got] = row.get(got, 0) + 1
    return {k: dict(sorted(v.items())) for k, v in sorted(out.items())}


def per_agent(confusion: dict) -> list[dict]:
    agents = sorted({*confusion, *(g for row in confusion.values() for g in row)})
    out = []
    for a in agents:
        tp = confusion.get(a, {}).get(a, 0)
        fp = sum(row.get(a, 0) for e, row in confusion.items() if e != a)
        fn = sum(n for g, n in confusion.get(a, {}).items() if g != a)
        p = round(tp / (tp + fp), 4) if tp + fp else None
        r = round(tp / (tp + fn), 4) if tp + fn else None
        f1 = round(2 * p * r / (p + r), 4) if p and r else (0.0 if p is not None and r is not None else None)
        out.append({'agent': a, 'tp': tp, 'fp': fp, 'fn': fn, 'precision': p, 'recall': r, 'f1': f1,
                    'support': tp + fn})
    return out


def stage_histogram(cases: list[dict]) -> dict[str, int]:
    """One stage per failing case: the stage of its first code that has one."""
    out: dict[str, int] = {}
    for c in cases:
        if c.get('pass') is not False:
            continue
        stage = next((x['stage'] for x in c.get('codes') or [] if x.get('stage')), None)
        if stage:
            out[stage] = out.get(stage, 0) + 1
    return dict(sorted(out.items(), key=lambda kv: (-kv[1], kv[0])))


def reliability(points: list[tuple[float, bool]], signal: str, bins: int = 10) -> dict:
    """Calibration of one signal: 10 equal-width bins of {lo, hi, n, accuracy, mean} and the expected calibration
    error (ECE)."""
    out, total, ece = [], len(points), 0.0
    for b in range(bins):
        lo, hi = b / bins, (b + 1) / bins
        inside = [(x, ok) for x, ok in points if lo <= x < hi or (b == bins - 1 and x == hi)]
        n = len(inside)
        acc = round(sum(ok for _, ok in inside) / n, 4) if n else None
        mean = round(sum(x for x, _ in inside) / n, 4) if n else round((lo + hi) / 2, 4)
        if n:
            ece += n / total * abs(acc - mean)
        out.append({'lo': round(lo, 2), 'hi': round(hi, 2), 'n': n, 'accuracy': acc, 'mean': mean})
    return {'signal': signal, 'bins': out, 'ece': round(ece, 4) if total else None}


def calibration_of(cases: list[dict]) -> list[dict]:
    conf, clear = [], []
    for c in cases:
        if c.get('unrecorded'):
            continue
        for s in c.get('steps') or []:
            e = s.get('expected')
            if not e:
                continue
            if isinstance(s.get('confidence'), (int, float)):
                conf.append((float(s['confidence']), s.get('agent') == e or fullmatch(e, s.get('agent'))))
            if isinstance(s.get('clear'), (int, float)):
                clear.append((float(s['clear']), e != 'clarify'))
    return [reliability(conf, 'confidence'), reliability(clear, 'clear')]


def summary(e: dict) -> dict:
    """The stored and returned form of an eval: its fields plus the counts. Private working keys (_ms) are dropped.
    Cases with pass None (unjudged or unrecorded) are left out of passed and total, and counted on their own."""
    cases = e['cases']
    passed = sum(c['pass'] is True for c in cases)
    unjudged = sum(1 for c in cases if c.get('unjudged') and c['pass'] is None)
    unrecorded = sum(1 for c in cases if c.get('unrecorded') and c['pass'] is None)
    total = e['total'] - sum(c['pass'] is None for c in cases)
    by_tag: dict[str, dict] = {}
    by_tag_route: dict[str, dict] = {}
    for c in cases:
        for t in c.get('tags', []):
            if c['pass'] is not None:
                s = by_tag.setdefault(t, {'passed': 0, 'total': 0})
                s['total'] += 1
                s['passed'] += c['pass'] is True
            if c.get('route_pass') is not None:
                s = by_tag_route.setdefault(t, {'passed': 0, 'total': 0})
                s['total'] += 1
                s['passed'] += c['route_pass'] is True
    ms = e.get('_ms') or [t['ms'] for c in cases for t in (c.get('turns') or [c])]
    judged = [c['judge']['mean'] for c in cases if c.get('judge')]
    tags_score = lambda key: {'passed': sum(c.get(key) is True for c in cases),
                              'total': sum(c.get(key) is not None for c in cases)}
    tokens = {k: sum((c.get('tokens') or {}).get(k, 0) for c in cases) for k in ('jev_in', 'llm_in', 'llm_out')}
    out = {**{k: v for k, v in e.items() if not k.startswith('_')}, 'done': len(cases), 'passed': passed,
           'total': total, 'accuracy': round(passed / total, 4) if total else 0.0,
           'silent_wrong': sum(1 for c in cases if c['pass'] is False and c['all_ok'] and not c.get('judge_error')),
           'by_tag': dict(sorted(by_tag.items())), 'p50_ms': percentile(ms, 50), 'p95_ms': percentile(ms, 95),
           'flaky': sum(1 for c in cases if c.get('flaky')),
           'judge_mean': round(sum(judged) / len(judged), 2) if judged else None,
           'judge_errors': [f"{c['id']}: {c['judge_error']}" for c in cases if c.get('judge_error')],
           'unjudged': unjudged, 'unrecorded': unrecorded,
           'route_pass': tags_score('route_pass'), 'answer_pass': tags_score('answer_pass'),
           'by_tag_route': dict(sorted(by_tag_route.items())), 'tokens': tokens,
           'stages': stage_histogram(cases)}
    if e.get('mode') == 'route' or any(c.get('steps') for c in cases):
        conf = confusion_of(cases)
        out.update(confusion=conf, per_agent=per_agent(conf), calibration=calibration_of(cases))
    warnings = [f"{c['id']}: {w}" for c in cases for w in c.get('budget_warnings', [])]
    if warnings:
        out['budget_warnings'] = warnings
    return out


# ---------- gates and baselines (D7) ----------

def wilson_lower(passed: int, total: int, z: float = 1.96) -> float:
    """The lower bound of the Wilson 95% interval for a pass rate."""
    if total <= 0:
        return 0.0
    p = passed / total
    d = 1 + z * z / total
    centre = p + z * z / (2 * total)
    spread = z * math.sqrt(p * (1 - p) / total + z * z / (4 * total * total))
    return max(0.0, (centre - spread) / d)


def suite_name(mode: str, engine_name: str | None) -> str:
    """The name gates and baselines are kept under: route, keyless, or the engine's name."""
    return 'route' if mode == 'route' else 'keyless' if engine_name in (None, 'none') else engine_name


def gate_tags(result: dict, rules: dict, baseline: dict | None) -> list[dict]:
    """TagGate per gated tag. A tag passes when every case passed, when the Wilson lower bound of its pass rate is at
    least `min`, or when its rate is at most `max_drop` below the baseline's. A tag with no baseline yet (a new tag, or
    an auto tag such as forced:weather) is judged on its pass rate against `min`, since the lower bound of a handful of
    cases can't reach it. A tag no case ran is not reported. `*` applies to every tag without its own rule."""
    by_tag = result.get('by_tag_route') if result.get('mode') == 'route' else result.get('by_tag')
    by_tag = by_tag or result.get('by_tag') or {}
    rates = (baseline or {}).get('rates') or {}
    default = rules.get('*')
    out = []
    for tag in sorted(by_tag):
        rule = rules.get(tag, default)
        if not isinstance(rule, dict):
            continue
        s = by_tag[tag]
        if not s['total']:
            continue
        rate = s['passed'] / s['total']
        lower = wilson_lower(s['passed'], s['total'])
        base = rates.get(tag)
        lo = float(rule.get('min', 0.0))
        ok = s['passed'] == s['total'] or lower >= lo or (base is None and rate >= lo)
        if not ok and base is not None and 'max_drop' in rule and lo < 1.0:
            ok = rate >= base - float(rule['max_drop']) - 1e-4  # baselines keep 4 decimals
        out.append({'tag': tag, 'min': lo, 'passed': s['passed'], 'total': s['total'], 'rate': round(rate, 4),
                    'lower': round(lower, 4), 'baseline': base, 'ok': ok})
    return out


def strict_unrecorded(result: dict, rules: dict) -> list[str]:
    """Gate failures for cases left unscored (a cassette miss) in a tag that allows no failure (min 1.0: safety,
    injection, control): the MAX_UNRECORDED share must never hide one of them."""
    strict = {t for t, r in rules.items() if t != '*' and isinstance(r, dict) and float(r.get('min', 0)) >= 1.0}
    out = []
    for tag in sorted(strict):
        ids = [c['id'] for c in result.get('cases') or [] if c.get('unrecorded') and tag in (c.get('tags') or [])]
        if ids:
            out.append(f"{tag} unrecorded ({', '.join(ids)})")
    return out


def read_json(path) -> dict:
    try:
        d = json.loads(Path(path).read_text())
    except (OSError, ValueError):
        return {}
    return d if isinstance(d, dict) else {}


def baseline_for(baseline: dict, suite: str) -> dict | None:
    b = baseline.get(suite)
    return b if isinstance(b, dict) and isinstance(b.get('rates'), dict) else None


def save_baseline(baseline: dict, suite: str, result: dict, sha: str) -> dict:
    """Records this run's per-tag pass rates for its suite. A --split/--tags run updates its own tags and keeps the
    others; a suite_sha change starts over."""
    by_tag = (result.get('by_tag_route') if result.get('mode') == 'route' else None) or result.get('by_tag') or {}
    old = baseline_for(baseline, suite) or {}
    kept = old.get('rates', {}) if old.get('suite_sha') == sha else {}
    rates = {**kept, **{t: round(s['passed'] / s['total'], 4) for t, s in by_tag.items() if s['total']}}
    totals = {**(old.get('totals', {}) if old.get('suite_sha') == sha else {}),
              **{t: s['total'] for t, s in by_tag.items() if s['total']}}
    return {**baseline, suite: {'suite_sha': sha, 'at': round(time.time()), 'engine': result.get('engine'),
                                'rates': dict(sorted(rates.items())), 'totals': dict(sorted(totals.items()))}}


def tag_drops(result: dict, base: dict, max_drop: float = 0.1) -> list[str]:
    """Tags whose pass rate is more than max_drop below the baseline's (the check when no gates file is given)."""
    by_tag = (result.get('by_tag_route') if result.get('mode') == 'route' else None) or result.get('by_tag') or {}
    rates = base.get('rates') or {}
    return [t for t, s in sorted(by_tag.items())
            if s['total'] and t in rates and s['passed'] / s['total'] < rates[t] - max_drop - 1e-4]


def regressions(result: dict, baseline: dict) -> list[str]:
    """Cases that pass in an older per-case baseline ({engine: {case id: pass}}) and fail now."""
    before = baseline.get(result['engine'] or 'none') or {}
    if not all(isinstance(v, bool) for v in before.values()):
        return []
    return [c['id'] for c in result['cases'] if before.get(c['id']) and c['pass'] is False]


# ---------- budgets (D8) ----------

def engine_class(engine) -> str:
    if engine is None:
        return 'keyless'
    return 'api' if getattr(engine, 'billing', '') == 'api' else 'cli'


def budget_for(case: dict, budgets: dict, eclass: str) -> dict:
    """The budget for one case: evals/budgets.json by the case's kind, then by each of its tags (a tag wins), then the
    case's own `budget`."""
    out = {}
    for key in (case_kind(case), *case.get('tags', [])):
        b = (budgets.get(key) or {}).get(eclass)
        if isinstance(b, dict):
            out.update(b)
    out.update(case.get('budget') or {})
    return out


def budget_misses(b: dict, turns: list[dict]) -> tuple[list[str], list[str]]:
    """(hard, soft) budget misses of one attempt. `planner` is hard (it is what the run did, not how long it took);
    time and tokens are warnings unless the eval is strict."""
    hard, soft = [], []
    ms = sum(t['ms'] or 0 for t in turns)
    for key in ('ms', 'p95_ms'):
        if key in b and ms > b[key]:
            soft.append(f'took {ms} ms, over the {b[key]} ms budget')
            break
    tok = {k: sum(t.get('tokens', {}).get(k, 0) for t in turns) for k in ('jev_in', 'llm_in', 'llm_out')}
    for k in ('jev_in', 'llm_in', 'llm_out'):
        if k in b and tok[k] > b[k]:
            soft.append(f'used {tok[k]} {k} tokens, over the {b[k]} budget')
    ft = sum(t.get('file_tokens', 0) for t in turns)
    if 'file_tokens' in b and ft > b['file_tokens']:
        soft.append(f"its files used {ft} tokens, over the {b['file_tokens']} budget")
    for t in turns:
        tm = t.get('timings') or {}
        if 'planner' in b and tm.get('planner') and tm['planner'] not in b['planner']:
            hard.append(f"planned with {tm['planner']}, expected {' or '.join(b['planner'])}")
        if 'plan_ms' in b and (tm.get('plan_ms') or 0) > b['plan_ms']:
            soft.append(f"planning took {tm['plan_ms']} ms, over the {b['plan_ms']} ms budget")
    return hard, soft


# ---------- storage (the evals table's `extra` column, added here on first use) ----------

_migrated: 'weakref.WeakSet' = weakref.WeakSet()


def migrate(store):
    """Adds the `extra` JSON column to an older database's evals table. Safe to call again."""
    if store in _migrated:
        return
    if 'extra' not in {r['name'] for r in store.q('PRAGMA table_info(evals)')}:
        store.x('ALTER TABLE evals ADD COLUMN extra TEXT')
    _migrated.add(store)


def save(store, s: dict):
    migrate(store)
    store.save_eval(s)
    store.x('UPDATE evals SET extra = ? WHERE id = ?', json.dumps({k: s[k] for k in EXTRA_KEYS if k in s}), s['eval_id'])


def extras_of(store, ids: list[str]) -> dict[str, dict]:
    migrate(store)
    if not ids:
        return {}
    rows = store.q(f'SELECT id, extra FROM evals WHERE id IN ({",".join("?" * len(ids))})', *ids)
    out = {}
    for r in rows:
        try:
            d = json.loads(r['extra']) if r['extra'] else {}
        except ValueError:
            d = {}
        out[r['id']] = d if isinstance(d, dict) else {}
    return out


def get(store, eid: str) -> dict | None:
    """A stored eval with its cases and extra fields (older evals simply have none)."""
    e = store.get_eval(eid)
    return {**e, **extras_of(store, [eid]).get(eid, {})} if e else None


def listed(store, limit: int = 50) -> list[dict]:
    evals = store.list_evals(limit)
    extra = extras_of(store, [e['eval_id'] for e in evals])
    return [{**e, **extra.get(e['eval_id'], {})} for e in evals]


# ---------- running ----------

def cassette_for(router, jev: str):
    """The Jev a route-mode eval asks: the router's own (live), or the committed cassette (replay, record)."""
    from .cassette import JevCassette
    if jev == 'live':
        return JevCassette(router.jev, None, 'live')
    return JevCassette(router.jev if jev == 'record' else None, CASSETTE, jev)


def start(router, engine, engine_name: str | None, cases: list[dict] | None = None, examples: bool | None = None, *,
          split: str = 'all', tags=None, repeat: int = 1, judge=None, mode: str = 'full', jev: str = 'live') -> str:
    """Starts an eval in the background and returns its id. engine: an Engine or None (keyless). examples forces route
    examples on or off for this eval's runs only; None uses the global switch as it is now. split and tags pick the
    cases; repeat runs every attempt that many times; judge is the Engine that scores rubric cases (None: unjudged).
    mode 'route' only routes (no agent runs), asking Jev live or from the cassette (jev)."""
    eid = uuid.uuid4().hex[:10]
    picked = select(load_cases() if cases is None else cases, split, tags)
    cassette = cassette_for(router, jev) if mode == 'route' else None
    task = asyncio.create_task(run_eval(router, eid, engine, engine_name, picked, examples, split=split, tags=tags,
                                        repeat=repeat, judge=judge, mode=mode, jev=jev, cassette=cassette))
    router.evals[eid] = task
    task.add_done_callback(lambda t: router.evals.pop(eid, None))
    return eid


class Fixtures:
    """Fixture files uploaded once per eval through the store, like a user's upload, and deleted when it ends."""

    def __init__(self, store):
        self.store, self.ids, self.lock = store, {}, asyncio.Lock()

    async def ids_for(self, names) -> list[str]:
        from .files import extract
        async with self.lock:
            for name in names or ():
                if name not in self.ids:
                    raw = (FIXTURES / name).read_bytes()
                    meta, text = await asyncio.to_thread(extract, name, raw)
                    self.store.add_file(meta, raw, text)
                    self.ids[name] = meta['id']
        return [self.ids[n] for n in names or ()]

    def remove(self):
        for fid in self.ids.values():
            self.store.delete_file(fid)


def chat_for(case: dict, turn: dict) -> dict:
    """The run's chat options: the case's, then the turn's own. '@create' and 'create' are the same agent."""
    chat = {**(case.get('chat') or {}), **(turn.get('chat') or {})}
    if chat.get('agent'):
        chat['agent'] = chat['agent'].lstrip('@')
    return chat


def attempt_engine(router, ctx: dict, chat: dict):
    """The engine a turn runs on. Route mode: none (nothing runs, and the heuristic planner plans). Research mode
    takes a web engine the way app.mode_engine does, but only when the eval already runs on an engine: a keyless
    suite never starts spending an engine's quota."""
    if ctx['mode'] == 'route':
        return None
    e = ctx['engine']
    if chat.get('mode') == 'research' and e is not None and not getattr(e, 'supports_web', False):
        e = router.web_engine() or e
    return e


def picked_agent_problem(router, chat: dict, engine, files: list[str], text: str) -> str | None:
    """Why the case's @agent can't be picked on this run (app.check_agent), else None."""
    if not chat.get('agent'):
        return None
    from .app import Bad, check_agent
    attached = router.store.list_files(files) if files else []
    try:
        check_agent(router, {'mode': chat.get('mode') or 'balanced', 'agent': chat['agent']}, engine, attached, text)
    except Bad as e:
        return str(e)
    return None


async def run_attempt(router, ctx: dict, case: dict, text: str, label: str) -> dict:
    """One attempt at a case: every turn in order (one session when there are several), then the judge and budget."""
    turns = case.get('turns') or [{'query': text, **{k: case[k] for k in (*EXPECT_KEYS, 'unordered', 'plan')
                                                     if k in case}}]
    session_id = f'eval-{ctx["eid"]}-{uuid.uuid4().hex[:8]}' if len(turns) > 1 else None
    files = await ctx['fixtures'].ids_for(case.get('files'))
    mode = ctx.get('mode', 'full')
    results, unrecorded = [], False
    read = lambda fid: read_created(router.store, fid)
    spec_of = spec_reader(router.store)
    try:
        for t in turns:
            chat = chat_for(case, t)
            engine = attempt_engine(router, ctx, chat)
            if problem := picked_agent_problem(router, chat, engine, files, t['query']):
                results.append({'query': t['query'], 'pass': False, 'reasons': [problem], 'answer': '', 'agents': [],
                                'ms': None, 'qid': None, 'all_ok': False, 'jev_tokens': 0, 'engines': [],
                                'codes': [{'code': 'run_status'}], 'steps': [], 'file': None, 'tokens': {},
                                'timings': {}, 'file_tokens': 0})
                continue
            extras = {'holdout': t['query']}
            if t.get('plan'):
                extras['plan'] = {'subtasks': list(t['plan']['subtasks']),
                                  'deps': [list(d) for d in t['plan'].get('deps') or [[] for _ in t['plan']['subtasks']]]}
            view = None
            if mode == 'route':
                extras['dry_run'] = 'route'
                if ctx.get('cassette') is not None:
                    view = ctx['cassette'].view()
                    extras['jev'] = view
            qid = router.submit(t['query'], 'eval', session_id=session_id, engine=engine, files=files,
                                examples=ctx['examples'], extras=extras,
                                **{k: v for k, v in chat.items() if k in ('mode', 'style', 'agent')})
            ctx['qids'].add(qid)
            await router.running[qid]
            ctx['qids'].discard(qid)
            rec = router.get_run(qid)
            ctx['ms'].append(rec.get('total_ms'))
            if view is not None and view.misses:
                unrecorded = True
            if mode == 'full' and ('expect_file' in t or created_files(rec)):  # reopening a file is blocking work
                results.append(await asyncio.to_thread(turn_result, t, rec, read, spec_of, mode, case))
            else:
                results.append(turn_result(t, rec, None, None, mode, case))
            if mode == 'route' and rec.get('dry_run') != 'route' and rec.get('status') == 'done':
                results[-1]['reasons'].append('the run did not stop after routing (the pipeline has no dry-run hook)')
                results[-1]['codes'].append({'code': 'run_status'})
    finally:
        if session_id:  # the runs stay (eval detail links to them); the chat list must not fill with eval sessions
            router.store.x('DELETE FROM sessions WHERE id = ?', session_id)
    reasons = [f'turn {i}: {r}' if len(turns) > 1 else r for i, t in enumerate(results, 1) for r in t['reasons']]
    codes = [c for t in results for c in t.get('codes', [])]
    over = [t['ms'] for t in results if case.get('max_ms') and (t['ms'] or 0) > case['max_ms']]
    if over and case.get('strict_ms') and ctx['engine'] is None:  # the budgets are keyless ones (the plan's "< 2 s keyless")
        reasons.append(f"took {max(over)} ms, over the {case['max_ms']} ms budget")
        codes.append({'code': 'budget', 'stage': 'budget'})
    hard, soft = budget_misses(budget_for(case, ctx.get('budgets') or {}, engine_class(ctx['engine'])), results)
    if ctx.get('strict_budgets'):
        hard, soft = hard + soft, []
    if mode == 'route':  # nothing ran, so only what the plan did can be over budget
        soft = []
    reasons += hard
    codes += [{'code': 'budget', 'stage': 'budget'}] * len(hard)
    score_ = judge_error = None
    unjudged = False
    if case.get('judge') and mode == 'full':
        if ctx['judge'] is None:
            unjudged = True
            codes.append({'code': 'unjudged', 'stage': 'judge'})
        else:
            question = '\n'.join(f'User: {t["query"]}' for t in turns)
            answered = {e for t in results for e in t['engines']}
            fs = next((t['file'] for t in reversed(results) if t.get('file')), None)
            try:
                score_ = await judge_mod.grade(judge_mod.avoiding(ctx['judge'], answered, router.engines), question,
                                               results[-1]['answer'], case['judge'], file_summary=file_summary(fs))
            except judge_mod.JudgeError as e:
                # the rubric could not be checked, so the case can't pass on its regexes alone
                judge_error = str(e)[:200]
                reasons.append(f'not judged: {judge_error}')
                codes.append({'code': 'judge', 'stage': 'judge'})
            if score_ and score_['mean'] < case.get('judge_min', judge_mod.JUDGE_MIN):
                reasons.append(f"judge mean {score_['mean']:g} is below {case.get('judge_min', judge_mod.JUDGE_MIN):g}"
                               + (f" ({score_['note']})" if score_['note'] else ''))
                codes.append({'code': 'judge', 'stage': 'judge'})
    if unrecorded:
        codes.append({'code': 'unrecorded'})
    passed = None if unrecorded else False if reasons else None if unjudged else True
    return {'label': label, 'pass': passed, 'reasons': reasons, 'codes': codes, 'turns': results, 'judge': score_,
            'judge_error': judge_error, 'over_budget': bool(over), 'unjudged': unjudged, 'unrecorded': unrecorded,
            'budget_warnings': soft}


async def run_case(router, ctx: dict, case: dict) -> dict:
    texts = [case_query(case), *case.get('paraphrases', [])] if not case.get('turns') else [case_query(case)]
    attempts = []
    for n in range(ctx['repeat']):
        for i, text in enumerate(texts):
            label = attempt_label(n, i, ctx['repeat'], len(texts))
            try:
                attempts.append({**await run_attempt(router, ctx, case, text, label), 'phrasing': i})
            except asyncio.CancelledError:
                raise
            except Exception as e:  # a broken fixture or harness bug fails this case, not the whole eval
                attempts.append({'label': label, 'pass': False, 'reasons': [f'eval harness error: {str(e)[:160]}'],
                                 'codes': [{'code': 'run_status'}],
                                 'turns': [{'query': text, 'pass': False, 'reasons': [], 'answer': '', 'agents': [],
                                            'ms': None, 'qid': None, 'all_ok': False, 'jev_tokens': 0, 'engines': [],
                                            'steps': [], 'file': None, 'tokens': {}, 'file_tokens': 0}],
                                 'judge': None, 'over_budget': False, 'phrasing': i})
    return combine(case, attempts, ctx.get('mode', 'full'))


async def run_eval(router, eid: str, engine, engine_name, cases: list[dict], examples: bool | None = None, *,
                   split: str = 'all', tags=None, repeat: int = 1, judge=None, mode: str = 'full', jev: str = 'live',
                   cassette=None, strict_budgets: bool = False, budgets: dict | None = None) -> dict:
    examples = router.route_examples if examples is None else bool(examples)  # fixed for the whole eval
    repeat = max(1, min(MAX_REPEAT, int(repeat)))
    if mode == 'route':
        cases = [c for c in cases if routable(c)]
        judge = None
    eclass = 'route' if mode == 'route' else engine_class(engine)
    cases = [c for c in cases if runs_on(c, eclass)]
    e = {'eval_id': eid, 'at': time.time(), 'engine': engine_name, 'status': 'running', 'total': len(cases), 'cases': [],
         'examples': examples, 'split': split, 'repeat': repeat, 'judge': judge.name if judge else None,
         'tags': list(tags) if tags else None, 'mode': mode, 'jev': jev if mode == 'route' else 'live',
         'suite_sha': suite_sha(), '_ms': []}
    save(router.store, summary(e))
    sem = asyncio.Semaphore(ROUTE_CONCURRENCY if mode == 'route' else CONCURRENCY)
    ctx = {'eid': eid, 'engine': None if mode == 'route' else engine, 'examples': examples, 'repeat': repeat,
           'judge': judge, 'qids': set(), 'ms': e['_ms'], 'fixtures': Fixtures(router.store), 'mode': mode,
           'cassette': cassette, 'strict_budgets': strict_budgets,
           'budgets': read_json(BUDGETS) if budgets is None else budgets}

    async def one(case):
        async with sem:  # cases start in file order, a few at a time
            e['cases'].append(await run_case(router, ctx, case))
            s = summary(e)
            save(router.store, s)
            router.bus.emit('eval_progress', eval_id=eid, done=s['done'], total=s['total'], passed=s['passed'])

    try:
        await asyncio.gather(*(one(c) for c in cases))
        e['status'] = 'done'
    except asyncio.CancelledError:
        for qid in list(ctx['qids']):
            router.cancel(qid)
        e['status'] = 'cancelled'
    except Exception:
        e['status'] = 'error'
    finally:
        ctx['fixtures'].remove()
    order = {c['id']: i for i, c in enumerate(cases)}
    e['cases'].sort(key=lambda c: order.get(c['id'], 0))
    s = summary(e)
    save(router.store, s)
    router.bus.emit('eval_done', eval_id=eid, passed=s['passed'], total=s['total'], accuracy=s['accuracy'], status=e['status'])
    return s


def side(e: dict) -> dict:
    tokens = [c['jev_tokens'] for c in e['cases'] if isinstance(c.get('jev_tokens'), (int, float))]
    return {'eval_id': e['eval_id'], 'engine': e['engine'], 'examples': bool(e.get('examples')), 'accuracy': e['accuracy'],
            'passed': e['passed'], 'total': e['total'], 'silent_wrong': e['silent_wrong'],
            'mean_jev_tokens': round(sum(tokens) / len(tokens), 1) if tokens else 0.0,
            'suite_sha': e.get('suite_sha'), 'mode': e.get('mode', 'full')}


def compare(a: dict, b: dict) -> dict:
    """Two evals side by side (GET /api/evals/compare): each side's summary and every case in either, a's order first.
    a_pass / b_pass is None where that eval didn't score the case. `warnings` says when the suites or a case's
    definition differ, so a change in score may be a change in the test."""
    pa, pb = ({c['id']: c for c in e['cases']} for e in (a, b))
    ids = list(dict.fromkeys([*pa, *pb]))
    warnings = []
    if a.get('suite_sha') and b.get('suite_sha') and a['suite_sha'] != b['suite_sha']:
        warnings.append('the two evals ran different versions of the case suite')
    changed = [i for i in ids if i in pa and i in pb and pa[i].get('case_sha') and pb[i].get('case_sha')
               and pa[i]['case_sha'] != pb[i]['case_sha']]
    if changed:
        warnings.append(f"{len(changed)} cases changed between the two evals: {', '.join(changed[:10])}")
    if a.get('mode', 'full') != b.get('mode', 'full'):
        warnings.append('one eval only routed and the other ran every agent')
    out = {'a': side(a), 'b': side(b),
           'cases': [{'id': i, 'query': (pa.get(i) or pb[i])['query'], 'a_pass': pa[i]['pass'] if i in pa else None,
                      'b_pass': pb[i]['pass'] if i in pb else None} for i in ids]}
    if warnings:
        out['warnings'] = warnings
    return out


# ---------- promoting a real run to a case (D6) ----------

class PromoteError(Exception):
    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


def case_from_run(rec: dict, file_names: list[str], suspects: list[dict]) -> dict:
    """A draft case from a stored run: its query, chat options, fixture names and pinned plan, with expect_steps drafted
    from the agents the run used (so edit them to what should have happened) and the run's suspects in the note."""
    case = {'id': f"run-{rec['qid']}", 'query': rec['text'], 'tags': ['real-traffic', 'draft']}
    chat = {k: rec[k] for k in ('mode', 'agent', 'style') if rec.get(k) and rec[k] not in ('balanced', 'default')}
    if chat:
        case['chat'] = chat
    if file_names:
        case['files'] = file_names
    subs = (rec.get('plan') or {}).get('subtasks') or []
    tasks = rec.get('tasks') or []
    if subs:
        index = {s['tid']: i for i, s in enumerate(subs)}
        case['plan'] = {'subtasks': [s['text'] for s in subs],
                        'deps': [[index[d] for d in s.get('depends_on') or [] if d in index] for s in subs]}
    steps = []
    for i, t in enumerate(tasks):
        if not t.get('agent') or t.get('tid') == 'merge':
            continue
        s = {'agent': re.escape(t['agent']), 'forced': bool(t.get('forced'))}
        deps = [step_index(d) for d in t.get('depends_on') or []]
        if any(d is not None for d in deps):
            s['depends_on'] = [d for d in deps if d is not None and d < i]
        steps.append(s)
    if steps:
        case['expect_steps'] = steps
    notes = [f"Drafted from run {rec['qid']} ({time.strftime('%Y-%m-%d', time.localtime(rec.get('at') or 0))}); "
             'expect_steps copy what the run did, so correct them before relying on this case.']
    if suspects:
        notes.append('Suspects: ' + '; '.join(f"{s.get('code')}: {s.get('note')}" for s in suspects))
    missing = [n for n in file_names if not (FIXTURES / n).is_file()]
    if missing:
        notes.append(f"Copy {', '.join(missing)} into evals/fixtures/ before running this case.")
    case['note'] = ' '.join(notes)
    return case


def suspects_of(rec: dict) -> list[dict]:
    """The run's stored suspects, else computed now (runs saved before suspects were stored)."""
    if isinstance(rec.get('suspects'), list):
        return rec['suspects']
    try:
        from .suspects import suspects
        return suspects(rec) or []
    except Exception:
        return []


def promote_run(router, qid: int) -> dict:
    """Writes run `qid` as a draft case to evals/cases.local.jsonl: {case_id, created, path}. 404 for an unknown or
    sandbox run, 409 when a case with its id already exists."""
    from .labels import PROMOTE_LOCK
    if type(qid) is not int or qid in router.sandbox:
        raise PromoteError(f'no run {qid}', 404)
    rec = router.get_run(qid)
    if rec is None:
        raise PromoteError(f'no run {qid}', 404)
    if rec.get('status') == 'running':
        raise PromoteError(f'run {qid} has not finished', 409)
    names = [f['name'] for f in router.store.list_files(rec['files'])] if rec.get('files') else []
    case = case_from_run(rec, names, suspects_of(rec))
    path = LOCAL_CASES
    with PROMOTE_LOCK:
        existing = {c['id'] for c in read_cases(CASES)}
        if path.exists():
            for line in path.read_text().splitlines():
                try:
                    existing.add(json.loads(line).get('id'))
                except (ValueError, AttributeError):
                    continue
        if case['id'] in existing:
            raise PromoteError(f"a case named {case['id']} already exists", 409)
        path.parent.mkdir(parents=True, exist_ok=True)
        gap = path.exists() and path.stat().st_size and not path.read_bytes().endswith(b'\n')
        with open(path, 'a') as f:
            f.write(('\n' if gap else '') + json.dumps(case, ensure_ascii=False) + '\n')
    return {'case_id': case['id'], 'created': True, 'path': shown_path(path)}


# ---------- read-only views of the app's database (harvest, flaky) ----------

def open_readonly(path):
    import sqlite3
    from urllib.parse import quote
    db = sqlite3.connect(f'file:{quote(str(Path(path).resolve()))}?mode=ro', uri=True)
    db.row_factory = sqlite3.Row
    return db


def parse_since(text: str) -> float:
    """'30d', '12h', '2w' or '90m' as seconds."""
    m = re.fullmatch(r'\s*(\d+)\s*([mhdw])\s*', text or '')
    if not m:
        raise ValueError('since must look like 30d, 12h, 2w or 90m')
    return int(m.group(1)) * {'m': 60, 'h': 3600, 'd': 86400, 'w': 604800}[m.group(2)]


def harvest(db_path, since: float, limit: int = 200, now: float | None = None) -> list[dict]:
    """Runs from the last `since` seconds with suspects, newest first: [{qid, at, text, suspects}]. Read-only."""
    now = time.time() if now is None else now
    db = open_readonly(db_path)
    try:
        rows = db.execute("SELECT record FROM runs WHERE at >= ? AND status != 'running' AND source != 'eval' "
                          'ORDER BY qid DESC', (now - since,)).fetchall()
    finally:
        db.close()
    out = []
    for r in rows:
        try:
            rec = json.loads(r['record'])
        except (ValueError, TypeError):
            continue
        found = suspects_of(rec)
        if found:
            out.append({'qid': rec['qid'], 'at': rec.get('at'), 'text': rec.get('text', ''), 'suspects': found})
            if len(out) >= limit:
                break
    return out


def flaky(db_path, last: int = 10) -> list[dict]:
    """Cases whose result flips across the newest `last` stored evals of the same suite_sha, engine and mode:
    [{id, suite_sha, engine, mode, passes, fails}]. Read-only."""
    db = open_readonly(db_path)
    try:
        cols = {r['name'] for r in db.execute('PRAGMA table_info(evals)')}
        extra = ', extra' if 'extra' in cols else ''
        rows = db.execute(f"SELECT id, engine, cases{extra} FROM evals WHERE status = 'done' ORDER BY at DESC LIMIT ?",
                          (last,)).fetchall()
    finally:
        db.close()
    groups: dict[tuple, dict[str, list]] = {}
    for r in rows:
        try:
            x = json.loads(r['extra']) if 'extra' in r.keys() and r['extra'] else {}
            cases = json.loads(r['cases'] or '[]')
        except (ValueError, TypeError):
            continue
        key = (x.get('suite_sha') or '', r['engine'] or 'none', x.get('mode') or 'full')
        for c in cases:
            if isinstance(c, dict) and c.get('pass') is not None:
                groups.setdefault(key, {}).setdefault(c['id'], []).append(bool(c['pass']))
    out = []
    for (sha, engine, mode), cases in groups.items():
        for cid, results in cases.items():
            if len(set(results)) > 1:
                out.append({'id': cid, 'suite_sha': sha[:12], 'engine': engine, 'mode': mode,
                            'passes': sum(results), 'fails': len(results) - sum(results)})
    return sorted(out, key=lambda d: (-min(d['passes'], d['fails']), d['id']))


# ---------- calibration sweep (D4) ----------

THRESHOLDS = {'MIN_CLEAR': [round(0.10 + 0.05 * i, 2) for i in range(7)],
              'CONFIRM_AT': [0.5, 0.6, 0.7, 0.8],
              'UNSUPPORTED_AT': [0.5, 0.6, 0.7, 0.8, 0.9],
              'DESCRIBED_AT': [0.4, 0.5, 0.6, 0.7, 0.8]}


class patched:
    """Sets config thresholds everywhere a module imported them by name, and restores them after."""

    def __init__(self, values: dict):
        self.values, self.saved = values, []

    def __enter__(self):
        import importlib
        for name in ('config', 'jev', 'policy', 'gate', 'pipeline', 'planner'):
            try:
                mod = importlib.import_module(f'jevrouter.{name}')
            except Exception:
                continue
            for k, v in self.values.items():
                if hasattr(mod, k):
                    self.saved.append((mod, k, getattr(mod, k)))
                    setattr(mod, k, v)
        return self

    def __exit__(self, *exc):
        for mod, k, v in reversed(self.saved):
            setattr(mod, k, v)
        return False


def step_samples(result: dict, cases_by_id: dict) -> list[dict]:
    """Every routed step of a route-mode result with an expected agent, with what the policy needs to decide again."""
    out = []
    for c in result['cases']:
        if c.get('unrecorded'):
            continue
        case = cases_by_id.get(c['id'], {})
        tags = set(case.get('tags', []))
        outs = [o for t in [case, *(case.get('turns') or [])] for o in as_outcomes(t.get('expect_outcome'))]
        # an honest "can't do that here" is wanted: the case says unsupported, or it is an unsupported, live-data or
        # future case whose answer must be a refusal (@cant, @unknowable)
        refusal = any(str(case.get(k) or '').startswith(('@cant', '@unknowable')) for k in ('must_match',))
        wants_unsupported = 'unsupported' in outs or (bool(tags & {'unsupported', 'live-data', 'future'}) and refusal)
        for s in c.get('steps') or []:
            if not s.get('expected') and not wants_unsupported:
                continue
            out.append({'case': c['id'], 'split': case.get('split', 'dev'), 'honesty': 'honesty' in tags,
                        'expected': s.get('expected'), 'step': s, 'unsupported_expected': wants_unsupported,
                        'mode': (case.get('chat') or {}).get('mode') or 'balanced',
                        'attached': bool(case.get('files')), 'multi_turn': bool(case.get('turns'))})
    return out


def redecide(sample: dict) -> str:
    """The final agent for a recorded step under the thresholds in force: jev.decide then policy.apply."""
    import dataclasses
    from types import SimpleNamespace as NS
    from . import jev as jev_mod
    from . import policy
    from .config import AGENTS
    s = sample['step']
    probs = s.get('probabilities') or {s.get('pick'): s.get('confidence') or 0}
    unsafe, clear, conf = float(s.get('unsafe') or 0), float(s.get('clear') or 0), float(s.get('confidence') or 0)
    agent, reason = jev_mod.decide(NS(choice=s.get('pick'), confidence=conf, probabilities=probs), unsafe, clear)
    d = {'agent': agent, 'pick': s.get('pick'), 'reason': reason, 'probabilities': probs, 'unsafe': unsafe,
         'clear': clear, 'confidence': conf, 'signals': s.get('signals') or {}}
    fields = {f.name for f in dataclasses.fields(policy.StepCtx)}
    values = {'text': s.get('input') or '', 'ctx': '', 'tid': s.get('tid') or '',
              'offered': {a: AGENTS.get(a, a) for a in probs}, 'runnable': set(probs),
              'forced': s.get('agent') if s.get('bound') or s.get('forced') else None,
              'mode': sample['mode'], 'has_engine': False, 'web': False, 'attached': sample['attached'],
              'has_context': sample['multi_turn'], 'depends_on': s.get('depends_on') or [], 'frame': None}
    ctx = policy.StepCtx(**{k: v for k, v in values.items() if k in fields},
                         **{k: None for k in fields - set(values)})
    try:
        return policy.apply(d, ctx).agent
    except Exception:
        return agent


def sweep_metrics(samples: list[dict]) -> dict:
    """Clarify rate, wrong-route rate, unsupported precision/recall and honesty pass rate of the samples."""
    got = [(x, redecide(x)) for x in samples]
    judged = [(x, a) for x, a in got if x['expected']]
    ok = lambda x, a: a == x['expected'] or fullmatch(x['expected'], a)
    pred = [(x, a) for x, a in got if a == 'unsupported']
    real = [(x, a) for x, a in got if x['unsupported_expected']]
    honest = [(x, a) for x, a in judged if x['honesty']]
    rate = lambda xs, f: round(sum(f(*p) for p in xs) / len(xs), 4) if xs else None
    return {'n': len(got), 'clarify_rate': rate(got, lambda x, a: a == 'clarify'),
            'wrong_route': rate(judged, lambda x, a: not ok(x, a)),
            'unsupported_precision': rate(pred, lambda x, a: x['unsupported_expected']),
            'unsupported_recall': rate(real, lambda x, a: a == 'unsupported'),
            'honesty_pass': rate(honest, ok)}


def calibrate(result: dict, cases: list[dict]) -> dict:
    """The threshold sweep: each threshold over its range with the others at their config values, on dev and holdout.
    The best point per threshold is chosen on dev (lowest wrong-route rate, then clarify rate) and reported on holdout."""
    from . import config
    by_id = {c['id']: c for c in cases}
    samples = step_samples(result, by_id)
    out = {'n': len(samples), 'current': {k: getattr(config, k, None) for k in THRESHOLDS}, 'sweeps': {}}
    for name, values in THRESHOLDS.items():
        if not hasattr(config, name):
            continue
        rows = []
        for v in values:
            with patched({name: v}):
                dev = sweep_metrics([x for x in samples if x['split'] == 'dev'])
                hold = sweep_metrics([x for x in samples if x['split'] == 'holdout'])
            rows.append({'value': v, 'dev': dev, 'holdout': hold})
        key = lambda r: ((r['dev']['wrong_route'] if r['dev']['wrong_route'] is not None else 1.0),
                         (r['dev']['clarify_rate'] or 0.0), abs(r['value'] - (getattr(config, name) or 0)))
        best = min(rows, key=key)
        out['sweeps'][name] = {'rows': rows, 'best_on_dev': best['value'], 'holdout_at_best': best['holdout']}
    return out


# ---------- CLI ----------

def mark(c: dict) -> str:
    return 'PASS' if c['pass'] is True else 'FAIL' if c['pass'] is False else 'SKIP'


def print_table(result: dict):
    for c in result['cases']:
        marks = ''.join(m for m, on in ((' flaky', c.get('flaky')), (' slow', c.get('over_budget')),
                                        (' unjudged', c.get('unjudged')), (' unrecorded', c.get('unrecorded'))) if on)
        j = f" judge {c['judge']['mean']:g}" if c.get('judge') else ''
        print(f"{mark(c)}  {c['id']:<22} {c['ms'] or 0:>6}ms  {','.join(c['agents'])[:22]:<22} "
              f"{c['query'][:44]:<44}{marks}{j}  {'; '.join(c['reasons'])[:110]}")
    print('\nby split:')  # dev and holdout are reported separately (B3), so a fix tuned on dev shows up honestly
    for split in ('dev', 'holdout'):
        cs = [c for c in result['cases'] if c.get('split', 'dev') == split and c['pass'] is not None]
        if cs:
            n = sum(c['pass'] is True for c in cs)
            print(f"  {split:<16} {n:>3}/{len(cs):<3} {n / len(cs):>5.0%}")
    print('by tag:')
    for tag, t in result['by_tag'].items():
        print(f"  {tag:<16} {t['passed']:>3}/{t['total']:<3} {t['passed'] / t['total']:>5.0%}")


def print_route_report(r: dict):
    conf = r.get('confusion') or {}
    if conf:
        cols = sorted({g for row in conf.values() for g in row})
        w = max(len(x) for x in [*conf, *cols, 'expected']) + 1
        print('\nconfusion (rows expected, columns got):')
        print('  ' + 'expected'.ljust(w) + ''.join(c[:9].rjust(10) for c in cols))
        for e, row in conf.items():
            print('  ' + e.ljust(w) + ''.join(str(row.get(c, '')).rjust(10) for c in cols))
    if r.get('per_agent'):
        print('\nper agent:            P       R      F1  support')
        for a in sorted(r['per_agent'], key=lambda x: (x['f1'] is None, x['f1'] or 0)):
            f = lambda v: '   -  ' if v is None else f'{v:6.2f}'
            print(f"  {a['agent']:<16} {f(a['precision'])}  {f(a['recall'])}  {f(a['f1'])}  {a['support']:>5}")
    if r.get('stages'):
        print('\nfailure stages: ' + ', '.join(f'{k} {v}' for k, v in r['stages'].items()))
    for cal in r.get('calibration') or []:
        ece = '-' if cal['ece'] is None else f"{cal['ece']:.3f}"
        cells = [f"[{b['lo']:.1f},{b['hi']:.1f}) n{b['n']} acc {b['accuracy']:.2f}" for b in cal['bins'] if b['n']]
        print(f"calibration {cal['signal']}: ECE {ece}; " + '; '.join(cells))
    robust = {t: s for t, s in (r.get('by_tag_route') or r.get('by_tag') or {}).items() if t.startswith('transform:')}
    if robust:
        print('robustness by transform: ' + ', '.join(f"{t[10:]} {s['passed']}/{s['total']}" for t, s in robust.items()))


def print_gates(gates: list[dict]):
    if not gates:
        return
    print('\ngates:  tag               min   pass rate  lower  baseline  ok')
    for g in gates:
        base = '-' if g['baseline'] is None else f"{g['baseline']:.2f}"
        print(f"  {g['tag']:<20} {g['min']:.2f}  {g['passed']:>3}/{g['total']:<3} {g['rate']:.2f}  {g['lower']:.2f}  "
              f"{base:>8}  {'yes' if g['ok'] else 'NO'}")


def headline(name: str, r: dict) -> str:
    j = f", judge mean {r['judge_mean']:g} ({r['judge']})" if r.get('judge_mean') is not None else ''
    extra = ''
    if r.get('unjudged'):
        extra += f", unjudged {r['unjudged']}"
    if r.get('unrecorded'):
        extra += f", unrecorded {r['unrecorded']}"
    rp, ap = r.get('route_pass') or {}, r.get('answer_pass') or {}
    if rp.get('total'):
        extra += f", route {rp['passed']}/{rp['total']}"
    if ap.get('total'):
        extra += f", answer {ap['passed']}/{ap['total']}"
    mode = f", mode {r['mode']} (Jev {r.get('jev')})" if r.get('mode') == 'route' else ''
    return (f"{name}: {r['passed']}/{r['total']} passed ({r['accuracy']:.0%}), silent_wrong {r['silent_wrong']}, "
            f"flaky {r['flaky']}, p50 {r['p50_ms']}ms, p95 {r['p95_ms']}ms, split {r['split']}, repeat {r['repeat']}, "
            f"examples {'on' if r['examples'] else 'off'}, mean Jev tokens {side(r)['mean_jev_tokens']:g}{extra}{mode}{j}")


def matrix_table(results: list[dict]) -> str:
    """Engines as columns: the headline numbers, then accuracy per tag."""
    names = [r['engine'] for r in results]
    rows = [('passed', [f"{r['passed']}/{r['total']}" for r in results]),
            ('accuracy', [f"{r['accuracy']:.0%}" for r in results]),
            ('silent_wrong', [str(r['silent_wrong']) for r in results]),
            ('flaky', [str(r['flaky']) for r in results]),
            ('p50 ms', [str(r['p50_ms']) for r in results]),
            ('p95 ms', [str(r['p95_ms']) for r in results]),
            ('judge mean', ['-' if r.get('judge_mean') is None else f"{r['judge_mean']:g}" for r in results])]
    for tag in sorted({t for r in results for t in r['by_tag']}):
        rows.append((f'tag {tag}', [f"{r['by_tag'][tag]['passed']}/{r['by_tag'][tag]['total']}" if tag in r['by_tag']
                                    else '-' for r in results]))
    w = max(len(k) for k, _ in rows) + 2
    cw = [max(len(n), *(len(v[i]) for _, v in rows)) + 2 for i, n in enumerate(names)]
    lines = [''.ljust(w) + ''.join(n.rjust(cw[i]) for i, n in enumerate(names))]
    lines += [k.ljust(w) + ''.join(v[i].rjust(cw[i]) for i in range(len(names))) for k, v in rows]
    return '\n'.join(lines)


def shown_path(path) -> str:
    """A path relative to the repo when it is inside it."""
    try:
        return str(Path(path).relative_to(ROOT))
    except ValueError:
        return str(path)


def cassette_problem(jev: str, sha: str, route: str | None = None) -> str | None:
    """Replay needs a cassette recorded for the cases route mode runs. The meta's route_sha is compared with `route`
    when both are known; a meta without one (recorded before route_sha existed) falls back to the whole suite_sha."""
    if jev != 'replay':
        return None
    if not CASSETTE.exists():
        return f'no Jev cassette at {shown_path(CASSETTE)}; record one with --mode route --jev record'
    meta = read_json(CASSETTE_META)
    if meta.get('route_sha') and route:
        if meta['route_sha'] != route:
            return ('the routed cases changed since the cassette was recorded (route_sha differs); refresh it with '
                    '--mode route --jev record')
        return None
    if meta.get('suite_sha') and meta['suite_sha'] != sha:
        return ('the cases changed since the cassette was recorded (suite_sha differs); refresh it with '
                '--mode route --jev record')
    return None


async def cli(engine_name: str | None, save: bool, examples: bool | None = None, *, split: str = 'all',
              tags=None, repeat: int = 1, judge: str | None = None, matrix: bool = False, as_json: bool = False,
              mode: str = 'full', jev: str = 'live', gates_path=None, strict_budgets: bool = False,
              cases_path=None) -> int:
    import aiohttp

    from .config import load_env
    from .engines import catalog, choose
    from .pipeline import Router
    from .store import DEFAULT_DB, Store, read_labels

    load_env()
    engines = catalog()
    if mode == 'route':
        if matrix or (engine_name not in (None, 'none')):
            print('route mode runs no engine; leave out --engine and --matrix', file=sys.stderr)
            return 2
        names = ['none']
    elif matrix:
        names = ['none', *(n for n, e in engines.items() if n != 'auto' and e.available()[0])]
        print(f"--matrix runs {len(names)} engines ({', '.join(names)}). Every engine other than none spends your "
              'subscription or API quota: about one to three calls per case.', file=sys.stderr)
    else:
        engine = None if engine_name == 'none' else choose(engines, engine_name) if engine_name else choose(engines)
        if engine_name not in (None, 'none') and (engine is None or engine.name != engine_name):
            print(f'engine {engine_name!r} is not available', file=sys.stderr)
            return 2
        names = [engine.name if engine else 'none']
    try:
        every = load_cases(cases_path) if cases_path else load_cases()
        cases = select(every, split, tags)
    except CaseError as e:
        print(f'bad eval case: {e}', file=sys.stderr)
        return 2
    if mode == 'route':
        cases = [c for c in cases if routable(c) and runs_on(c, 'route')]
    if not cases:
        print('no cases match that split and those tags', file=sys.stderr)
        return 2
    sha = suite_sha()
    rsha = route_sha()
    if problem := cassette_problem(jev if mode == 'route' else 'live', sha, rsha):
        print(problem, file=sys.stderr)
        return 2
    baseline = read_json(BASELINE)
    gates = read_json(gates_path) if gates_path else {}
    if gates_path and not gates:
        print(f'no gates in {gates_path}', file=sys.stderr)
        return 2
    results, code = [], 0
    async with aiohttp.ClientSession() as http:
        for name in names:
            engine = None if name == 'none' else engines[name]
            try:
                judge_engine, note = judge_mod.pick(engines, judge, engine) if mode == 'full' else (None, '')
            except judge_mod.JudgeError as e:
                print(str(e), file=sys.stderr)
                return 2
            if judge and note:
                print(note, file=sys.stderr)
            cassette = None
            if mode == 'route':
                from .cassette import JevCassette
                inner = None
                if jev != 'replay':
                    from typesafe_sdk import AsyncTypeSafeClient
                    inner = AsyncTypeSafeClient()
                cassette = JevCassette(inner, None if jev == 'live' else CASSETTE, jev)
                # nothing runs in route mode: no engine, no HTTP (the lone-term lookup needs a session)
                router = Router(cassette, None, None, engines=engines, store=Store())
            else:
                from typesafe_sdk import AsyncTypeSafeClient
                router = Router(AsyncTypeSafeClient(), http, engine, engines=engines, store=Store())  # in-memory: CLI evals don't touch the app DB
            ex = router.route_examples if examples is None else examples
            if ex:  # route examples come from the app's labels, read without opening its DB for writing
                router.labels = read_labels(os.environ.get('TG_DB') or DEFAULT_DB)
            result = await run_eval(router, 'cli', engine, name, cases, ex, split=split, tags=tags, repeat=repeat,
                                    judge=judge_engine, mode=mode, jev=jev, cassette=cassette,
                                    strict_budgets=strict_budgets)
            suite = suite_name(mode, name)
            if mode == 'route' and jev == 'record' and cassette is not None:
                CASSETTE_META.parent.mkdir(parents=True, exist_ok=True)
                CASSETTE_META.write_text(json.dumps({'suite_sha': sha, 'route_sha': rsha, 'at': round(time.time()),
                                                     'entries': len(cassette)},
                                                    indent=1) + '\n')
            rules = gates.get(suite) if isinstance(gates.get(suite), dict) else None
            if rules:
                result['gates'] = gate_tags(result, rules, baseline_for(baseline, suite))
            results.append(result)
            if not as_json:
                print_table(result)
                if mode == 'route' or result.get('confusion'):
                    print_route_report(result)
                print_gates(result.get('gates') or [])
                print('\n' + headline(name if mode == 'full' else 'route', result))
                rubric = sum(1 for c in cases if c.get('judge'))
                if rubric and judge_engine is None and mode == 'full':
                    print(f'{rubric} rubric cases were not judged: {note}; pass --judge NAME or --judge auto to score them')
                for err in result.get('judge_errors', []):
                    print(f'judge error: {err}')
                for w in result.get('budget_warnings', [])[:20]:
                    print(f'over budget: {w}')
            if save:
                baseline = save_baseline(baseline, suite, result, sha)
                BASELINE.write_text(json.dumps(baseline, indent=1, sort_keys=True) + '\n')
                if not as_json:
                    print(f'baseline for {suite} saved to {BASELINE}')
                continue
            failed = []
            if rules:
                failed = [g['tag'] for g in result['gates'] if not g['ok']]
                if mode == 'route' and result['total'] + result['unrecorded'] and \
                        result['unrecorded'] / (result['total'] + result['unrecorded']) > MAX_UNRECORDED:
                    failed.append(f"unrecorded {result['unrecorded']}")
                failed += strict_unrecorded(result, rules)
            else:
                failed = regressions(result, baseline)
                base = baseline_for(baseline, suite)
                if base and base.get('suite_sha') == sha:
                    failed += tag_drops(result, base)
            if failed:
                code = 1
                if not as_json:
                    print(f"{'gates failed' if rules else 'regressions against baseline'}: {', '.join(failed)}")
    if as_json:
        print(json.dumps(results if matrix else results[0], indent=1))
    elif matrix:
        print('\n' + matrix_table(results))
    return code


def parse_args(argv=None) -> argparse.Namespace:
    ap = argparse.ArgumentParser(prog='python -m jevrouter.evals', description='Run the TraceGraph eval cases through the real pipeline.')
    ap.add_argument('--engine', help='engine name, or none for keyless (default: TG_ENGINE / auto)')
    ap.add_argument('--examples', choices=('on', 'off'), help='route examples from your labels on or off (default: TG_ROUTE_EXAMPLES)')
    ap.add_argument('--split', choices=SPLITS, default='all', help='dev, holdout or all cases (default: all)')
    ap.add_argument('--tags', type=lambda s: [t.strip() for t in s.split(',') if t.strip()],
                    help='only cases with any of these comma-separated tags (all: every case)')
    ap.add_argument('--repeat', type=int, default=1, choices=range(1, MAX_REPEAT + 1), metavar='1-5',
                    help='attempts per case; a case is flaky when they disagree (default 1)')
    ap.add_argument('--judge', help='engine that scores rubric cases (a name, or auto for any engine other than the one under test)')
    ap.add_argument('--matrix', action='store_true', help='run keyless and then every available engine, and print a table (spends quota)')
    ap.add_argument('--json', action='store_true', dest='as_json', help='print the result as JSON instead of a table')
    ap.add_argument('--save-baseline', action='store_true', help='record this run as the baseline for its suite')
    ap.add_argument('--mode', choices=EVAL_MODES, default='full', help='full runs every agent; route stops after routing')
    ap.add_argument('--jev', choices=JEV_SOURCES, default='live',
                    help='route mode: ask Jev live, replay the cassette, or record what the cassette lacks')
    ap.add_argument('--gates', dest='gates_path', help='fail when a tag misses its gate in this file (evals/gates.json)')
    ap.add_argument('--strict-budgets', action='store_true', help='a case over its time or token budget fails')
    ap.add_argument('--cases', dest='cases_path', help='a cases file to run instead of the suite (e.g. evals/paraphrases.gen.jsonl)')
    args = ap.parse_args(argv)
    if args.matrix and args.engine:
        ap.error('--matrix runs every engine; leave out --engine')
    if args.jev != 'live' and args.mode != 'route':
        ap.error('--jev replay and record are for --mode route')
    return args


SUBCOMMANDS = ('calibrate', 'harvest', 'flaky', 'paraphrase', 'judge-calibrate')


def sub_args(argv: list[str]) -> argparse.Namespace:
    ap = argparse.ArgumentParser(prog='python -m jevrouter.evals')
    sp = ap.add_subparsers(dest='command', required=True)
    c = sp.add_parser('calibrate', help='sweep the routing thresholds over recorded Jev answers')
    c.add_argument('--jev', choices=('replay', 'record', 'live'), default='replay')
    c.add_argument('--tags', type=lambda s: [t.strip() for t in s.split(',') if t.strip()])
    c.add_argument('--json', action='store_true', dest='as_json')
    h = sp.add_parser('harvest', help='list real runs with suspects, newest first (read-only)')
    h.add_argument('--since', default='30d')
    h.add_argument('--db')
    h.add_argument('--limit', type=int, default=200)
    h.add_argument('--json', action='store_true', dest='as_json')
    f = sp.add_parser('flaky', help='cases whose result flips across stored evals of the same suite (read-only)')
    f.add_argument('--last', type=int, default=10)
    f.add_argument('--db')
    f.add_argument('--json', action='store_true', dest='as_json')
    p = sp.add_parser('paraphrase', help='print deterministic paraphrases of the cases as JSON lines')
    p.add_argument('--seed', type=int, default=7)
    p.add_argument('--n', type=int, default=3)
    p.add_argument('--tags', type=lambda s: [t.strip() for t in s.split(',') if t.strip()])
    p.add_argument('--split', choices=SPLITS, default='all')
    p.add_argument('--transforms', type=lambda s: [t.strip() for t in s.split(',') if t.strip()])
    j = sp.add_parser('judge-calibrate', help="score the judge against the gold set (spends the judge engine's quota)")
    j.add_argument('--judge', required=True)
    j.add_argument('--gold', default=str(JUDGE_GOLD))
    j.add_argument('--json', action='store_true', dest='as_json')
    return ap.parse_args(argv)


def db_path(arg) -> str:
    from .store import DEFAULT_DB
    return arg or os.environ.get('TG_DB') or str(DEFAULT_DB)


def cmd_harvest(a) -> int:
    try:
        since = parse_since(a.since)
    except ValueError as e:
        print(str(e), file=sys.stderr)
        return 2
    rows = harvest(db_path(a.db), since, a.limit)
    if a.as_json:
        print(json.dumps(rows, indent=1))
        return 0
    for r in rows:
        when = time.strftime('%Y-%m-%d %H:%M', time.localtime(r['at'] or 0))
        print(f"{r['qid']:>6}  {when}  {','.join(s['code'] for s in r['suspects']):<34} {r['text'][:70]}")
    print(f'{len(rows)} runs with suspects in the last {a.since}')
    return 0


def cmd_flaky(a) -> int:
    rows = flaky(db_path(a.db), a.last)
    if a.as_json:
        print(json.dumps(rows, indent=1))
        return 0
    for r in rows:
        print(f"{r['id']:<30} {r['engine']:<12} {r['mode']:<6} suite {r['suite_sha'] or '-':<12} "
              f"passed {r['passes']}, failed {r['fails']}")
    print(f'{len(rows)} flaky cases over the last {a.last} evals')
    return 0


def cmd_paraphrase(a) -> int:
    from .paraphrase import TRANSFORMS, generate
    try:
        cases = select(load_cases(CASES), a.split, a.tags)
        variants = generate(cases, a.seed, a.n, a.transforms or TRANSFORMS)
    except (CaseError, ValueError) as e:
        print(str(e), file=sys.stderr)
        return 2
    for v in variants:
        sys.stdout.write(json.dumps(v, ensure_ascii=False) + '\n')
    return 0


def read_gold(path) -> list[dict]:
    rows = []
    for n, line in enumerate(Path(path).read_text().splitlines(), 1):
        if not line.strip():
            continue
        d = json.loads(line)
        if not (isinstance(d.get('question'), str) and isinstance(d.get('answer'), str) and isinstance(d.get('rubric'), str)
                and isinstance(d.get('scores'), dict) and all(c in d['scores'] for c in judge_mod.CRITERIA)):
            raise CaseError(f'{Path(path).name} line {n}: needs question, answer, rubric and scores for '
                            f'{", ".join(judge_mod.CRITERIA)}')
        rows.append(d)
    return rows


async def judge_agreement(engine, rows: list[dict]) -> dict:
    """The judge against human scores: mean absolute error per criterion and pass/fail agreement at JUDGE_MIN."""
    errors = {c: [] for c in judge_mod.CRITERIA}
    agree, failed, results = 0, 0, []
    for r in rows:
        try:
            s = await judge_mod.grade(engine, r['question'], r['answer'], r['rubric'], file_summary=r.get('file'))
        except judge_mod.JudgeError as e:
            failed += 1
            results.append({'question': r['question'][:60], 'error': str(e)})
            continue
        gold_mean = sum(r['scores'][c] for c in judge_mod.CRITERIA) / len(judge_mod.CRITERIA)
        for c in judge_mod.CRITERIA:
            errors[c].append(abs(s[c] - r['scores'][c]))
        same = (s['mean'] >= judge_mod.JUDGE_MIN) == (gold_mean >= judge_mod.JUDGE_MIN)
        agree += same
        results.append({'question': r['question'][:60], 'judge': s['mean'], 'gold': round(gold_mean, 2), 'agree': same})
    scored = len(rows) - failed
    agreement = round(agree / scored, 4) if scored else None
    drafts = sum(1 for r in rows if r.get('scored_by') == 'draft')
    # scores a person hasn't reviewed can't vouch for the judge, however well it agrees with them
    return {'rows': len(rows), 'failed': failed, 'drafts': drafts,
            'mae': {c: round(sum(v) / len(v), 3) if v else None for c, v in errors.items()},
            'agreement': agreement, 'trusted': agreement is not None and agreement >= 0.85 and not drafts,
            'results': results}


async def cmd_judge_calibrate(a) -> int:
    from .config import load_env
    from .engines import catalog
    load_env()
    try:
        rows = read_gold(a.gold)
        engine, note = judge_mod.pick(catalog(), a.judge, None)
    except (OSError, ValueError, CaseError, judge_mod.JudgeError) as e:
        print(str(e), file=sys.stderr)
        return 2
    if engine is None:
        print(note or 'no judge engine', file=sys.stderr)
        return 2
    out = await judge_agreement(engine, rows)
    if a.as_json:
        print(json.dumps(out, indent=1))
        return 0
    print(f"judge {engine.name} on {out['rows']} gold rows ({out['failed']} not scored)")
    print('mean absolute error: ' + ', '.join(f'{c} {v}' for c, v in out['mae'].items()))
    print(f"pass/fail agreement at {judge_mod.JUDGE_MIN}: {out['agreement']}; "
          f"{'trusted for gating' if out['trusted'] else 'not trusted for gating (needs 0.85)'}")
    if out['drafts']:
        print(f"{out['drafts']} gold rows are unreviewed drafts (scored_by: draft); have a person check their scores and "
              'set scored_by to their name before the judge can be trusted')
    return 0


async def cmd_calibrate(a) -> int:
    """Runs the route suite once (replay by default) and sweeps the thresholds over its recorded steps."""
    import io
    import contextlib
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        code = await cli(None, False, mode='route', jev=a.jev, tags=a.tags, as_json=True)
    if code == 2:
        return 2
    try:
        result = json.loads(buf.getvalue())
    except ValueError:
        print('the route run printed no result', file=sys.stderr)
        return 2
    t0 = time.perf_counter()
    out = calibrate(result, load_cases())
    out['sweep_ms'] = round((time.perf_counter() - t0) * 1000)
    if a.as_json:
        print(json.dumps(out, indent=1))
        return 0
    print(f"{out['n']} recorded steps; sweep took {out['sweep_ms']} ms; current: " +
          ', '.join(f'{k} {v}' for k, v in out['current'].items()))
    for name, s in out['sweeps'].items():
        print(f'\n{name}:  value  clarify  wrong  unsup P  unsup R  honesty   (dev | holdout)')
        for r in s['rows']:
            f = lambda m: ' '.join('   -  ' if m[k] is None else f'{m[k]:6.2f}' for k in
                                   ('clarify_rate', 'wrong_route', 'unsupported_precision', 'unsupported_recall', 'honesty_pass'))
            print(f"  {r['value']:>12}  {f(r['dev'])} | {f(r['holdout'])}")
        print(f"  best on dev: {s['best_on_dev']} (change config only when holdout confirms it)")
    return 0


def main(argv=None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)
    if argv and argv[0] in SUBCOMMANDS:
        a = sub_args(argv)
        if a.command == 'harvest':
            return cmd_harvest(a)
        if a.command == 'flaky':
            return cmd_flaky(a)
        if a.command == 'paraphrase':
            return cmd_paraphrase(a)
        if a.command == 'judge-calibrate':
            return asyncio.run(cmd_judge_calibrate(a))
        return asyncio.run(cmd_calibrate(a))
    a = parse_args(argv)
    return asyncio.run(cli(a.engine, a.save_baseline, None if a.examples is None else a.examples == 'on', split=a.split,
                           tags=a.tags, repeat=a.repeat, judge=a.judge, matrix=a.matrix, as_json=a.as_json,
                           mode=a.mode, jev=a.jev, gates_path=a.gates_path, strict_budgets=a.strict_budgets,
                           cases_path=a.cases_path))


if __name__ == '__main__':
    sys.exit(main())
