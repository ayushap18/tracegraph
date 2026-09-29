"""The PPTX painter (docs/PLAN-designer.md 2 and 9.6): draws a positioned DesignPlan with python-pptx and makes no
layout decisions. Text boxes use the plan's sizes and wrapped lines (auto-fit off), shapes and icons are native vector
shapes, diagrams are native shapes (create/diagram.py), charts are native charts, images are the workspace crops. Page
notes become speaker notes. Rules F3/A1-A4 and X1-X3 hold exactly as in create/render._pptx.

Owner: builder L.
"""
from __future__ import annotations

import importlib
import io

from .plan import DesignPlan
from .workspace import Workspace

EMU_PER_PT = 12700
AUTHOR = 'TraceGraph'
BODY_SLOTS = ('body', 'left', 'right', 'context', 'takeaway', 'points', 'col1', 'col2', 'body2')


def _emu(pt: float) -> int:
    return int(round(float(pt) * EMU_PER_PT))


def _render():
    return importlib.import_module('jevrouter.create.render')


def title_box(page, spec):
    """The box that becomes the slide's real title (A1): the file title on the cover, the section heading elsewhere,
    the closing title on the closing slide."""
    from .layout import parse_ref
    for bx in page.boxes:
        if bx.kind != 'text':
            continue
        r = parse_ref(bx.content)
        if r is not None and ((r.top == 'title' and r.field != 'footer') or r.part == 'heading'):
            return bx
    if page.layout == 'closing':
        return next((bx for bx in page.boxes if bx.kind == 'text' and bx.slot == 'title'), None)
    return None


def image_bytes(bx, ws, *, mono: bool = False) -> bytes:
    """The picture a box shows: a workspace crop, or an asset-cache image cropped here (fractions of the source)."""
    ref = bx.content
    if ref.startswith('ws:'):
        data = ws.get(ref)
        if not data:
            raise FileNotFoundError(f'{ref} is not in the workspace')
        return data
    if not ref.startswith('asset:'):
        raise ValueError(f'not an image ref: {ref[:40]}')
    from PIL import Image
    data = _render().asset_path(ref[6:]).read_bytes()
    crop = bx.crop
    if not mono and (crop is None or tuple(crop) == (0.0, 0.0, 1.0, 1.0)):
        return data
    with Image.open(io.BytesIO(data)) as im:
        im = im.convert('L' if mono else 'RGB')
        if crop is not None and tuple(crop) != (0.0, 0.0, 1.0, 1.0):
            W, H = im.size
            x, y, w, h = crop
            box = (round(x * W), round(y * H), max(round(x * W) + 1, round((x + w) * W)),
                   max(round(y * H) + 1, round((y + h) * H)))
            im = im.crop(box)
        buf = io.BytesIO()
        im.save(buf, 'PNG')
        return buf.getvalue()


class _Painter:
    def __init__(self, plan: DesignPlan, spec: dict, ws: Workspace):
        from pptx.dml.color import RGBColor
        from .tokens import DesignSystem
        self.plan, self.spec, self.ws = plan, spec, ws
        self.ds = DesignSystem.from_dict(plan.system)
        self.RGB = RGBColor

    def rgb(self, token: str | None, default: str = 'text'):
        return self.RGB.from_string(self.ds.color(token or default).upper())

    # ----- text -----
    def fill_frame(self, tf, bx, *, title: bool = False):
        from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
        from pptx.util import Pt
        from ..create.spec import runs
        from .layout import box_text, bullet_indent, bullet_of
        tf.clear()
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {'top': MSO_ANCHOR.TOP, 'middle': MSO_ANCHOR.MIDDLE,
                              'bottom': MSO_ANCHOR.BOTTOM}.get(bx.valign, MSO_ANCHOR.TOP)
        text = box_text(bx, self.spec)
        marker = bullet_of(bx, self.spec)
        size = float(bx.size or self.ds.size(bx.step or 'body', 'pptx'))
        color = self.rgb(bx.style.text_color)
        align = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT,
                 'justify': PP_ALIGN.JUSTIFY}.get(bx.align, PP_ALIGN.LEFT)
        tracking = self.ds.scale.display_tracking if bx.step in ('display', 'h1') else 0.0
        paras = text.split('\n') if text else ['']
        for k, line in enumerate(paras):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = Pt(size * bx.line_height)
            p.space_before = p.space_after = Pt(0)
            if marker and k == 0:
                self.bullet(p, marker, bullet_indent(size, marker != '•'))
            elif marker or self.continues_bullet(bx):
                ppr = p._p.get_or_add_pPr()
                ppr.set('marL', str(_emu(bullet_indent(size, bool(marker and marker != '•')))))
            for piece, b, it in (runs(line) or [(line, False, False)]):
                r = p.add_run()
                r.text = piece
                f = r.font
                f.size = Pt(size)
                f.name = bx.font or self.ds.families.body
                f.bold = True if (bx.bold or b) else (False if title else None)
                f.italic = True if (bx.italic or it) else None
                f.color.rgb = color
                if tracking:
                    r._r.get_or_add_rPr().set('spc', str(int(round(tracking * size * 100))))

    def continues_bullet(self, bx) -> bool:
        from .layout import block_of, parse_ref
        r = parse_ref(bx.content)
        blk = block_of(bx.content, self.spec)
        return bool(r and r.items and r.words and blk and blk.get('type') == 'bullets')

    @staticmethod
    def bullet(p, marker: str, indent_pt: float):
        from lxml import etree
        from pptx.oxml.ns import qn
        ppr = p._p.get_or_add_pPr()
        ppr.set('marL', str(_emu(indent_pt)))
        ppr.set('indent', str(-_emu(indent_pt)))
        for tag in ('a:buNone', 'a:buChar', 'a:buAutoNum'):
            for el in ppr.findall(qn(tag)):
                ppr.remove(el)
        if marker == '•':
            el = etree.SubElement(ppr, qn('a:buChar'))
            el.set('char', marker)
        else:
            el = etree.SubElement(ppr, qn('a:buAutoNum'))
            el.set('type', 'arabicPeriod')
            n = int(marker.rstrip('.') or 1)
            if n > 1:
                el.set('startAt', str(n))

    def text(self, shapes, bx):
        from pptx.util import Emu
        shape = shapes.add_textbox(Emu(_emu(bx.x)), Emu(_emu(bx.y)), Emu(_emu(bx.w)), Emu(_emu(bx.h)))
        size = float(bx.size or 0)
        if bx.slot in BODY_SLOTS and size >= 18:
            shape.name = 'Body'
        elif bx.content.startswith('spec:') and '#' in bx.content:
            shape.name = bx.content.rsplit('#', 1)[1].capitalize()
        else:
            shape.name = (bx.slot or 'Text').replace('_', ' ').capitalize()
        if bx.style.fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = self.rgb(bx.style.fill)
            self.alpha(shape.fill, bx.style.opacity)
        self.fill_frame(shape.text_frame, bx)
        return shape

    # ----- shapes -----
    @staticmethod
    def alpha(fill, opacity: float):
        """A solid fill's transparency (opacity 0..1)."""
        if opacity is None or opacity >= 0.999:
            return
        from lxml import etree
        from pptx.oxml.ns import qn
        try:
            clr = fill._xPr.find(qn('a:solidFill'))
            srgb = clr.find(qn('a:srgbClr')) if clr is not None else None
            if srgb is not None:
                for el in srgb.findall(qn('a:alpha')):
                    srgb.remove(el)
                a = etree.SubElement(srgb, qn('a:alpha'))
                a.set('val', str(int(max(0.0, min(1.0, opacity)) * 100000)))
        except Exception:
            pass

    def line_alpha(self, line, opacity: float):
        if opacity is None or opacity >= 0.999:
            return
        from lxml import etree
        from pptx.oxml.ns import qn
        try:
            ln = line._get_or_add_ln()
            srgb = ln.find(qn('a:solidFill')).find(qn('a:srgbClr'))
            a = etree.SubElement(srgb, qn('a:alpha'))
            a.set('val', str(int(max(0.0, min(1.0, opacity)) * 100000)))
        except Exception:
            pass

    def rect(self, shapes, x, y, w, h, *, fill=None, stroke=None, stroke_w=0.0, opacity=1.0, radius=0.0,
             ellipse=False, name='Shape'):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Emu, Pt
        kind = MSO_SHAPE.OVAL if ellipse else MSO_SHAPE.ROUNDED_RECTANGLE if radius > 0 else MSO_SHAPE.RECTANGLE
        shp = shapes.add_shape(kind, Emu(_emu(x)), Emu(_emu(y)), Emu(max(1, _emu(w))), Emu(max(1, _emu(h))))
        shp.name = name
        if radius > 0 and not ellipse:
            shp.adjustments[0] = min(0.5, radius / max(min(w, h), 1))
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = self.rgb(fill)
            self.alpha(shp.fill, opacity)
        else:
            shp.fill.background()
        if stroke and stroke_w:
            shp.line.color.rgb = self.rgb(stroke)
            shp.line.width = Pt(stroke_w)
            self.line_alpha(shp.line, opacity)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        if shp.has_text_frame:
            shp.text_frame.text = ''
        return shp

    def poly(self, shapes, points, *, stroke=None, stroke_w=0.0, fill=None, closed=False, opacity=1.0, name='Line'):
        from pptx.util import Pt
        pts = [(float(a), float(b)) for a, b in points]
        if len(pts) < 2:
            return None
        builder = shapes.build_freeform(_emu(pts[0][0]), _emu(pts[0][1]), scale=1.0)
        builder.add_line_segments([(_emu(a), _emu(b)) for a, b in pts[1:]], close=bool(closed))
        shp = builder.convert_to_shape()
        shp.name = name
        if fill and closed:
            shp.fill.solid()
            shp.fill.fore_color.rgb = self.rgb(fill)
            self.alpha(shp.fill, opacity)
        else:
            shp.fill.background()
        if stroke and stroke_w:
            shp.line.color.rgb = self.rgb(stroke)
            shp.line.width = Pt(stroke_w)
            self.line_alpha(shp.line, opacity)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def gradient(self, shapes, prim, ox, oy, opacity):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Emu
        x, y, w, h = (float(prim.get(k, 0)) for k in ('x', 'y', 'w', 'h'))
        shp = shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(_emu(ox + x)), Emu(_emu(oy + y)), Emu(max(1, _emu(w))),
                               Emu(max(1, _emu(h))))
        shp.name = 'Gradient'
        shp.fill.gradient()
        shp.fill.gradient_angle = float(prim.get('angle', 0)) % 360
        stops = shp.fill.gradient_stops
        stops[0].color.rgb = self.rgb(prim.get('from') or 'accent')
        stops[-1].color.rgb = self.rgb(prim.get('to') or 'accent2')
        op = float(prim.get('opacity', 1.0)) * opacity
        if op < 0.999:
            from lxml import etree
            from pptx.oxml.ns import qn
            for gs in shp.fill._xPr.iter(qn('a:srgbClr')):
                a = etree.SubElement(gs, qn('a:alpha'))
                a.set('val', str(int(op * 100000)))
        shp.line.fill.background()
        shp.shadow.inherit = False

    def art(self, shapes, bx):
        _, kind, seed = (bx.content.split(':') + ['', ''])[:3]
        try:
            from . import art
            prims = art.draw(kind, int(seed or 0), bx.w, bx.h, self.ds)
        except Exception:
            prims = [{'type': 'ellipse', 'x': bx.w * 0.35, 'y': bx.h * 0.15, 'w': bx.w * 0.6, 'h': bx.w * 0.6,
                      'fill': 'accent', 'opacity': 0.25},
                     {'type': 'ellipse', 'x': bx.w * 0.1, 'y': bx.h * 0.6, 'w': bx.w * 0.3, 'h': bx.w * 0.3,
                      'fill': 'accent2', 'opacity': 0.25}]
        group = shapes.add_group_shape()
        group.name = f'Art: {kind}'
        self.decorative(group)
        op = float(bx.style.opacity if bx.style.opacity is not None else 1.0)
        for pr in prims:
            t = pr.get('type')
            o = float(pr.get('opacity', 1.0)) * op
            if t in ('rect', 'ellipse'):
                self.rect(group.shapes, bx.x + float(pr.get('x', 0)), bx.y + float(pr.get('y', 0)),
                          float(pr.get('w', 0)), float(pr.get('h', 0)), fill=pr.get('fill'), stroke=pr.get('stroke'),
                          stroke_w=float(pr.get('stroke_w') or 0), opacity=o, radius=float(pr.get('radius') or 0),
                          ellipse=t == 'ellipse', name='Art')
            elif t in ('line', 'path'):
                pts = [(bx.x + float(a), bx.y + float(b)) for a, b in pr.get('points') or []]
                self.poly(group.shapes, pts, stroke=pr.get('stroke'), stroke_w=float(pr.get('stroke_w') or 0),
                          fill=pr.get('fill'), closed=bool(pr.get('closed')), opacity=o, name='Art')
            elif t == 'gradient':
                self.gradient(group.shapes, pr, bx.x, bx.y, op)
        if not len(group.shapes):
            group._element.getparent().remove(group._element)

    @staticmethod
    def decorative(shape, alt: str | None = None):
        """Alt text on a shape (descr); decorative shapes get an empty one."""
        try:
            cnv = shape._element.xpath('./*[1]/p:cNvPr')[0]
            cnv.set('descr', (alt or '')[:1000])
        except Exception:
            pass

    def icon(self, shapes, bx):
        name = bx.content[5:]
        color = bx.style.stroke or 'accent'
        try:
            from . import icons
            ic = icons.get(name)
            if ic is None:
                raise LookupError(name)
            lines = icons.polylines(ic, min(bx.w, bx.h))
            group = shapes.add_group_shape()
            group.name = f'Icon: {name}'
            width = max(0.75, float(getattr(icons, 'STROKE', 2.0)) * min(bx.w, bx.h) / float(getattr(icons, 'VIEWBOX',
                                                                                                        24)))
            for pl in lines:
                pts = [(bx.x + a, bx.y + b) for a, b in pl]
                closed = len(pts) > 2 and pts[0] == pts[-1]
                shp = self.poly(group.shapes, pts[:-1] if closed else pts, stroke=color, stroke_w=width,
                                closed=closed, name='Icon')
                if shp is not None:
                    self.round_caps(shp)
            self.decorative(group, bx.alt)
            return
        except Exception:
            pass
        if name == 'quote':   # the icon set is not available: a large quotation mark in the accent colour
            from pptx.util import Emu, Pt
            tb = shapes.add_textbox(Emu(_emu(bx.x)), Emu(_emu(bx.y - bx.h * 0.2)), Emu(_emu(bx.w * 1.5)),
                                    Emu(_emu(bx.h * 1.4)))
            tb.name = 'Icon: quote'
            tf = tb.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            r = tf.paragraphs[0].add_run()
            r.text = '“'
            r.font.size = Pt(max(24.0, bx.h * 1.6))
            r.font.bold = True
            r.font.color.rgb = self.rgb(color)
            self.decorative(tb, bx.alt)

    @staticmethod
    def round_caps(shp):
        from pptx.oxml.ns import qn
        try:
            ln = shp.line._get_or_add_ln()
            ln.set('cap', 'rnd')
            from lxml import etree
            if ln.find(qn('a:round')) is None:
                etree.SubElement(ln, qn('a:round'))
        except Exception:
            pass

    # ----- pictures, charts, tables, diagrams -----
    def picture(self, shapes, bx):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Emu, Pt
        data = image_bytes(bx, self.ws, mono=bool(self.ds.grey_images) and bx.content.startswith('asset:'))
        pic = shapes.add_picture(io.BytesIO(data), Emu(_emu(bx.x)), Emu(_emu(bx.y)), Emu(_emu(bx.w)), Emu(_emu(bx.h)))
        pic.name = 'Image'
        pic._element.nvPicPr.cNvPr.set('descr', (bx.alt or '')[:1000])
        if bx.style.radius and bx.style.radius > 0:
            try:
                pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
            except Exception:
                pass
        if bx.style.stroke and bx.style.stroke_w:
            pic.line.color.rgb = self.rgb(bx.style.stroke)
            pic.line.width = Pt(bx.style.stroke_w)
        return pic

    def chart(self, shapes, bx):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
        from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_PATTERN
        from pptx.util import Emu, Pt
        from .layout import block_of
        b = block_of(bx.content, self.spec)
        if not b or b.get('type') != 'chart':
            raise ValueError(f'{bx.content} is not a chart')
        render = _render()
        kind = {'bar': XL_CHART_TYPE.COLUMN_CLUSTERED, 'line': XL_CHART_TYPE.LINE_MARKERS,
                'pie': XL_CHART_TYPE.PIE}.get(b.get('kind'), XL_CHART_TYPE.COLUMN_CLUSTERED)
        data = CategoryChartData()
        data.categories = [render.safe_cell(lab) for lab in b['labels']]   # X2: never a live formula
        series = b['series'][:1] if b.get('kind') == 'pie' else b['series']
        for s in series:
            data.add_series(render.safe_cell(s['name']), s['values'])
        frame = shapes.add_chart(kind, Emu(_emu(bx.x)), Emu(_emu(bx.y)), Emu(_emu(bx.w)), Emu(_emu(bx.h)), data)
        frame.name = 'Chart'
        frame._element.nvGraphicFramePr.cNvPr.set('descr', (bx.alt or '')[:1000])
        chart = frame.chart
        text_c = self.rgb('text')
        small = max(12.0, self.ds.size('caption', 'pptx'))
        chart.has_title = True
        chart.chart_title.text_frame.text = b.get('title') or 'Chart'
        for r in chart.chart_title.text_frame.paragraphs[0].runs:
            r.font.size, r.font.bold, r.font.color.rgb = Pt(small + 2), True, text_c
            r.font.name = self.ds.families.heading
        chart.font.size, chart.font.color.rgb = Pt(small), text_c
        chart.font.name = self.ds.families.body
        chart.has_legend = len(series) > 1 or b.get('kind') == 'pie'
        if chart.has_legend:
            chart.legend.position, chart.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
        palette = [self.RGB.from_string(p.upper()) for p in (self.ds.chart_palette or [self.ds.color('accent')])]
        hatch = bool(self.ds.hatch)
        plot = chart.plots[0]
        if b.get('kind') == 'pie':
            plot.has_data_labels = True
            plot.data_labels.show_percentage, plot.data_labels.show_value = True, False
            plot.data_labels.number_format, plot.data_labels.number_format_is_linked = '0%', False
            for i, point in enumerate(plot.series[0].points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = palette[i % len(palette)]
        else:
            markers = [XL_MARKER_STYLE.CIRCLE, XL_MARKER_STYLE.SQUARE, XL_MARKER_STYLE.DIAMOND,
                       XL_MARKER_STYLE.TRIANGLE, XL_MARKER_STYLE.X, XL_MARKER_STYLE.STAR]
            patterns = [MSO_PATTERN.WIDE_UPWARD_DIAGONAL, MSO_PATTERN.PERCENT_20, MSO_PATTERN.WIDE_DOWNWARD_DIAGONAL,
                        MSO_PATTERN.NARROW_HORIZONTAL, MSO_PATTERN.CROSS, MSO_PATTERN.NARROW_VERTICAL]
            dashes = [MSO_LINE_DASH_STYLE.SOLID, MSO_LINE_DASH_STYLE.DASH, MSO_LINE_DASH_STYLE.ROUND_DOT,
                      MSO_LINE_DASH_STYLE.DASH_DOT, MSO_LINE_DASH_STYLE.SQUARE_DOT, MSO_LINE_DASH_STYLE.LONG_DASH]
            for i, s in enumerate(plot.series):
                if b.get('kind') == 'line':
                    s.format.line.color.rgb = text_c if hatch else palette[i % len(palette)]
                    if hatch:
                        s.format.line.dash_style = dashes[i % len(dashes)]
                    s.marker.style = markers[i % len(markers)]
                    s.marker.format.fill.solid()
                    s.marker.format.fill.fore_color.rgb = text_c if hatch else palette[i % len(palette)]
                elif hatch:
                    s.format.fill.patterned()
                    s.format.fill.pattern = patterns[i % len(patterns)]
                    s.format.fill.fore_color.rgb = text_c
                    s.format.fill.back_color.rgb = self.rgb('bg')
                    s.format.line.color.rgb = text_c
                else:
                    s.format.fill.solid()
                    s.format.fill.fore_color.rgb = palette[i % len(palette)]
            if b.get('kind') == 'bar' and len(b['labels']) * len(series) <= 24:
                plot.has_data_labels = True
                plot.data_labels.font.size, plot.data_labels.font.color.rgb = Pt(small), text_c
        return frame

    def table(self, shapes, bx):
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Emu, Pt
        from .layout import block_of, parse_ref, table_cells, table_geometry
        b = block_of(bx.content, self.spec)
        if not b:
            raise ValueError(f'{bx.content} is not a block')
        r = parse_ref(bx.content)
        cols, rows = table_cells(b)
        rng = r.rows or (0, len(rows))
        size = float(bx.size or self.ds.size('body', 'pptx'))
        g = table_geometry(b, rng, bx.w, size, self.ds, 'pptx', height=bx.h)
        body = rows[rng[0]:rng[1]]
        n = max(1, len(cols))
        frame = shapes.add_table(len(body) + 1, n, Emu(_emu(bx.x)), Emu(_emu(bx.y)), Emu(_emu(bx.w)),
                                 Emu(_emu(sum(g.heights))))
        frame.name = 'Table'
        frame._element.nvGraphicFramePr.cNvPr.set('descr', (bx.alt or '')[:1000])
        tbl = frame.table
        tbl.first_row = True          # A2: the header row is marked as a header
        tbl.horz_banding = False
        for j, w in enumerate(g.cols):
            tbl.columns[j].width = Emu(_emu(w))
        for i, h in enumerate(g.heights):
            tbl.rows[i].height = Emu(_emu(h))
        is_num = _render().is_num
        raw = [b.get('columns') or cols] + [(b.get('rows') or [])[rng[0] + k] if b.get('type') == 'table' else row
                                            for k, row in enumerate(body)]
        for i, data in enumerate([cols] + body):
            for j in range(n):
                v = data[j] if j < len(data) else ''
                cell = tbl.cell(i, j)
                cell.fill.solid()
                cell.fill.fore_color.rgb = self.rgb('header_bg' if i == 0 else 'stripe' if i % 2 == 0 else 'bg')
                cell.margin_left = cell.margin_right = Emu(_emu(g.pad))
                cell.margin_top = cell.margin_bottom = Emu(_emu(g.vpad or g.pad))
                cell.vertical_anchor = MSO_ANCHOR.TOP
                p = cell.text_frame.paragraphs[0]
                src = raw[i][j] if i and j < len(raw[i]) else None
                if i and is_num(src):
                    p.alignment = PP_ALIGN.RIGHT
                run = p.add_run()
                run.text = str(v)
                run.font.size = Pt(size)
                run.font.bold = True if i == 0 else None
                run.font.name = g.family
                run.font.color.rgb = self.rgb('header_text' if i == 0 else 'text')
        return frame

    def diagram(self, slide, bx):
        from ..create import diagram
        from .layout import block_of, theme_dict
        b = block_of(bx.content, self.spec)
        if not b:
            raise ValueError(f'{bx.content} is not a block')
        diagram.pptx_draw(slide, b, theme_dict(self.ds), (bx.x / 72, bx.y / 72, bx.w / 72, bx.h / 72),
                          font=self.ds.families.body)
        grp = slide.shapes[-1]
        self.decorative(grp, bx.alt)

    # ----- the file -----
    def paint(self) -> bytes:
        from pptx import Presentation
        from pptx.util import Emu
        plan, spec = self.plan, self.spec
        if plan.format != 'pptx' or not plan.pages:
            raise ValueError('a pptx plan with pages is needed')
        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(_emu(plan.pages[0].w)), Emu(_emu(plan.pages[0].h))   # 16:9 (F3)
        prs.core_properties.title = spec.get('title') or ''
        prs.core_properties.author = AUTHOR
        prs.core_properties.subject = _render().subject(spec.get('subtitle'))
        for page in plan.pages:
            self.slide(prs, page)
        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()

    def slide(self, prs, page):
        from pptx.util import Emu
        from .layout import block_of
        tb = title_box(page, self.spec)
        slide = prs.slides.add_slide(prs.slide_layouts[5] if tb is not None else prs.slide_layouts[6])
        fill = slide.background.fill
        fill.solid()
        bg = page.background if page.background in self.ds.colors or str(page.background).startswith('chart') \
            else 'bg'
        fill.fore_color.rgb = self.rgb(bg)
        tree = slide.shapes._spTree
        group = None
        order = sorted(range(len(page.boxes)), key=lambda k: (page.boxes[k].z, k))
        for k in order:
            bx = page.boxes[k]
            shapes = slide.shapes
            if bx.slot == 'cards' and page.layout == 'stat-cards' or (bx.slot == 'cards' and group is not None):
                if group is None:
                    blk = next((block_of(x.content, self.spec) for x in page.boxes
                                if x.slot == 'cards' and x.content.startswith('spec:')), None)
                    group = slide.shapes.add_group_shape()
                    title = (blk or {}).get('title') or 'Key figures'
                    group.name = f'Diagram: {title}' if (blk or {}).get('type') == 'stat-cards' else 'Stat cards'
                shapes = group.shapes
            if bx is tb:
                ph = slide.shapes.title
                ph.left, ph.top, ph.width, ph.height = (Emu(_emu(v)) for v in (bx.x, bx.y, bx.w, bx.h))
                self.fill_frame(ph.text_frame, bx, title=True)
                el = ph._element
                el.getparent().remove(el)
                tree.insert_element_before(el, 'p:extLst')
                continue
            kind = bx.kind
            if kind == 'text':
                self.text(shapes, bx)
            elif kind == 'shape':
                if bx.content.startswith('art:'):
                    self.art(shapes, bx)
                else:
                    shp = self.rect(shapes, bx.x, bx.y, bx.w, bx.h, fill=bx.style.fill, stroke=bx.style.stroke,
                                    stroke_w=bx.style.stroke_w, opacity=bx.style.opacity, radius=bx.style.radius,
                                    ellipse=bool(bx.style.radius and bx.style.radius >= min(bx.w, bx.h) / 2 - 0.01
                                                 and abs(bx.w - bx.h) < 0.5),
                                    name=(bx.slot or 'Shape').capitalize())
                    self.decorative(shp, bx.alt)
            elif kind == 'image':
                self.picture(shapes, bx)
            elif kind == 'chart':
                self.chart(shapes, bx)
            elif kind == 'table':
                self.table(shapes, bx)
            elif kind == 'diagram':
                self.diagram(slide, bx)
            elif kind == 'icon':
                self.icon(shapes, bx)
        if tb is None and slide.shapes.title is not None:
            el = slide.shapes.title._element
            el.getparent().remove(el)
        if page.notes:
            slide.notes_slide.notes_text_frame.text = page.notes   # overflow text lives in the notes (L3, F3)


def paint(plan: DesignPlan, spec: dict, ws: Workspace) -> bytes:
    """The .pptx bytes for a plan of format 'pptx' (one slide per page, 960 x 540 pt). Raises on a library failure;
    the caller turns that into a fallback, never a crash."""
    return _Painter(plan, spec, ws).paint()
