"""Seeded fuzz of created files (docs/PLAN-files-robust.md section 7): every spec a model could plausibly send, and many
it should not, becomes a file in all five formats, and no text is thrown away on the way.

The generator starts from one of six valid bases and applies 1 to 4 mutations from a catalog (the 24 probe shapes of the
run 2750 forensics, the string series, the 27 odd shapes of the fuzz audit, number strings, lone surrogates, long
headings and column names, long subtitles, 41 sections, 31 blocks, page-break-only and figure-only specs, CJK, Arabic,
Hindi and emoji text, None or a wrong type at any key, depth 60 nesting, NaN, infinity, booleans, link keys everywhere,
and the type, field, section and top-level aliases of the repair algorithm). Every text field carries a unique marker
word (`zq` + 6 letters), so preservation can be checked.

Properties, for every spec and every format:
  P1  normalize raises nothing but SpecError, and only with rule S2, L6 or X1.
  P2  when the oracle finds text in a section, normalize does not raise.
  P3  render_safe(norm, fmt) gives bytes that verify reopens (V1 ok).
  P4  normalize(normalize(spec, fmt)[0], fmt) does not raise and keeps the section count.
  P5  every marker in a block's text survives into the text of normalize(spec, 'md').
  P6  a mutation that drops or converts content leaves at least one `fix` result with ok false.

TG_FUZZ_N sets the number of specs (default 300); TG_FUZZ_SHARD=k/m runs every m-th spec starting at k, so the 5,000
spec run of the acceptance table can be split across processes:
  TG_FUZZ_N=5000 TG_FUZZ_SHARD=0/4 .venv/bin/pytest -q tests/test_create_fuzz.py   (and 1/4, 2/4, 3/4)
Everything here is offline: no engine, no network, fonts and the image cache are patched to temporary ones.
"""
import copy
import functools
import io
import json
import os
import random
import re
from pathlib import Path

import pytest

import jevrouter.create as cf
from jevrouter import evals
from jevrouter.agents import create as ca
from jevrouter.create import spec as spec_mod
from jevrouter.create.brief import parse_brief
from jevrouter.create.longdoc import OUTLINE_SCHEMA
from jevrouter.engines import Reply

ROOT = Path(__file__).resolve().parent.parent
FIXTURES = ROOT / 'evals' / 'fixtures'
FORMATS = ('pdf', 'docx', 'pptx', 'xlsx', 'md')
SEED = 20260928
N = int(os.environ.get('TG_FUZZ_N', '300'))
REFUSALS = {'S2', 'L6', 'X1'}   # the only rules that may refuse a whole file (plan 0.1)
MARK = re.compile(r'zq[a-z]{6}')

# plan 2.3 step 5: keys whose strings are never content, and the block types that never hold text
SKIP_KEYS = {'type', 'kind', 'lang', 'ordered', 'level', 'id', 'parent', 'from', 'to', 'formats', 'asset', 'credit',
             'query'} | set(spec_mod.FETCH_KEYS)
FIGURE_TYPES = {'figure', 'picture', 'image', 'photo', 'img', 'illustration'}
BREAK_TYPES = {'page_break', 'hr', 'divider', 'break', 'pagebreak'}
KNOWN_TYPES = set(spec_mod.BLOCKS) | {
    'list', 'ul', 'bullet', 'bulleted_list', 'points', 'ol', 'numbered_list', 'numbered', 'text', 'para', 'p', 'body',
    'markdown', 'description', 'blockquote', 'pre', 'snippet', 'source_code', 'grid', 'matrix', 'data_table', 'graph',
    'plot', 'diagram', 'heading', 'header', 'h1', 'h2', 'h3', 'subheading', 'title', 'donut', 'doughnut',
    *spec_mod.CHART_KINDS, *spec_mod.CHART_ALIASES} | FIGURE_TYPES | BREAK_TYPES


def selected() -> list[int]:
    """The spec indices this process runs (TG_FUZZ_SHARD=k/m)."""
    shard = os.environ.get('TG_FUZZ_SHARD', '').strip()
    k, m = 0, 1
    if shard:
        a, b = shard.split('/')
        k, m = int(a), max(1, int(b))
    return [i for i in range(N) if i % m == k % m]


@pytest.fixture(autouse=True)
def offline(monkeypatch, tmp_path):
    """Fonts see no installed faces and no TRACEGRAPH_BODY_FONT (the result does not depend on the machine) and images
    go to an empty temporary cache."""
    from jevrouter.create import assets, fonts
    monkeypatch.setattr(fonts, 'system_index', lambda: {})
    monkeypatch.delenv('TRACEGRAPH_BODY_FONT', raising=False)
    monkeypatch.setattr(assets, 'CACHE', tmp_path / 'assets')


def needs(obj, name: str, who: str):
    if not hasattr(obj, name):
        pytest.skip(f'waiting for builder {who}: {getattr(obj, "__name__", obj)}.{name} (docs/PLAN-files-robust.md)')


# ---------- markers ----------

class Ctx:
    """One spec's generation state: its random source, a marker counter, the markers a mutation may lose (unsafe), and
    whether any applied mutation drops or converts content (P6)."""

    def __init__(self, rng: random.Random):
        self.rng, self.i = rng, rng.randrange(26 ** 5)
        self.unsafe: set[str] = set()
        self.p6: list[str] = []
        self.names: list[str] = []
        self.cut_tail = False   # L1: sections past 40 are cut from the end

    def m(self) -> str:
        """A marker word unique within the spec: 7919 is prime to 26, so the counter maps one to one."""
        self.i += 1
        n = (self.i * 7919 + 12345) % 26 ** 6
        s = ''
        for _ in range(6):
            n, r = divmod(n, 26)
            s += chr(97 + r)
        return 'zq' + s

    def t(self, words: int = 5) -> str:
        """A marker followed by a few plain words."""
        return ' '.join([self.m(), *self.rng.sample(VOCAB, words)])

    def lose(self, obj):
        self.unsafe |= markers(obj)


VOCAB = ('phones got smaller and cheaper every decade while batteries lasted longer networks moved from analog to '
         'digital then data came first and voice became an app cameras replaced film maps replaced paper people '
         'carried the internet in a pocket').split()


def markers(obj) -> set[str]:
    try:
        return set(MARK.findall(json.dumps(obj, ensure_ascii=False, default=str)))
    except (ValueError, RecursionError):
        return set(MARK.findall(repr(obj)))


# ---------- the oracle: where a spec holds text (plan 2.3, read conservatively) ----------

def strings(x, skip=SKIP_KEYS):
    if isinstance(x, str):
        yield x
    elif isinstance(x, dict):
        for k, v in x.items():
            if k not in skip:
                yield from strings(v, skip)
    elif isinstance(x, list):
        for v in x:
            yield from strings(v, skip)


def infer(b: dict) -> str | None:
    """The type repair infers for a block without a known one (plan 2.3 step 4, first match wins); None: salvaged."""
    has = lambda k: b.get(k) is not None
    if has('items'):
        return 'bullets'
    if has('rows') or has('columns'):
        return 'table'
    if (has('series') or has('data')) and has('labels'):
        return 'chart'
    if has('events'):
        return 'timeline'
    if has('edges'):
        return 'flow'
    if has('nodes'):
        return 'tree'
    if has('text') or has('content'):
        return 'paragraph'
    if has('code') or has('lang'):
        return 'code'
    if has('query') or has('caption'):
        return 'figure'
    return None


def block_texts(b):
    if isinstance(b, str):
        yield b   # a block given as a string keeps its text (salvage)
        return
    if not isinstance(b, dict):
        return
    t = b.get('type')
    t = t.strip().lower() if isinstance(t, str) else None
    if t not in KNOWN_TYPES:
        t = infer(b) or 'salvage'
    if t in FIGURE_TYPES or t in BREAK_TYPES:
        return
    yield from strings(b)


def container(sec: dict):
    if sec.get('blocks') is not None:
        return sec['blocks']
    return next((sec[k] for k in ('content', 'body', 'items') if sec.get(k) is not None), None)


def section_texts(sec, notes: bool = True, heading: bool = True):
    if isinstance(sec, str):
        yield sec
        return
    if not isinstance(sec, dict):
        return
    if heading:
        if isinstance(sec.get('heading'), str):
            yield sec['heading']
        elif sec.get('heading') is None:
            yield from (sec[k] for k in ('title', 'name') if isinstance(sec.get(k), str))
    if notes and isinstance(sec.get('notes'), str):
        yield sec['notes']
    c = container(sec)
    if isinstance(c, list):
        for b in c:
            yield from block_texts(b)
    elif isinstance(c, (dict, str)):
        yield from block_texts(c)
    for k in ('bullets', 'points'):
        if isinstance(sec.get(k), (list, str)):
            yield from strings(sec[k])


def sections_of(spec) -> list:
    if isinstance(spec, list):
        return spec
    if isinstance(spec, str):
        return [spec]
    if not isinstance(spec, dict):
        return []
    v = spec['sections'] if spec.get('sections') is not None else next(
        (spec[k] for k in ('slides', 'pages', 'parts') if spec.get(k) is not None), None)
    if isinstance(v, list):
        return v
    if isinstance(v, dict):
        return list(v.values())
    return [v] if isinstance(v, str) else []


def visible(s: str) -> bool:
    try:
        return bool(spec_mod.plain(s).strip())
    except Exception:
        return False


def oracle_has_text(spec) -> bool:
    """P2: headings, notes and block strings (outside the skipped keys and figure blocks) of any section."""
    return any(visible(s) for sec in sections_of(spec) for s in section_texts(sec))


def required(spec, ctx: Ctx) -> set[str]:
    """P5: the markers in block text (and headings) that must reach the Markdown spec."""
    secs = sections_of(spec)
    if ctx.cut_tail:
        for sec in secs[25:]:
            ctx.lose(sec)
    found = {m for sec in secs for s in section_texts(sec, notes=False) for m in MARK.findall(s)}
    return found - ctx.unsafe


def norm_text(norm: dict) -> str:
    """The words of a normalized spec: the create agent's own spec_text plus diagram labels (which it leaves out)."""
    out = [ca.spec_text(norm)]
    for s in norm.get('sections') or []:
        for b in s.get('blocks') or []:
            for e in b.get('events') or []:
                out += [str(e.get('date') or ''), str(e.get('label') or '')]
            for n in b.get('nodes') or []:
                out.append(str(n.get('label') or ''))
            for e in b.get('edges') or []:
                if isinstance(e, dict):
                    out.append(str(e.get('label') or ''))
    return '\n'.join(out)


# ---------- the six valid bases ----------

def b_para(c):
    return {'type': 'paragraph', 'text': c.t(8) + '.'}


def b_bullets(c, n=3, ordered=False):
    return {'type': 'bullets', 'items': [c.t(4) for _ in range(n)], 'ordered': ordered}


def b_table(c, rows=4):
    return {'type': 'table', 'title': c.t(2), 'columns': [c.t(1), c.t(1), c.t(1)],
            'rows': [[c.t(1), 1990 + i, round(c.rng.uniform(1, 500), 1)] for i in range(rows)]}


def b_chart(c, kind='bar'):
    labels = [c.t(1) for _ in range(4)]
    series = [{'name': c.t(1), 'values': [c.rng.randint(1, 90) for _ in labels]}]
    if kind != 'pie':
        series.append({'name': c.t(1), 'values': [c.rng.randint(1, 90) for _ in labels]})
    return {'type': 'chart', 'kind': kind, 'title': c.t(2), 'labels': labels, 'series': series}


def b_timeline(c, n=5):
    return {'type': 'timeline', 'title': c.t(2),
            'events': [{'date': str(1946 + 9 * i), 'label': c.t(3)} for i in range(n)]}


def b_tree(c):
    nodes = [{'id': 'root', 'parent': '', 'label': c.t(1)}]
    for i in range(4):
        nodes.append({'id': f'n{i}', 'parent': 'root' if i < 2 else 'n0', 'label': c.t(1)})
    return {'type': 'tree', 'title': c.t(2), 'nodes': nodes}


def b_flow(c):
    nodes = [{'id': f's{i}', 'label': c.t(1)} for i in range(4)]
    return {'type': 'flow', 'title': c.t(2), 'nodes': nodes,
            'edges': [{'from': f's{i}', 'to': f's{i + 1}', 'label': ''} for i in range(3)]}


def sec(c, *blocks, notes=''):
    return {'heading': c.t(2), 'level': 1, 'blocks': list(blocks), 'notes': notes}


def base_short(c):
    return {'title': c.t(3), 'subtitle': c.t(4), 'sections': [
        sec(c, b_para(c), b_para(c)), sec(c, b_bullets(c)),
        sec(c, b_para(c), {'type': 'quote', 'text': c.t(6), 'by': c.t(1)})]}


def base_deck(c):
    secs = [sec(c, b_bullets(c, c.rng.randint(2, 4)), notes=c.t(6)) for _ in range(11)]
    secs[2]['blocks'] = [b_bullets(c, 2), b_table(c, 3)]
    secs[5]['blocks'] = [b_bullets(c, 2), b_timeline(c, 4)]
    return {'title': c.t(3), 'subtitle': '', 'sections': secs}


def base_tables(c):
    return {'title': c.t(3), 'subtitle': c.t(3), 'sections': [
        sec(c, b_para(c), b_table(c, 6)), sec(c, b_table(c, 12)), sec(c, b_para(c), b_table(c, 3), b_para(c)),
        sec(c, b_bullets(c))]}


def base_charts(c):
    return {'title': c.t(3), 'subtitle': '', 'sections': [
        sec(c, b_para(c), b_chart(c, 'bar')), sec(c, b_chart(c, 'line')), sec(c, b_para(c), b_chart(c, 'pie')),
        sec(c, b_table(c, 4))]}


def base_diagrams(c):
    return {'title': c.t(3), 'subtitle': c.t(2), 'sections': [
        sec(c, b_para(c), b_timeline(c)), sec(c, b_tree(c)), sec(c, b_para(c), b_flow(c))]}


def base_mixed(c):
    return {'title': c.t(3), 'subtitle': c.t(3), 'sections': [
        sec(c, b_para(c), b_bullets(c, 3, ordered=True)),
        sec(c, b_table(c, 3), b_chart(c, 'bar')),
        {**sec(c, {'type': 'quote', 'text': c.t(6), 'by': c.t(1)}, {'type': 'code', 'lang': 'python',
                                                                    'text': 'print("' + c.m() + '")'}), 'level': 2},
        sec(c, b_timeline(c, 4), {'type': 'figure', 'query': 'Motorola DynaTAC', 'caption': 'A DynaTAC 8000X'}),
        sec(c, {'type': 'page_break'}, b_para(c))]}


BASES = [('short doc', base_short), ('11-section deck', base_deck), ('tables report', base_tables),
         ('charts', base_charts), ('diagrams', base_diagrams), ('mixed', base_mixed)]


# ---------- helpers for mutations ----------

def dict_sections(spec) -> list[dict]:
    return [s for s in sections_of(spec) if isinstance(s, dict)]


def section_list(spec) -> list | None:
    if isinstance(spec, dict) and isinstance(spec.get('sections'), list):
        return spec['sections']
    return spec if isinstance(spec, list) else None


def insert_block(c: Ctx, spec, block) -> bool:
    """Puts a block in a random section that has a list of blocks, or in a new section at the end."""
    homes = [s for s in dict_sections(spec) if isinstance(s.get('blocks'), list)]
    if homes:
        s = c.rng.choice(homes)
        s['blocks'].insert(c.rng.randint(0, len(s['blocks'])), block)
        return True
    secs = section_list(spec)
    if secs is None:
        return False
    secs.append({'heading': c.t(2), 'level': 1, 'blocks': [block]})
    return True


def blocks_of(spec, types=None) -> list[tuple[dict, int, dict]]:
    out = []
    for s in dict_sections(spec):
        if isinstance(s.get('blocks'), list):
            for i, b in enumerate(s['blocks']):
                if isinstance(b, dict) and (types is None or b.get('type') in types):
                    out.append((s, i, b))
    return out


def pick_block(c: Ctx, spec, types=None):
    found = blocks_of(spec, types)
    return c.rng.choice(found) if found else None


def junk(c: Ctx):
    return c.rng.choice([None, 0, -1, 1.5, True, False, '', '   ', [], {}, [None], 'text', float('nan'),
                         float('inf'), {'a': 1}, [[]], 10 ** 30])


# ---------- the catalog: (name, phase, p6, function(ctx, spec) -> applied) ----------
# phase orders the mutations of one spec: 0 replaces the whole spec, 1 breaks existing objects, 2 adds or changes
# blocks, 3 changes sections and the spec, 4 changes the top-level shape. p6: the mutation drops or converts content.

CATALOG: list[tuple[str, int, bool, object]] = []


def mut(name: str, phase: int = 2, p6: bool = False):
    def deco(fn):
        CATALOG.append((name, phase, p6, fn))
        return fn
    return deco


def adds(name: str, make, p6: bool = False, lose: bool = False):
    """A mutation that inserts one block made by make(ctx). lose: its markers may be dropped (a risky shape)."""
    def fn(c, spec):
        b = make(c)
        if lose:
            c.lose(b)
        return insert_block(c, spec, b)
    CATALOG.append((name, 2, p6, fn))


# the 24 probe shapes of the run 2750 forensics (one bad block in a valid deck; 21 of them discarded the whole file)
adds('probe: bullets without items', lambda c: {'type': 'bullets'}, p6=True)
adds('probe: bullets items null', lambda c: {'type': 'bullets', 'items': None}, p6=True)
adds('probe: bullets text not items', lambda c: {'type': 'bullets', 'text': c.t(3) + '\n' + c.t(3)}, p6=True)
adds('probe: bullets points key', lambda c: {'type': 'bullets', 'points': [c.t(3), c.t(3)]})
adds('probe: bullets items string', lambda c: {'type': 'bullets', 'items': c.t(2) + ', ' + c.t(2)})
adds('probe: bullets items empty', lambda c: {'type': 'bullets', 'items': []}, p6=True)
adds('probe: paragraph content key', lambda c: {'type': 'paragraph', 'content': c.t(8)})
adds('probe: table without rows', lambda c: {'type': 'table', 'columns': [c.t(1), c.t(1)]})
adds('probe: chart radar', lambda c: {'type': 'chart', 'kind': 'radar', 'title': c.t(2), 'labels': [c.t(1), c.t(1)],
                                      'series': [{'name': c.t(1), 'values': [3, 4]}]}, p6=True)
adds('probe: chart series without name', lambda c: {'type': 'chart', 'kind': 'bar', 'title': c.t(2),
                                                    'labels': [c.t(1), c.t(1)], 'series': [{'values': [1, 2]}]})
adds('probe: chart values strings', lambda c: {'type': 'chart', 'kind': 'bar', 'title': c.t(2),
                                               'labels': [c.t(1), c.t(1)],
                                               'series': [{'name': c.t(1), 'values': ['1', '2']}]}, p6=True)
adds('probe: timeline without events', lambda c: {'type': 'timeline', 'title': c.t(3)}, p6=True)
adds('probe: timeline items key', lambda c: {'type': 'timeline', 'title': c.t(2), 'items': [
    {'date': '1973', 'label': c.t(2)}, {'date': '1983', 'label': c.t(2)}]})
adds('probe: tree nodes dict', lambda c: {'type': 'tree', 'title': c.t(2), 'nodes': {'a': c.t(1)}}, lose=True)
adds('probe: flow edges dict', lambda c: {'type': 'flow', 'title': c.t(2), 'nodes': [{'id': 'a', 'label': c.t(1)}],
                                          'edges': {'a': 'b'}}, p6=True, lose=True)
adds('probe: heading block', lambda c: {'type': 'heading', 'text': c.t(3)}, p6=True)
adds('probe: diagram block', lambda c: {'type': 'diagram', 'nodes': []}, p6=True)
adds('probe: picture block', lambda c: {'type': 'picture', 'query': 'Nokia 3310'}, lose=True)
adds('probe: block is a string', lambda c: c.t(6))
adds('probe: figure with url', lambda c: {'type': 'figure', 'query': 'phone', 'caption': 'c',
                                          'url': 'https://example.com/x.png'}, p6=True)
adds('probe: figure with image key', lambda c: {'type': 'figure', 'query': 'phone', 'caption': 'c',
                                                'image': 'dynatac'}, p6=True)
adds('probe: bullets with link key', lambda c: {'type': 'bullets', 'items': [c.t(3)], 'link': 'https://example.com'},
     p6=True)


@mut('probe: section blocks dict', phase=1)
def _(c, spec):
    secs = dict_sections(spec)
    if not secs:
        return False
    c.rng.choice(secs)['blocks'] = {'type': 'bullets', 'items': [c.t(3), c.t(2)]}
    return True


@mut('probe: section is a string', phase=1, p6=True)
def _(c, spec):
    secs = section_list(spec)
    if not secs:
        return False
    secs[c.rng.randrange(len(secs))] = c.t(2)
    return True


# the string series
adds('string series', lambda c: {'type': 'chart', 'kind': 'pie', 'title': c.t(2), 'labels': [c.t(1), c.t(1)],
                                 'series': 'x'}, p6=True)

# the 27 odd shapes of the fuzz audit
adds('odd: ragged table', lambda c: {'type': 'table', 'columns': [c.t(1), c.t(1), c.t(1)],
                                     'rows': [[c.t(1)], [c.t(1), c.t(1), c.t(1), 'extra'], []]}, p6=True)
adds('odd: table without columns', lambda c: {'type': 'table', 'columns': [], 'rows': [[c.t(1), c.t(1)]]}, p6=True)
adds('odd: table rows of strings', lambda c: {'type': 'table', 'columns': [c.t(1)], 'rows': [c.t(1), c.t(1)]})
adds('odd: chart mismatched', lambda c: {'type': 'chart', 'kind': 'bar', 'title': c.t(2),
                                         'labels': [c.t(1), c.t(1), c.t(1)],
                                         'series': [{'name': c.t(1), 'values': [1]}]})
adds('odd: chart series string', lambda c: {'type': 'chart', 'kind': 'bar', 'title': c.t(2), 'labels': [c.t(1)],
                                            'series': ['x']}, p6=True, lose=True)
adds('odd: chart empty', lambda c: {'type': 'chart', 'kind': 'line', 'title': c.t(2), 'labels': [], 'series': []},
     p6=True)
adds('odd: chart negative pie', lambda c: {'type': 'chart', 'kind': 'pie', 'title': c.t(2), 'labels': [c.t(1), c.t(1)],
                                           'series': [{'name': c.t(1), 'values': [-1, 0]}]}, p6=True, lose=True)
adds('odd: timeline events strings', lambda c: {'type': 'timeline', 'title': c.t(2), 'events': [
    '1973 ' + c.t(2), '1983: ' + c.t(2), '2007 - ' + c.t(2)]})
adds('odd: timeline event without date', lambda c: {'type': 'timeline', 'title': c.t(2),
                                                    'events': [{'label': c.t(1)}]}, p6=True, lose=True)
adds('odd: timeline 40 events', lambda c: {'type': 'timeline', 'title': c.t(2), 'events': [
    {'date': str(1900 + i), 'label': c.t(2) + ' ' + 'e' * 30} for i in range(40)]}, p6=True, lose=True)
adds('odd: tree cycle', lambda c: {'type': 'tree', 'title': c.t(2), 'nodes': [
    {'id': 'a', 'label': c.t(1), 'parent': 'b'}, {'id': 'b', 'label': c.t(1), 'parent': 'a'}]}, lose=True)
adds('odd: tree without root', lambda c: {'type': 'tree', 'title': c.t(2), 'nodes': [
    {'id': 'a', 'label': c.t(1), 'parent': 'zz'}]}, lose=True)
adds('odd: tree nodes strings', lambda c: {'type': 'tree', 'title': c.t(2), 'nodes': [c.t(1), c.t(1)]}, lose=True)
adds('odd: tree 200 nodes', lambda c: {'type': 'tree', 'title': c.t(2), 'nodes': [{'id': 'r', 'label': 'Root'}] + [
    {'id': f'n{i}', 'label': f'N{i}', 'parent': 'r'} for i in range(200)]}, p6=True, lose=True)
adds('odd: flow dangling edge', lambda c: {'type': 'flow', 'title': c.t(2), 'nodes': [{'id': 'a', 'label': c.t(1)}],
                                           'edges': [{'from': 'a', 'to': 'zz'}]}, p6=True, lose=True)
adds('odd: flow nodes strings', lambda c: {'type': 'flow', 'title': c.t(2), 'nodes': [c.t(1), c.t(1), c.t(1)],
                                           'edges': []}, p6=True)
adds('odd: flow self loop', lambda c: {'type': 'flow', 'title': c.t(2), 'nodes': [
    {'id': 'a', 'label': c.t(1)}, {'id': 'b', 'label': c.t(1)}], 'edges': [
    {'from': 'a', 'to': 'a'}, {'from': 'b', 'to': 'a'}, {'from': 'a', 'to': 'b'}]})
adds('odd: nested bullets', lambda c: {'type': 'bullets', 'items': [[c.t(2), c.t(2)], {'text': c.t(2)}, None, 5]},)
adds('odd: paragraph text dict', lambda c: {'type': 'paragraph', 'text': {'a': 1}}, p6=True)
adds('odd: paragraph text number', lambda c: {'type': 'paragraph', 'text': 12345})
adds('odd: 5,000-character word', lambda c: {'type': 'paragraph', 'text': c.m() + ' ' + 'x' * 5000})
adds('odd: emoji, CJK and RTL', lambda c: {'type': 'paragraph',
                                           'text': c.m() + ' \U0001F4F1 手机 الهاتف '
                                                           '\U0001F1EE\U0001F1F3 ‍'})
adds('odd: huge code', lambda c: {'type': 'code', 'lang': 'py', 'text': '# ' + c.m() + '\n' + 'print(1)\n' * 800})
adds('odd: quote without by', lambda c: {'type': 'quote', 'text': c.t(5)})
adds('odd: figure without query', lambda c: {'type': 'figure'})
adds('odd: image with a bad asset', lambda c: {'type': 'image', 'asset': 'nothex', 'caption': 'x'}, lose=True)
adds('odd: page break block', lambda c: {'type': 'page_break'})


# number strings
@mut('number strings in a chart', p6=True)
def _(c, spec):
    got = pick_block(c, spec, {'chart'})
    if not got or not isinstance(got[2].get('series'), list) or not got[2]['series']:
        return insert_block(c, spec, {'type': 'chart', 'kind': 'bar', 'title': c.t(2), 'labels': [c.t(1), c.t(1)],
                                      'series': [{'name': c.t(1), 'values': [c.rng.choice(['', ' ', '%', '()', '(%)',
                                                                                           '$']), 5]}]})
    s = got[2]['series'][0]
    if not isinstance(s, dict) or not isinstance(s.get('values'), list) or not s['values']:
        return False
    # a chart left with no numbers becomes bullets of its labels, without series names; a pie drops a slice with no
    # value, label and all
    c.lose(got[2] if str(got[2].get('kind')).lower() in ('pie', 'donut') else got[2].get('series'))
    s['values'][c.rng.randrange(len(s['values']))] = c.rng.choice(['', ' ', '%', '()', '(%)', '$'])
    return True


@mut('number strings in a table')
def _(c, spec):
    got = pick_block(c, spec, {'table'})
    if not got or not isinstance(got[2].get('rows'), list) or not got[2]['rows']:
        return insert_block(c, spec, {'type': 'table', 'columns': [c.t(1), c.t(1)],
                                      'rows': [['', '%'], ['()', '(%)'], [' ', '$']]})
    rows = [r for r in got[2]['rows'] if isinstance(r, list) and r]
    if not rows:
        return False
    r = c.rng.choice(rows)
    j = c.rng.randrange(len(r))
    c.lose(r[j])
    r[j] = c.rng.choice(['', ' ', '%', '()', '(%)', '$'])
    return True


@mut('lone surrogate')
def _(c, spec):
    return insert_block(c, spec, {'type': 'paragraph', 'text': c.m() + ' emoji \ud83d cut and \udc00 more'})


@mut('lone surrogate in a heading', phase=3)
def _(c, spec):
    secs = dict_sections(spec)
    if not secs:
        return False
    c.rng.choice(secs)['heading'] = c.m() + ' Phones \ud83d'
    return True


@mut('130-character heading', phase=3)
def _(c, spec):
    secs = dict_sections(spec)
    if not secs:
        return False
    h = c.t(3)
    c.rng.choice(secs)['heading'] = (h + ' ' + ' '.join(['mobile'] * 30))[:130]
    return True


@mut('5,000-character column name', p6=True)
def _(c, spec):
    got = pick_block(c, spec, {'table'})
    if not got or not isinstance(got[2].get('columns'), list) or not got[2]['columns']:
        return insert_block(c, spec, {'type': 'table', 'columns': [c.m() + ' ' + 'x' * 5000, c.t(1)],
                                      'rows': [[1, 2]]})
    got[2]['columns'][0] = c.m() + ' ' + 'x' * 5000
    return True


@mut('5,000-character column name of words', p6=True)
def _(c, spec):
    return insert_block(c, spec, {'type': 'table', 'columns': [c.m() + ' ' + 'word ' * 1000, c.t(1), c.t(1)],
                                  'rows': [[1, 2, 3]]})


@mut('subtitle of 256 characters', phase=3)
def _(c, spec):
    if not isinstance(spec, dict):
        return False
    spec['subtitle'] = (c.m() + ' ' + 'history of the phone ' * 20)[:256]
    return True


@mut('subtitle of 300 characters', phase=3)
def _(c, spec):
    if not isinstance(spec, dict):
        return False
    spec['subtitle'] = (c.m() + ' ' + 'word ' * 80)[:300]
    return True


@mut('41 sections', phase=3, p6=True)
def _(c, spec):
    secs = section_list(spec)
    if secs is None:
        return False
    while len(secs) < 41:
        secs.append({'heading': c.t(2), 'level': 1, 'blocks': [b_para(c)]})
    c.cut_tail = True
    return True


@mut('31 blocks', p6=True)
def _(c, spec):
    secs = [s for s in dict_sections(spec) if isinstance(s.get('blocks'), list)]
    if not secs:
        return False
    s = c.rng.choice(secs)
    while len(s['blocks']) < 31:
        s['blocks'].append({'type': 'paragraph', 'text': c.t(3)})
    return True


@mut('page-break-only spec', phase=0)
def _(c, spec):
    spec.clear()
    spec.update({'title': c.t(2), 'sections': [{'heading': '', 'blocks': [{'type': 'page_break'}] * 3}]})
    return True


@mut('figure-only spec', phase=0)
def _(c, spec):
    spec.clear()
    spec.update({'title': c.t(2), 'sections': [{'heading': '', 'blocks': [
        {'type': 'figure', 'query': 'Motorola DynaTAC', 'caption': 'A phone'}]}]})
    return True


@mut('CJK, Arabic, Hindi and emoji text')
def _(c, spec):
    text = c.rng.choice(['携帯電話の歴史 and 手机',
                         'تاريخ الهاتف المحمول',
                         'मोबाइल फोन का इतिहास',
                         'Phones \U0001F4F1 are great \U0001F680 → 5G ≈ 2 Mbps'])
    return insert_block(c, spec, c.rng.choice([{'type': 'paragraph', 'text': c.m() + ' ' + text},
                                               {'type': 'bullets', 'items': [c.m() + ' ' + text, c.t(2)]}]))


@mut('non-Latin heading', phase=3)
def _(c, spec):
    secs = dict_sections(spec)
    if not secs:
        return False
    c.rng.choice(secs)['heading'] = c.m() + ' 携帯電話 मोबाइल \U0001F4F1'
    return True


def objects(spec, rows: bool = True):
    """(object, its block or section) for every dict in the spec. rows: include dict table rows (their keys are cells,
    so a link key there is content, not a fetch)."""
    out = [(spec, None)] if isinstance(spec, dict) else []
    for s in dict_sections(spec):
        out.append((s, s))
        if isinstance(s.get('blocks'), list):
            for b in s['blocks']:
                if isinstance(b, dict):
                    out.append((b, b))
                    for k in ('events', 'nodes', 'edges', 'series', 'items') + (('rows',) if rows else ()):
                        for x in b.get(k) if isinstance(b.get(k), list) else []:
                            if isinstance(x, dict):
                                out.append((x, b))
    return out


@mut('None or a wrong type at a key', phase=1)
def _(c, spec):
    objs = objects(spec)
    if not objs:
        return False
    obj, owner = c.rng.choice(objs)
    k = c.rng.choice(list(obj) or ['type'])
    c.lose(owner if owner is not None else obj.get(k))   # a broken block may be re-typed, salvaged or left out
    obj[k] = junk(c)
    return True


@mut('depth 60 nesting', phase=1)
def _(c, spec):
    got = pick_block(c, spec)
    deep = [c.t(2)]
    for i in range(60):
        deep = [deep] if i % 2 else {'k': deep}
    c.lose(deep)
    if not got:
        return insert_block(c, spec, {'type': 'paragraph', 'text': deep})
    c.lose(got[2])
    key = c.rng.choice([k for k in got[2] if k != 'type'] or ['text'])
    got[2][key] = deep
    return True


@mut('NaN and infinity in a chart', p6=True)
def _(c, spec):
    return insert_block(c, spec, {'type': 'chart', 'kind': 'line', 'title': c.t(2), 'labels': [c.t(1), c.t(1), c.t(1)],
                                  'series': [{'name': c.t(1), 'values': [float('nan'), float('inf'), 4]}]})


@mut('NaN and infinity in a table')
def _(c, spec):
    return insert_block(c, spec, {'type': 'table', 'columns': [c.t(1), c.t(1)],
                                  'rows': [[float('nan'), float('-inf')], [c.t(1), 10 ** 400]]})


@mut('infinite and NaN levels', phase=3)
def _(c, spec):
    secs = dict_sections(spec)
    if not secs:
        return False
    c.rng.choice(secs)['level'] = c.rng.choice([float('inf'), float('-inf'), float('nan'), 10 ** 30])
    return True


@mut('booleans')
def _(c, spec):
    return insert_block(c, spec, c.rng.choice([
        {'type': 'table', 'columns': [c.t(1), c.t(1)], 'rows': [[True, False], [c.t(1), True]]},
        {'type': 'bullets', 'items': [c.t(2), True, False], 'ordered': 'yes'},
        {'type': 'paragraph', 'text': c.t(3), 'ordered': True}]))


@mut('booleans in a chart')
def _(c, spec):
    b = {'type': 'chart', 'kind': 'bar', 'title': c.t(2), 'labels': [c.t(1), c.t(1)],
         'series': [{'name': c.t(1), 'values': [True, 3]}]}
    c.lose(b['series'])
    return insert_block(c, spec, b)


@mut('link keys everywhere', phase=3, p6=True)
def _(c, spec):
    keys = sorted(spec_mod.FETCH_KEYS)
    placed = 0
    for obj, _owner in objects(spec, rows=False):
        if c.rng.random() < 0.5 or not placed:
            obj[c.rng.choice(keys)] = c.rng.choice(['https://example.com/a.png', '/etc/passwd', 'file:///tmp/x',
                                                    ['https://a', 'https://b'], {'href': 'https://c'}])
            placed += 1
    return placed > 0


# the repair algorithm's aliases (plan 2.3): content is converted, never lost
TYPE_ALIASES = {'bullets': ['list', 'ul', 'bullet', 'bulleted_list', 'points', 'ol', 'numbered_list', 'numbered'],
                'paragraph': ['text', 'para', 'p', 'body', 'markdown', 'description'],
                'quote': ['blockquote'], 'code': ['pre', 'snippet', 'source_code'],
                'table': ['grid', 'matrix', 'data_table'], 'chart': ['graph', 'plot']}


@mut('block type alias')
def _(c, spec):
    kind = c.rng.choice(sorted(TYPE_ALIASES))
    alias = c.rng.choice(TYPE_ALIASES[kind])
    make = {'bullets': b_bullets, 'paragraph': b_para, 'table': b_table, 'chart': lambda c: b_chart(c, 'line'),
            'quote': lambda c: {'type': 'quote', 'text': c.t(5), 'by': c.t(1)},
            'code': lambda c: {'type': 'code', 'lang': 'sh', 'text': 'echo ' + c.m()}}[kind]
    b = make(c)
    b['type'] = alias
    return insert_block(c, spec, b)


@mut('block type in another case')
def _(c, spec):
    b = c.rng.choice([b_bullets, b_para, b_table])(c)
    b['type'] = c.rng.choice([' ' + b['type'].capitalize() + ' ', b['type'].upper()])
    return insert_block(c, spec, b)


@mut('chart kind as the type')
def _(c, spec):
    b = b_chart(c, 'bar')
    b['type'] = c.rng.choice(['bar', 'line', 'pie', 'column', 'donut'])
    del b['kind']
    if b['type'] in ('pie', 'donut'):
        c.lose(b['series'][1:])   # a pie shows its first series only
    return insert_block(c, spec, b)


@mut('page break aliases')
def _(c, spec):
    return insert_block(c, spec, {'type': c.rng.choice(['hr', 'divider', 'break', 'pagebreak'])})


@mut('heading block aliases', p6=True)
def _(c, spec):
    t = c.rng.choice(['heading', 'header', 'h1', 'h2', 'h3', 'subheading', 'title'])
    return insert_block(c, spec, {'type': t, 'text': c.t(3)})


@mut('diagram by its keys')
def _(c, spec):
    return insert_block(c, spec, c.rng.choice([
        {'type': 'diagram', 'title': c.t(2), 'events': [{'date': '1973', 'label': c.t(2)},
                                                        {'date': '2007', 'label': c.t(2)}]},
        {'type': 'diagram', 'title': c.t(2), 'nodes': [{'id': 'a', 'label': c.t(1)}, {'id': 'b', 'label': c.t(1)}],
         'edges': [{'from': 'a', 'to': 'b'}]},
        {'type': 'diagram', 'title': c.t(2), 'nodes': [{'id': 'a', 'label': c.t(1), 'parent': ''},
                                                       {'id': 'b', 'label': c.t(1), 'parent': 'a'},
                                                       {'id': 'c', 'label': c.t(1), 'parent': 'a'}]}]))


@mut('missing block type')
def _(c, spec):
    got = pick_block(c, spec, {'bullets', 'table', 'chart', 'timeline', 'flow', 'tree', 'paragraph', 'code'})
    if not got:
        return False
    b = got[2]
    b.pop('type', None)
    if 'text' in b and 'lang' in b:
        b.pop('lang')   # a code block without its type reads as a paragraph (text comes first)
    return True


@mut('missing quote type')
def _(c, spec):
    b = {'text': c.t(5), 'by': c.t(1)}
    c.unsafe |= markers(b['by'])   # read as a paragraph: the author may go
    return insert_block(c, spec, b)


@mut('unknown block type kept as text', p6=True)
def _(c, spec):
    return insert_block(c, spec, {'type': c.rng.choice(['callout', 'note', 'aside', 'card']),
                                  'heading_text': c.t(3), 'body_text': c.t(4)})


FIELD_ALIASES = [
    lambda c: {'type': 'paragraph', c.rng.choice(['body', 'value', 'paragraph', 'description']): c.t(6)},
    lambda c: {'type': 'paragraph', 'text': [c.t(3), c.t(3)]},
    lambda c: {'type': 'bullets', c.rng.choice(['list', 'bullets', 'lines', 'content', 'text']): [c.t(2), c.t(2)]},
    lambda c: {'type': 'bullets', 'items': [{'label': c.t(2)}, {'title': c.t(2)}, {'other': c.t(2)}]},
    lambda c: {'type': 'table', c.rng.choice(['headers', 'header', 'cols']): [c.t(1), c.t(1)],
               c.rng.choice(['data', 'cells', 'body']): [[c.t(1), 2], [c.t(1), 3]]},
    lambda c: {'type': 'table', 'columns': ['a', 'b'], 'rows': [{'a': c.t(1), 'b': 2}, {'a': c.t(1), 'b': 3}]},
    lambda c: {'type': 'table', 'rows': [{'name': c.t(1), 'year': 1973}, {'name': c.t(1), 'year': 2007}]},
    lambda c: {'type': 'chart', 'kind': 'bar', 'title': c.t(2),
               c.rng.choice(['categories', 'x', 'xlabels']): [c.t(1), c.t(1)],
               c.rng.choice(['datasets', 'values']): [{'name': c.t(1), 'values': [1, 2]}]},
    lambda c: {'type': 'chart', 'kind': 'bar', 'title': c.t(2), 'labels': [c.t(1), c.t(1)], 'series': [3, 4]},
    lambda c: {'type': 'chart', 'kind': 'bar', 'title': c.t(2), 'labels': [c.t(1), c.t(1)],
               'series': {'name': c.t(1), 'values': [3, 4]}},
    lambda c: {'type': 'chart', 'kind': 'line', 'title': c.t(2), 'labels': [c.t(1), c.t(1)], 'series': [[3, 4]]},
    lambda c: {'type': 'chart', 'kind': 'pie', 'title': c.t(2), 'data': [{'label': c.t(1), 'value': 3},
                                                                        {'label': c.t(1), 'value': 5}]},
    lambda c: {'type': 'timeline', 'title': c.t(2), c.rng.choice(['milestones', 'entries']): [
        {'date': '1973', 'label': c.t(2)}, {'date': '1983', 'label': c.t(2)}]},
    lambda c: {'type': 'timeline', 'title': c.t(2), 'events': [
        {c.rng.choice(['year', 'when', 'time']): '1973',
         c.rng.choice(['title', 'text', 'event', 'description']): c.t(2)},
        {'year': 1983, 'title': c.t(2)}]},
    lambda c: {'type': 'flow', 'title': c.t(2), 'nodes': [{'id': 'a', 'label': c.t(1)}, {'id': 'b', 'label': c.t(1)},
                                                          {'id': 'c', 'label': c.t(1)}],
               'edges': ['a->b', 'b → c']},
    lambda c: {'type': 'flow', 'title': c.t(2), 'nodes': [{'id': 'a', 'label': c.t(1)}, {'id': 'b', 'label': c.t(1)}],
               'edges': ['a to b']},
    lambda c: {'type': 'quote', c.rng.choice(['quote', 'content']): c.t(5),
               c.rng.choice(['author', 'source']): c.t(1)},
    lambda c: {'type': 'code', c.rng.choice(['code', 'content', 'source']): 'print("' + c.m() + '")'},
]


# content that changes shape on the way: a string split into items or rows, a chart without a number, a flow chained
FIELD_CONVERSIONS = [
    lambda c: {'type': 'bullets', 'items': '- ' + c.t(2) + '\n- ' + c.t(2) + '\n1. ' + c.t(2)},
    lambda c: {'type': 'bullets', 'items': c.t(2) + '\n' + c.t(2)},
    lambda c: {'type': 'bullets', 'items': c.t(2) + '; ' + c.t(2) + '; ' + c.t(2)},
    lambda c: {'type': 'table', 'columns': [c.t(1), c.t(1)], 'rows': c.t(1) + ' | 2\n' + c.t(1) + ' | 3'},
    lambda c: {'type': 'chart', 'kind': 'bar', 'title': c.t(2), 'labels': [c.t(1), c.t(1)],
               'series': [{'values': ['n/a', 'none']}]},   # no number at all: bullets of its title and labels
    lambda c: {'type': 'flow', 'title': c.t(2), 'nodes': [c.t(1), c.t(1), c.t(1)]},
]


@mut('block field alias')
def _(c, spec):
    return insert_block(c, spec, c.rng.choice(FIELD_ALIASES)(c))


@mut('block field conversion', p6=True)
def _(c, spec):
    return insert_block(c, spec, c.rng.choice(FIELD_CONVERSIONS)(c))


@mut('block type inferred from keys')
def _(c, spec):
    return insert_block(c, spec, c.rng.choice([
        lambda: {'items': [c.t(2), c.t(2)]}, lambda: {'rows': [[c.t(1), 1]], 'columns': [c.t(1), c.t(1)]},
        lambda: {'labels': [c.t(1), c.t(1)], 'series': [{'name': c.t(1), 'values': [1, 2]}], 'title': c.t(2)},
        lambda: {'events': [{'date': '1973', 'label': c.t(2)}, {'date': '2007', 'label': c.t(2)}], 'title': c.t(2)},
        lambda: {'nodes': [{'id': 'a', 'label': c.t(1)}, {'id': 'b', 'label': c.t(1)}],
                 'edges': [{'from': 'a', 'to': 'b'}], 'title': c.t(2)},
        lambda: {'text': c.t(6)}, lambda: {'content': c.t(6)},
    ])())


# sections
@mut('section aliases', phase=3)
def _(c, spec):
    secs = section_list(spec)
    if secs is None:
        return False
    secs.insert(c.rng.randint(0, len(secs)), c.rng.choice([
        lambda: {c.rng.choice(['title', 'name']): c.t(2), 'blocks': [b_para(c)]},
        lambda: {'heading': c.t(2), c.rng.choice(['content', 'body']): [b_para(c), b_bullets(c, 2)]},
        lambda: {'heading': c.t(2), 'blocks': c.t(8)},
        lambda: {'heading': c.t(2), 'blocks': [c.t(3), c.t(3)]},
        lambda: {'heading': c.t(2), 'items': [c.t(3), c.t(3)]},
        lambda: {'heading': c.t(2), c.rng.choice(['bullets', 'points']): [c.t(3), c.t(3)]},
    ])())
    return True


@mut('section without a heading', phase=3, p6=True)
def _(c, spec):
    secs = section_list(spec)
    if not secs:
        return False
    # the first section may stay without one (the text under the title); any later one is named (S8)
    secs.insert(c.rng.randint(1, len(secs)), {'blocks': [b_para(c), b_bullets(c, 2)]})
    return True


@mut('long string section', phase=3, p6=True)
def _(c, spec):
    secs = section_list(spec)
    if secs is None:
        return False
    secs.insert(c.rng.randint(0, len(secs)), c.t(20))
    return True


# the top-level shape
@mut('top level is a list', phase=4)
def _(c, spec):
    if not isinstance(spec, dict) or not isinstance(spec.get('sections'), list):
        return False
    secs = spec['sections']
    spec.clear()
    spec['__list__'] = secs
    return True


@mut('top level alias', phase=4)
def _(c, spec):
    if not isinstance(spec, dict) or 'sections' not in spec:
        return False
    spec[c.rng.choice(['slides', 'pages', 'parts'])] = spec.pop('sections')
    return True


@mut('top level dict of sections', phase=4)
def _(c, spec):
    if not isinstance(spec, dict) or not isinstance(spec.get('sections'), list):
        return False
    spec['sections'] = {f'part_{i}': s for i, s in enumerate(spec['sections'])}
    return True


@mut('top level is a string', phase=4)
def _(c, spec):
    if not isinstance(spec, dict):
        return False
    spec.clear()
    spec['__str__'] = c.t(12)
    return True


CATALOG.sort(key=lambda m: m[1])  # stable: the catalog order within a phase stays as written
NAMES = [m[0] for m in CATALOG]


def gen_spec(rng: random.Random, first: int | None = None) -> tuple[object, dict]:
    """(spec, info): a valid base with 1 to 4 mutations. info has the base, the applied mutation names, the markers
    that must survive (P5) and whether a mutation dropped or converted content (P6). `first` forces one catalog entry,
    so a run of len(CATALOG) specs covers every mutation at least once."""
    c = Ctx(rng)
    base_name, make = rng.choice(BASES)
    spec = make(c)
    forced = [CATALOG[first]] if first is not None else []
    picks = forced + rng.sample(CATALOG, rng.randint(1 - len(forced), 3))
    picks = sorted({p[0]: p for p in picks}.values(), key=lambda m: m[1])
    for name, _phase, p6, fn in picks:
        if fn(c, spec):
            c.names.append(name)
            if p6:
                c.p6.append(name)
    if '__list__' in spec:
        spec = spec['__list__']
    elif '__str__' in spec:
        spec = spec['__str__']
    return spec, {'base': base_name, 'mutations': c.names, 'p6': c.p6, 'markers': required(spec, c)}


@functools.lru_cache(maxsize=None)
def corpus() -> list[tuple[int, str, dict]]:
    """(index, the spec as JSON, info) for this process's shard; JSON so every test gets a fresh copy."""
    out = []
    for i in selected():
        spec, info = gen_spec(random.Random(SEED + i), first=i % len(CATALOG))
        out.append((i, json.dumps(spec, ensure_ascii=False), info))
    return out


def fresh(text: str):
    return json.loads(text)


def label(i: int, info: dict, fmt: str = '') -> str:
    return f'seed {SEED + i} ({info["base"]}: {", ".join(info["mutations"]) or "none"}){" " + fmt if fmt else ""}'


def report(failures: list[str], total: int) -> str:
    head = f'{len(failures)} of {total} checks failed; first {min(12, len(failures))}:'
    return '\n'.join([head, *failures[:12]])


# ---------- the generator itself ----------

def test_catalog_covers_the_plan():
    assert len(CATALOG) == len(set(NAMES))
    assert sum(n.startswith('probe: ') for n in NAMES) == 24
    assert sum(n.startswith('odd: ') for n in NAMES) == 27
    for need in ('string series', 'number strings in a chart', 'lone surrogate', '130-character heading',
                 '5,000-character column name', 'subtitle of 256 characters', 'subtitle of 300 characters',
                 '41 sections', '31 blocks', 'page-break-only spec', 'figure-only spec',
                 'CJK, Arabic, Hindi and emoji text', 'None or a wrong type at a key', 'depth 60 nesting',
                 'NaN and infinity in a chart', 'booleans', 'link keys everywhere'):
        assert need in NAMES, need
    assert len(BASES) == 6


def test_the_generator_is_seeded_and_every_mutation_applies():
    a = [gen_spec(random.Random(SEED + i), first=i % len(CATALOG)) for i in range(len(CATALOG))]
    b = [gen_spec(random.Random(SEED + i), first=i % len(CATALOG)) for i in range(len(CATALOG))]
    dump = lambda xs: [json.dumps(s, ensure_ascii=False, sort_keys=True) + repr(sorted(i['markers'])) for s, i in xs]
    assert dump(a) == dump(b)
    applied = {n for _, info in a for n in info['mutations']}
    assert applied == set(NAMES), sorted(set(NAMES) - applied)
    words = [m for _, info in a for m in info['markers']]
    assert len(words) > 20 * len(CATALOG)


@pytest.mark.parametrize('name,make', BASES, ids=[b[0] for b in BASES])
def test_every_base_is_valid_and_clean(name, make):
    """The bases hold only valid content, so a P6 note can only come from a mutation, and every base builds."""
    c = Ctx(random.Random(SEED))
    spec = make(c)
    marks = markers(spec)
    for fmt in FORMATS:
        norm, results = cf.normalize(copy.deepcopy(spec), fmt)
        assert not [r for r in results if not r.ok and r.severity in ('fix', 'block')], (fmt, results)
        data = cf.render(norm, fmt)
        assert next(r for r in cf.verify(norm, fmt, data) if r.id == 'V1').ok
    norm, _ = cf.normalize(copy.deepcopy(spec), 'md')
    text = norm_text(norm)
    lost = {m for m in marks if m not in text} - markers([s.get('notes') for s in spec['sections']]) - \
        markers([spec.get('title'), spec.get('subtitle')])
    assert not lost, lost


# ---------- P1, P2, P4, P5, P6: normalize ----------

def test_normalize_properties():
    needs(spec_mod, 'repair', 'S')
    failures, checks = [], 0
    for i, text, info in corpus():
        has_text = oracle_has_text(fresh(text))
        for fmt in FORMATS:
            checks += 1
            where = label(i, info, fmt)
            try:
                norm, results = cf.normalize(fresh(text), fmt)
            except cf.SpecError as e:
                if e.rule_id not in REFUSALS:
                    failures.append(f'P1 {where}: SpecError {e.rule_id}: {str(e)[:120]}')
                elif has_text:
                    failures.append(f'P2 {where}: refused with {e.rule_id} although a section holds text')
                continue
            except Exception as e:  # noqa: BLE001 (the property is that nothing else escapes)
                failures.append(f'P1 {where}: {type(e).__name__}: {str(e)[:120]}')
                continue
            try:
                again, _ = cf.normalize(copy.deepcopy(norm), fmt)
                if len(again['sections']) != len(norm['sections']):
                    failures.append(f'P4 {where}: {len(norm["sections"])} sections became {len(again["sections"])}')
            except Exception as e:  # noqa: BLE001
                failures.append(f'P4 {where}: normalizing again raised {type(e).__name__}: {str(e)[:120]}')
            if info['p6'] and not any(r.severity == 'fix' and not r.ok for r in results):
                failures.append(f'P6 {where}: {", ".join(info["p6"])} changed content but no fix note was left')
            if fmt == 'md':
                got = norm_text(norm)
                lost = sorted(m for m in info['markers'] if m not in got)
                if lost:
                    failures.append(f'P5 {where}: {len(lost)} text marker(s) lost, e.g. {lost[:3]}')
    assert not failures, report(failures, checks)


def test_repair_never_raises_and_agrees_with_the_oracle():
    needs(spec_mod, 'repair', 'S')
    needs(spec_mod, 'has_text', 'S')
    failures = []
    for i, text, info in corpus():
        try:
            fixed, results = spec_mod.repair(fresh(text))
            assert isinstance(fixed, dict) and isinstance(fixed.get('sections'), list)
            assert all(isinstance(r, cf.RuleResult) for r in results)
            if oracle_has_text(fresh(text)) and not spec_mod.has_text(fresh(text)):
                failures.append(f'{label(i, info)}: has_text is False although a section holds text')
        except Exception as e:  # noqa: BLE001
            failures.append(f'{label(i, info)}: {type(e).__name__}: {str(e)[:120]}')
    for garbage in (None, 5, 'x', [], {}, [None], {'sections': 5}, float('nan'), [[[]]], {'sections': [None, 3]}):
        try:
            fixed, _ = spec_mod.repair(garbage)
            assert isinstance(fixed.get('sections'), list)
        except Exception as e:  # noqa: BLE001
            failures.append(f'garbage {garbage!r}: {type(e).__name__}: {e}')
    assert not failures, report(failures, len(corpus()))


# ---------- P3: every format renders and reopens ----------

def test_render_safe_reopens_in_every_format():
    needs(cf, 'render_safe', 'S')
    failures, checks = [], 0
    for i, text, info in corpus():
        for fmt in FORMATS:
            try:
                norm, _ = cf.normalize(fresh(text), fmt)
            except cf.SpecError:
                continue   # a refusal is P1 and P2's business
            except Exception:  # noqa: BLE001
                continue
            checks += 1
            where = label(i, info, fmt)
            try:
                data, laid = cf.render_safe(norm, fmt)
                assert isinstance(data, bytes) and data
                v1 = next(r for r in cf.verify(norm, fmt, data) if r.id == 'V1')
                if not v1.ok:
                    failures.append(f'P3 {where}: V1 {v1.note[:120]}')
                if any(r.id == 'V11' and r.severity != 'fix' for r in laid):
                    failures.append(f'P3 {where}: V11 must be a fix, got {laid}')
            except Exception as e:  # noqa: BLE001
                failures.append(f'P3 {where}: {type(e).__name__}: {str(e)[:160]}')
    assert not failures, report(failures, checks)


# ---------- a model's reply, mangled the ways replies come back ----------

def mangle(rng: random.Random, text: str) -> str:
    return rng.choice([
        lambda t: '```json\n' + t + '\n```',
        lambda t: 'Here is the file:\n' + t + '\nLet me know if you want changes.',
        lambda t: re.sub(r'\]', ',]', t, count=3),     # trailing commas
        lambda t: t[:max(1, int(len(t) * rng.uniform(0.3, 0.95)))],   # a reply cut off mid-way
        lambda t: t,
    ])(text)


def test_parse_spec_reads_mangled_replies():
    needs(spec_mod, 'parse_spec', 'S')
    failures = []
    for i, text, info in corpus():
        rng = random.Random(SEED + i)
        reply = mangle(rng, text)
        try:
            got = spec_mod.parse_spec(reply)
            assert got is None or isinstance(got, dict)
            if got is not None:
                cf.normalize(got, 'md')
        except cf.SpecError as e:
            if e.rule_id not in REFUSALS:
                failures.append(f'{label(i, info)}: SpecError {e.rule_id}: {str(e)[:100]}')
        except Exception as e:  # noqa: BLE001
            failures.append(f'{label(i, info)}: {type(e).__name__}: {str(e)[:120]}')
    with pytest.raises(cf.SpecError) as e:
        spec_mod.parse_spec('{"sections": ["' + 'x' * (2 * 1024 * 1024 + 10) + '"]}')
    assert e.value.rule_id == 'L6'
    assert spec_mod.parse_spec('no json here at all') is None
    assert not failures, report(failures, len(corpus()))


# ---------- the fixtures of section 7.3 ----------

PLAN = ROOT / 'docs' / 'PLAN-files-robust.md'
Q2750 = ('create the ppt on the how mobile phone is being evolved history past present everything a ppt of 12 slides '
         'using the multiple pictured diagrams and also use the design.md for the design')


def shape_2750() -> dict:
    return json.loads((FIXTURES / 'spec_2750_shape.json').read_text())


def test_the_design_fixture_is_the_plans_text():
    plan = PLAN.read_text() if PLAN.exists() else None
    if plan is None:
        pytest.skip('docs/PLAN-files-robust.md is not in this checkout')
    block = re.search(r'`evals/fixtures/design_system.md` \(exact text\):\n\n```markdown\n(.*?)\n```', plan, re.S)
    assert block, 'the plan no longer holds the design fixture text'
    assert (FIXTURES / 'design_system.md').read_text() == block.group(1) + '\n'


def test_the_2750_shape_fixture():
    f = shape_2750()
    assert f['query'] == Q2750 and f['format'] == 'pptx' and f['slides'] == [12, 12]
    outline = f['outline']
    assert outline['title'] == 'The Evolution of Mobile Phones'
    heads = [s['heading'] for s in outline['sections']]
    assert len(heads) == 11 and heads[0] == 'Before mobile phones' and heads[-1] == 'What comes next'
    first, second = f['batches']
    assert [s['heading'] for s in first['sections']] == heads[:6]
    assert [s['heading'] for s in second['sections']] == heads[6:]
    assert second['sections'][0] == {'heading': 'Smartphones arrive', 'level': 1,
                                     'blocks': [{'type': 'bullets', 'items': None}], 'notes': ''}
    assert f['repair']['sections'][0]['heading'] == 'Smartphones arrive'
    kinds = {b['type'] for s in [*first['sections'], *second['sections']] for b in s['blocks']}
    assert {'timeline', 'tree', 'flow'} <= kinds


def test_the_estimate_runs_fixture():
    rows = [json.loads(line) for line in (FIXTURES / 'estimate_runs.jsonl').read_text().splitlines() if line.strip()]
    assert [r['qid'] for r in rows] == [1629, 1662, 1717, 2346, 2381, 2401, 2402, 2403, 2405, 2749, 2750]
    for r in rows:
        assert {'qid', 'query', 'mode', 'engine', 'web_engine', 'planner', 'merger', 'files', 'actual'} <= set(r)
        assert r['planner'] in ('heuristic', 'llm') and r['merger'] in ('single', 'llm')
        assert set(r['actual']) == {'tokens_in', 'tokens_out', 'seconds'}
        assert all(isinstance(f['name'], str) and type(f['chars']) is int for f in r['files'])
    by = {r['qid']: r for r in rows}
    assert by[2750]['actual'] == {'tokens_in': 250_046, 'tokens_out': 30_492, 'seconds': 272.5}
    assert by[2750]['engine'] == 'agy' and by[2750]['web_engine'] == 'claude-code' and by[2750]['mode'] == 'research'
    assert by[2750]['files'] == [{'name': 'DESIGN-lovable.md', 'chars': 17_284}]
    assert by[2749]['engine'] == 'claude-code' and by[2749]['actual']['tokens_in'] == 104_690
    assert by[1717]['files'] == [{'name': 'release_notes.txt', 'chars': 774}]
    assert by[2403]['files'] == [{'name': 'meeting_notes.txt', 'chars': 604}]
    assert by[2405]['has_answer'] is True and sum(r['has_answer'] for r in rows) == 1


# ---------- the run 2750 replay (fake engine, no model) ----------

class Replay2750:
    """Answers the long writer's calls with the recorded shape of run 2750: the outline, batch 1 (parts 1 to 6), batch 2
    (parts 7 to 11, with section 7 as a bullets block without items), and for any other call (a repair of the missing
    section) the fixture's rewrite of section 7."""
    name, label, billing, supports_web = 'agy', 'Antigravity', 'subscription', False

    def __init__(self):
        self.f = shape_2750()
        self.calls = []

    def available(self):
        return True, ''

    async def stream(self, *, system, prompt, effort='medium', emit_delta=None, max_tokens=2048, web=False, schema=None,
                     exec=False, **kw):
        self.calls.append({'system': system, 'prompt': prompt, 'schema': schema})
        heads = [s['heading'] for s in self.f['outline']['sections']]
        if schema is OUTLINE_SCHEMA:
            body = self.f['outline']
        elif f'- "{heads[0]}"' in prompt:
            body = self.f['batches'][0]
        elif f'- "{heads[7]}"' in prompt:
            body = self.f['batches'][1]
        else:
            body = self.f['repair']
        return Reply(json.dumps(body), 41_000, 3_000, engine='agy')


def design_doc() -> tuple[dict, str]:
    from jevrouter.files import extract
    return extract('design_system.md', (FIXTURES / 'design_system.md').read_bytes())


async def test_the_2750_replay_builds_a_12_slide_deck():
    needs(spec_mod, 'repair', 'S')
    from pptx import Presentation
    research = ('Mobile phones began with the 1946 Mobile Telephone Service. Martin Cooper made the first handheld '
                'call on 3 April 1973. 1G arrived in 1979, GSM in 1991, 3G in 2001, the iPhone in 2007, 4G in 2009 '
                'and 5G in 2019.')
    eng = Replay2750()
    job = ca.Job(Q2750, deps=[('Research how mobile phones evolved', research)], docs=[design_doc()],
                 brief=parse_brief(Q2750))
    made = await ca.make(job, eng, None)
    assert made.ok, made.answer
    assert made.file['format'] == 'pptx' and made.file['slides'] == 12
    prs = Presentation(io.BytesIO(made.data))
    assert len(prs.slides) == 12
    notes = ' '.join([*(r['note'] for r in made.file['rules'] if r['id'] == 'S1' and not r['ok']), *made.caveats])
    assert re.search(r'\b7\b', notes), (made.file['rules'], made.caveats)
    assert made.llm_in > 0 and made.file['tokens'] == made.llm_in + made.llm_out
    writer = [c for c in eng.calls if not c['prompt'].startswith('Plan the layout')]   # Studio's art direction
    assert len(writer) <= 4   # outline, two batches and at most one call that rewrites section 7
    assert len(eng.calls) - len(writer) <= 1   # and at most one design call
    text = ' '.join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
    assert 'Smartphones arrive' in text and 'Before mobile phones' in text and 'What comes next' in text
    try:
        from jevrouter.create import design  # noqa: F401 (builder A's module)
    except ImportError:
        return
    assert not any('Color Palette' in c['prompt'] or '#f7f4ed' in c['prompt'] for c in eng.calls)
    assert (made.file.get('design') or {}).get('colors', {}).get('bg', '').upper() == 'F7F4ED', made.file.get('design')
    for slide in prs.slides:
        fill = slide.background.fill
        assert str(fill.fore_color.rgb).upper() == 'F7F4ED'


# ---------- the eval file checks of section 7.2 ----------

def md_file(fid: str, **meta) -> dict:
    return {'id': fid, 'name': f'{fid}.md', 'format': 'md', 'size': 1, 'created': 0, 'qid': 1, 'title': fid,
            'tokens': 0, 'source': 'llm', 'rules': [], **meta}


def scored(expect: dict, f: dict) -> list[str]:
    rec = {'qid': 1, 'status': 'done', 'merged': {'answer': 'Made it.', 'engine': 'single'}, 'total_ms': 5,
           'tasks': [{'agent': 'create', 'ok': True, 'created_files': [f]}]}
    return evals.check_file(expect, rec, lambda fid: b'# Tea\n\nGreen and black tea.\n')


def test_eval_keys_design_bg_and_partial_ok():
    assert {'design_bg', 'partial_ok'} <= evals.FILE_KEYS
    ok = {'id': 'x', 'query': 'q', 'expect_file': {'format': 'md', 'design_bg': 'F7F4ED', 'partial_ok': False}}
    assert evals.validate_case(ok, strict=True) == []
    bad = {'id': 'y', 'query': 'q', 'expect_file': {'format': 'md', 'design_bg': 'cream', 'partial_ok': 'no'}}
    errors = evals.validate_case(bad)
    assert any('design_bg' in e for e in errors) and any('partial_ok' in e for e in errors)
    styled = md_file('a', design={'name': 'design_system.md', 'colors': {'bg': 'f7f4ed', 'text': '1C1C1C'}})
    assert scored({'format': 'md', 'design_bg': '#F7F4ED'}, styled) == []
    assert scored({'format': 'md', 'design_bg': 'FFFFFF'}, styled) == [
        'a.md has background F7F4ED from its design, expected FFFFFF']
    assert scored({'format': 'md', 'design_bg': 'F7F4ED'}, md_file('b')) == [
        'b.md has no design applied (expected background F7F4ED)']
    partial = md_file('c', partial={'planned': 11, 'written': 10, 'missing': ['Smartphones arrive'], 'resume': None})
    assert scored({'format': 'md', 'partial_ok': False}, partial) == [
        'c.md is partial: 10 of 11 written (missing Smartphones arrive)']
    assert scored({'format': 'md', 'partial_ok': True}, partial) == []
    assert scored({'format': 'md', 'partial_ok': False}, md_file('d', partial=None)) == []


def test_the_new_eval_cases():
    cases = {c['id']: c for c in evals.load_cases(evals.CASES)}
    new = ['cr-2750-deck', 'cr-design-pdf', 'cr-design-restyle', 'cr-pictured', 'cr-12-slides']
    for cid in new:
        assert evals.validate_case(cases[cid], strict=True) == [], cid
    deck = cases['cr-2750-deck']
    assert deck['query'] == Q2750 and deck['files'] == ['design_system.md'] and deck['chat']['mode'] == 'research'
    assert deck['expect_agents'] == ['research', 'create']
    assert {'create', 'pptx', 'design', 'long'} <= set(deck['tags'])
    assert deck['expect_file'] == {'format': 'pptx', 'slides': [12, 12], 'diagrams_min': 2, 'design_bg': 'F7F4ED'}
    restyle = cases['cr-design-restyle']['turns'][-1]['expect_file']
    assert restyle['source'] == 'convert' and restyle['design_bg'] == 'F7F4ED'
    assert cases['cr-pictured']['expect_file']['images_min'] == 1
    assert cases['cr-12-slides']['expect_file']['slides'] == [12, 12]
    assert all(cases[c].get('run_on') == ['cli', 'api'] for c in new)   # they need an engine; route replay skips them
