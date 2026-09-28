"""The ruleset: RULES matches docs/RULES-files.md, and verify() catches files that break V1-V4, L4, X1-X3, A3, A4."""
import io
import re
import zipfile
from pathlib import Path

import pytest

from jevrouter.create import FORMATS, RULES, RuleResult, SpecError, preview, render, verify
from jevrouter.create.themes import THEMES, contrast, get, worst_contrast

SPEC = {'title': 'Quarterly report', 'subtitle': 'Q3', 'sections': [
    {'heading': 'Summary', 'blocks': [{'type': 'paragraph', 'text': 'Revenue grew.'},
                                      {'type': 'bullets', 'items': ['North up', 'South flat']}]},
    {'heading': 'Numbers', 'blocks': [
        {'type': 'table', 'columns': ['Region', 'Revenue'], 'rows': [['North', 120], ['South', 80]]},
        {'type': 'chart', 'kind': 'bar', 'title': 'Revenue by region', 'labels': ['North', 'South'],
         'series': [{'name': 'Revenue', 'values': [120, 80]}]}]}]}


def by_id(results) -> dict:
    return {r.id: r for r in results}


def rewrite_zip(data: bytes, change) -> bytes:
    """A copy of an OOXML package with change(name, bytes) -> bytes | None applied (None drops the entry), plus any
    extra entries change.extra holds."""
    src, out = zipfile.ZipFile(io.BytesIO(data)), io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        for name in src.namelist():
            body = change(name, src.read(name))
            if body is not None:
                z.writestr(name, body)
        for name, body in getattr(change, 'extra', {}).items():
            z.writestr(name, body)
    return out.getvalue()


# ---------- RULES ----------

def test_rules_match_the_ruleset_document():
    doc = (Path(__file__).resolve().parent.parent / 'docs' / 'RULES-files.md').read_text()
    rows = re.findall(r'^\| ([SLFXVA]\d+) \|.*\| (block|fix|warn|—) \| (yes|no) \|$', doc, re.M)
    assert rows and [r['id'] for r in RULES] == [r[0] for r in rows]
    for rule, (rid, sev, enforced) in zip(RULES, rows):
        assert rule['severity'] == (None if sev == '—' else sev), rid
        assert rule['enforced'] == (enforced == 'yes')


def test_rules_shape_and_copy():
    groups = {'Content spec', 'Size limits', 'Structure and style', 'Safety', 'Verification', 'Accessibility'}
    assert len({r['id'] for r in RULES}) == len(RULES)
    for r in RULES:
        assert set(r) == {'id', 'group', 'text', 'severity', 'enforced'}
        assert r['group'] in groups and r['severity'] in ('block', 'fix', 'warn', None)
        assert r['text'] and '—' not in r['text'] and '–' not in r['text'], r['id']


def test_rule_result_is_plain_data():
    r = RuleResult('V2', 'warn', False, 'missing')
    assert r.to_dict() == {'id': 'V2', 'severity': 'warn', 'ok': False, 'note': 'missing'}


# ---------- verify on good files ----------

EXPECTED = {
    'pdf': {'V1', 'V2', 'V3', 'V4', 'L4', 'F1', 'F6', 'X1', 'X3', 'A1', 'A2', 'A3', 'A4'},
    'docx': {'V1', 'V2', 'V3', 'V4', 'L4', 'F2', 'F6', 'X1', 'X3', 'A1', 'A2', 'A3', 'A4'},
    'pptx': {'V1', 'V2', 'V3', 'V4', 'L4', 'F3', 'F6', 'X1', 'X2', 'X3', 'A1', 'A2', 'A3', 'A4'},
    'xlsx': {'V1', 'V2', 'V3', 'V4', 'L4', 'F4', 'X1', 'X2', 'X3', 'A2', 'A3', 'A4'},
    'md': {'V1', 'V2', 'V3', 'V4', 'L4', 'F5', 'X1', 'X3', 'A1', 'A3', 'A4'},
}


@pytest.mark.parametrize('fmt', FORMATS)
def test_verify_good_file(fmt):
    res = verify(SPEC, fmt, render(SPEC, fmt))
    assert {r.id for r in res} == EXPECTED[fmt]
    assert all(r.ok for r in res), [(r.id, r.note) for r in res if not r.ok]
    rules = {r['id']: r['severity'] for r in RULES}
    assert all(r.severity == rules[r.id] for r in res)


# ---------- verify catches broken files ----------

@pytest.mark.parametrize('fmt', FORMATS)
def test_v1_file_that_does_not_reopen_blocks(fmt):
    garbage = b'\xff\xfe not a document' if fmt == 'md' else b'PK\x03\x04 broken'
    with pytest.raises(SpecError) as e:
        verify(SPEC, fmt, garbage if fmt != 'pdf' else b'%PDF-1.4 truncated')
    assert e.value.rule_id == 'V1'


def test_l4_oversized_file_blocks():
    with pytest.raises(SpecError) as e:
        verify(SPEC, 'md', b'x' * (15 * 1024 * 1024 + 1))
    assert e.value.rule_id == 'L4'


@pytest.mark.parametrize('fmt', FORMATS)
def test_v2_v3_wrong_spec_warns(fmt):
    other = {**SPEC, 'title': 'Different title', 'sections': SPEC['sections'] + [
        {'heading': 'Missing section', 'blocks': [{'type': 'paragraph', 'text': 'x'},
                                                  {'type': 'table', 'columns': ['a'], 'rows': [[1]]}]}]}
    res = by_id(verify(other, fmt, render(SPEC, fmt)))
    assert not res['V2'].ok and 'Different title' in res['V2'].note and res['V2'].severity == 'warn'
    assert not res['V3'].ok


@pytest.mark.parametrize('fmt', ['pptx', 'xlsx', 'docx', 'md', 'pdf'])
def test_v4_missing_chart_warns(fmt):
    without = {**SPEC, 'sections': [SPEC['sections'][0], {**SPEC['sections'][1],
                                                          'blocks': SPEC['sections'][1]['blocks'][:1]}]}
    extra = {**SPEC, 'sections': SPEC['sections'] + [{'heading': 'More', 'blocks': [
        {'type': 'chart', 'kind': 'pie', 'title': 'Share of profit', 'labels': ['a', 'b'],
         'series': [{'name': 's', 'values': [1, 2]}]}]}]}
    res = by_id(verify(extra, fmt, render(without, fmt)))
    assert not res['V4'].ok


def test_x1_macro_in_package_blocks():
    data = render(SPEC, 'xlsx')

    def keep(name, body):
        return body
    keep.extra = {'xl/vbaProject.bin': b'\x00macro'}
    with pytest.raises(SpecError) as e:
        verify(SPEC, 'xlsx', rewrite_zip(data, keep))
    assert e.value.rule_id == 'X1'

    def macro_type(name, body):
        return body.replace(b'sheet.main+xml', b'sheet.macroEnabled.main+xml') if name == '[Content_Types].xml' else body
    with pytest.raises(SpecError):
        verify(SPEC, 'xlsx', rewrite_zip(data, macro_type))


def test_x1_script_in_pdf_blocks():
    data = render(SPEC, 'pdf').replace(b'/PageMode /UseOutlines', b'/PageMode /UseOutlines /JavaScript 1 0 R', 1)
    with pytest.raises(SpecError) as e:
        verify(SPEC, 'pdf', data)
    assert e.value.rule_id == 'X1'


def test_x2_formula_left_in_a_sheet_is_reported():
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(render(SPEC, 'xlsx')))
    ws = wb['Numbers']
    ws['A2'] = "=cmd|' /C calc'!A0"
    ws['B3'] = '=SUM(B2:B2)'
    ws['A3'] = '@SUM(1)'
    buf = io.BytesIO()
    wb.save(buf)
    res = by_id(verify(SPEC, 'xlsx', buf.getvalue()))
    assert not res['X2'].ok and 'A2' in res['X2'].note and 'A3' in res['X2'].note and 'B3' not in res['X2'].note


def test_x3_external_reference_is_reported():
    data = render(SPEC, 'docx')

    def external(name, body):
        if name == 'word/_rels/document.xml.rels':
            return body.replace(b'</Relationships>', b'<Relationship Id="rIdX" Type="http://schemas.openxmlformats.org/'
                                b'officeDocument/2006/relationships/image" Target="http://evil.example/p.png" '
                                b'TargetMode="External"/></Relationships>')
        return body
    res = by_id(verify(SPEC, 'docx', rewrite_zip(data, external)))
    assert not res['X3'].ok and res['X3'].severity == 'fix'
    md = render(SPEC, 'md') + b'\n![x](https://evil.example/p.png)\n'
    assert not by_id(verify(SPEC, 'md', md))['X3'].ok


def test_f5_markdown_with_two_titles_or_skipped_levels():
    md = b'# Quarterly report\n\n#### Summary\n\n# Numbers\n'
    res = by_id(verify(SPEC, 'md', md))
    assert not res['F5'].ok


def test_f4_unfrozen_header_is_reported():
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(render(SPEC, 'xlsx')))
    wb['Numbers'].freeze_panes = None
    buf = io.BytesIO()
    wb.save(buf)
    assert not by_id(verify(SPEC, 'xlsx', buf.getvalue()))['F4'].ok


# ---------- A3, A4 ----------

def test_a3_chart_summary_next_to_every_chart():
    for fmt in FORMATS:
        res = by_id(verify(SPEC, fmt, render(SPEC, fmt)))
        assert res['A3'].ok, fmt
    md = render(SPEC, 'md').replace(b'*Highest North (120), lowest South (80).*', b'')
    assert not by_id(verify(SPEC, 'md', md))['A3'].ok


def test_a4_theme_contrast():
    assert round(contrast('000000', 'FFFFFF'), 1) == 21.0 and contrast('777777', '777777') == 1.0
    for name in THEMES:
        ratio, pair = worst_contrast(name)
        assert ratio >= 4.5, (name, pair, ratio)
    assert get('nope') is THEMES['clean'] and get('dark', paper=True)['bg'] == 'FFFFFF'
    assert get('clean', paper=True) is THEMES['clean']


# ---------- the brief rules (docs/PLAN-accuracy-v2.md C6, C7) ----------

def test_brief_rules_are_listed_and_only_run_with_a_brief():
    from jevrouter.create.brief import parse_brief
    ids = [r['id'] for r in RULES]
    assert ids.index('V5') == ids.index('V4') + 1 and ids[ids.index('V5'):ids.index('V9') + 1] == \
        ['V5', 'V6', 'V7', 'V8', 'V9']
    assert ids.index('X6') == ids.index('X5') + 1 and by_id([RuleResult(**{k: r[k] for k in ('id', 'severity')},
                                                                        ok=True, note='') for r in RULES])['X6']
    assert next(r for r in RULES if r['id'] == 'X3')['text'].startswith('No remote references; images only as '
                                                                        'embedded bytes from the asset cache.')
    data = render(SPEC, 'pdf')
    assert not {'V5', 'V6', 'V7', 'V8', 'V9', 'X6'} & set(by_id(verify(SPEC, 'pdf', data)))
    res = by_id(verify(SPEC, 'pdf', data, brief=parse_brief('a 12-13 page PDF')))
    assert {'V5', 'X6'} <= set(res) and not {'V6', 'V7', 'V8', 'V9'} & set(res)


def test_v5_pages_and_slides():
    from jevrouter.create.brief import parse_brief
    res = by_id(verify(SPEC, 'pdf', render(SPEC, 'pdf'), brief=parse_brief('12-13 pages')))
    assert not res['V5'].ok and res['V5'].note == 'asked for 12-13 pages, made 1'
    assert by_id(verify(SPEC, 'pdf', render(SPEC, 'pdf'), brief=parse_brief('a one page pdf')))['V5'].note == \
        '1 pages, asked for 1'
    n = preview('pptx', render(SPEC, 'pptx'))['slides']
    res = by_id(verify(SPEC, 'pptx', render(SPEC, 'pptx'), brief=parse_brief(f'a {n} slide deck')))
    assert res['V5'].ok and res['V5'].note == f'{n} slides, asked for {n}'
    res = by_id(verify(SPEC, 'pptx', render(SPEC, 'pptx'), brief=parse_brief('a 10-12 slide deck')))
    assert not res['V5'].ok and res['V5'].note == f'asked for 10-12 slides, made {n}'
    res = by_id(verify(SPEC, 'docx', render(SPEC, 'docx'), brief=parse_brief('10 pages')))
    assert not res['V5'].ok and res['V5'].note.startswith('asked for 10 pages, made about 1 (')


def test_v8_theme_and_the_grey_scan():
    from jevrouter.create.brief import parse_brief
    from jevrouter.create.rules import grey_scan
    assert worst_contrast('mono')[0] >= 7
    mono = {**SPEC, 'theme': 'mono'}
    ok, note = grey_scan(render(mono, 'pdf'))
    assert ok and note.startswith('only greys drawn')
    ok, note = grey_scan(render(SPEC, 'pdf'))
    assert not ok and 'coloured fills or strokes' in note
    bw = parse_brief('in black and white')
    assert by_id(verify(mono, 'pdf', render(mono, 'pdf'), brief=bw))['V8'].ok
    res = by_id(verify(SPEC, 'pdf', render(SPEC, 'pdf'), brief=bw))['V8']
    assert not res.ok and res.note == 'theme clean, asked for mono'
    assert by_id(verify(mono, 'docx', render(mono, 'docx'), brief=bw))['V8'].ok


def test_v6_images_asked_for_but_none():
    from jevrouter.create.brief import parse_brief
    res = by_id(verify(SPEC, 'pdf', render(SPEC, 'pdf'), brief=parse_brief('with pictures')))
    assert not res['V6'].ok and res['V6'].note == 'no images embedded (asked for images)'
    res = by_id(verify(SPEC, 'md', render(SPEC, 'md'), brief=parse_brief('with pictures')))
    assert res['V6'].note == "no images embedded; a MD file can't hold them"
