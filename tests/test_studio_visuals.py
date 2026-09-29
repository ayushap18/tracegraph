"""Studio visuals (docs/PLAN-designer.md 3.4, builder V): the bundled Lucide icons, seeded decorative art, smart crop,
and the ten new native diagram kinds with their DocSpec repair. Offline: no network, a temporary asset cache."""
import hashlib
import io
import json
import re
import zipfile

import pytest
from PIL import Image, ImageDraw

from jevrouter.create import FORMATS, diagram, normalize, render, verify
from jevrouter.create import spec as spec_mod
from jevrouter.create.brief import parse_brief, studio_kinds
from jevrouter.create.rules import grey_scan
from jevrouter.studio import art, assets, icons, plan, tokens
from jevrouter.studio.assets import pixels


# ---------- helpers ----------

def ds_for(preset: str = 'minimal', dark: bool = False) -> tokens.DesignSystem:
    colors = {r: '888888' for r in tokens.COLOR_ROLES}
    colors.update(bg='111827' if dark else 'FFFFFF', text='F3F4F6' if dark else '1F2328', accent='2563EB',
                  accent2='DB2777', border='D0D7DE', surface='F3F6FA')
    return tokens.DesignSystem(id=preset, name=preset, dark=dark, colors=colors,
                               chart_palette=['2563EB', 'F59E0B', '059669', 'DC2626', '7C3AED', '0891B2'])


def png_bytes(img: Image.Image) -> bytes:
    buf = io.BytesIO()
    img.save(buf, 'PNG')
    return buf.getvalue()


def subject_image(cx: float, cy: float, size=(1200, 600)) -> bytes:
    """A flat sky with a busy subject around (cx, cy) fractions."""
    img = Image.new('RGB', size, (200, 220, 240))
    d = ImageDraw.Draw(img)
    x, y = cx * size[0], cy * size[1]
    for i in range(0, 120, 6):
        d.line([(x - 60 + i, y - 60), (x + 60 - i, y + 60)], fill=(10, 10, 10), width=3)
    d.ellipse([x - 70, y - 70, x + 70, y + 70], outline=(0, 0, 0), width=5)
    return png_bytes(img)


@pytest.fixture
def cache(monkeypatch, tmp_path):
    from jevrouter.create import assets as cache_mod
    monkeypatch.setattr(cache_mod, 'CACHE', tmp_path / 'assets')
    (tmp_path / 'assets').mkdir()

    def put(data: bytes) -> str:
        sha = hashlib.sha256(data).hexdigest()
        (tmp_path / 'assets' / f'{sha}.png').write_bytes(data)
        return sha
    return put


class FakeWorkspace:
    def __init__(self):
        self.files = {}

    def put(self, kind, name, data):
        ref = f'ws:{kind}/{name}'
        self.files[ref] = data
        return ref

    def get(self, ref):
        return self.files.get(ref)


# ---------- icons ----------

def test_bundled_icons_are_a_curated_lucide_set_with_its_licence():
    raw = json.loads(icons.ICONS_JSON.read_text())
    assert raw['source'] == 'lucide' and raw['license'] == 'ISC' and raw['viewbox'] == 24 and raw['version']
    names = icons.names()
    assert 140 <= len(names) <= 180 and names == sorted(names)
    for want in ('atom', 'dna', 'battery', 'orbit', 'landmark', 'sigma', 'globe', 'cpu', 'arrow-right', 'check'):
        assert want in names
    lic = icons.LICENSE.read_text()
    assert 'ISC License' in lic and 'Lucide' in lic and 'Feather' in lic
    for n in names:
        ic = icons.get(n)
        assert ic.nodes and all(k in icons.NODE_TYPES for k, _ in ic.nodes)


def test_every_icon_flattens_inside_its_square():
    for n in icons.names():
        lines = icons.polylines(icons.get(n), 48)
        assert lines, n
        pts = [p for ln in lines for p in ln]
        assert all(-1 <= x <= 49 and -1 <= y <= 49 for x, y in pts), n
        assert all(len(ln) >= 2 for ln in lines)


def test_path_parser_curves_arcs_and_closed_shapes():
    circle = icons.path_polylines('M12 2a10 10 0 1 0 0 20a10 10 0 1 0 0-20Z', 0.05)
    pts = circle[0]
    assert pts[0] == pts[-1]
    assert all(abs(((x - 12) ** 2 + (y - 12) ** 2) ** 0.5 - 10) < 0.1 for x, y in pts)
    packed = icons.path_polylines('M2 12a10 10 0 011 1', 0.1)     # arc flags written without spaces
    assert packed and abs(packed[0][-1][0] - 3) < 1e-6 and abs(packed[0][-1][1] - 13) < 1e-6
    rect = icons.node_polylines('rect', {'x': '2', 'y': '2', 'width': '20', 'height': '10', 'rx': '2'}, 0.1)[0]
    assert rect[0] == rect[-1] and min(x for x, _ in rect) == pytest.approx(2)
    # a finer tolerance gives more points
    atom = icons.get('atom')
    assert sum(map(len, icons.polylines(atom, 24, tolerance=0.05))) > sum(map(len, icons.polylines(atom, 24)))


def _reference_lines(icon) -> list:
    """The icon flattened by fontTools' own SVG path parser (arcs become cubics), independent of icons.path_polylines;
    the simple shapes (circle, rect, line, polyline) use icons.node_polylines."""
    from fontTools.pens.basePen import BasePen
    from fontTools.svgLib.path import parse_path

    class Flat(BasePen):
        def __init__(self):
            super().__init__(None)
            self.lines, self.cur = [], None

        def _moveTo(self, p):
            self.cur = [p]
            self.lines.append(self.cur)

        def _lineTo(self, p):
            self.cur.append(p)

        def _curveToOne(self, p1, p2, p3):
            p0 = self.cur[-1]
            for i in range(1, 17):
                t = i / 16
                a, b, c, d = (1 - t) ** 3, 3 * t * (1 - t) ** 2, 3 * t * t * (1 - t), t ** 3
                self.cur.append((a * p0[0] + b * p1[0] + c * p2[0] + d * p3[0],
                                 a * p0[1] + b * p1[1] + c * p2[1] + d * p3[1]))

        def _qCurveToOne(self, p1, p2):
            p0 = self.cur[-1]
            for i in range(1, 17):
                t = i / 16
                a, b, c = (1 - t) ** 2, 2 * t * (1 - t), t * t
                self.cur.append((a * p0[0] + b * p1[0] + c * p2[0], a * p0[1] + b * p1[1] + c * p2[1]))

        def _closePath(self):
            if self.cur:
                self.cur.append(self.cur[0])
    out = []
    for kind, attrs in icon.nodes:
        if kind == 'path':
            pen = Flat()
            parse_path(attrs['d'], pen)
            out += [ln for ln in pen.lines if len(ln) > 1]
        else:
            out += icons.node_polylines(kind, attrs, 0.05)
    return out


def _ink(img) -> tuple:
    """(ink pixel count, bbox) of an L or RGBA image (alpha for RGBA)."""
    band = img.getchannel('A') if img.mode == 'RGBA' else img
    mask = band.point(lambda v: 255 if v >= 96 else 0)
    return mask.histogram()[255], mask.getbbox()


@pytest.mark.parametrize('name', icons.names())
def test_every_icon_draws_all_of_itself(name):
    """icons.png (built on polylines) draws the whole icon: its ink spans the SVG's bounding box and has about as much
    ink as the same SVG flattened by fontTools. A path parser that stops early or misreads packed arc flags ("0010-10"
    is flags 0, 0 and x 10) drew only part of leaf and book-open."""
    px = 96
    icon = icons.get(name)
    ref = _reference_lines(icon)
    xs, ys = [x for ln in ref for x, _ in ln], [y for ln in ref for _, y in ln]
    s = px / icons.VIEWBOX
    ref_img = Image.new('L', (px, px), 0)
    d = ImageDraw.Draw(ref_img)
    for ln in ref:
        d.line([(x * s, y * s) for x, y in ln], fill=255, width=int(icons.STROKE * s), joint='curve')
    want, _ = _ink(ref_img)
    got, box = _ink(Image.open(io.BytesIO(icons.png(name, '000000', px))))
    assert box is not None, name
    # the drawn ink covers the SVG's bounding box (less a stroke's width of slack on each side)
    slack = icons.STROKE * s
    assert box[0] <= min(xs) * s + slack and box[2] >= max(xs) * s - slack, (name, box, (min(xs) * s, max(xs) * s))
    assert box[1] <= min(ys) * s + slack and box[3] >= max(ys) * s - slack, (name, box, (min(ys) * s, max(ys) * s))
    # and is not mostly empty inside it
    assert got >= 0.8 * want, (name, got, want)
    # the polylines themselves span the same box as fontTools' reading of the SVG
    pl = icons.polylines(icon, icons.VIEWBOX)
    gx, gy = [x for ln in pl for x, _ in ln], [y for ln in pl for _, y in ln]
    assert max(abs(min(gx) - min(xs)), abs(max(gx) - max(xs)), abs(min(gy) - min(ys)), abs(max(gy) - max(ys))) < 0.3


def test_packed_arc_flags_touching_the_next_number():
    # "a10 10 0 0010-10": flags 0 and 0, then x 10 (the tokenizer reads "0010" as one number)
    got = icons.path_polylines('M11 20a10 10 0 0010-10', 0.05)[0]
    assert got[-1] == pytest.approx((21, 10))
    # "a5 5 0 012.9-4.5": flags 0 and 1, then x 2.9
    got = icons.path_polylines('M2 21a5 5 0 012.9-4.5', 0.05)[0]
    assert got[-1] == pytest.approx((4.9, 16.5))
    # "a1 1 0 01.5.5": flags 0 and 1, then .5 .5
    got = icons.path_polylines('M0 0a1 1 0 01.5.5', 0.05)[0]
    assert got[-1] == pytest.approx((0.5, 0.5))


@pytest.mark.parametrize('text,want', [('battery', 'battery'), ('planet', 'orbit'), ('dna', 'dna'),
                                       ('The water cycle', 'droplet'), ('History of Rome', 'landmark'),
                                       ('Growth in 2020', 'trending-up'), ('photosynthesis', 'leaf'),
                                       ('rocket', 'rocket'), ('arrow-right', 'arrow-right')])
def test_pick_by_keyword(text, want):
    assert icons.pick(text) == want


def test_pick_is_deterministic_skips_used_and_admits_no_match():
    assert icons.pick('battery') == icons.pick('battery')
    other = icons.pick('battery', used={'battery'})
    assert other and other != 'battery'
    assert icons.pick('xyzzy qwerty') is None and icons.pick('') is None


def test_icon_png_is_a_transparent_stroked_picture():
    img = Image.open(io.BytesIO(icons.png('atom', '1D4ED8', 64)))
    assert img.size == (64, 64) and img.mode == 'RGBA'
    alphas = [a for *_, a in pixels(img)]
    assert alphas[0] == 0 and max(alphas) == 255
    drawn = [p for p in pixels(img) if p[3] > 200]
    assert drawn and all(abs(p[2] - 0xD8) < 40 for p in drawn)
    assert Image.open(io.BytesIO(icons.png('no-such-icon', '000000', 32))).getextrema()[3] == (0, 0)


# ---------- art ----------

@pytest.mark.parametrize('kind', art.ART_KINDS)
def test_art_is_seeded_bounded_and_uses_token_colours(kind):
    ds = ds_for('bold-dark', dark=True)
    a = art.draw(kind, 7, 960, 540, ds)
    assert a == art.draw(kind, 7, 960, 540, ds) and 0 < len(a) <= art.MAX_PRIMITIVES
    assert art.draw(kind, 8, 960, 540, ds) != a or kind == 'grid'
    for p in a:
        assert p['type'] in ('rect', 'ellipse', 'line', 'path', 'gradient')
        for key in ('fill', 'stroke', 'from', 'to'):
            if p.get(key) is not None:
                assert p[key] in plan.TOKEN_COLORS, (kind, key, p[key])
        if p['type'] in ('rect', 'ellipse'):
            assert 0 <= p.get('opacity', 1) <= 1
    img = Image.open(io.BytesIO(art.png(kind, 7, 480, 270, ds)))
    assert img.size == (480, 270) and len(set(pixels(img))) > 1


def test_art_png_matches_across_sizes_and_mono_has_no_gradient():
    ds = ds_for()
    small = Image.open(io.BytesIO(art.png('blobs', 3, 240, 135, ds)))
    big = Image.open(io.BytesIO(art.png('blobs', 3, 960, 540, ds))).resize((240, 135))
    diff = sum(abs(a - b) for pa, pb in zip(pixels(small), pixels(big)) for a, b in zip(pa, pb))
    assert diff / (240 * 135 * 3) < 12   # the same picture at two resolutions
    mono = ds_for('mono')
    assert not any(p['type'] == 'gradient' for p in art.draw('gradient-mesh', 1, 960, 540, mono))
    assert art.draw('nope', 1, 10, 10, ds) == art.draw('dots', 1, 10, 10, ds) and art.draw('dots', 1, 0, 10, ds) == []


def test_art_pick_follows_preset_and_mood():
    assert art.pick(ds_for('bold-dark', True), ['space'], 1) == 'orbit-rings'
    assert art.pick(ds_for('pastel'), ['calm'], 1) == 'waves'
    assert art.pick(ds_for('vibrant'), [], 5) in art.PRESET_KINDS['vibrant']
    assert art.pick(ds_for('mono'), ['sunset'], 2) in art.PRESET_KINDS['mono']
    assert art.pick(ds_for('academic'), [], 9) == art.pick(ds_for('academic'), [], 9)


# ---------- smart crop ----------

def test_focal_point_finds_the_subject_and_centres_flat_pictures():
    fx, fy = assets.focal_point(subject_image(0.8, 0.55))
    assert abs(fx - 0.8) < 0.08 and abs(fy - 0.55) < 0.12
    fx, fy = assets.focal_point(subject_image(0.2, 0.3))
    assert abs(fx - 0.2) < 0.08 and abs(fy - 0.3) < 0.12
    assert assets.focal_point(png_bytes(Image.new('RGB', (300, 200), (40, 40, 40)))) == (0.5, 0.5)
    assert assets.focal_point(b'not an image') == (0.5, 0.5)


def test_cover_crops_keep_the_focal_point_on_a_third():
    size = (1200, 600)
    for focal in ((0.8, 0.5), (0.1, 0.5), (0.5, 0.5), (0.97, 0.9)):
        c = assets.crop_for(size, focal, 300, 540)      # a tall slot from a wide picture
        x, y, w, h = c
        assert assets.focal_in(c, focal) and 0 <= x <= 1 - w + 1e-9 and h == 1
        assert abs((w * size[0]) / (h * size[1]) - 300 / 540) < 0.01
    x, y, w, h = assets.crop_for(size, (0.8, 0.5), 300, 540)
    assert abs((0.8 - x) / w - 2 / 3) < 0.01       # an off-centre subject sits on the right third
    assert assets.crop_for(size, (0.8, 0.5), 300, 540, fit='contain') == (0, 0, 1, 1)
    assert assets.crop_for(size, (0.5, 0.5), 0, 540) == (0, 0, 1, 1)
    tall = assets.crop_for((600, 1200), (0.5, 0.2), 960, 540)
    assert tall[2] == 1 and assets.focal_in(tall, (0.5, 0.2))


def test_upscale_against_the_minimum_resolution():
    # a full-bleed slide (960 x 540 pt) at 96 dpi needs 1280 x 720 px
    assert assets.upscale((1280, 720), (0, 0, 1, 1), 960, 540, 'pptx') == pytest.approx(1.0)
    assert assets.upscale((640, 360), (0, 0, 1, 1), 960, 540, 'pptx') == pytest.approx(2.0)
    assert assets.upscale((640, 360), (0, 0, 1, 1), 960, 540, 'pptx') > assets.MAX_UPSCALE
    assert assets.upscale((1280, 720), (0, 0, 0.5, 1), 480, 540, 'pdf') > assets.upscale(
        (1280, 720), (0, 0, 0.5, 1), 480, 540, 'pptx')


def test_prepare_stores_a_crop_in_the_workspace(cache):
    sha = cache(subject_image(0.75, 0.5))
    ws = FakeWorkspace()
    ref, crop, focal = assets.prepare(sha, 300, 540, ws, fmt='pptx')
    assert ref.startswith('ws:images/') and assets.focal_in(crop, focal)
    img = Image.open(io.BytesIO(ws.get(ref)))
    assert abs(img.width / img.height - 300 / 540) < 0.02
    assert img.width <= round(300 / 72 * 96 * assets.KEEP_DPI)
    again = assets.prepare(f'asset:{sha}', 300, 540, ws, fmt='pptx')
    assert again[0] == ref and len(ws.files) == 1                 # reused, not re-encoded
    grey, _, _ = assets.prepare(sha, 300, 540, ws, fmt='pptx', mono=True)
    assert all(r == g == b for r, g, b in pixels(Image.open(io.BytesIO(ws.get(grey))).convert('RGB')))
    duo, _, _ = assets.prepare(sha, 300, 540, ws, fmt='pptx', duotone=('1E3A8A', 'FDE68A'))
    assert duo not in (ref, grey)
    with pytest.raises(FileNotFoundError):
        assets.prepare('0' * 64, 100, 100, ws, fmt='pdf')


# ---------- the new diagram kinds ----------

BLOCKS = {
    'cycle': {'type': 'cycle', 'title': 'The water cycle',
              'steps': ['Evaporation', 'Condensation', 'Precipitation', 'Collection']},
    'venn': {'type': 'venn', 'title': 'Plants and animals', 'shared': ['Cells', 'DNA'],
             'sets': [{'label': 'Plants', 'items': ['Photosynthesis', 'Cell walls']},
                      {'label': 'Animals', 'items': ['Move around', 'Eat food']}]},
    'pyramid': {'type': 'pyramid', 'title': 'Needs',
                'levels': ['Self-actualisation', 'Esteem', 'Belonging', 'Safety', 'Physiological needs']},
    'matrix': {'type': 'matrix', 'title': 'Priorities', 'x_axis': 'Effort', 'y_axis': 'Impact', 'quadrants': [
        {'label': 'Quick wins', 'items': ['Recycle paper']}, {'label': 'Big projects', 'items': ['Solar roof']},
        {'label': 'Fill-ins', 'items': ['Posters']}, {'label': 'Avoid', 'items': ['New car park']}]},
    'mindmap': {'type': 'mindmap', 'title': 'Energy', 'nodes': [
        {'id': 'e', 'label': 'Energy', 'parent': None}, {'id': 'r', 'label': 'Renewable', 'parent': 'e'},
        {'id': 'n', 'label': 'Fossil fuels', 'parent': 'e'}, {'id': 's', 'label': 'Solar', 'parent': 'r'},
        {'id': 'w', 'label': 'Wind', 'parent': 'r'}, {'id': 'c', 'label': 'Coal', 'parent': 'n'}]},
    'process': {'type': 'process', 'title': 'Scientific method', 'steps': [
        {'label': 'Question', 'detail': 'Ask what you want to find out'}, {'label': 'Hypothesis', 'detail': ''},
        {'label': 'Experiment', 'detail': 'Test it fairly'}, {'label': 'Conclusion', 'detail': 'Explain it'}]},
    'comparison': {'type': 'comparison', 'title': 'Cell division', 'columns': [
        {'label': 'Mitosis', 'items': ['Two cells', 'Identical']}, {'label': 'Meiosis', 'items': ['Four cells']}]},
    'stat-cards': {'type': 'stat-cards', 'title': 'Key numbers', 'stats': [
        {'value': '71%', 'label': 'of Earth is water', 'icon': 'droplet'}, {'value': '8', 'label': 'planets'},
        {'value': '4.5 bn', 'label': 'years old'}]},
    'scatter': {'type': 'scatter', 'title': 'Height and arm span', 'x_label': 'Height', 'y_label': 'Arm span',
                'points': [{'x': 150, 'y': 148, 'label': 'Ana'}, {'x': 160, 'y': 161}, {'x': 172, 'y': 170,
                                                                                         'label': 'Ben'}]},
}


def labelled_block(cache) -> dict:
    sha = cache(subject_image(0.5, 0.5, (800, 600)))
    return {'type': 'labelled', 'title': 'A plant cell', 'image': sha, 'callouts': [
        {'label': 'Nucleus', 'x': 0.5, 'y': 0.5}, {'label': 'Cell wall', 'x': 0.1, 'y': 0.2},
        {'label': 'Vacuole', 'x': '70%', 'y': 0.8}]}


def spec(*blocks, theme='clean') -> dict:
    return {'title': 'Science', 'theme': theme, 'sections': [
        {'heading': 'Overview', 'level': 1, 'blocks': [{'type': 'paragraph', 'text': 'Some words first.'}, *blocks]}]}


def last_block(s, fmt='pdf'):
    return normalize(s, fmt)[0]['sections'][-1]['blocks'][-1]


def squash(text: str) -> str:
    return re.sub(r'\s+', '', text).lower()


def all_blocks(cache) -> dict:
    return {**BLOCKS, 'labelled': labelled_block(cache)}


def test_the_kinds_match_the_contract():
    assert diagram.NEW_KINDS == plan.DIAGRAM_KINDS_NEW
    assert set(diagram.NEW_KINDS) <= set(spec_mod.BLOCKS) and set(diagram.ALL_KINDS) == set(spec_mod.DIAGRAMS)
    assert diagram.KINDS == ('timeline', 'tree', 'flow')


@pytest.mark.parametrize('kind', diagram.NEW_KINDS)
def test_normalized_blocks_follow_the_block_schema_and_are_idempotent(kind, cache):
    block = all_blocks(cache)[kind]
    out = last_block(spec(block))
    assert out['type'] == kind
    assert plan.validate(out, plan.DIAGRAM_BLOCK_SCHEMAS[kind]) == [], out
    norm = normalize(spec(block), 'pdf')[0]
    assert normalize(norm, 'pdf')[0]['sections'] == norm['sections']
    lay = diagram.layout(out)
    assert diagram.overlaps(lay) == [] and lay.height > 20
    assert all(b['x'] >= -0.5 and b['x'] + b['w'] <= lay.width + 0.5 for b in lay.boxes)


@pytest.mark.parametrize('kind', diagram.NEW_KINDS)
@pytest.mark.parametrize('fmt', ['pdf', 'docx', 'md', 'xlsx'])
def test_every_new_kind_is_drawn_by_the_legacy_renderers(kind, fmt, cache):
    s = spec(all_blocks(cache)[kind])
    data = render(s, fmt)
    results = {r.id: r for r in verify(s, fmt, data)}
    assert results['V4'].ok and '1 of 1 diagrams drawn' in results['V4'].note, results['V4']
    block = last_block(s, fmt)
    labels = diagram.labels_of(block)
    assert labels
    if fmt == 'pdf':
        from pypdf import PdfReader
        text = squash(' '.join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages))
        drawn = labels if kind != 'scatter' else [x for x in labels if x in ('Ana', 'Ben', 'Height', 'Arm span')]
        assert all(squash(lab) in text for lab in drawn), [lab for lab in drawn if squash(lab) not in text]
    elif fmt == 'docx':
        from docx import Document
        alts = [x._inline.docPr.get('descr') for x in Document(io.BytesIO(data)).inline_shapes]
        assert len(alts) == 1 and alts[0].startswith('Diagram: ')
        assert all(lab in alts[0] for lab in labels)
    elif fmt == 'md':
        text = data.decode()
        assert '```mermaid\n' in text and all(lab in text for lab in labels)
    else:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data))
        cells = {str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None}
        assert all(lab in cells for lab in labels), [lab for lab in labels if lab not in cells]


def test_every_new_kind_draws_as_native_editable_pptx_shapes(cache):
    from pptx import Presentation
    from jevrouter.create import themes
    prs = Presentation()
    blocks = all_blocks(cache)
    for kind, block in blocks.items():
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        diagram.pptx_draw(slide, last_block(spec(block)), themes.get('dark'), (0.5, 0.5, 9, 6.5), font='Inter')
    buf = io.BytesIO()
    prs.save(buf)
    prs = Presentation(io.BytesIO(buf.getvalue()))

    def walk(shapes):
        for s in shapes:
            yield s
            if s.shape_type == 6:
                yield from walk(s.shapes)
    for slide, (kind, block) in zip(prs.slides, blocks.items()):
        shapes = list(walk(slide.shapes))
        groups = [s for s in shapes if s.shape_type == 6]
        assert len(groups) == 1 and groups[0].name.startswith('Diagram: ')
        texts = squash(' '.join(s.text_frame.text for s in shapes if s.has_text_frame))
        norm = last_block(spec(block))
        drawn = diagram.labels_of(norm) if kind != 'scatter' else ['Ana', 'Ben', 'Height', 'Arm span']
        assert all(squash(lab) in texts for lab in drawn), (kind, drawn)
        if kind == 'labelled':
            assert any(s.shape_type == 13 and s.name == 'Labelled figure' for s in shapes)
    xml = ''.join(zipfile.ZipFile(buf).read(n).decode() for n in zipfile.ZipFile(buf).namelist()
                  if n.startswith('ppt/slides/slide'))
    assert 'prst="chevron"' in xml and 'prst="ellipse"' in xml and '<a:custGeom>' in xml and '<a:alpha ' in xml


@pytest.mark.parametrize('kind', diagram.NEW_KINDS)
def test_legacy_pptx_renderer_draws_every_new_kind_as_a_diagram(kind, cache, monkeypatch):
    """The legacy PPTX slide builder sends every kind in DIAGRAMS to diagram_shape (native grouped shapes)."""
    from pptx import Presentation
    monkeypatch.setenv('TG_STUDIO', 'off')
    s = spec(all_blocks(cache)[kind])
    data = render(s, 'pptx')
    prs = Presentation(io.BytesIO(data))
    groups = [sh for slide in prs.slides for sh in slide.shapes if sh.shape_type == 6]
    assert len(groups) == 1 and groups[0].name.startswith('Diagram: '), [sh.name for sh in groups]
    assert any(sh.name == 'Diagram title' for slide in prs.slides for sh in slide.shapes)
    results = {r.id: r for r in verify(s, 'pptx', data)}
    assert results['V4'].ok and '1 of 1 diagrams drawn' in results['V4'].note


def test_new_kinds_in_mono_draw_only_greys(cache):
    blocks = [b for k, b in all_blocks(cache).items() if k != 'labelled']
    ok, note = grey_scan(render(spec(*blocks, theme='mono'), 'pdf'))
    assert ok, note


# ---------- repair, not block ----------

def notes(res) -> str:
    return ' | '.join(r.note for r in res if not r.ok)


@pytest.mark.parametrize('raw,kind', [
    ({'type': 'venn diagram', 'sets': {'Cats': ['Purr'], 'Dogs': ['Bark']}, 'overlap': 'Pets'}, 'venn'),
    ({'type': 'swot', 'quadrants': {'Strengths': 'Fast', 'Weaknesses': 'Costly', 'Opportunities': 'New users',
                                    'Threats': 'Rivals'}}, 'matrix'),
    ({'type': 'Life Cycle', 'stages': 'Egg\nLarva\nPupa\nAdult'}, 'cycle'),
    ({'type': 'mind-map', 'centre': 'Rome', 'branches': [{'label': 'Army', 'children': ['Legions']}, 'Roads']},
     'mindmap'),
    ({'type': 'process', 'steps': ['Plan: decide the aim', 'Do', 'Review']}, 'process'),
    ({'type': 'pros and cons', 'columns': {'Pros': ['Cheap'], 'Cons': ['Slow']}}, 'comparison'),
    ({'type': 'stats', 'items': ['72% walk to school', '3 buses a day']}, 'stat-cards'),
    ({'type': 'scatter', 'points': [[1, 2], [2, '3.5'], {'x': 3, 'y': 5, 'label': 'c'}], 'x_axis': 'Hours'},
     'scatter'),
    ({'type': 'diagram', 'kind': 'pyramid', 'levels': ['Top', 'Middle', 'Base']}, 'pyramid'),
    ({'type': 'diagram', 'kind': 'stat_cards', 'stats': [{'value': 1, 'label': 'a'}, {'value': 2, 'label': 'b'}]},
     'stat-cards'),
    ({'title': 'No type', 'quadrants': [{'label': 'A'}, {'label': 'B'}, {'label': 'C'}, {'label': 'D'}]}, 'matrix'),
    ({'sets': [{'label': 'A', 'items': ['x']}, {'label': 'B'}]}, 'venn'),
])
def test_models_other_spellings_are_read(raw, kind):
    out, res = normalize(spec(raw), 'pdf')
    b = out['sections'][0]['blocks'][-1]
    assert b['type'] == kind, (b, notes(res))
    assert plan.validate(b, plan.DIAGRAM_BLOCK_SCHEMAS[kind]) == [], b


def test_old_meanings_of_process_matrix_and_scatter_are_kept():
    flow = {'type': 'process', 'nodes': ['Plan', 'Build'], 'edges': [['Plan', 'Build']]}
    assert last_block(spec(flow))['type'] == 'flow'
    table = {'type': 'matrix', 'columns': ['a', 'b'], 'rows': [[1, 2]]}
    assert last_block(spec(table))['type'] == 'table'
    chart = {'type': 'scatter', 'labels': ['a', 'b'], 'series': [{'name': 's', 'values': [1, 2]}]}
    got = last_block(spec(chart))
    assert got['type'] == 'chart' and got['kind'] == 'line'
    bullets = {'type': 'bullets', 'points': ['one', 'two']}
    assert last_block(spec(bullets))['type'] == 'bullets'


def test_limits_are_repaired_with_a_note():
    ten = {'type': 'cycle', 'steps': [f'Step {i}' for i in range(10)]}
    out, res = normalize(spec(ten), 'pdf')
    assert len(out['sections'][0]['blocks'][-1]['steps']) == 8 and 'cycles cut to 8' in notes(res)
    long = {'type': 'pyramid', 'levels': ['x' * 90, 'b', 'c']}
    out, res = normalize(spec(long), 'pdf')
    assert len(out['sections'][0]['blocks'][-1]['levels'][0]) == 60 and 'shortened to 60' in notes(res)
    three = {'type': 'matrix', 'quadrants': [{'label': 'A'}, {'label': 'B'}, {'label': 'C'}]}
    out, res = normalize(spec(three), 'pdf')
    assert [q['label'] for q in out['sections'][0]['blocks'][-1]['quadrants']] == ['A', 'B', 'C', '']
    deep = {'type': 'mindmap', 'nodes': [{'id': 'a', 'label': 'A'}, {'id': 'b', 'label': 'B', 'parent': 'a'},
                                         {'id': 'c', 'label': 'C', 'parent': 'b'}, {'id': 'd', 'label': 'D', 'parent': 'c'}]}
    out, res = normalize(spec(deep), 'pdf')
    nodes = {n['id']: n['parent'] for n in out['sections'][0]['blocks'][-1]['nodes']}
    assert nodes == {'a': None, 'b': 'a', 'c': 'b', 'd': 'b'} and 'joined to their branch' in notes(res)
    pts = {'type': 'scatter', 'points': [[1, 2], ['a', 3], [4, 5]]}
    out, res = normalize(spec(pts), 'pdf')
    assert len(out['sections'][0]['blocks'][-1]['points']) == 2 and 'points without two numbers' in notes(res)


@pytest.mark.parametrize('raw,word', [
    ({'type': 'cycle', 'title': 'Loop', 'steps': ['Only', 'Two']}, 'Two'),
    ({'type': 'venn', 'sets': [{'label': 'Alone', 'items': ['x']}]}, 'Alone'),
    ({'type': 'stat-cards', 'stats': [{'value': '42', 'label': 'answers'}]}, '42 answers'),
    ({'type': 'labelled', 'image': 'f' * 64, 'callouts': [{'label': 'Nucleus', 'x': 0.5, 'y': 0.5}]}, 'Nucleus'),
])
def test_too_little_to_draw_keeps_the_words_as_a_list(raw, word, cache):
    for fmt in FORMATS:
        out, res = normalize(spec(raw), fmt)
        b = out['sections'][0]['blocks'][-1]
        assert b['type'] == 'bullets' and any(word in it for it in b['items']), (fmt, b)
        assert 'shown as a list' in notes(res)


def test_labelled_pictures_come_only_from_code(cache):
    block = labelled_block(cache)
    # code wrote it: S7 keeps the asset id (it is not an address) and normalize keeps the figure
    out, res = normalize(spec(block), 'pdf')
    assert out['sections'][0]['blocks'][-1]['image'] == block['image'] and 'S7' not in {r.id for r in res if not r.ok}
    # a model wrote it: strip_internal drops the picture id, so its callouts stay as words
    from jevrouter.create.spec import strip_internal
    stripped = strip_internal(spec(block))
    b = last_block(stripped)
    assert b['type'] == 'bullets' and 'Nucleus' in b['items']
    # an address instead of an asset id is removed by S7 like any other
    web = {**block, 'image': 'https://example.com/cell.png'}
    out, res = normalize(spec(web), 'pdf')
    assert out['sections'][0]['blocks'][-1]['type'] == 'bullets' and 'nothing is fetched' in notes(res)


def test_brief_reads_the_studio_diagram_words():
    assert studio_kinds('a deck with a venn diagram and a SWOT analysis') == ['venn', 'matrix']
    assert studio_kinds('a labelled diagram of a plant cell and a scatter plot') == ['labelled', 'scatter']
    assert studio_kinds('explain the water cycle with a cycle diagram') == ['cycle']
    assert studio_kinds('a report about cycling') == []
    # the writer can name the Studio kinds, so the brief asks for them by name, not for a legacy stand-in
    assert parse_brief('a pdf with a water cycle diagram').diagram_kinds == ['cycle']
    assert parse_brief('slides with a mind map').diagram_kinds == ['mindmap']
    assert parse_brief('a deck with a venn diagram and a timeline').diagram_kinds == ['timeline', 'venn']
    assert parse_brief('slides with a flowchart and a SWOT diagram').diagram_kinds == ['flow', 'matrix']
    # a labelled picture comes only from code, so it is never asked of the writer
    assert parse_brief('a pdf with a labelled diagram of a cell').diagram_kinds == []
    # the words alone, without asking for a diagram, ask for nothing
    assert parse_brief('a report about the water cycle').diagram_kinds == []
    # a stored brief with a kind this build doesn't know drops it instead of failing later
    from jevrouter.create.brief import from_dict, to_dict
    b = from_dict({**to_dict(parse_brief('slides with a venn diagram')), 'diagram_kinds': ['venn', 'hexagon']})
    assert b.diagram_kinds == ['venn']


# ---------- the writer's schema offers the Studio kinds (one generic diagram block) ----------

def _item(label='', detail='', parent='', value='', x=None, y=None, items=()):
    return {'label': label, 'detail': detail, 'parent': parent, 'value': value, 'x': x, 'y': y, 'items': list(items)}


GENERIC = {
    'cycle': {'items': [_item('Evaporation'), _item('Condensation'), _item('Precipitation'), _item('Collection')]},
    'venn': {'items': [_item('Plants', items=['Cell walls']), _item('Animals', items=['Move around'])],
             'shared': ['Cells', 'DNA']},
    'pyramid': {'items': [_item('Self-actualisation'), _item('Esteem'), _item('Safety')]},
    'matrix': {'items': [_item('Quick wins', items=['Recycle']), _item('Big projects', items=['Solar']),
                         _item('Fill-ins'), _item('Avoid', items=['Car park'])], 'x_label': 'Effort',
               'y_label': 'Impact'},
    'mindmap': {'items': [_item('Energy'), _item('Renewable', parent='Energy'), _item('Solar', parent='Renewable'),
                          _item('Coal', parent='Energy')]},
    'process': {'items': [_item('Question', 'Ask it'), _item('Test', 'Fairly'), _item('Conclude')]},
    'comparison': {'items': [_item('Mitosis', items=['Two cells']), _item('Meiosis', items=['Four cells'])]},
    'stat-cards': {'items': [_item('of Earth is water', value='71%'), _item('planets', value='8')]},
    'scatter': {'items': [_item('Ana', x=150, y=148), _item('', x=160, y=161), _item('Ben', x=172, y=170)],
                'x_label': 'Height', 'y_label': 'Arm span'},
}


def generic_block(kind: str) -> dict:
    base = {'type': 'diagram', 'kind': kind, 'title': f'A {kind}', 'items': [], 'shared': [], 'x_label': '',
            'y_label': ''}
    return {**base, **GENERIC[kind]}


def test_the_schema_offers_every_writable_kind_once():
    blocks = spec_mod.DOCSPEC_SCHEMA['properties']['sections']['items']['properties']['blocks']['items']['anyOf']
    generic = [b for b in blocks if b['properties']['type']['enum'] == ['diagram']]
    assert len(generic) == 1
    assert set(generic[0]['properties']['kind']['enum']) == set(diagram.NEW_KINDS) - {'labelled'}
    assert set(GENERIC) == set(spec_mod.DIAGRAM_WRITABLE)


def strict_ok(v, sch) -> bool:
    """The JSON Schema subset DOCSPEC_SCHEMA uses (type, enum, anyOf, properties, required, additionalProperties,
    items), checked the way an engine that enforces the schema would."""
    if 'anyOf' in sch:
        return any(strict_ok(v, s) for s in sch['anyOf'])
    t = sch.get('type')
    kinds = {'object': dict, 'array': list, 'string': str, 'boolean': bool, 'null': type(None)}
    if t in kinds and not isinstance(v, kinds[t]):
        return False
    if t in ('number', 'integer') and (isinstance(v, bool) or not isinstance(v, (int, float) if t == 'number' else int)):
        return False
    if 'enum' in sch and v not in sch['enum']:
        return False
    if t == 'object':
        props = sch.get('properties', {})
        if set(sch.get('required', ())) - set(v) or (sch.get('additionalProperties') is False and set(v) - set(props)):
            return False
        return all(strict_ok(v[k], props[k]) for k in v if k in props)
    if t == 'array':
        return all(strict_ok(x, sch.get('items', {})) for x in v)
    return True


@pytest.mark.parametrize('kind', sorted(GENERIC))
def test_a_generic_diagram_block_becomes_its_kind(kind):
    raw = generic_block(kind)
    doc = {'title': 'Science', 'subtitle': '', 'sections': [{'heading': 'Overview', 'level': 1, 'notes': '',
                                                             'blocks': [raw]}]}
    assert strict_ok(doc, spec_mod.DOCSPEC_SCHEMA)   # what an engine that enforces the schema may send
    assert not strict_ok({**doc, 'sections': [{**doc['sections'][0], 'blocks': [{**raw, 'kind': 'labelled'}]}]},
                         spec_mod.DOCSPEC_SCHEMA)
    for fmt in ('pdf', 'pptx'):
        out, res = normalize(json.loads(json.dumps(doc)), fmt)
        b = out['sections'][0]['blocks'][-1]
        assert b['type'] == kind, (b, notes(res))
        assert plan.validate(b, plan.DIAGRAM_BLOCK_SCHEMAS[kind]) == [], b
    labels = diagram.labels_of(b)
    for item in GENERIC[kind]['items']:
        if item['label']:
            assert item['label'] in labels, (item['label'], labels)
    data = render(doc, 'pdf')
    assert {r.id: r for r in verify(doc, 'pdf', data)}['V4'].ok
