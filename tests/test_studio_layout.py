"""Studio builder L: the layout library, the fitting pipeline, the art director (keyless and model) and the PPTX/PDF
painters (docs/PLAN-designer.md 3.1, 3.5, 9.4-9.6). No network and no paid models: engines are fakes."""
import asyncio
import hashlib
import io
import json
from types import SimpleNamespace as NS

import pytest

from jevrouter.create.spec import normalize
from jevrouter.studio import direct, layout, library, paint_pdf, paint_pptx, plan, tokens

# ---------- fixtures ----------


def _ds(preset='minimal', dark=False, **kw) -> tokens.DesignSystem:
    colors = {r: '000000' for r in tokens.COLOR_ROLES}
    colors.update(bg='FFFFFF', surface='F3F4F6', text='1F2328', heading='111827', muted='57606A', accent='1D4ED8',
                  accent2='0891B2', border='D0D7DE', header_bg='1D4ED8', header_text='FFFFFF', stripe='F6F8FA',
                  code_bg='F6F8FA', overlay='F8FAFC', on_accent='FFFFFF')
    return tokens.DesignSystem(id=preset, name=preset, dark=dark, colors=colors,
                               chart_palette=['2563EB', 'F59E0B', '059669', 'DC2626', '7C3AED', '0891B2'], **kw)


class FakeWS:
    """The Workspace surface the layout engine and painters use."""

    def __init__(self):
        self.files = {}

    def put(self, kind, name, data):
        self.files[f'ws:{kind}/{name}'] = data
        return f'ws:{kind}/{name}'

    def get(self, ref):
        return self.files.get(ref)

    def log(self, event):
        pass


@pytest.fixture
def cache(tmp_path, monkeypatch):
    """A private asset cache with real PNGs (create/assets.CACHE)."""
    from PIL import Image, ImageDraw
    from jevrouter.create import assets
    monkeypatch.setattr(assets, 'CACHE', tmp_path)

    def make(color=(30, 60, 120), size=(1200, 800)):
        im = Image.new('RGB', size, color)
        ImageDraw.Draw(im).ellipse((400, 200, 800, 600), fill=(250, 220, 90))
        buf = io.BytesIO()
        im.save(buf, 'PNG')
        data = buf.getvalue()
        sha = hashlib.sha256(data).hexdigest()
        (tmp_path / f'{sha}.png').write_bytes(data)
        return sha
    return make


def _deck(a1, a2) -> dict:
    return {'title': 'Photosynthesis', 'subtitle': 'How plants turn light into food', 'sections': [
        {'heading': 'Why it matters', 'level': 1, 'blocks': [
            {'type': 'image', 'asset': a1, 'caption': 'A leaf in sunlight', 'credit': 'Photo: J. Doe, CC BY-SA 4.0'},
            {'type': 'bullets', 'items': ['Plants make their own food', 'Oxygen for us to breathe']}]},
        {'heading': 'The big number', 'level': 1, 'blocks': [
            {'type': 'paragraph', 'text': 'About 71% of Earth is covered by water where algae make oxygen.'}]},
        {'heading': 'Key figures', 'level': 1, 'blocks': [
            {'type': 'bullets', 'items': ['120 billion tonnes of carbon fixed a year', '50% of oxygen comes from the '
                                          'ocean', '1.5% of sunlight energy captured']}]},
        {'heading': 'Timeline of discovery', 'level': 1, 'blocks': [
            {'type': 'timeline', 'title': 'Discoveries', 'events': [
                {'date': '1648', 'label': 'Van Helmont weighs a willow'}, {'date': '1771', 'label': 'Priestley'},
                {'date': '1779', 'label': 'Ingenhousz: light needed'}]}]},
        {'heading': 'Light vs dark reactions', 'level': 1, 'blocks': [
            {'type': 'paragraph', 'text': 'Light reactions'}, {'type': 'bullets', 'items': ['In the thylakoids',
                                                                                              'Need light']},
            {'type': 'paragraph', 'text': 'Calvin cycle'}, {'type': 'bullets', 'items': ['In the stroma',
                                                                                          'Fix carbon dioxide']}]},
        {'heading': 'Rate by light level', 'level': 1, 'blocks': [
            {'type': 'chart', 'kind': 'bar', 'title': 'Oxygen output', 'labels': ['Low', 'Medium', 'High'],
             'series': [{'name': 'ml per hour', 'values': [2, 5, 9]}]},
            {'type': 'paragraph', 'text': 'More light means more oxygen, up to a limit.'}]},
        {'heading': 'In their words', 'level': 1, 'blocks': [
            {'type': 'quote', 'text': 'Nature does nothing uselessly.', 'by': 'Aristotle'}]},
        {'heading': 'The equation', 'level': 1, 'notes': 'Say the equation out loud.', 'blocks': [
            {'type': 'paragraph', 'text': 'Carbon dioxide and water become glucose and oxygen, using light energy.'},
            {'type': 'image', 'asset': a2, 'caption': 'Chloroplasts', 'credit': 'Photo: A. Smith, CC0'}]},
        {'heading': 'Results table', 'level': 1, 'blocks': [
            {'type': 'table', 'columns': ['Light', 'Bubbles per minute'],
             'rows': [['Low', 4], ['Medium', 11], ['High', 19]]}]},
        {'heading': 'Summary', 'level': 1, 'blocks': [
            {'type': 'bullets', 'items': ['Light becomes chemical energy', 'Two stages work together']}]},
    ]}


def _laid(spec_raw, fmt, ds=None, **kw):
    spec, _ = normalize(spec_raw, fmt)
    ds = ds or _ds()
    outline = direct.outline_of(spec, fmt)
    d = direct.direct_keyless(outline, ds)
    ws = FakeWS()
    p = layout.lay_out(spec, fmt, ds, d, ws, file_id='0123456789ab', **kw)
    return spec, p, ws, d


# ---------- the library ----------


def test_next_best_stays_in_family_then_text():
    assert library.next_best('image-left-text') == 'image-right-text'
    assert library.next_best('image-left-text', {'image-right-text'}) == 'title-bullets'
    assert library.next_best('image-left-text', {'image-right-text'}, needs=('image',)) is None
    assert library.next_best('full-width-diagram', needs=('chart',)) == 'chart-focus'
    assert library.next_best('two-column-text') == 'pull-quote'
    assert library.next_best('cover-hero') == 'cover-type'


def test_grid_rects_follow_margins_gutters_and_bleed():
    ds = _ds()
    x, y, w, h = layout.grid_rect((0, 0, 12, 12), layout.SLIDE, ds)
    assert (x, y) == (48, 48) and abs(w - 864) < 0.01 and abs(h - 444) < 0.01
    assert layout.grid_rect((0, 0, 6, 12), layout.SLIDE, ds, bleed=True) == (0, 0, 480, 540)
    edges = layout.column_edges(layout.SLIDE, ds, 'pptx')
    assert len(edges) == 12 and edges[0] == 48 and abs(edges[1] - edges[0] - (864 - 11 * 16) / 12 - 16) < 0.01
    pw, ph = layout.page_size('pdf', 'letter')
    assert (pw, ph) == (612.0, 792.0)
    x, y, w, h = layout.grid_rect((0, 0, 12, 16), (pw, ph), ds, fmt='pdf')
    assert abs(x - 56.7) < 0.01 and abs(x + w - (pw - 56.7)) < 0.01


def test_compact_variant_moves_slots_up_and_steps_down():
    d = layout.slot_areas('title-bullets', 'compact')
    assert d['title'] == ((0, 0, 12, 1), 'h3') and d['body'] == ((0, 1, 12, 11), 'body')
    assert layout.slot_areas('image-left-text', 'compact')['caption'] == ((7, 10, 5, 2), 'caption')


# ---------- fitting ----------


def test_fit_text_wraps_with_real_metrics_and_shrinks_within_the_scale(monkeypatch):
    from jevrouter.studio import fonts
    calls = []

    def fake_measure(text, family, size, *, bold=False, italic=False, fmt='pptx'):
        calls.append(family)
        return len(text) * size * 0.5
    monkeypatch.setattr(fonts, 'measure', fake_measure)
    ds = _ds()
    r = layout.fit_text('Photosynthesis in green plants', 'Inter', 'h2', ds, 400, 60, 'pptx')
    assert r.fits and calls and set(calls) == {'Inter'}
    assert all(len(ln) * r.size * 0.5 <= 400 + 0.01 for ln in r.lines)
    assert r.size in {ds.size(s, 'pptx') for s in plan.TYPE_STEPS}          # on the type scale (D6)
    # too much text: shrinks down the scale but never below the slide body minimum, and reports the overflow
    r = layout.fit_text('word ' * 300, 'Inter', 'h1', ds, 300, 80, 'pptx')
    assert not r.fits and r.size == ds.size('body', 'pptx') == 18 and r.overflow_words > 0
    # a caption never goes below the caption minimum
    r = layout.fit_text('tiny ' * 200, 'Inter', 'caption', ds, 200, 20, 'pdf')
    assert r.size >= 8


def test_wrap_keeps_emphasis_runs_and_breaks_long_words():
    lines = layout.wrap_runs('A **bold** word and *italic* text', 'Inter', 18, 1000)
    segs = {(t.strip(), b, i) for t, b, i in lines[0]}
    assert len(lines) == 1 and ('bold', True, False) in segs and ('italic', False, True) in segs
    assert layout.plain_line(lines[0]) == 'A bold word and italic text'
    lines = layout.wrap_runs('x' * 400, 'Inter', 18, 100)
    assert len(lines) > 3 and all(layout.line_width(ln, 'Inter', 18, 'pptx') <= 100.01 for ln in lines)


def test_stats_are_found_but_years_are_not():
    assert layout.stat_of('72% of students prefer it') == ('72%', 'of students prefer it')
    assert layout.stat_of('About $3.2 billion was spent')[0] == '$3.2 billion'
    assert layout.stat_of('In 1914 the war began') is None
    assert layout.stat_of('There are 3 states') is None


# ---------- the art director ----------


def test_keyless_direction_by_content_shape(cache):
    spec, _ = normalize(_deck(cache(), cache((20, 110, 60))), 'pptx')
    o = direct.outline_of(spec, 'pptx')
    assert o.pages[0].section is None and o.pages[-1].section is None and len(o.pages) == 12   # cover + 10 + closing
    d = direct.direct_keyless(o, _ds())
    got = [p.layout for p in d.pages]
    assert got[0] == 'cover-hero' and got[-1] == 'closing' and d.source == 'keyless'
    assert got[1] in ('image-left-text', 'image-right-text')
    assert got[2] == 'big-number' and got[3] == 'stat-cards' and got[4] == 'timeline-strip'
    assert got[5] == 'comparison' and got[6] == 'chart-focus' and got[7] == 'quote'
    assert got[9] == 'full-width-diagram'
    assert plan.validate(d.to_dict(), plan.art_schema(llm=False)) == []
    assert len(set(got)) >= 4 and _max_run(got) <= 3                       # D7


def _max_run(xs):
    best = run = 1
    for a, b in zip(xs, xs[1:]):
        run = run + 1 if a == b else 1
        best = max(best, run)
    return best


def test_keyless_variety_on_a_text_only_deck():
    spec = {'title': 'Notes', 'sections': [{'heading': f'Part {i}', 'level': 1, 'blocks': [
        {'type': 'paragraph', 'text': 'Some words about this part of the topic.'},
        {'type': 'bullets', 'items': ['One point', 'Another point']}]} for i in range(10)]}
    spec, _ = normalize(spec, 'pptx')
    got = [p.layout for p in direct.direct_keyless(direct.outline_of(spec, 'pptx'), _ds()).pages]
    assert _max_run(got) <= 3 and len(set(got)) >= 4


def test_dividers_only_in_long_hierarchical_decks():
    secs = []
    for i in range(4):
        secs.append({'heading': f'Chapter {i}', 'level': 1, 'blocks': [{'type': 'paragraph', 'text': 'Intro text.'}]})
        secs += [{'heading': f'Topic {i}.{k}', 'level': 2, 'blocks': [{'type': 'paragraph', 'text': 'Detail.'}]}
                 for k in range(2)]
    spec, _ = normalize({'title': 'Course', 'sections': secs}, 'pptx')
    got = [p.layout for p in direct.direct_keyless(direct.outline_of(spec, 'pptx'), _ds()).pages]
    assert 'section-divider' in got


class FakeEngine:
    supports_vision = False

    def __init__(self, reply=None, error=None):
        self.reply, self.error, self.calls = reply, error, []

    async def stream(self, *, system, prompt, effort='medium', emit_delta=None, max_tokens=2048, web=False,
                     schema=None, exec=False):
        self.calls.append({'system': system, 'prompt': prompt, 'max_tokens': max_tokens, 'schema': schema,
                           'effort': effort})
        if self.error:
            raise self.error
        return NS(text=self.reply, input_tokens=900, output_tokens=300)


def test_llm_direction_is_validated_against_the_library(cache):
    spec, _ = normalize(_deck(cache(), cache((20, 110, 60))), 'pptx')
    o = direct.outline_of(spec, 'pptx')
    pages = [{'layout': 'cover-type'}, {'layout': 'hero-banner'}, {'layout': 'big-number', 'focus': 'block:0'},
             {'layout': 'quote'}]                                   # unknown id; a quote layout with no quote; short
    eng = FakeEngine(reply=json.dumps({'preset': 'bold-dark', 'mood': ['bold', 'cinematic'], 'dark': True,
                                       'pages': pages + [{'layout': 'title-bullets', 'freeform': True}] * 8}))
    d, usage = asyncio.run(direct.direct_llm(o, _ds(), eng, mood=['bold']))
    assert len(eng.calls) == 1 and eng.calls[0]['max_tokens'] <= 2000 and eng.calls[0]['schema'] == plan.art_schema()
    assert usage['calls'] == 1 and usage['llm_in'] == 900 and usage['llm_out'] == 300 and usage['ms'] >= 0
    assert d.source == 'llm' and d.preset == 'bold-dark' and d.dark is True and len(d.pages) == len(o.pages)
    got = [p.layout for p in d.pages]
    assert got[0] == 'cover-type' and got[1] != 'hero-banner' and got[2] == 'big-number' and got[3] != 'quote'
    assert got[-1] == 'closing'                                   # the closing page keeps its layout
    assert sum(p.freeform for p in d.pages) <= 2 and _max_run(got) <= 3
    assert all(direct.fits_layout(p.layout, op, o) or p.layout == 'section-divider' for p, op in zip(d.pages, o.pages))
    assert d.notes


@pytest.mark.parametrize('engine', [FakeEngine(reply='not json at all'), FakeEngine(error=RuntimeError('boom')),
                                    FakeEngine(reply='{"pages": []}'), None])
def test_llm_direction_falls_back_to_keyless(engine, cache):
    spec, _ = normalize(_deck(cache(), cache((20, 110, 60))), 'pptx')
    o = direct.outline_of(spec, 'pptx')
    d, usage = asyncio.run(direct.direct_llm(o, _ds(), engine))
    keyless = direct.direct_keyless(o, _ds())
    assert [p.layout for p in d.pages] == [p.layout for p in keyless.pages] and d.source == 'keyless'
    assert any('rules' in n for n in d.notes) and set(usage) == {'calls', 'llm_in', 'llm_out', 'ms'}
    assert usage['calls'] == (0 if engine is None else 1)


# ---------- the layout engine ----------


def test_lay_out_a_deck_positions_every_box_on_the_grid(cache):
    spec, p, ws, _ = _laid(_deck(cache(), cache((20, 110, 60))), 'pptx')
    d = json.loads(json.dumps(p.to_dict()))
    assert plan.validate(d, plan.design_plan_schema()) == [] and plan.DesignPlan.from_dict(d) == p
    edges = layout.column_edges(layout.SLIDE, _ds(), 'pptx')
    ds = _ds()
    sizes = {ds.size(s, 'pptx') for s in plan.TYPE_STEPS}
    for page in p.pages:
        assert page.w == 960 and page.h == 540 and page.direction is not None
        ids = [b.id for b in page.boxes]
        assert len(ids) == len(set(ids)) and all(i.startswith(f'p{page.index}.') for i in ids)
        for b in page.boxes:
            assert 0 <= b.x and 0 <= b.y and b.x + b.w <= 960.01 and b.y + b.h <= 540.01, (page.layout, b.id)
            if b.kind == 'text':
                assert b.size in sizes and b.lines, b.id                              # D6: on the scale
                assert len(b.lines) * b.size * b.line_height <= b.h + 0.5, b.id       # D1: measured fit
                assert b.size >= 12 and (b.step == 'caption' or b.size >= 18), b.id   # D3 minimums
                if not b.bleed and b.slot is not None:
                    assert any(abs(b.x - e) <= 2 for e in edges) or b.slot == 'cards', b.id   # snapped
            if b.kind in ('image', 'chart', 'diagram', 'table', 'icon'):
                assert b.alt, b.id
    layouts = [pg.layout for pg in p.pages]
    assert layouts[0] == 'cover-hero' and layouts[-1] == 'closing'
    cover = p.pages[0]
    img = next(b for b in cover.boxes if b.kind == 'image')
    assert img.bleed and img.fit == 'cover' and img.crop is not None and (img.w, img.h) == (960, 540)
    title = next(b for b in cover.boxes if b.slot == 'title')
    assert title.content == 'spec:title' and title.style.text_color == 'text' and title.overlay_ok
    assert [b.z for b in cover.boxes] == sorted(b.z for b in cover.boxes)
    # reading order: title before body on content slides
    tb = next(pg for pg in p.pages if pg.layout == 'title-bullets')
    order = [b.slot for b in tb.boxes if b.kind == 'text']
    assert order.index('title') < order.index('body')
    # the closing slide lists the image credits
    closing = p.pages[-1]
    assert any(b.content.endswith('#credit') for b in closing.boxes)
    assert {a.kind for a in p.assets} >= {'image', 'art'} and p.fonts
    # notes: the section's speaker notes survive
    eq = next(pg for pg in p.pages if pg.section == 7 and not pg.continued)
    assert 'Say the equation' in eq.notes


def test_box_text_resolves_every_ref_kind(cache):
    spec, _ = normalize(_deck(cache(), cache((20, 110, 60))), 'pptx')
    t = lambda ref: layout.box_text(plan.Box(id='x', kind='text', x=0, y=0, w=1, h=1, content=ref), spec)  # noqa
    assert t('spec:title') == 'Photosynthesis' and t('spec:0/heading#continued') == 'Why it matters (continued)'
    kinds = [b['type'] for b in spec['sections'][0]['blocks']]
    bi, ii = kinds.index('bullets'), kinds.index('image')
    assert t(f'spec:0/{bi}/items[1:2]') == 'Oxygen for us to breathe' and t(f'spec:0/{ii}#credit').startswith('Photo')
    assert t('spec:1/0#stat') == '71%' and t('spec:1/0/words[0:2]') == 'About 71%'
    ci = [b['type'] for b in spec['sections'][5]['blocks']].index('chart')
    assert t(f'spec:5/{ci}#summary').startswith('Highest') and t('text:Page 3') == 'Page 3'
    assert t('spec:6/0#by') == 'Aristotle' and t('spec:99/0') == '' and t('icon:quote') == ''
    b = plan.Box(id='x', kind='text', x=0, y=0, w=1, h=1, content=f'spec:0/{bi}/items[1:2]')
    assert layout.bullet_of(b, spec) == '•'


def test_overflow_is_rebalanced_split_evenly_and_never_near_empty():
    long_items = [f'Point {i} with a few more words to make it wrap onto a second line of the slide body' for i in
                  range(6)]
    raw = {'title': 'T', 'sections': [
        {'heading': 'Many points', 'level': 1, 'blocks': [{'type': 'bullets', 'items': long_items}]},
        {'heading': 'Essay', 'level': 1, 'blocks': [{'type': 'paragraph', 'text': ' '.join(['word'] * 38) + '.'},
                                                    {'type': 'paragraph', 'text': ' '.join(['more'] * 38) + '.'}]}]}
    spec, p, _, _ = _laid(raw, 'pptx')
    first = next(pg for pg in p.pages if pg.section == 0)
    body = [b for b in first.boxes if b.slot == 'body']
    words = sum(len(layout.box_text(b, spec).split()) for b in body)
    assert words <= 40 and first.notes            # over the slot's word budget: to the speaker notes (rebalance)
    for pg in p.pages:
        if pg.continued:
            text = [b for b in pg.boxes if b.kind == 'text' and b.slot != 'title']
            words = sum(len(layout.box_text(b, spec).split()) for b in text)
            assert words > layout.SMALL_REST_WORDS or any(b.kind == 'table' for b in pg.boxes), pg.index


def test_a_long_table_is_split_into_even_parts():
    rows = [[f'Row {i}', i] for i in range(12)]
    raw = {'title': 'T', 'sections': [{'heading': 'Data', 'level': 1, 'blocks': [
        {'type': 'table', 'columns': ['Item with a longer header', 'Value'], 'rows': rows}]}]}
    spec, p, _, _ = _laid(raw, 'pdf')
    spec2, p2, _, _ = _laid({**raw, 'sections': [{**raw['sections'][0], 'blocks': [
        {'type': 'table', 'columns': ['A', 'B'], 'rows': [[f'r{i}', i] for i in range(12)]}]}]}, 'pptx')
    tables = [layout.parse_ref(b.content).rows or (0, 12) for pg in p2.pages for b in pg.boxes if b.kind == 'table']
    sizes = [j - i for i, j in tables]
    assert sum(sizes) == 12 and max(sizes) - min(sizes) <= 1
    # every table part keeps its header row and the continued page says so
    assert all(pg.continued for pg in p2.pages if any(b.kind == 'table' and layout.parse_ref(b.content).rows
                                                      and layout.parse_ref(b.content).rows[0] > 0 for b in pg.boxes))


def test_image_contain_keeps_aspect_and_caps_upscaling(cache):
    small = cache(size=(200, 100))
    raw = {'title': 'T', 'sections': [{'heading': 'Pic', 'level': 1, 'blocks': [
        {'type': 'image', 'asset': small, 'caption': 'Small', 'credit': 'Photo: X, CC0'}]}]}
    spec, p, _, _ = _laid(raw, 'pdf')
    img = next(b for pg in p.pages for b in pg.boxes if b.kind == 'image' and not b.bleed)
    assert img.fit == 'contain' and abs(img.w / img.h - 2) < 0.01
    assert img.w <= 200 * 1.5 * 72 / 150 + 0.01                         # never past 1.5x at 150 dpi
    assert layout.crop_fractions((1200, 800), (0.9, 0.5), 400, 400) == (0.6667, 0.0, 0.6667, 1.0) or \
        layout.crop_fractions((1200, 800), (0.9, 0.5), 400, 400)[0] > 0.3   # focal crop leans to the subject


def test_print_layout_packs_sections_and_keeps_the_outline(cache):
    raw = {'title': 'Lab report', 'subtitle': 'Enzymes', 'sections': [
        {'heading': f'Part {i}', 'level': 1, 'blocks': [{'type': 'paragraph', 'text': 'Text ' * 80}]}
        for i in range(6)]}
    spec, p, ws, _ = _laid(raw, 'pdf')
    assert p.pages[0].layout == 'cover' and len(p.pages) < 1 + 6          # short sections share pages
    heads = [b for pg in p.pages for b in pg.boxes if b.content.endswith('/heading')]
    assert len(heads) == 6
    for i, pg in enumerate(p.pages):
        assert any(b.content == f'text:Page {i + 1}' for b in pg.boxes)
    data = paint_pdf.paint(p, spec, ws)
    from jevrouter.create.rules import verify
    bad = [r for r in verify(spec, 'pdf', data) if not r.ok]
    assert not bad, [(r.id, r.note) for r in bad]


# ---------- code fixes and restyle ----------


def test_refit_actions(cache):
    spec, p, ws, _ = _laid(_deck(cache(), cache((20, 110, 60))), 'pptx')
    ds = _ds()
    n = len(p.pages)
    i = next(k for k, pg in enumerate(p.pages) if pg.layout == 'title-bullets')
    title = next(b for b in p.pages[i].boxes if b.slot == 'title')
    size = title.size
    assert layout.refit(p, spec, i, 'shrink', ds, ws, box=title.id) and title.size < size
    assert layout.refit(p, spec, i, 'compact', ds, ws)
    assert ArtDirection_pages(p)[p.pages[i].direction].variant == 'compact'
    k = next(k for k, pg in enumerate(p.pages) if pg.layout == 'image-left-text')
    assert layout.refit(p, spec, k, 'swap_layout', ds, ws)
    assert p.pages[k].layout == 'image-right-text' and len(p.pages) == n
    assert not layout.refit(p, spec, 0, 'nonsense', ds, ws)
    before = len(p.pages)
    j = next(k for k, pg in enumerate(p.pages) if len([b for b in pg.boxes if b.slot == 'body']) >= 2)
    assert layout.refit(p, spec, j, 'split', ds, ws) and len(p.pages) == before + 1 and p.pages[j + 1].continued
    assert [pg.index for pg in p.pages] == list(range(len(p.pages)))
    for pg in p.pages:
        assert all(b.id.startswith(f'p{pg.index}.') for b in pg.boxes)
    assert plan.validate(json.loads(json.dumps(p.to_dict())), plan.design_plan_schema()) == []


def ArtDirection_pages(p):
    return plan.ArtDirection.from_dict(p.direction).pages


def test_relayout_restyles_with_overrides(cache):
    spec, p, ws, _ = _laid(_deck(cache(), cache((20, 110, 60))), 'pptx')
    dark = _ds('bold-dark', dark=True)
    q = layout.relayout(p, spec, dark, ws, layouts={2: 'title-bullets'})
    assert q.preset == 'bold-dark' and q.system['dark'] is True
    assert plan.ArtDirection.from_dict(q.direction).source == 'restyle'
    assert q.pages[2].layout == 'title-bullets'
    with pytest.raises(ValueError):
        layout.relayout(p, spec, dark, ws, layouts={1: 'hero-banner'})


def test_flow_styles_for_flowing_formats(cache):
    spec, _ = normalize(_deck(cache(), cache((20, 110, 60))), 'docx')
    f = layout.flow_styles(spec, 'docx', _ds())
    assert plan.validate(f, plan.FLOW_SCHEMA) == []
    assert set(f['styles']) == set(plan.TYPE_STEPS) and f['styles']['body']['size'] == 11
    assert any(x['placement'] == 'full' for x in f['figures']) and len(f['palette']) == 6
    p = layout.lay_out(spec, 'docx', _ds(), direct.direct_keyless(direct.outline_of(spec, 'docx'), _ds()), FakeWS(),
                       file_id='0123456789ab')
    assert p.pages == [] and p.flow == f


# ---------- painters ----------


def test_pptx_painter_draws_the_plan_natively(cache):
    from pptx import Presentation
    from jevrouter.create.rules import verify
    spec, p, ws, _ = _laid(_deck(cache(), cache((20, 110, 60))), 'pptx')
    data = paint_pptx.paint(p, spec, ws)
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == len(p.pages) and (prs.slide_width, prs.slide_height) == (12192000, 6858000)
    first = prs.slides[0]
    assert first.shapes.title.text_frame.text == 'Photosynthesis'
    names = [s.name for sl in prs.slides for s in sl.shapes]
    assert 'Image' in names and 'Chart' in names and 'Table' in names
    assert any(n.startswith('Diagram:') for n in names)
    # the title sits above the photo and its overlay (z order), in the plan's size and colour
    shapes = list(first.shapes)
    assert shapes.index(first.shapes.title) > [s.name for s in shapes].index('Image')
    run = first.shapes.title.text_frame.paragraphs[0].runs[0]
    box = next(b for b in p.pages[0].boxes if b.slot == 'title')
    assert run.font.size.pt == box.size and str(run.font.color.rgb) == _ds().color('text')
    notes = [sl.notes_slide.notes_text_frame.text for sl in prs.slides if sl.has_notes_slide]
    assert any('Say the equation' in n for n in notes)
    report = {'layouts': {}}
    for pg in p.pages:
        report['layouts'][pg.layout] = report['layouts'].get(pg.layout, 0) + 1
    bad = [r for r in verify(spec, 'pptx', data, design_report=report) if not r.ok]
    assert not bad, [(r.id, r.note) for r in bad]


def test_pdf_painter_embeds_outline_and_draws_the_same_geometry(cache):
    from pypdf import PdfReader
    from jevrouter.create.rules import verify
    spec, p, ws, _ = _laid(_deck(cache(), cache((20, 110, 60))), 'pdf')
    data = paint_pdf.paint(p, spec, ws)
    r = PdfReader(io.BytesIO(data))
    assert len(r.pages) == len(p.pages) and abs(float(r.pages[0].mediabox.width) - 595.28) < 0.1
    titles = [str(o.title) for o in r.outline if not isinstance(o, list)]
    assert titles[0] == 'Photosynthesis'
    text = ''.join(pg.extract_text() for pg in r.pages)
    assert 'Oxygen output' in text and 'Page 1' in text
    bad = [x for x in verify(spec, 'pdf', data) if not x.ok]
    assert not bad, [(x.id, x.note) for x in bad]


def test_pdf_painter_subsets_a_truetype_family(cache, monkeypatch):
    """A family studio/fonts resolves to a TrueType file is measured with it and subset-embedded."""
    import os
    import reportlab
    from jevrouter.studio import fonts
    vera = os.path.join(os.path.dirname(reportlab.__file__), 'fonts', 'Vera.ttf')
    verabd = os.path.join(os.path.dirname(reportlab.__file__), 'fonts', 'VeraBd.ttf')
    if not os.path.exists(vera):
        pytest.skip('reportlab ships without Vera')
    face = lambda path, st: fonts.FontFace('Vera', st, path, 'builtin', 'user')  # noqa: E731
    res = fonts.FontResolution('Vera', 'Vera', {'regular': face(vera, 'regular'), 'bold': face(verabd, 'bold')},
                               'Calibri', 'builtin', 'user', True)
    monkeypatch.setattr(fonts, 'resolve_local', lambda family: res if family == 'Vera' else None)
    ds = _ds()
    ds.families = tokens.Families(display='Vera', heading='Vera', body='Vera', caption='Vera')
    spec, p, ws, _ = _laid(_deck(cache(), cache((20, 110, 60))), 'pdf', ds=ds)
    assert any(f.family == 'Vera' and f.embedded for f in p.fonts)
    data = paint_pdf.paint(p, spec, ws)
    from jevrouter.create.rules import pdf_scan
    from pypdf import PdfReader
    used = pdf_scan(b'', PdfReader(io.BytesIO(data)))['fonts']
    assert any('+' in f and 'Vera' in f for f in used)            # a subset (ABCDEF+Name)


def test_painters_raise_for_the_wrong_format(cache):
    spec, p, ws, _ = _laid(_deck(cache(), cache((20, 110, 60))), 'pptx')
    with pytest.raises(ValueError):
        paint_pdf.paint(p, spec, ws)
    p.format = 'pdf'
    with pytest.raises(ValueError):
        paint_pptx.paint(p, spec, ws)


@pytest.mark.parametrize('seed', range(6))
def test_fuzz_specs_never_crash_and_always_fit(seed, cache):
    import random
    rnd = random.Random(seed)
    img = cache()
    kinds = ['paragraph', 'bullets', 'table', 'chart', 'quote', 'code', 'timeline', 'image']
    sections = []
    for s in range(rnd.randint(1, 14)):
        blocks = []
        for _ in range(rnd.randint(0, 4)):
            k = rnd.choice(kinds)
            words = ' '.join(rnd.choice(['alpha', 'beta', '42%', 'gamma', 'delta', '1,200', 'epsilon'])
                             for _ in range(rnd.randint(1, 90)))
            blocks.append({'paragraph': {'type': 'paragraph', 'text': words},
                           'bullets': {'type': 'bullets', 'items': [words[:rnd.randint(5, 120)]] * rnd.randint(1, 9)},
                           'table': {'type': 'table', 'columns': ['A', 'B', 'C'],
                                     'rows': [[words[:20], i, 'x'] for i in range(rnd.randint(1, 30))]},
                           'chart': {'type': 'chart', 'kind': rnd.choice(['bar', 'line', 'pie']), 'title': 'C',
                                     'labels': ['a', 'b', 'c'], 'series': [{'name': 's', 'values': [1, 2, 3]}]},
                           'quote': {'type': 'quote', 'text': words, 'by': 'Someone'},
                           'code': {'type': 'code', 'lang': 'py', 'text': '\n'.join(['x = 1'] * rnd.randint(1, 40))},
                           'timeline': {'type': 'timeline', 'title': 'T', 'events': [
                               {'date': str(1900 + i), 'label': 'event'} for i in range(rnd.randint(2, 9))]},
                           'image': {'type': 'image', 'asset': img, 'caption': 'cap', 'credit': 'Photo: X, CC0'}}[k])
        sections.append({'heading': f'Section {s}', 'level': rnd.choice([1, 1, 2]), 'blocks': blocks})
    raw = {'title': 'Fuzz ' * rnd.randint(1, 12), 'subtitle': 'sub', 'sections': sections}
    for fmt in ('pptx', 'pdf'):
        try:
            spec, p, ws, _ = _laid(raw, fmt)
        except Exception as e:
            from jevrouter.create.rules import SpecError
            if isinstance(e, SpecError):
                return
            raise
        assert plan.validate(json.loads(json.dumps(p.to_dict())), plan.design_plan_schema()) == []
        for pg in p.pages:
            for b in pg.boxes:
                if b.kind == 'text' and b.lines and 'overflow' not in ' '.join(p.notes):
                    assert len(b.lines) * b.size * b.line_height <= b.h + 0.5, (fmt, pg.layout, b.id)
        data = (paint_pptx if fmt == 'pptx' else paint_pdf).paint(p, spec, ws)
        assert len(data) > 1000


# ---------- visuals fill their slot (the keyless CSV regressions) ----------

def _csv_spec(n_rows: int, chart: bool = False) -> dict:
    cols = ['date', 'region', 'product', 'units', 'unit_price', 'revenue']
    rows = [[f'2025-01-{k + 1:02d}', ['North', 'South'][k % 2], ['Widget', 'Gadget', 'Gizmo'][k % 3], 5 + k, 25.0,
             125.0 + k] for k in range(n_rows)]
    blocks = [{'type': 'table', 'title': 'Sales q1', 'columns': cols, 'rows': rows}]
    if chart:
        blocks.append({'type': 'chart', 'kind': 'bar', 'title': 'Total revenue by status',
                       'labels': ['paid', 'refunded', 'pending'], 'series': [{'name': 'revenue', 'values': [5, 1, 2]}]})
    return {'title': 'Sales q1', 'subtitle': f'{n_rows} rows', 'sections': [
        {'heading': 'Sales q1', 'level': 1, 'blocks': blocks}]}


def _checks(p, cid):
    from jevrouter.studio import qa
    return [r for r in qa.run_check(cid, p) if not r.ok]


def test_a_short_slide_table_grows_toward_its_slot_and_is_painted_that_tall():
    from pptx import Presentation
    spec, p, ws, _ = _laid(_csv_spec(4), 'pptx')
    page = next(pg for pg in p.pages if any(b.kind == 'table' for b in pg.boxes))
    tb = next(b for b in page.boxes if b.kind == 'table')
    natural = sum(layout.table_geometry(layout.block_of(tb.content, spec), None, tb.w, tb.size, _ds(), 'pptx').heights)
    assert natural < tb.h <= natural * layout.TABLE_STRETCH + 0.5       # grown, but within the stretch cap
    assert not _checks(p, 'D4')                                           # no longer "looks empty"
    # the table is titled like its section, so the title is not said twice above it
    assert not any(b.content.endswith('#title') for b in page.boxes)
    prs = Presentation(io.BytesIO(paint_pptx.paint(p, spec, ws)))
    frame = next(sh for sh in prs.slides[page.index].shapes if sh.has_table)
    assert abs(sum(r.height for r in frame.table.rows) / 12700 - tb.h) < 1.0


def test_a_twelve_row_table_fits_one_slide_without_a_repeated_title():
    spec, p, _, _ = _laid(_csv_spec(12), 'pptx')
    tables = [b for pg in p.pages for b in pg.boxes if b.kind == 'table']
    assert len(tables) == 1 and '/rows[' not in tables[0].content        # not split 6 + 6
    assert not _checks(p, 'D4')


def test_print_figures_of_one_section_share_a_page_and_it_is_balanced():
    spec, p, _, _ = _laid(_csv_spec(8, chart=True), 'pdf')
    figure_pages = [pg for pg in p.pages if any(b.kind in ('table', 'chart') for b in pg.boxes)]
    assert len(figure_pages) == 1
    kinds = {b.kind for b in figure_pages[0].boxes}
    assert {'table', 'chart'} <= kinds
    assert not _checks(p, 'D5')
    # a table too long for the page still gets pages of its own and the chart is never squeezed under it
    spec2, p2, _, _ = _laid(_csv_spec(60, chart=True), 'pdf')
    charts = [b for pg in p2.pages for b in pg.boxes if b.kind == 'chart']
    assert len(charts) == 1 and charts[0].h >= 200


def test_the_enlarge_fix_grows_a_small_visual_without_adding_anything():
    from jevrouter.studio import qa
    spec, p, ws, _ = _laid(_csv_spec(4), 'pptx')
    page = next(pg for pg in p.pages if any(b.kind == 'table' for b in pg.boxes))
    tb = next(b for b in page.boxes if b.kind == 'table')
    tb.h = round(tb.h / layout.TABLE_STRETCH, 2)                          # as small as the old layout left it
    assert any(r.threshold == qa.THRESHOLDS['white_max'] for r in _checks(p, 'D4'))
    n_boxes, h0 = len(page.boxes), tb.h
    assert layout.refit(p, spec, page.index, 'enlarge', _ds(), ws) is True
    assert len(page.boxes) == n_boxes and tb.h > h0
    x, y, w, h = layout.grid_rect(layout.slot_areas(page.layout, page.variant)[tb.slot][0], (page.w, page.h), _ds())
    assert tb.y + tb.h <= y + h + 0.5                                     # never past its slot
    assert not _checks(p, 'D4')
    assert 'enlarge' in layout.FIX_ACTIONS and qa.EMPTY_FIXES == ('enlarge',)
