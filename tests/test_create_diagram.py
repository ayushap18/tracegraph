"""Native diagrams (docs/PLAN-accuracy-v2.md C4): timeline, tree and flow blocks drawn in every format from one
layout, normalized to their limits, counted by V4 and V9, and black and white charts that never rely on colour."""
import io
import re
import zipfile

import pytest

from jevrouter.create import FORMATS, normalize, render, verify
from jevrouter.create import diagram
from jevrouter.create.brief import parse_brief
from jevrouter.create.rules import grey_scan

TIMELINE = {'type': 'timeline', 'title': 'Milestones of AI', 'events': [
    {'date': '1950', 'label': 'Turing asks whether machines can think'},
    {'date': '1956', 'label': 'The Dartmouth workshop names the field'},
    {'date': '1997', 'label': 'Deep Blue beats Kasparov'},
    {'date': '2012', 'label': 'Deep learning wins ImageNet'}]}
TREE = {'type': 'tree', 'title': 'Fields of AI', 'nodes': [
    {'id': 'ai', 'parent': '', 'label': 'Artificial intelligence'},
    {'id': 'ml', 'parent': 'ai', 'label': 'Machine learning'},
    {'id': 'dl', 'parent': 'ml', 'label': 'Deep learning'},
    {'id': 'nlp', 'parent': 'ai', 'label': 'Language processing'},
    {'id': 'cv', 'parent': 'ai', 'label': 'Computer vision'}]}
FLOW = {'type': 'flow', 'title': 'Training a model', 'nodes': [
    {'id': 'a', 'label': 'Collect data'}, {'id': 'b', 'label': 'Train model'}, {'id': 'c', 'label': 'Evaluate'}],
    'edges': [{'from': 'a', 'to': 'b', 'label': 'clean'}, {'from': 'b', 'to': 'c', 'label': ''}]}
BLOCKS = {'timeline': TIMELINE, 'tree': TREE, 'flow': FLOW}


def spec(*blocks, theme='clean') -> dict:
    return {'title': 'About AI', 'theme': theme, 'sections': [
        {'heading': 'Overview', 'level': 1, 'blocks': [{'type': 'paragraph', 'text': 'Some words first.'}, *blocks]}]}


def squash(text: str) -> str:
    return re.sub(r'\s+', '', text).lower()


def pptx_texts(data: bytes) -> list[str]:
    from pptx import Presentation

    def walk(shapes):
        for s in shapes:
            yield s
            if s.shape_type == 6:
                yield from walk(s.shapes)
    prs = Presentation(io.BytesIO(data))
    return [s.text_frame.text for sl in prs.slides for s in walk(sl.shapes) if s.has_text_frame]


@pytest.mark.parametrize('kind', diagram.KINDS)
@pytest.mark.parametrize('fmt', FORMATS)
def test_every_kind_in_every_format(kind, fmt):
    block = BLOCKS[kind]
    data = render(spec(block), fmt)
    results = {r.id: r for r in verify(spec(block), fmt, data)}
    assert results['V4'].ok and '1 of 1 diagrams drawn' in results['V4'].note, results['V4']
    labels = diagram.labels_of(normalize(spec(block), fmt)[0]['sections'][-1]['blocks'][-1])
    if fmt == 'pdf':
        from pypdf import PdfReader
        text = squash(' '.join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages))
        assert all(squash(lab) in text for lab in labels)
    elif fmt == 'pptx':
        texts = squash(' '.join(pptx_texts(data)))
        assert all(squash(lab) in texts for lab in labels)
    elif fmt == 'docx':
        from docx import Document
        doc = Document(io.BytesIO(data))
        alts = [s._inline.docPr.get('descr') for s in doc.inline_shapes]
        assert len(alts) == 1 and alts[0].startswith('Diagram: ') and all(lab in alts[0] for lab in labels)
        assert any(n.startswith('word/media/') and n.endswith('.png') for n in zipfile.ZipFile(io.BytesIO(data)).namelist())
    elif fmt == 'md':
        text = data.decode()
        assert '```mermaid\n' in text and all(lab in text for lab in labels)
    else:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data))
        cells = {str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None}
        assert any(c.startswith('Diagram: ') for c in cells) and all(lab in cells for lab in labels)


def test_mermaid_text_is_safe():
    tricky = {'type': 'flow', 'title': 'x', 'nodes': [{'id': 'a', 'label': 'a"] --> evil["b'}, {'id': 'b', 'label': 'b'}],
              'edges': [{'from': 'a', 'to': 'b', 'label': 'x|y'}]}
    out = diagram.mermaid(normalize(spec(tricky), 'md')[0]['sections'][0]['blocks'][1])
    assert out.count('-->') == 1 and '"]' not in out.split('\n')[1][:-2] and '|x y|' in out


def test_a_40_node_tree_has_no_overlapping_boxes():
    nodes = [{'id': 'root', 'parent': '', 'label': 'Everything we know about the subject'}]
    for i in range(6):
        nodes.append({'id': f'b{i}', 'parent': 'root', 'label': f'Branch number {i} of the tree'})
        for j in range(6):
            nodes.append({'id': f'b{i}l{j}', 'parent': f'b{i}', 'label': f'Leaf {j} under branch {i}'})
    nodes = nodes[:40]
    block = normalize(spec({'type': 'tree', 'title': 'Big', 'nodes': nodes}), 'pdf')[0]['sections'][0]['blocks'][1]
    assert len(block['nodes']) == 40
    for width in (480.0, 300.0, 900.0):
        lay = diagram.layout(block, width)
        assert diagram.overlaps(lay) == [] and all(b['x'] >= 0 and b['x'] + b['w'] <= width + 0.5 for b in lay.boxes)
    small = diagram.layout(normalize(spec(TREE), 'pdf')[0]['sections'][0]['blocks'][1])
    assert diagram.overlaps(small) == [] and len(small.boxes) == 5
    for block in (TIMELINE, FLOW, {**TIMELINE, 'events': TIMELINE['events'] * 5}):
        assert diagram.overlaps(diagram.layout(normalize(spec(block), 'pdf')[0]['sections'][0]['blocks'][1])) == []


def test_normalize_trims_and_fixes_with_notes():
    long = {'type': 'timeline', 'title': 'Long', 'events': [{'date': str(1900 + i), 'label': f'e{i}'} for i in range(31)]}
    cyclic = {'type': 'flow', 'title': 'Loop', 'nodes': [{'id': x, 'label': x.upper()} for x in 'abc'],
              'edges': [{'from': 'a', 'to': 'b', 'label': ''}, {'from': 'b', 'to': 'c', 'label': ''},
                        {'from': 'c', 'to': 'a', 'label': 'again'}, {'from': 'a', 'to': 'zzz', 'label': ''}]}
    out, res = normalize(spec(long, cyclic), 'pdf')
    t, f = out['sections'][0]['blocks'][1:]
    notes = {r.id: r.note for r in res}
    assert len(t['events']) == 30 and 'timelines cut to 30 events' in notes['L2']
    assert [(e['from'], e['to']) for e in f['edges']] == [('a', 'b'), ('b', 'c')]
    assert 'a flow with a cycle had its back edges removed' in notes['S6']


def test_normalize_repairs_trees():
    nodes = [{'id': 'a', 'parent': '', 'label': 'A'}, {'id': 'b', 'parent': '', 'label': 'B'},
             {'id': 'c', 'parent': 'd', 'label': 'C'}, {'id': 'd', 'parent': 'c', 'label': 'D'},
             {'id': 'a', 'parent': 'b', 'label': 'dup'}, {'id': 'e', 'parent': 'nowhere', 'label': 'x' * 60}]
    deep = [{'id': f'n{i}', 'parent': f'n{i - 1}' if i else '', 'label': f'N{i}'} for i in range(7)]
    out, res = normalize(spec({'type': 'tree', 'title': 'T', 'nodes': nodes}, {'type': 'tree', 'title': 'D', 'nodes': deep}),
                         'pdf')
    t, d = out['sections'][0]['blocks'][1:]
    parents = {n['id']: n['parent'] for n in t['nodes']}
    assert list(parents) == ['a', 'b', 'c', 'd', 'e'] and parents['a'] == '' and parents['b'] == 'a'
    assert parents['e'] == 'a' and len(next(n for n in t['nodes'] if n['id'] == 'e')['label']) <= 40
    assert sum(1 for p in parents.values() if not p) == 1  # one root, the loop c <-> d cut
    assert len(d['nodes']) == 4 and 'trees cut to 4 levels' in {r.id: r.note for r in res}['L2']


def test_a_diagram_with_too_little_is_left_out():
    out, res = normalize(spec({'type': 'timeline', 'title': 'x', 'events': [{'date': '1', 'label': 'only'}]}), 'md')
    assert [b['type'] for b in out['sections'][0]['blocks']] == ['paragraph']
    assert 'fewer than 2 events' in {r.id: r.note for r in res}['S6']


def test_v9_counts_diagrams_against_the_brief():
    two = parse_brief('a pdf with multiple diagrams')
    one = spec(TIMELINE)
    data = render(one, 'pdf')
    v9 = {r.id: r for r in verify(one, 'pdf', data, brief=two)}['V9']
    assert not v9.ok and v9.note == '1 diagram drawn, asked for at least 2'
    both = spec(TIMELINE, TREE)
    v9 = {r.id: r for r in verify(both, 'pdf', render(both, 'pdf'), brief=two)}['V9']
    assert v9.ok and v9.note == '2 diagrams drawn, asked for at least 2'
    chart = spec({'type': 'chart', 'kind': 'bar', 'title': 'Sales', 'labels': ['a', 'b'],
                  'series': [{'name': 's', 'values': [1, 2]}]})
    assert {r.id: r for r in verify(chart, 'md', render(chart, 'md'), brief=parse_brief('with charts'))}['V9'].ok


def test_diagrams_in_mono_draw_only_greys():
    s = spec(TIMELINE, TREE, FLOW, theme='mono')
    assert grey_scan(render(s, 'pdf'))[0]


@pytest.mark.parametrize('kind', ['bar', 'line', 'pie'])
def test_mono_charts_use_patterns_not_colour(kind):
    chart = {'type': 'chart', 'kind': kind, 'title': 'Funding', 'labels': ['A', 'B', 'C'],
             'series': [{'name': 'x', 'values': [1, 2, 3]}] + ([{'name': 'y', 'values': [2, 3, 1]}] if kind != 'pie' else [])}
    mono, clean = spec(chart, theme='mono'), spec(chart)
    ok, note = grey_scan(render(mono, 'pdf'))
    assert ok, note
    assert not grey_scan(render(clean, 'pdf'))[0]
    if kind == 'bar':  # hatching: many short strokes inside the bars, and the pattern named in the legend
        from pypdf import PdfReader
        text = PdfReader(io.BytesIO(render(mono, 'pdf'))).pages[0].extract_text()
        assert 'diagonal: x' in text
    pptx = zipfile.ZipFile(io.BytesIO(render(mono, 'pptx')))
    xml = ''.join(pptx.read(n).decode() for n in pptx.namelist() if n.startswith('ppt/charts/chart'))
    assert ('<a:pattFill' in xml) == (kind == 'bar') and ('prstDash' in xml) == (kind == 'line')
    xlsx = zipfile.ZipFile(io.BytesIO(render(mono, 'xlsx')))
    xml = ''.join(xlsx.read(n).decode() for n in xlsx.namelist() if n.startswith('xl/charts/chart'))
    assert ('pattFill' in xml) == (kind == 'bar')


def test_page_breaks():
    s = {'title': 'T', 'sections': [{'heading': 'One', 'blocks': [{'type': 'paragraph', 'text': 'a'}]},
                                    {'heading': 'Two', 'blocks': [{'type': 'page_break'}, {'type': 'page_break'},
                                                                  {'type': 'paragraph', 'text': 'b'}]}]}
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(render(s, 'pdf')))
    assert len(reader.pages) == 2 and 'Two' in reader.pages[1].extract_text()
    assert normalize(s, 'pptx')[0]['sections'][1]['blocks'] == [{'type': 'paragraph', 'text': 'b'}]
    assert b'w:br w:type="page"' in zipfile.ZipFile(io.BytesIO(render(s, 'docx'))).read('word/document.xml')
