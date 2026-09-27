"""The eval harness for created files (docs/PLAN-files.md, "Evals"): the expect_file schema, reopening real PDF, DOCX,
PPTX, XLSX and Markdown files, scoring them, and a run through run_eval with a fake create step."""
import io

import pytest

from jevrouter import evals
from jevrouter.agents import AgentResult
from tests.test_evals_harness import registry, router, run


# ---------- real small files ----------

def pdf(pages=('Solid-state batteries', 'Why they matter'), title='Solid-state batteries') -> bytes:
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.setTitle(title)
    for text in pages:
        c.drawString(72, 720, text)
        c.showPage()
    c.save()
    return buf.getvalue()


def docx(paragraphs=('Compound interest', 'Interest earns interest.'), table=(('Year', 'Balance'), ('1', '1050'))) -> bytes:
    import docx as d
    doc = d.Document()
    doc.core_properties.title = paragraphs[0]
    doc.add_heading(paragraphs[0], 1)
    for p in paragraphs[1:]:
        doc.add_paragraph(p)
    t = doc.add_table(rows=len(table), cols=len(table[0]))
    for i, row in enumerate(table):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def pptx(titles=('Sales', 'Units by product'), chart=True, notes='Gizmo sold least') -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    prs = Presentation()
    for t in titles:
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = t
    if chart:
        data = CategoryChartData()
        data.categories = ['Widget', 'Gadget', 'Gizmo']
        data.add_series('Units', (52, 21, 40))
        s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(4), data)
    s.notes_slide.notes_text_frame.text = notes
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def xlsx(rows=(('region', 'revenue'), ('North', 300), ('West', 480)), sheet='Sales', chart=True, extra=()) -> bytes:
    import openpyxl
    from openpyxl.chart import BarChart, Reference
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for r in (*rows, *extra):
        ws.append(list(r))
    if chart:
        c = BarChart()
        c.add_data(Reference(ws, min_col=2, min_row=1, max_row=len(rows)), titles_from_data=True)
        ws.add_chart(c, 'D2')
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def rec(files, answer='Created the file.', agent='create'):
    return {'qid': 1, 'status': 'done', 'merged': {'answer': answer, 'engine': 'single'}, 'total_ms': 5,
            'tasks': [{'agent': agent, 'ok': True, 'created_files': files}]}


def made(fid, fmt, rules=(), **kw):
    return {'id': fid, 'name': f'{fid}.{fmt}', 'format': fmt, 'size': 1, 'created': 0, 'qid': 1, 'title': fid,
            'tokens': 0, 'source': 'answer', 'rules': list(rules), **kw}


def scored(expect, files, blobs):
    return evals.check_file(expect, rec(files), blobs.__getitem__)


# ---------- schema ----------

def test_expect_file_schema():
    ok = {'id': 'a', 'query': 'q', 'expect_file': {'format': 'pdf', 'contains': ['TCP'], 'min_pages': 1, 'rules_ok': True}}
    assert evals.validate_case(ok, strict=True) == []
    assert evals.validate_case({'id': 'b', 'query': 'q', 'expect_file': False}, strict=True) == []
    bad = {'id': 'c', 'query': 'q', 'expect_file': {'format': 'odt', 'contains': ['('], 'slides_min': -1, 'rules_ok': 1,
                                                    'colour': 'red'}}
    errors = evals.validate_case(bad)
    assert any('unknown key' in e for e in errors) and any('format must be one of' in e for e in errors)
    assert any('bad regex' in e for e in errors) and any('slides_min' in e for e in errors)
    assert any('rules_ok' in e for e in errors)
    assert evals.validate_case({'id': 'd', 'query': 'q', 'expect_file': 'pdf'}) == ['expect_file must be an object or false']
    turn = {'id': 'e', 'turns': [{'query': 'hi'}, {'query': 'as pdf', 'expect_file': {'format': 'pdf'}}]}
    assert evals.validate_case(turn, strict=True) == []


def test_committed_create_cases():
    cases = [c for c in evals.load_cases(evals.CASES) if 'create' in c['tags']]
    assert len(cases) >= 20
    assert 0.2 <= sum(c.get('split') == 'holdout' for c in cases) / len(cases) <= 0.35
    expects = [e for c in cases for e in [c.get('expect_file'), *(t.get('expect_file') for t in c.get('turns', []))]
               if e is not None]
    assert {e['format'] for e in expects if e} == set(evals.FILE_FORMATS)
    assert sum(e is False for e in expects) >= 3  # blocked and honest "no file" cases
    for c in cases:
        assert set(c['tags']) & set(evals.FILE_FORMATS), c['id']


# ---------- reopening and scoring each format ----------

def test_pdf_pages_and_text():
    blobs = {'p': pdf()}
    files = [made('p', 'pdf')]
    assert scored({'format': 'pdf', 'contains': ['solid.state', 'matter'], 'min_pages': 2, 'max_pages': 2}, files, blobs) == []
    r = scored({'format': 'pdf', 'contains': ['lithium'], 'min_pages': 3, 'max_pages': 1}, files, blobs)
    assert r == ['p.pdf does not contain /lithium/', 'p.pdf has 2 pages, fewer than 3', 'p.pdf has 2 pages, more than 1']


def test_docx_text_includes_tables():
    blobs = {'d': docx()}
    assert scored({'format': 'docx', 'contains': ['compound interest', 'Balance', '1050']}, [made('d', 'docx')], blobs) == []


def test_pptx_slides_charts_and_notes():
    blobs = {'s': pptx()}
    files = [made('s', 'pptx')]
    assert scored({'format': 'pptx', 'contains': ['Units by product', 'Gizmo', 'sold least'], 'slides_min': 2,
                   'charts_min': 1}, files, blobs) == []
    blobs['n'] = pptx(titles=('Only',), chart=False)
    r = scored({'format': 'pptx', 'slides_min': 3, 'charts_min': 1}, [made('n', 'pptx')], blobs)
    assert r == ['n.pptx has 1 slides, fewer than 3', 'n.pptx has 0 charts, fewer than 1']


def test_xlsx_sheets_charts_and_values():
    blobs = {'x': xlsx()}
    files = [made('x', 'xlsx')]
    assert scored({'format': 'xlsx', 'contains': ['North', '480'], 'sheets': ['^sales$'], 'charts_min': 1}, files, blobs) == []
    r = scored({'format': 'xlsx', 'sheets': ['Summary']}, files, blobs)
    assert r == ['x.xlsx has no sheet like /Summary/ (sheets: Sales)']


def test_xlsx_formula_injection_must_stay_text():
    """X2: an injected formula kept as text passes; the same cell written as a live formula fails any xlsx case, and
    formulas from the allow-list are fine."""
    safe = xlsx(extra=[('evil', "'=HYPERLINK(\"http://evil.example\")"), ('total', '=SUM(B2:B3)')], chart=False)
    live = xlsx(extra=[('evil', '=HYPERLINK("http://evil.example")')], chart=False)
    pipe = xlsx(extra=[('evil', "+SUM(1+1)*cmd|' /C calc'!A0")], chart=False)
    blobs = {'safe': safe, 'live': live, 'pipe': pipe}
    assert scored({'format': 'xlsx', 'contains': ['HYPERLINK']}, [made('safe', 'xlsx')], blobs) == []
    assert scored({'format': 'xlsx'}, [made('live', 'xlsx')], blobs) == ["live.xlsx has an unsafe formula '=HYPERLINK(\"http://evil.example\")'"]
    assert scored({'format': 'xlsx'}, [made('pipe', 'xlsx')], blobs) == []  # openpyxl keeps a leading + as text
    assert not evals.SAFE_FORMULA.match("=cmd|' /C calc'!A0") and evals.SAFE_FORMULA.match('=ROUND(AVERAGE(B2:B9), 2)')


def test_markdown():
    blobs = {'m': b'# Notes\n\n- 100 USD is about 92 EUR\n- Lisbon: sunny\n'}
    assert scored({'format': 'md', 'contains': ['USD', 'Lisbon']}, [made('m', 'md')], blobs) == []


# ---------- which file, rules, and failures ----------

def test_picks_the_file_in_the_expected_format():
    blobs = {'p': pdf(), 'd': docx()}
    files = [made('p', 'pdf'), made('d', 'docx')]
    assert scored({'format': 'docx', 'contains': ['interest']}, files, blobs) == []
    assert scored({'format': 'xlsx'}, files, blobs) == ['expected a xlsx file, got p.pdf, d.docx']
    assert scored({'format': 'pdf'}, [], blobs) == ['expected a pdf file, got none']


def test_expect_no_file():
    assert evals.check_file(False, rec([]), {}.__getitem__) == []
    assert evals.check_file(False, rec([made('p', 'pdf')]), {}.__getitem__) == ['expected no file, got p.pdf']


def test_rules_ok_fails_on_block_or_warn_only():
    rules = [{'id': 'S4', 'severity': 'fix', 'ok': False, 'note': 'markdown removed'},
             {'id': 'V2', 'severity': 'warn', 'ok': True, 'note': ''}]
    blobs = {'p': pdf()}
    assert scored({'format': 'pdf', 'rules_ok': True}, [made('p', 'pdf', rules)], blobs) == []
    warned = [*rules, {'id': 'V3', 'severity': 'warn', 'ok': False, 'note': 'page count'}]
    assert scored({'format': 'pdf', 'rules_ok': True}, [made('p', 'pdf', warned)], blobs) == ['p.pdf failed rules V3']
    assert scored({'format': 'pdf'}, [made('p', 'pdf', warned)], blobs) == []


def test_a_file_that_does_not_reopen_fails():
    r = scored({'format': 'pptx'}, [made('bad', 'pptx')], {'bad': b'not a zip'})
    assert len(r) == 1 and r[0].startswith('bad.pptx does not reopen')
    r = scored({'format': 'pdf'}, [made('gone', 'pdf')], {})
    assert r[0].startswith('gone.pdf does not reopen')


def test_turn_result_without_a_reader_cannot_pass():
    t = {'query': 'as pdf', 'expect_file': {'format': 'pdf'}}
    out = evals.turn_result(t, rec([made('p', 'pdf')]))
    assert not out['pass'] and out['reasons'] == ['created files cannot be checked here']


def test_read_created_uses_the_store_by_id(tmp_path):
    class Plain:
        files_dir = tmp_path / 'files'
    (tmp_path / 'created').mkdir()
    (tmp_path / 'created' / 'abc123').write_bytes(b'# hi')
    assert evals.read_created(Plain(), 'abc123') == b'# hi'
    with pytest.raises(ValueError):
        evals.read_created(Plain(), '../secret')

    class WithReader(Plain):
        def created_raw(self, fid):
            return b'from store ' + fid.encode()
    assert evals.read_created(WithReader(), 'x1') == b'from store x1'


# ---------- through run_eval ----------

async def test_run_eval_scores_created_files(monkeypatch):
    """A fake create step: the run record carries created_files (as the pipeline stores them) and the store serves the
    bytes. The harness reopens them per turn and fails the turn whose file misses. The step runs on an agent Jev is
    already offered (math), so this test does not depend on how the create agent is registered."""
    blobs = {'f1': pdf(pages=('TCP and UDP', 'UDP has no handshake')), 'f2': b'# Notes\n\nnothing about it\n'}

    async def create(text, emit):
        return AgentResult('Created **notes.pdf**, 2 pages', True)

    def route(text):
        return 'math', 0.9

    from tests.fakes import FakeJev
    from jevrouter.pipeline import Router
    r = Router(FakeJev(route_for=route), registry=registry(math=create))
    r.store.created_raw = blobs.__getitem__
    get_run = r.get_run

    def with_files(qid):
        out = get_run(qid)
        fid = 'f1' if 'pdf' in out['text'].lower() else 'f2'
        fmt = 'pdf' if fid == 'f1' else 'md'
        for t in out['tasks']:
            if t.get('agent') == 'math':
                t['created_files'] = [made(fid, fmt)]
        return out
    monkeypatch.setattr(r, 'get_run', with_files)
    case = {'id': 'cr', 'tags': ['create'], 'turns': [
        {'query': 'Put TCP vs UDP in a PDF', 'expect_agents': ['math'],
         'expect_file': {'format': 'pdf', 'contains': ['handshake'], 'min_pages': 2}},
        {'query': 'now as markdown', 'expect_file': {'format': 'md', 'contains': ['TCP']}}]}
    single = {'id': 'none', 'tags': ['create'], 'query': 'Put that in a PDF', 'expect_file': False}
    s = await run(r, [case, single])
    c, n = s['cases']
    assert [t['pass'] for t in c['turns']] == [True, False] and not c['pass']
    assert c['reasons'] == ['turn 2: f2.md does not contain /TCP/']
    assert not n['pass'] and n['reasons'] == ['expected no file, got f1.pdf']
