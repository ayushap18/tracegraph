"""Created files in the pipeline and the API (docs/PLAN-files.md): the create agent's zero-token paths and its one spec
call, the X4 safety check, block rules, storage by id, the /api/created endpoints, /api/rules, sandbox files kept in
memory only, and the research -> create plan."""
import asyncio
import io
import json
import sqlite3

import pytest
from aiohttp.test_utils import TestClient, TestServer

from jevrouter import app as appmod
from jevrouter import create as cf
from jevrouter.agents import AgentResult
from jevrouter.agents import create as create_agent
from jevrouter.config import CREATE_CONTEXT_CHARS, CREATE_MAX_TOKENS
from jevrouter.engines import Reply
from jevrouter.files import extract
from jevrouter.pipeline import Router
from jevrouter.planner import is_file_request, plan, split_file_request
from jevrouter.store import Store
from tests.fakes import FakeJev, ScriptEngine

SID = 'sandbox-create-1'
ANSWER = ('Ada Lovelace was an English mathematician.\n\n## Work\n\nShe wrote the first published algorithm for '
          "Babbage's Analytical Engine.\n\n- Born 1815\n- Died 1852")
RESEARCH = ('## Why they matter\n\nSolid-state batteries swap the liquid electrolyte for a solid one.\n\n'
            '## Trade-offs\n\n- Higher energy density\n- Harder to manufacture')
SPEC = {'title': 'Solar power', 'subtitle': '',
        'sections': [{'heading': 'How it works', 'level': 1, 'notes': '',
                      'blocks': [{'type': 'paragraph', 'text': 'Panels turn sunlight into electricity.'},
                                 {'type': 'bullets', 'items': ['Clean', 'Cheap to run'], 'ordered': False}]},
                     {'heading': 'Output', 'level': 1, 'notes': '',
                      'blocks': [{'type': 'table', 'columns': ['Year', 'TWh'], 'rows': [[2022, 1300], [2023, 1600]]}]}]}
FORMAT_WORDS = ('pdf', 'slides', 'deck', 'spreadsheet', 'excel', 'word document', 'markdown', 'docx', 'as ')


@pytest.fixture(autouse=True)
def legacy_renderer(monkeypatch):
    """These tests pin the standard renderer's behaviour (page and slide counts, calls to stub engines), so Studio is
    off here; tests/test_studio_api.py covers the same paths with Studio on."""
    monkeypatch.setenv('TG_STUDIO', 'off')


def route(text):
    t = text.lower()
    if any(w in t for w in FORMAT_WORDS) and ('put' in t or 'make' in t or 'turn' in t or 'now' in t or 'write' in t):
        return 'create', 0.9
    if 'research' in t or 'ada' in t:
        return 'knowledge', 0.9
    return 'chat', 0.9


def registry():
    async def knowledge(text, emit):
        answer = RESEARCH if 'research' in text.lower() else ANSWER
        emit(answer)
        return AgentResult(answer, True, 'wikipedia.org')

    async def chat(text, emit):
        emit('Hello there.')
        return AgentResult('Hello there.', True)
    return {'knowledge': knowledge, 'chat': chat}


class SpecEngine(ScriptEngine):
    """ScriptEngine that answers the create agent's spec call (schema DOCSPEC_SCHEMA) with `spec`, recording it."""

    def __init__(self, spec=None, merged=None, **kw):
        super().__init__(**kw)
        self.spec, self.merged, self.spec_calls = SPEC if spec is None else spec, merged, []

    async def stream(self, *, system, prompt, effort='medium', emit_delta=None, max_tokens=2048, web=False, schema=None,
                     exec=False):
        if schema is cf.DOCSPEC_SCHEMA:
            self.spec_calls.append({'system': system, 'prompt': prompt, 'effort': effort, 'max_tokens': max_tokens})
            return Reply(json.dumps(self.spec), 120, 80)
        if self.merged is not None and system.startswith('Do not use tools. You combine'):
            return Reply(self.merged, 3, 2)
        return await super().stream(system=system, prompt=prompt, effort=effort, emit_delta=emit_delta,
                                    max_tokens=max_tokens, web=web, schema=schema, exec=exec)


class UnsafeJev(FakeJev):
    """The safety check on a text on its own (not routing) flags anything that mentions `word`."""

    def __init__(self, word='forbidden', **kw):
        super().__init__(route_for=route, **kw)
        self.word = word

    async def system_one(self, state, qs):
        r = await super().system_one(state, qs)
        if 'route' not in qs and 'multi' not in qs and self.word in state:
            r.answers['unsafe'].noul = 0.95
        return r


def router(engine=None, jev=None, store=None) -> Router:
    return Router(jev or FakeJev(route_for=route), None, engine, registry=registry(), store=store or Store(),
                  engines={engine.name: engine} if engine else None)


async def ask(r: Router, query: str, session='s1', **kw) -> dict:
    events = []
    r.bus.taps.append(events.append)
    qid = next(r.ids)
    await r.handle(query, 'chat', qid, session_id=session, **kw)
    r.bus.taps.remove(events.append)
    return {'qid': qid, 'events': events, 'rec': r.get_run(qid)}


def made(run: dict) -> list[dict]:
    return [f for t in run['rec']['tasks'] for f in t.get('created_files') or []]


def answered(run: dict) -> list[dict]:
    return [e for e in run['events'] if e['type'] == 'answered']


def same_meta(got: dict, want: dict) -> bool:
    """The stored CreatedFile equals the one made, docs/PLAN-accuracy-v2.md C8 keys included."""
    return want == got


def spec_calls(engine) -> list:
    return engine.spec_calls if engine is not None else []


def add_csv(store: Store, name='sales.csv', rows=8) -> str:
    raw = ('region,revenue\n' + '\n'.join(f'{"North" if i % 2 else "West"},{100 + i}' for i in range(rows))).encode()
    meta, text = extract(name, raw)
    store.add_file(meta, raw, text)
    return meta['id']


def stored(store: Store, spec=SPEC, fmt='pdf', **kw) -> dict:
    meta, spec, data = create_agent.build(spec, fmt, source=kw.pop('source', 'llm'), **kw)
    store.add_created(meta, spec, data)
    return meta


# ---------- zero-token paths ----------

@pytest.mark.parametrize('engine', [None, SpecEngine()], ids=['keyless', 'engine'])
async def test_put_that_in_a_pdf_uses_the_previous_answer_with_no_spec_call(engine):
    r = router(engine)
    await ask(r, 'Who was Ada Lovelace')
    run = await ask(r, 'put that in a PDF')
    files = made(run)
    assert [f['format'] for f in files] == ['pdf'] and files[0]['source'] == 'answer' and files[0]['tokens'] == 0
    assert files[0]['qid'] == run['qid'] and files[0]['pages'] >= 1 and files[0]['sandbox'] is None
    assert spec_calls(engine) == []
    step = answered(run)[-1]
    assert step['agent'] == 'create' and step['created_files'] == files and step['engine'] == 'keyless'
    assert step['answer'].startswith(f"Created **{files[0]['name']}**, {files[0]['pages']} page")
    assert run['rec']['merged']['answer'].startswith('Created **')
    data = r.store.created_raw(files[0]['id'])
    text = '\n'.join(p.extract_text() for p in __import__('pypdf').PdfReader(io.BytesIO(data)).pages)
    assert 'Babbage' in text and 'Analytical Engine' in text
    assert r.store.created_spec(files[0]['id'])['title'] == 'Ada Lovelace'
    ids = [x['id'] for x in files[0]['rules']]
    assert {'S1', 'X4', 'V1', 'L4'} <= set(ids) and all(x['ok'] for x in files[0]['rules'] if x['severity'] == 'block')


@pytest.mark.parametrize('engine', [None, SpecEngine()], ids=['keyless', 'engine'])
async def test_now_as_slides_rerenders_the_last_file_at_zero_tokens(engine):
    r = router(engine)
    await ask(r, 'Who was Ada Lovelace')
    first = made(await ask(r, 'put that in a PDF'))[0]
    run = await ask(r, 'now as slides')
    files = made(run)
    assert [f['format'] for f in files] == ['pptx'] and files[0]['source'] == 'convert'
    assert files[0]['from_id'] == first['id'] and files[0]['tokens'] == 0 and files[0]['slides'] >= 2
    assert r.store.created_spec(files[0]['id']) == r.store.created_spec(first['id'])
    assert any(x['id'] == 'X4' and x['ok'] for x in files[0]['rules'])  # carried over from the checked source
    assert spec_calls(engine) == []
    # asking for the format it already is gives that file back, and makes nothing new
    again = await ask(r, 'now as slides')
    assert made(again)[0]['id'] == files[0]['id'] and 'already' in answered(again)[-1]['answer']
    assert len(r.store.list_created()) == 2


@pytest.mark.parametrize('engine', [None, SpecEngine()], ids=['keyless', 'engine'])
async def test_attached_table_becomes_a_spreadsheet_with_a_chart_at_zero_tokens(engine):
    r = router(engine)
    fid = add_csv(r.store)
    run = await ask(r, 'turn this into an Excel spreadsheet', files=[fid])
    f = made(run)[0]
    assert f['format'] == 'xlsx' and f['source'] == 'table' and f['tokens'] == 0 and f['sheets']
    assert spec_calls(engine) == []
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.store.created_raw(f['id'])))
    ws = wb.worksheets[0]
    assert (ws['A1'].value, ws['B1'].value, ws['B2'].value) == ('region', 'revenue', 100) and ws._charts


async def test_nothing_to_build_from_keyless_is_an_honest_answer():
    r = router()
    run = await ask(r, 'make slides about solar power', session=None)
    step = answered(run)[-1]
    assert step['agent'] == 'create' and not step['ok'] and 'created_files' not in step
    assert 'needs an LLM engine' in step['answer'] and r.store.list_created() == []
    run = await ask(r, 'put that in a PDF', session='fresh')
    assert 'no earlier answer' in answered(run)[-1]['answer'] and not made(run)


async def test_forced_create_works_keyless():
    r = router(jev=FakeJev(route_for=lambda t: ('chat', 0.9)))
    await ask(r, 'Who was Ada Lovelace')
    run = await ask(r, 'save it please as markdown', agent='create')
    assert made(run)[0]['format'] == 'md' and run['rec']['tasks'][0]['forced']
    assert r.store.created_raw(made(run)[0]['id']).decode().startswith('# Ada Lovelace')


async def test_a_vague_file_request_with_context_is_not_sent_to_clarify():
    r = router()
    await ask(r, 'Who was Ada Lovelace')
    r.jev.clear = 0.1  # Jev finds every request from here on unclear
    run = await ask(r, 'put that in a PDF')
    assert run['rec']['tasks'][0]['agent'] == 'create' and made(run)
    run = await ask(r, 'now as slides')
    assert run['rec']['tasks'][0]['agent'] == 'create' and made(run)[0]['format'] == 'pptx'
    fresh = await ask(r, 'put that in a PDF', session='other')  # nothing to refer to: Jev's clarify stands
    assert fresh['rec']['tasks'][0]['agent'] == 'clarify'


async def test_turns_that_did_not_answer_are_skipped():
    """A clarify between the file and "now as slides" doesn't hide the file, and is never put in a file itself."""
    r = router()
    await ask(r, 'Who was Ada Lovelace')
    pdf = made(await ask(r, 'put that in a PDF'))[0]
    r.jev.clear = 0.1
    # a knowledge pick Jev finds unclear (a chat pick is exempt from the clarity veto, docs/PLAN-accuracy-v2.md A6)
    assert (await ask(r, 'research hmm'))['rec']['tasks'][0]['agent'] == 'clarify'
    r.jev.clear = 0.9
    f = made(await ask(r, 'now as slides'))[0]
    assert f['format'] == 'pptx' and f['from_id'] == pdf['id']
    r2 = router()
    r2.jev.clear = 0.1
    await ask(r2, 'research hmm')
    r2.jev.clear = 0.9
    run = await ask(r2, 'put that in a PDF')
    assert not made(run) and 'no earlier answer' in answered(run)[-1]['answer']


# ---------- the one spec call ----------

async def test_fresh_content_costs_one_spec_call_with_the_schema_and_records_tokens():
    e = SpecEngine()
    r = router(e)
    run = await ask(r, 'make slides about solar power', session=None)
    assert len(e.spec_calls) == 1
    call = e.spec_calls[0]
    assert call['max_tokens'] == CREATE_MAX_TOKENS['pptx'] and call['effort'] == 'low'
    assert 'Request: make slides about solar power' in call['prompt'] and 'Format: PowerPoint' in call['prompt']
    f = made(run)[0]
    assert f['format'] == 'pptx' and f['source'] == 'llm' and f['tokens'] == 200 and f['title'] == 'Solar power'
    step = answered(run)[-1]
    assert step['engine'] == e.name and step['checks']['effort'] == 'low'
    assert run['rec']['tokens']['llm_in'] >= 120 and run['rec']['tokens']['llm_out'] >= 80


@pytest.mark.parametrize('query, fmt, effort', [
    ('make a spreadsheet of solar output', 'xlsx', 'low'), ('make slides about solar power', 'pptx', 'low'),
    ('make a PDF about solar power', 'pdf', 'low'), ('write a markdown file on solar power', 'md', 'low'),
    ('make a word document about solar power', 'docx', 'low'),
    ('make a detailed research report PDF on solar power', 'pdf', 'high')])
async def test_token_caps_per_format(query, fmt, effort):
    e = SpecEngine()
    run = await ask(router(e), query, session=None)
    assert e.spec_calls[0]['max_tokens'] == CREATE_MAX_TOKENS[fmt] and e.spec_calls[0]['effort'] == effort
    assert made(run)[0]['format'] == fmt


async def test_no_format_named_picks_one_and_says_so():
    e = SpecEngine()
    run = await ask(router(e), 'write up solar power for me', agent='create', session=None)
    f = made(run)[0]
    assert f['format'] == 'md' and 'No format was named, so this is Markdown' in answered(run)[-1]['answer']
    run = await ask(router(SpecEngine()), 'write a report on solar power', agent='create', session=None)
    assert made(run)[0]['format'] == 'pdf'


async def test_spec_call_context_is_trimmed_and_tables_send_schema_and_first_rows_only():
    e = SpecEngine()
    r = router(e)
    fid = add_csv(r.store, rows=40)
    long_answer = 'word ' * 20_000

    async def knowledge(text, emit):
        return AgentResult(long_answer, True)
    r.fixed_registry['knowledge'] = knowledge
    await ask(r, 'Who was Ada Lovelace')
    await ask(r, 'write an analysis report of this table as a PDF', files=[fid])
    prompt = e.spec_calls[0]['prompt']
    assert len(prompt) <= CREATE_CONTEXT_CHARS + 200
    assert 'columns: region, revenue' in prompt and 'West,100' in prompt and 'North,105' not in prompt


async def test_blocked_spec_makes_no_file():
    e = SpecEngine(spec={**SPEC, 'title': 'A forbidden recipe'})
    r = router(e, jev=UnsafeJev())
    run = await ask(r, 'make a PDF about chemistry', session=None)
    step = answered(run)[-1]
    assert not step['ok'] and 'Rule X4 blocked it' in step['answer'] and 'created_files' not in step
    assert r.store.list_created() == [] and list(r.store.created_dir.iterdir()) == []


async def test_a_block_rule_is_named_and_no_file_is_made():
    e = SpecEngine(spec={'title': 'Empty', 'subtitle': '', 'sections': []})
    r = router(e)
    run = await ask(r, 'make a PDF about nothing much', session=None)
    step = answered(run)[-1]
    assert not step['ok'] and step['answer'].startswith('No file was made. Rule S2 blocked it')
    assert r.store.list_created() == []


# ---------- research -> create ----------

async def test_keyless_splitter_plans_research_then_create():
    assert split_file_request('research solid-state batteries and make slides about it') == (
        'research solid-state batteries', 'make slides about it')
    assert split_file_request('weather in Paris, then put it in a PDF') == ('weather in Paris', 'put it in a PDF')
    assert split_file_request('salt and pepper') is None and split_file_request('make slides about solar') is None
    assert split_file_request('How do I open a file and save it as a PDF?') is None
    assert is_file_request('put that in a PDF') and not is_file_request('what is the word count')
    p = await plan('research solid-state batteries and make slides about it', FakeJev(multi=0.1))
    assert p['subtasks'] == ['research solid-state batteries', 'make slides about it'] and p['deps'] == [[], [0]]
    p = await plan('weather in Paris and time in Tokyo and put both into a PDF', FakeJev(multi=0.9))
    assert len(p['subtasks']) == 3 and p['deps'][-1] == [0, 1]


@pytest.mark.parametrize('engine', [None, 'engine'])
async def test_research_then_slides_builds_the_file_from_the_research_answer(engine):
    plan_ = {'subtasks': [{'text': 'research solid-state batteries', 'depends_on': []},
                          {'text': 'make slides from it', 'depends_on': [0]}]}
    e = SpecEngine(plan=plan_) if engine else None
    r = router(e)
    run = await ask(r, 'research solid-state batteries and make slides about it', session=None)
    subtasks = run['rec']['plan']['subtasks']
    assert len(subtasks) == 2 and subtasks[1]['depends_on'] == [subtasks[0]['tid']]
    f = made(run)[0]
    assert f['format'] == 'pptx' and f['source'] == 'answer' and f['tokens'] == 0
    assert spec_calls(e) == []
    if e:  # the file step is not rewritten by the engine either
        assert not any(c['system'].startswith('Do not use tools. Rewrite') for c in e.calls)
    spec = r.store.created_spec(f['id'])
    assert spec['title'] == 'Solid-state batteries'
    assert [s['heading'] for s in spec['sections']] == ['Why they matter', 'Trade-offs']
    assert f['name'] in run['rec']['merged']['answer']


async def test_llm_merge_keeps_the_created_file_named():
    plan_ = {'subtasks': [{'text': 'research solid-state batteries', 'depends_on': []},
                          {'text': 'make slides from it', 'depends_on': [0]}]}
    e = SpecEngine(plan=plan_, merged='Here is what I found.')

    async def knowledge(text, emit):
        return AgentResult(RESEARCH, True, engine='claude-code')  # an LLM answer, so the LLM merger runs
    r = router(e)
    r.fixed_registry['knowledge'] = knowledge
    run = await ask(r, 'research solid-state batteries and make slides about it', session=None)
    f = made(run)[0]
    # however the answer is merged (docs/PLAN-accuracy-v2.md B2 leads with the file), it names the file exactly once
    assert run['rec']['merged']['answer'].count(f'**{f["name"]}**') == 1


# ---------- HTTP ----------

@pytest.fixture
async def client():
    box = {}

    def factory(http):
        box['r'] = Router(FakeJev(route_for=route), http, None, store=Store(), registry=registry())
        return box['r']
    async with TestClient(TestServer(appmod.create_app(factory))) as c:
        c.router = box['r']
        yield c


async def finish(client, qid):
    task = client.router.running.get(qid)
    if task:
        await asyncio.wait_for(task, 10)


async def json_of(resp, status=200):
    assert resp.status == status, await resp.text()
    return await resp.json()


async def test_keyless_ask_then_pdf_then_download_and_convert(client):
    body = await json_of(await client.post('/ask', json={'query': 'Who was Ada Lovelace', 'source': 'chat'}))
    await finish(client, body['qid'])
    body = await json_of(await client.post('/ask', json={'query': 'put that in a PDF', 'source': 'chat',
                                                          'session_id': body['session_id']}))
    await finish(client, body['qid'])
    run = await json_of(await client.get(f'/api/runs/{body["qid"]}'))
    f = run['tasks'][0]['created_files'][0]
    assert f['format'] == 'pdf' and f['tokens'] == 0 and run['tokens']['llm_in'] == run['tokens']['llm_out'] == 0
    resp = await client.get(f'/api/created/{f["id"]}/download')
    data = await resp.read()
    assert resp.status == 200 and data.startswith(b'%PDF') and len(data) == f['size']
    assert resp.headers['Content-Type'] == 'application/pdf'
    assert resp.headers['Content-Disposition'] == (f'attachment; filename="{f["name"]}"; '
                                                   f"filename*=UTF-8''{f['name']}")
    assert resp.headers['X-Content-Type-Options'] == 'nosniff'
    docx = await json_of(await client.post(f'/api/created/{f["id"]}/convert', json={'format': 'docx'}), 201)
    assert docx['format'] == 'docx' and docx['tokens'] == 0 and docx['source'] == 'convert' and docx['from_id'] == f['id']
    resp = await client.get(f'/api/created/{docx["id"]}/download')
    assert resp.headers['Content-Type'] == cf.MIME['docx'] and (await resp.read())[:2] == b'PK'


async def test_created_endpoints(client):
    store = client.router.store
    a = stored(store, fmt='pdf')
    b = stored(store, fmt='md')
    files = (await json_of(await client.get('/api/created')))['files']
    assert [f['id'] for f in files] == [b['id'], a['id']]
    assert [f['id'] for f in (await json_of(await client.get(f'/api/created?limit=1&before={b["created"]}')))['files']] == [a['id']]
    assert (await client.get('/api/created?limit=x')).status == 400
    assert (await client.get('/api/created?before=soon')).status == 400
    assert same_meta(await json_of(await client.get(f'/api/created/{a["id"]}')), a)
    for bad in ('0123456789ab', 'nope', 'ABCDEF123456'):
        assert (await client.get(f'/api/created/{bad}')).status == 404
        assert (await client.get(f'/api/created/{bad}/download')).status == 404
        assert (await client.get(f'/api/created/{bad}/preview')).status == 404
        assert (await client.delete(f'/api/created/{bad}')).status == 404
    pv = await json_of(await client.get(f'/api/created/{a["id"]}/preview'))
    assert pv['kind'] == 'outline' and pv['pages'] == a['pages'] and pv['items'][0]['text'] == 'Solar power'
    pv = await json_of(await client.get(f'/api/created/{b["id"]}/preview'))
    assert pv['kind'] == 'markdown' and pv['text'].startswith('# Solar power')
    resp = await client.get(f'/api/created/{b["id"]}/download')
    assert resp.headers['Content-Type'] == 'text/markdown; charset=utf-8'
    assert (await json_of(await client.delete(f'/api/created/{b["id"]}'))) == {'ok': True}
    assert (await client.get(f'/api/created/{b["id"]}')).status == 404
    assert not (store.created_dir / b['id']).exists()


async def test_convert_endpoint_errors_and_chain(client):
    store = client.router.store
    src = stored(store, fmt='pdf')
    url = f'/api/created/{src["id"]}/convert'
    assert (await client.post(url, json={'format': 'exe'})).status == 400
    assert (await client.post(url, json={})).status == 400
    assert (await client.post(url, data='not json')).status == 400
    assert (await client.post(url, json={'format': 'pdf'})).status == 409
    assert (await client.post('/api/created/0123456789ab/convert', json={'format': 'md'})).status == 404
    chain, cur = [], src
    for fmt in ('docx', 'pptx', 'md'):
        cur = await json_of(await client.post(f'/api/created/{cur["id"]}/convert', json={'format': fmt}), 201)
        chain.append(cur)
    assert [f['format'] for f in chain] == ['docx', 'pptx', 'md']
    assert [f['from_id'] for f in chain] == [src['id'], chain[0]['id'], chain[1]['id']]
    assert all(f['tokens'] == 0 and f['source'] == 'convert' for f in chain)
    assert chain[1]['slides'] >= 3
    md = store.created_raw(chain[2]['id']).decode()
    assert md.startswith('# Solar power') and '| Year | TWh |' in md
    assert store.created_spec(chain[2]['id']) == store.created_spec(src['id'])


async def test_rules_and_agents(client):
    rules = (await json_of(await client.get('/api/rules')))['rules']
    ids = [r['id'] for r in rules]
    assert len(ids) == len(set(ids)) == len(cf.RULES) and {'S1', 'L5', 'X4', 'V4', 'A4'} <= set(ids)
    assert all(set(r) == {'id', 'group', 'text', 'severity', 'enforced'} for r in rules)
    agents = {a['name']: a for a in (await json_of(await client.get('/api/agents')))['agents']}
    assert agents['create']['available'] and not agents['create']['engine_required']
    assert (await client.post('/api/agents', json={'name': 'create', 'description': 'x' * 20, 'prompt': 'y' * 20})).status == 400


def test_disposition_is_safe_for_any_name():
    d = appmod.disposition('a"b\r\nX-Evil: 1;.pdf')
    assert '\r' not in d and '\n' not in d and d.count('"') == 2
    assert d.startswith('attachment; filename="a_b__X-Evil__1_.pdf"; filename*=UTF-8\'\'a%22b%0D%0AX-Evil%3A%201%3B.pdf')
    assert appmod.disposition('..hidden.pdf').startswith('attachment; filename="hidden.pdf"')


async def test_download_rejects_path_traversal(client):
    store = client.router.store
    secret = store.created_dir.parent / 'secret.txt'
    secret.write_text('top secret')
    try:
        for path in ('/api/created/..%2Fsecret.txt/download', '/api/created/%2E%2E%2Fsecret.txt/download',
                     '/api/created/../secret.txt/download', '/api/created/..%5Csecret.txt/download',
                     f'/api/sandbox/{SID}/created/..%2F..%2Fsecret.txt/download',
                     '/api/created/%2Fetc%2Fpasswd/download'):
            resp = await client.get(path)
            assert resp.status in (400, 404), path
            assert b'top secret' not in await resp.read() and b'root:' not in await resp.read()
        with pytest.raises(ValueError):
            store.created_path('../secret.txt')
        with pytest.raises(ValueError):
            store.created_raw('/etc/passwd')
    finally:
        secret.unlink()


async def test_sandbox_created_files_live_in_memory_only(client):
    store = client.router.store
    for q in ('Who was Ada Lovelace', 'put that in a PDF'):
        body = await json_of(await client.post('/ask', json={'query': q, 'source': 'sandbox', 'sandbox_id': SID}))
        await finish(client, body['qid'])
    mem = client.router.sandboxes.peek(SID)
    (fid, (meta, spec, data)), = mem.created.items()
    assert meta['sandbox'] == SID and meta['format'] == 'pdf' and meta['tokens'] == 0
    assert mem.thread[-1]['record']['tasks'][0]['created_files'] == [meta]
    # nothing in the store or on disk
    assert store.list_created() == [] and list(store.created_dir.iterdir()) == []
    assert (await client.get(f'/api/created/{fid}')).status == 404
    assert (await client.get(f'/api/created/{fid}/download')).status == 404
    resp = await client.get(f'/api/sandbox/{SID}/created/{fid}/download')
    assert resp.status == 200 and await resp.read() == data and resp.headers['X-Content-Type-Options'] == 'nosniff'
    assert (await json_of(await client.get(f'/api/sandbox/{SID}/created/{fid}/preview')))['kind'] == 'outline'
    assert (await client.get(f'/api/sandbox/sandbox-someone-else/created/{fid}/download')).status == 404
    assert (await client.get(f'/api/sandbox/x/created/{fid}/download')).status == 400
    # "now as markdown" in the sandbox converts from its memory, still off disk
    body = await json_of(await client.post('/ask', json={'query': 'now as markdown', 'source': 'sandbox',
                                                          'sandbox_id': SID}))
    await finish(client, body['qid'])
    assert [m['format'] for m, _, _ in mem.created.values()] == ['pdf', 'md'] and store.list_created() == []
    assert (await client.delete(f'/api/sandbox/{SID}')).status == 200
    assert (await client.get(f'/api/sandbox/{SID}/created/{fid}/download')).status == 404
    assert (await client.get(f'/api/sandbox/{SID}/created/{fid}/preview')).status == 404


async def test_keeping_a_sandbox_saves_its_files():
    r = router()
    for q in ('Who was Ada Lovelace', 'put that in a PDF'):
        qid = r.submit(q, 'sandbox', sandbox=SID)
        await r.running[qid]
    (fid, (_, spec, data)), = r.sandboxes.peek(SID).created.items()
    session, qids = r.keep_sandbox(SID)
    kept = r.store.get_created(fid)
    assert kept['qid'] == qids[1] and kept['sandbox'] is None and r.store.created_raw(fid) == data
    assert r.store.created_spec(fid) == spec
    assert r.get_run(qids[1])['tasks'][0]['created_files'][0]['sandbox'] is None


def test_sandbox_keeps_at_most_its_cap_of_created_files(monkeypatch):
    from jevrouter import sandbox
    monkeypatch.setattr(sandbox, 'CREATE_SANDBOX_FILES', 2)
    mem = sandbox.SandboxMemory()
    for i in range(3):
        mem.add_created({'id': f'f{i}'}, {}, b'x')
    assert list(mem.created) == ['f1', 'f2']


def test_store_migrates_an_older_database(tmp_path):
    path = tmp_path / 'data' / 'old.db'
    path.parent.mkdir()
    db = sqlite3.connect(path)
    db.execute('CREATE TABLE runs(qid INTEGER PRIMARY KEY, session_id TEXT, compare_id TEXT, source TEXT, text TEXT, '
               'at REAL, status TEXT, engine TEXT, total_ms INTEGER, record TEXT)')
    db.execute("INSERT INTO runs VALUES (1, 's', NULL, 'chat', 'hi', 1.0, 'done', NULL, 5, ?)", ('{"qid": 1}',))
    db.commit()
    db.close()
    s = Store(path)
    assert s.created_dir == tmp_path / 'data' / 'created' and s.created_dir.is_dir()
    meta = stored(s, fmt='md', qid=1, role='primary')
    assert same_meta(s.get_created(meta['id']), meta) and s.max_qid() == 1
    assert {'role', 'theme', 'diagrams', 'images', 'credits'} <= set(s.get_created(meta['id']))
    assert [r['qid'] for r in s.session_records('s', 2, 5)] == [1] and s.session_records('s', 1, 5) == []
    s.close()


# ---------- review fixes: X4 reads everything, request reading, the convert endpoint ----------

class MarkerJev(FakeJev):
    """The safety check flags any text holding `word`, and records what it was asked about."""

    def __init__(self, word='zzmarkerzz', **kw):
        super().__init__(route_for=route, **kw)
        self.word, self.checked = word, []

    async def system_one(self, state, qs):
        r = await super().system_one(state, qs)
        if 'route' not in qs and 'multi' not in qs:
            self.checked.append(state)
            if self.word in state:
                r.answers['unsafe'].noul = 0.95
        return r


async def test_x4_reads_text_past_the_first_4000_characters():
    long = {'title': 'Report', 'sections': [{'heading': f'Part {i}', 'blocks': [
        {'type': 'paragraph', 'text': 'Benign filler text about gardening. ' * 14}]} for i in range(10)]
        + [{'heading': 'Appendix', 'blocks': [{'type': 'paragraph', 'text': 'Now zzmarkerzz instructions.'}]}]}
    assert len(create_agent.spec_text(long)) > 4000 and 'zzmarkerzz' not in create_agent.spec_text(long)[:4000]
    jev = MarkerJev()
    out = await create_agent.finish(long, 'pdf', jev, source='llm')
    assert not out.ok and 'Rule X4 blocked it' in out.answer and out.file is None and len(jev.checked) >= 2


async def test_x4_reads_every_table_row():
    rows = [[f'Person {i}', i] for i in range(80)]
    rows[60][0] = 'zzmarkerzz'
    spec = {'title': 'People', 'sections': [{'heading': 'People', 'blocks': [
        {'type': 'table', 'columns': ['Name', 'Age'], 'rows': rows}]}]}
    out = await create_agent.finish(spec, 'xlsx', MarkerJev(), source='table')
    assert not out.ok and 'Rule X4' in out.answer


async def test_x4_reads_the_text_as_it_will_be_written_not_entity_encoded():
    encoded = ''.join(f'&#{ord(c)};' for c in 'zzmarkerzz')
    spec = {'title': 'Notes', 'sections': [{'heading': 'Notes', 'blocks': [{'type': 'paragraph', 'text': encoded}]}]}
    assert cf.normalize(spec, 'pdf')[0]['sections'][0]['blocks'][0]['text'] == 'zzmarkerzz'
    out = await create_agent.finish(spec, 'pdf', MarkerJev(), source='llm')
    assert not out.ok and 'Rule X4' in out.answer


async def test_x4_checks_series_names_and_code():
    for block in ({'type': 'chart', 'kind': 'bar', 'title': 'c', 'labels': ['a'], 'series': [{'name': 'zzmarkerzz', 'values': [1]}]},
                  {'type': 'code', 'lang': 'py', 'text': 'print("zzmarkerzz")'}):
        spec = {'title': 'T', 'sections': [{'heading': 'H', 'blocks': [block]}]}
        out = await create_agent.finish(spec, 'md', MarkerJev(), source='llm')
        assert not out.ok and 'Rule X4' in out.answer, block['type']


class PlanEngine(SpecEngine):
    """An engine whose LLM planner rewrites file follow-ups into self-contained steps, as its prompt used to ask."""

    async def stream(self, *, system, prompt, **kw):
        if system.startswith('You are a query planner'):
            q = prompt.rsplit('Current query:', 1)[-1]
            if 'put that in a PDF' in q:
                return Reply(json.dumps({'subtasks': [{'text': 'Put the information about Ada Lovelace in a PDF',
                                                       'depends_on': []}]}), 50, 20)
            if 'now as slides' in q:
                return Reply(json.dumps({'subtasks': [{'text': 'Make slides about Ada Lovelace', 'depends_on': []}]}),
                             50, 20)
        return await super().stream(system=system, prompt=prompt, **kw)


async def test_a_planner_rewrite_does_not_hide_a_zero_token_file_request():
    e = PlanEngine()
    r = router(e)
    await ask(r, 'Who was Ada Lovelace')
    pdf = await ask(r, 'put that in a PDF')
    assert pdf['rec']['tasks'][0]['text'] == 'Put the information about Ada Lovelace in a PDF'  # the rewrite happened
    assert [(f['format'], f['source'], f['tokens']) for f in made(pdf)] == [('pdf', 'answer', 0)]
    slides = await ask(r, 'now as slides')
    assert [(f['format'], f['source'], f['tokens']) for f in made(slides)] == [('pptx', 'convert', 0)]
    assert e.spec_calls == []


@pytest.mark.parametrize('engine', [None, SpecEngine()], ids=['keyless', 'engine'])
async def test_an_earlier_step_wins_over_the_last_file(engine):
    r = router(engine)
    await ask(r, 'Who was Ada Lovelace')
    pdf = made(await ask(r, 'put that in a PDF'))[0]
    run = await ask(r, 'research solid-state batteries and make slides about it')
    files = made(run)
    assert [f['format'] for f in files] == ['pptx'] and files[0]['source'] == 'answer'
    assert files[0]['from_id'] != pdf['id'] and 'Ada' not in files[0]['title']


async def test_an_attached_table_wins_over_the_last_file():
    r = router(SpecEngine())
    await ask(r, 'Who was Ada Lovelace')
    pdf = made(await ask(r, 'put that in a PDF'))[0]
    fid = add_csv(r.store)
    files = made(await ask(r, 'turn this into an Excel spreadsheet', files=[fid]))
    assert [(f['format'], f['source']) for f in files] == [('xlsx', 'table')] and files[0]['from_id'] != pdf['id']


@pytest.mark.parametrize('query,fmt', [('I also want an Excel spreadsheet', 'xlsx'), ('make a word document too', 'docx'),
                                       ('can I get slides', 'pptx')])
async def test_another_format_right_after_a_file_converts_it_not_its_reply(query, fmt):
    r = router()
    await ask(r, 'Who was Ada Lovelace')
    pdf = made(await ask(r, 'put that in a PDF'))[0]
    files = made(await ask(r, query, agent='create'))
    assert [(f['format'], f['source'], f['from_id']) for f in files] == [(fmt, 'convert', pdf['id'])]
    spec = r.store.created_spec(files[0]['id'])
    assert 'Created **' not in json.dumps(spec)


async def test_a_turn_that_only_made_a_file_gives_no_answer_to_build_on():
    r = router()
    await ask(r, 'Who was Ada Lovelace')
    rec = (await ask(r, 'put that in a PDF'))['rec']
    turns, last = r.chat_files({'session_id': 's1'}, None, rec['qid'] + 1)
    assert [t['answer'][:12] for t in turns] == ['Ada Lovelace', ''] and last[2] is True
    assert create_agent.only_created(rec)


@pytest.mark.parametrize('engine', [None, SpecEngine()], ids=['keyless', 'engine'])
async def test_notes_of_our_conversation_cover_every_turn(engine):
    ctx = [{'query': 'Convert 100 USD to EUR', 'answer': '100 USD is 92.10 EUR.'},
           {'query': "What's the weather in Lisbon?", 'answer': 'Lisbon: 22 C, sunny.'}]
    out = await create_agent.make(create_agent.Job('Write markdown notes of our conversation', context=ctx), engine, None)
    text = out.data.decode()
    assert out.ok and out.file['source'] == 'answer' and 'USD' in text and 'Lisbon' in text
    assert spec_calls(engine) == []


async def test_convert_endpoint_turns_an_unexpected_failure_into_422(client, monkeypatch):
    src = stored(client.router.store, fmt='pdf')

    def boom(*a, **kw):
        raise RuntimeError('renderer broke')
    monkeypatch.setattr(create_agent, 'build', boom)
    resp = await client.post(f'/api/created/{src["id"]}/convert', json={'format': 'docx'})
    assert resp.status == 422 and 'renderer broke' in (await resp.json())['error']


async def test_convert_endpoint_makes_a_pdf_of_a_table_with_huge_rows(client):
    cells = [' '.join(['word'] * 60)] * 15
    spec = {'title': 'Survey', 'sections': [{'heading': 'Answers', 'blocks': [
        {'type': 'table', 'columns': [f'Q{i}' for i in range(15)], 'rows': [cells] * 3}]}]}
    src = stored(client.router.store, spec=spec, fmt='xlsx', source='table')
    pdf = await json_of(await client.post(f'/api/created/{src["id"]}/convert', json={'format': 'pdf'}), 201)
    assert pdf['format'] == 'pdf' and pdf['pages'] >= 1


async def test_excel_from_an_attached_table_keeps_the_data_after_an_analysis_step():
    """"Turn this sales data into an Excel sheet with a chart": the planner may add an analysis step first; the sheet
    still holds the table (with a chart), and the analysis rides along as notes."""
    import io
    import openpyxl
    meta = {'id': 'f1', 'name': 'sales.csv', 'columns': ['region', 'product', 'revenue']}
    rows = [['North', 'Widget', 1200], ['West', 'Gadget', 900], ['South', 'Widget', 700]]
    job = create_agent.Job('Turn this sales data into an Excel sheet with a chart of revenue by region',
                           deps=[('analyze the sales data', 'North leads with 1,200 in revenue.')],
                           tables=[(meta, meta['columns'], rows)])
    out = await create_agent.make(job, None, None)
    assert out.ok and out.file['format'] == 'xlsx' and out.file['source'] == 'table' and out.file['tokens'] == 0
    wb = openpyxl.load_workbook(io.BytesIO(out.data))
    cells = {str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None}
    assert {'North', 'West', 'Widget'} <= cells and any('North leads' in c for c in cells)
    assert any(ws._charts for ws in wb.worksheets)


# ---------- resume and checkpoints (docs/PLAN-files-robust.md 3.4, 3.5; builder E) ----------

def resume_state(written: int = 1, file_id=None) -> dict:
    parts = [{'heading': f'Part {i}', 'level': 1, 'words': 45, 'hints': [], 'diagrams': [], 'figures': 0} for i in range(3)]
    return {'v': 1, 'kind': 'longdoc', 'format': 'pptx', 'request': 'make a 4 slide deck about tea', 'brief': None,
            'theme': None, 'font': None, 'design': None, 'ctx': 'Tea is a drink.', 'title': 'Tea', 'subtitle': '',
            'parts': parts,
            'written': {str(i): {'heading': f'Part {i}', 'level': 1, 'blocks': [{'type': 'paragraph', 'text': 'Tea.'}],
                                 'notes': ''} for i in range(written)},
            'failed': [], 'phase': 'sections', 'engine': 'claude-code', 'tokens_in': 500, 'tokens_out': 200, 'calls': 2,
            'at': 1.0, 'file_id': file_id}


def saved_run(store: Store, qid: int, state: dict, status='timeout') -> dict:
    rec = {'qid': qid, 'text': 'make a 4 slide deck about tea', 'source': 'chat', 'at': 1.0, 'plan': None, 'tasks': [],
           'merged': None, 'total_ms': 10, 'error': None, 'status': status, 'engine': None, 'session_id': 'sess-resume',
           'compare_id': None, 'files': [], 'tokens': {'jev_in': 0, 'llm_in': 0, 'llm_out': 0}, 'mode': 'research',
           'style': 'default', 'agent': 'create', 'group_id': None, 'chosen': True,
           'checkpoints_state': {f'{qid}.2': state}}
    store.touch_session('sess-resume', rec['text'])
    store.save_run(rec)
    return rec


async def test_resume_endpoint(client, monkeypatch):
    seen = []

    async def make(job, engine, jev, mode='balanced'):
        seen.append(job)
        return create_agent.Made('No file was made: keyless test.', False)
    monkeypatch.setattr(create_agent, 'make', make)
    store = client.router.store
    saved_run(store, 41, resume_state())
    saved_run(store, 42, resume_state(written=3, file_id='0123456789ab'))
    url = '/api/created/resume'
    assert (await client.post(url, json={'qid': 41, 'tid': '41.9'})).status == 404
    assert (await json_of(await client.post(url, json={'qid': 99, 'tid': '99.1'}), 404))['error'] == \
        'no checkpoint for that file step'
    assert (await client.post(url, json={'qid': '41', 'tid': '41.2'})).status == 400
    assert (await client.post(url, json={'qid': 41, 'tid': ''})).status == 400
    assert (await client.post(url, json={'qid': 41, 'tid': '41.2', 'confirm_cost': 1})).status == 400
    assert (await json_of(await client.post(url, json={'qid': 42, 'tid': '42.2'}), 409))['error'] == \
        'That file is already complete.'
    got = await json_of(await client.post(url, json={'qid': 41, 'tid': '41.2'}), 202)
    assert got['ok'] and got['qid'] not in (41, 42) and got['session_id'] == 'sess-resume'
    assert got['estimate']['keyless'] and got['estimate']['calls'] == 0  # this client is keyless
    await finish(client, got['qid'])
    # the job gets {qid, tid, state} so the new file's resumed_from names the source run and step
    assert seen and seen[-1].resume == {'qid': 41, 'tid': '41.2', 'state': resume_state()}
    new = await json_of(await client.get(f'/api/runs/{got["qid"]}'))
    assert new['agent'] == 'create' and new['text'] == 'make a 4 slide deck about tea' and new['mode'] == 'balanced'
    assert new['tasks'][0]['reason'] == 'resume a partial file'


async def test_checkpoint_state_never_leaves_the_server(client):
    store = client.router.store
    saved_run(store, 51, resume_state())
    client.router.history.append(store.get_run(51))
    run = await json_of(await client.get('/api/runs/51'))
    assert 'checkpoints_state' not in run and run['checkpoints'][0]['tid'] == '51.2'
    assert run['checkpoints'][0]['missing'] == ['Part 1', 'Part 2'] and run['checkpoints'][0]['resumable']
    runs = (await json_of(await client.get('/api/runs')))['runs']
    assert runs and all('checkpoints_state' not in r for r in runs)
    session = await json_of(await client.get('/api/sessions/sess-resume'))
    assert all('checkpoints_state' not in r for r in session['runs'])
    cps = await json_of(await client.get('/api/runs/51/checkpoints'))
    assert cps == {'checkpoints': run['checkpoints']}
    assert (await client.get('/api/runs/999/checkpoints')).status == 404
    hello = client.router.hello()
    assert all('checkpoints_state' not in r for r in hello['history'])
    assert store.get_run(51)['checkpoints_state']  # kept on the server for Resume


async def test_resume_is_priced_and_guarded(monkeypatch):
    monkeypatch.setenv('TG_COST_CONFIRM', '1')
    seen = []

    async def make(job, engine, jev, mode='balanced'):
        seen.append(job)
        return create_agent.Made('No file was made: test.', False)
    monkeypatch.setattr(create_agent, 'make', make)
    box = {}

    def factory(http):
        e = SpecEngine(name='agy', label='Antigravity')
        box['r'] = Router(FakeJev(route_for=route), None, e, store=Store(), registry=registry(), engines={'agy': e})
        return box['r']
    async with TestClient(TestServer(appmod.create_app(factory))) as c:
        store = box['r'].store
        state = resume_state(written=0)
        state['parts'] = state['parts'] * 7  # 21 missing slides on Antigravity
        saved_run(store, 61, state)
        before = store.max_qid()
        got = await json_of(await c.post('/api/created/resume', json={'qid': 61, 'tid': '61.2'}), 409)
        assert got['needs_confirmation'] and got['estimate']['needs_confirmation'] and store.max_qid() == before
        assert not seen and got['estimate']['summary'].startswith('Resuming this file needs about')
        got = await json_of(await c.post('/api/created/resume', json={'qid': 61, 'tid': '61.2', 'confirm_cost': True}),
                            202)
        task = box['r'].running.get(got['qid'])
        if task:
            await asyncio.wait_for(task, 10)
        assert seen and seen[-1].resume['state']['parts'] == state['parts'] and seen[-1].resume['qid'] == 61


async def test_a_design_file_from_an_earlier_turn_is_found_when_the_request_names_one(monkeypatch):
    """H5: "use the design.md" with nothing attached to this message finds the design attached earlier in the chat."""
    from pathlib import Path
    seen = []

    async def make(job, engine, jev, mode='balanced'):
        seen.append(job)
        return create_agent.Made('No file was made: test.', False)
    monkeypatch.setattr(create_agent, 'make', make)
    store = Store()
    text = (Path(__file__).resolve().parent.parent / 'evals' / 'fixtures' / 'design_system.md').read_text()
    design, _ = extract('DESIGN-lovable.md', text.encode())
    store.add_file(design, text.encode(), text)
    notes, _ = extract('notes.txt', b'meeting notes: we agreed to ship on Friday')
    store.add_file(notes, b'meeting notes: we agreed to ship on Friday', 'meeting notes: we agreed to ship on Friday')
    earlier = {'qid': 5, 'text': 'here are my files', 'source': 'chat', 'at': 1.0, 'plan': None,
               'tasks': [{'tid': '5.1', 'agent': 'chat', 'ok': True, 'answer': 'Got them.'}], 'merged': {'answer': 'Got them.'},
               'total_ms': 1, 'error': None, 'status': 'done', 'engine': None, 'session_id': 's-design', 'compare_id': None,
               'files': [notes['id'], design['id']], 'tokens': {'jev_in': 0, 'llm_in': 0, 'llm_out': 0}}
    store.touch_session('s-design', earlier['text'])
    store.save_run(earlier)
    r = router(store=store)
    await ask(r, 'make slides about tea using the design.md', session='s-design')
    assert [m['name'] for m, _ in seen[-1].design_docs] == ['DESIGN-lovable.md']
    assert 'Color Palette' in seen[-1].design_docs[0][1]
    await ask(r, 'make slides about tea', session='s-design')  # no design named: nothing is looked up
    assert seen[-1].design_docs == []
